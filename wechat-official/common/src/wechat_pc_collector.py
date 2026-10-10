from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import signal
import sqlite3
import subprocess
import sys
import time
import shutil
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
import win32clipboard
from PIL import ImageGrab
from pywinauto import Desktop, keyboard, mouse

# Windows consoles often default to GBK; force UTF-8 so real WeChat titles
# containing Japanese punctuation/emoji cannot terminate a long collection.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


WINDOW_TITLE = "微信"
DEFAULT_START_DATE = "2026-01-01"
STOP_REQUESTED = False
SEARCH_LAYOUT: str | None = None


class SearchWindowUnavailableError(RuntimeError):
    """The WeChat 搜一搜 surface disappeared while Chromium was navigating."""


class ResultPositionError(RuntimeError):
    """A search result exists but Chromium could not make it safely clickable."""


class ArticleMetadataUnavailableError(RuntimeError):
    """The article and its search card expose no trustworthy publish date."""


def request_graceful_stop(_signum=None, _frame=None) -> None:
    global STOP_REQUESTED
    STOP_REQUESTED = True
    print("Graceful stop requested. Finishing the current article, then exporting and importing committed rows.", flush=True)
ODS_FIELDS = [
    "id", "文章图片链接", "文章图片信息", "作者名字", "文章标题", "文章内容", "文章标签",
    "文章时间地点", "文章点赞数", "文章收藏数", "文章评论数", "文章链接", "文章id", "作者主页链接",
    "文章视频链接", "文章视频内容", "关键词", "作者id", "作者粉丝数", "作者点赞数", "文章分享数",
    "文章视频时长", "平台", "文章获取时间", "文章标注", "分类", "清洗标注", "项目", "地点",
    "内容类型", "是否广告", "广告关键词",
]


@dataclass
class ResultRow:
    project: str
    keyword: str
    title: str
    url: str
    author_name: str | None = None
    published_at: str | None = None
    location: str | None = None
    content: str | None = None
    image_count: int = 0
    video_count: int = 0
    like_count: int | None = None
    comment_count: int | None = None
    favorite_count: int | None = None
    share_count: int | None = None
    captured_at: str = ""

    def __post_init__(self) -> None:
        if not self.captured_at:
            self.captured_at = datetime.now().isoformat(timespec="seconds")


def load_config(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(f"配置文件不存在：{path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _wechat_main_window_candidates():
    # The PC WeChat main window is Qt-based. The separate 搜一搜 article
    # window has the same title but is Chromium-based, so do not count it as
    # a second main window.
    candidates = []
    for candidate in Desktop(backend="uia").windows():
        try:
            class_name = candidate.element_info.class_name or ""
            rect = candidate.rectangle()
            # Legacy accounts use a Qt51514... top-level window; newer account
            # rollouts use mmui::MainWindow.  Do not depend on the title: some
            # Windows locale combinations expose 微信 as mojibake through UIA.
            ui_family = class_name.startswith("Qt51514") or class_name.startswith("mmui::MainWindow")
            main_sized = rect.width() >= 600 and rect.height() >= 500
            if ui_family and candidate.is_visible() and main_sized:
                candidates.append(candidate)
        except Exception:
            continue
    return candidates


def _wechat_hidden_main_candidates():
    candidates = []
    for candidate in Desktop(backend="uia").windows(visible_only=False):
        try:
            class_name = candidate.element_info.class_name or ""
            if class_name.startswith("Qt51514") or class_name.startswith("mmui::MainWindow"):
                if not class_name.startswith("mmui::LoginWindow"):
                    candidates.append(candidate)
        except Exception:
            continue
    return candidates


def _wechat_login_windows():
    matches = []
    for candidate in Desktop(backend="uia").windows():
        try:
            if (candidate.element_info.class_name or "").startswith("mmui::LoginWindow"):
                matches.append(candidate)
        except Exception:
            continue
    return matches


def wake_wechat_main_window() -> bool:
    """Wake WeChat when all background processes remain but its UI is hidden.

    Launching the already-running executable is WeChat's supported single-
    instance activation path and restores a tray-hidden main window.  The
    executable is supplied by the environment probe, with standard install
    locations only as portable fallbacks.
    """
    # First restore an existing tray-hidden instance.  Some WeChat builds do
    # not react to starting Weixin.exe again, but pywinauto restore does.
    hidden = _wechat_hidden_main_candidates()
    for candidate in hidden:
        try:
            candidate.restore()
            candidate.set_focus()
            time.sleep(1.0)
            if _wechat_main_window_candidates():
                print(json.dumps({
                    "event": "wechat_window_recovered",
                    "method": "restore_existing_window",
                }, ensure_ascii=False), flush=True)
                return True
        except Exception:
            continue

    candidates = [
        os.environ.get("WECHAT_CLIENT_PATH", ""),
        os.path.join(os.environ.get("ProgramFiles", ""), "Tencent", "Weixin", "Weixin.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", ""), "Tencent", "WeChat", "WeChat.exe"),
    ]
    executable = next((path for path in candidates if path and Path(path).is_file()), None)
    if not executable:
        print(json.dumps({
            "event": "wechat_window_recovery_failed",
            "reason": "client executable path not found",
        }, ensure_ascii=False), flush=True)
        return False
    print(json.dumps({
        "event": "wechat_window_recovery_started",
        "executable": executable,
    }, ensure_ascii=False), flush=True)
    try:
        subprocess.Popen(
            [executable],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            close_fds=True,
        )
    except Exception as exc:
        print(json.dumps({
            "event": "wechat_window_recovery_failed",
            "reason": str(exc),
        }, ensure_ascii=False), flush=True)
        return False
    deadline = time.time() + 10.0
    while time.time() < deadline:
        time.sleep(0.5)
        if _wechat_main_window_candidates():
            print(json.dumps({
                "event": "wechat_window_recovered",
                "method": "single_instance_launch",
            }, ensure_ascii=False), flush=True)
            return True
    print(json.dumps({
        "event": "wechat_window_recovery_failed",
        "reason": "main window did not appear within 10 seconds",
    }, ensure_ascii=False), flush=True)
    return False


def get_wechat_window():
    candidates = _wechat_main_window_candidates()
    if not candidates:
        wake_wechat_main_window()
        candidates = _wechat_main_window_candidates()
    if not candidates and _wechat_login_windows():
        raise RuntimeError("微信客户端已打开，但仍停留在‘需在手机上完成登录’界面；请先完成登录")
    if len(candidates) == 1:
        try:
            rect = candidates[0].rectangle()
            if rect.width() <= 10 or rect.height() <= 10:
                candidates[0].restore()
                time.sleep(0.5)
        except Exception:
            pass
    matches = []
    for candidate in candidates:
        try:
            rect = candidate.rectangle()
            if rect.width() > 10 and rect.height() > 10:
                matches.append(candidate)
        except Exception:
            continue
    if len(matches) != 1:
        raise RuntimeError(f"需要且只能有一个可见标题为‘{WINDOW_TITLE}’的窗口，当前找到 {len(matches)} 个")
    win = matches[0]
    rect = win.rectangle()
    if rect.width() <= 10 or rect.height() <= 10:
        win.restore()
        time.sleep(0.5)
    try:
        win.maximize()
        time.sleep(0.4)
    except Exception:
        pass
    win.set_focus()
    return win


def web_windows():
    matches = []
    for win in Desktop(backend="uia").windows():
        try:
            title = (win.window_text() or "").strip()
            class_name = win.element_info.class_name or ""
            if class_name == "Chrome_WidgetWin_0" and (
                title == WINDOW_TITLE or "搜一搜" in title
            ):
                rect = win.rectangle()
                if rect.width() > 100 and rect.height() > 100:
                    matches.append(win)
        except Exception:
            continue
    return matches


def get_web_window():
    matches = web_windows()
    if len(matches) != 1:
        raise RuntimeError(f"需要且只能有一个微信搜一搜网页窗口，当前找到 {len(matches)} 个")
    win = matches[0]
    try:
        win.maximize()
        time.sleep(0.4)
    except Exception:
        pass
    win.set_focus()
    return win


def _is_embedded_search_surface(win) -> bool:
    """Recognize the account-specific 搜一搜 page embedded in the Qt main window."""
    names = set()
    has_search_action = False
    has_search_edit = False
    has_search_document = False
    try:
        for control in win.descendants():
            try:
                name = (control.window_text() or "").strip()
                if name:
                    names.add(name)
                if (control.element_info.automation_id or "") == "weixin-search-input":
                    has_search_edit = True
                if (
                    control.element_info.control_type == "Document"
                    and (control.element_info.automation_id or "") == "RootWebArea"
                    and "搜一搜" in name
                ):
                    has_search_document = True
                if name == "搜索" and control.element_info.control_type in {"Button", "Text", "Hyperlink"}:
                    has_search_action = True
            except Exception:
                continue
    except Exception:
        return False
    # Requiring category labels avoids confusing the global chat-list search
    # box with the full embedded 搜一搜 page.
    return (
        has_search_document and has_search_edit and has_search_action
    ) or (
        has_search_action and "文章" in names and "公众号" in names
        and ("搜一搜" in names or "全部" in names)
    )


def try_get_search_window():
    """Return either the detached Chromium window or the embedded Qt surface."""
    global SEARCH_LAYOUT
    detached = web_windows()
    if len(detached) == 1:
        if SEARCH_LAYOUT != "detached":
            SEARCH_LAYOUT = "detached"
            print(json.dumps({"event": "wechat_search_layout", "layout": SEARCH_LAYOUT}, ensure_ascii=False), flush=True)
        # Re-enumerating inside get_web_window introduces a race while the
        # detached Chromium window navigates.  Use the unique wrapper from the
        # enumeration that just succeeded.
        win = detached[0]
        try:
            win.maximize()
            time.sleep(0.2)
        except Exception:
            pass
        win.set_focus()
        return win
    if len(detached) > 1:
        raise RuntimeError(f"找到多个微信搜一搜网页窗口：{len(detached)}")
    main = get_wechat_window()
    if SEARCH_LAYOUT == "embedded" or _is_embedded_search_surface(main):
        if SEARCH_LAYOUT != "embedded":
            SEARCH_LAYOUT = "embedded"
            print(json.dumps({"event": "wechat_search_layout", "layout": SEARCH_LAYOUT}, ensure_ascii=False), flush=True)
        return main
    return None


def get_search_window(wait_seconds: float = 2.0):
    """Return a stable 搜一搜 surface, tolerating short Chromium rebuilds.

    Detached 搜一搜 windows are briefly destroyed and recreated during some
    navigations.  A single UIA enumeration in that interval used to abort the
    complete worker.  Keep this wait deliberately short; the higher-level
    recovery routine is responsible for reopening a genuinely closed page.
    """
    deadline = time.time() + max(0.0, float(wait_seconds))
    last_error = None
    while True:
        try:
            surface = try_get_search_window()
            if surface is not None:
                return surface
        except Exception as exc:
            last_error = exc
        if time.time() >= deadline:
            detail = f"：{last_error}" if last_error else ""
            raise SearchWindowUnavailableError(
                f"未检测到独立或内嵌的微信搜一搜页面{detail}"
            ) from last_error
        time.sleep(0.25)


def has_search_window() -> bool:
    try:
        return try_get_search_window() is not None
    except Exception:
        return False


def display_scale_factor() -> float:
    try:
        percent = int(os.environ.get("WECHAT_DISPLAY_SCALE", "100"))
    except ValueError:
        percent = 100
    return 1.5 if percent == 150 else 1.0


def scale_pixels(value: int | float) -> int:
    return int(round(float(value) * display_scale_factor()))


def click_relative(win, point: list[int]) -> None:
    rect = win.rectangle()
    mouse.click(coords=(rect.left + scale_pixels(point[0]), rect.top + scale_pixels(point[1])))


def set_clipboard(text: str) -> None:
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardText(text, win32clipboard.CF_UNICODETEXT)
    finally:
        win32clipboard.CloseClipboard()


def get_clipboard() -> str:
    win32clipboard.OpenClipboard()
    try:
        if win32clipboard.IsClipboardFormatAvailable(win32clipboard.CF_UNICODETEXT):
            return win32clipboard.GetClipboardData(win32clipboard.CF_UNICODETEXT)
        return ""
    finally:
        win32clipboard.CloseClipboard()


def paste_text(text: str) -> None:
    set_clipboard(text)
    keyboard.send_keys("^a{BACKSPACE}^v", pause=0.02)


def save_window_screenshot(win, path: Path) -> None:
    rect = win.rectangle()
    path.parent.mkdir(parents=True, exist_ok=True)
    ImageGrab.grab(
        bbox=(rect.left, rect.top, rect.right, rect.bottom),
        all_screens=True,
    ).save(path)


def load_keywords(path: Path, limit: int | None = None) -> list[tuple[str, str]]:
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        entries = payload.get("keywords", payload) if isinstance(payload, dict) else payload
        rows = []
        for item in entries:
            if isinstance(item, str):
                keyword, project = item, ""
            else:
                keyword, project = item.get("keyword"), item.get("project", "")
            if keyword:
                rows.append((str(keyword).strip(), str(project or "").strip()))
            if limit and len(rows) >= limit:
                break
        return rows
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["关键词"]
    rows: list[tuple[str, str]] = []
    for keyword, project, *_ in sheet.iter_rows(min_row=2, values_only=True):
        if keyword is None:
            continue
        rows.append((str(keyword).strip(), str(project or "").strip()))
        if limit and len(rows) >= limit:
            break
    return rows


def normalize_article_key(url: str) -> str:
    query = parse_qs(urlparse(url).query)
    parts = [query.get("__biz", [""])[0], query.get("mid", [""])[0], query.get("idx", [""])[0]]
    stable = "|".join(parts) if all(parts) else url.strip()
    return hashlib.sha256(stable.encode("utf-8")).hexdigest()


def parse_count(text: str | None) -> int | None:
    if text is None:
        return None
    value = text.strip().replace(",", "")
    if not value or value in {"--", "不可见", "未显示"}:
        return None
    match = re.search(r"(\d+(?:\.\d+)?)\s*([万wW]?)", value)
    if not match:
        return None
    number = float(match.group(1))
    if match.group(2):
        number *= 10000
    return int(number)


_STATE_BACKUP_TIMES: dict[str, float] = {}


def _state_backup_path(db_path: Path) -> Path:
    return db_path.with_name(db_path.name + ".lastgood")


def _verify_state_connection(connection: sqlite3.Connection) -> None:
    result = connection.execute("PRAGMA quick_check").fetchone()
    if not result or str(result[0]).lower() != "ok":
        raise sqlite3.DatabaseError(
            f"SQLite quick_check failed: {result[0] if result else 'no result'}"
        )


def _archive_corrupt_state(db_path: Path, reason: Exception) -> Path:
    """Keep every damaged SQLite artefact for diagnosis; never delete it."""
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    archive_dir = db_path.parent / "corrupt-state" / f"{db_path.stem}-{stamp}"
    archive_dir.mkdir(parents=True, exist_ok=True)
    for suffix in ("", "-wal", "-shm", "-journal"):
        source = Path(str(db_path) + suffix)
        if source.exists():
            os.replace(source, archive_dir / source.name)
    (archive_dir / "reason.txt").write_text(str(reason), encoding="utf-8")
    return archive_dir


def checkpoint_state_backup(connection: sqlite3.Connection, *, force: bool = False) -> None:
    """Create a consistent rotating backup without copying a live SQLite file."""
    db_row = connection.execute("PRAGMA database_list").fetchone()
    if not db_row or not db_row[2] or db_row[2] == ":memory:":
        return
    db_path = Path(db_row[2])
    key = str(db_path.resolve())
    now = time.monotonic()
    if not force and now - _STATE_BACKUP_TIMES.get(key, 0.0) < 20.0:
        return
    target = _state_backup_path(db_path)
    temporary = target.with_name(target.name + ".tmp")
    temporary.unlink(missing_ok=True)
    backup_connection = sqlite3.connect(temporary)
    try:
        connection.backup(backup_connection)
        _verify_state_connection(backup_connection)
    finally:
        backup_connection.close()
    os.replace(temporary, target)
    _STATE_BACKUP_TIMES[key] = now


def init_state(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = None
    recovery_event = None
    try:
        connection = sqlite3.connect(db_path, timeout=30)
        connection.execute("PRAGMA busy_timeout=30000")
        _verify_state_connection(connection)
    except sqlite3.DatabaseError as exc:
        if connection is not None:
            connection.close()
        archive_dir = _archive_corrupt_state(db_path, exc)
        backup_path = _state_backup_path(db_path)
        restored = False
        if backup_path.exists():
            try:
                candidate = sqlite3.connect(backup_path, timeout=30)
                try:
                    _verify_state_connection(candidate)
                finally:
                    candidate.close()
                shutil.copy2(backup_path, db_path)
                restored = True
            except (OSError, sqlite3.DatabaseError):
                restored = False
        recovery_event = {
            "event": "sqlite_state_recovered",
            "database": str(db_path),
            "archive": str(archive_dir),
            "restored_lastgood": restored,
            "reason": str(exc),
        }
        connection = sqlite3.connect(db_path, timeout=30)
        connection.execute("PRAGMA busy_timeout=30000")
    # DELETE journalling is materially safer than WAL when a portable package
    # was placed in OneDrive.  FULL sync plus SQLite's backup API gives us a
    # recoverable checkpoint without copying a live database file.
    connection.execute("PRAGMA journal_mode=DELETE")
    connection.execute("PRAGMA synchronous=FULL")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS captured_article (
            article_key TEXT NOT NULL,
            project TEXT NOT NULL,
            keyword TEXT NOT NULL,
            url TEXT NOT NULL,
            like_count INTEGER,
            comment_count INTEGER,
            favorite_count INTEGER,
            share_count INTEGER,
            captured_at TEXT NOT NULL,
            PRIMARY KEY (article_key, keyword)
        )
        """
    )
    # Idempotent migration for databases created before title capture was added.
    columns = {row[1] for row in connection.execute("PRAGMA table_info(captured_article)")}
    migrations = {
        "title": "TEXT", "author_name": "TEXT", "published_at": "TEXT",
        "location": "TEXT", "content": "TEXT", "image_count": "INTEGER DEFAULT 0",
        "video_count": "INTEGER DEFAULT 0",
    }
    for column, definition in migrations.items():
        if column not in columns:
            connection.execute(f"ALTER TABLE captured_article ADD COLUMN {column} {definition}")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS crawl_task (
            keyword TEXT PRIMARY KEY,
            project TEXT NOT NULL,
            next_result_index INTEGER NOT NULL DEFAULT 0,
            old_streak INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'pending',
            last_error TEXT,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS collection_issue (
            keyword TEXT NOT NULL,
            result_index INTEGER NOT NULL,
            issue_type TEXT NOT NULL,
            title TEXT,
            fingerprint TEXT,
            detail TEXT,
            observed_at TEXT NOT NULL,
            PRIMARY KEY (keyword, result_index, issue_type)
        )
        """
    )
    issue_columns = {
        row[1] for row in connection.execute("PRAGMA table_info(collection_issue)")
    }
    if "fingerprint" not in issue_columns:
        connection.execute("ALTER TABLE collection_issue ADD COLUMN fingerprint TEXT")
    connection.commit()
    checkpoint_state_backup(connection, force=True)
    if recovery_event:
        print(json.dumps(recovery_event, ensure_ascii=False), flush=True)
    return connection


def load_task_state(connection: sqlite3.Connection, keyword: str, project: str) -> tuple[int, int, str]:
    row = connection.execute(
        "SELECT next_result_index, old_streak, status FROM crawl_task WHERE keyword=?",
        (keyword,),
    ).fetchone()
    if row:
        return int(row[0]), int(row[1]), str(row[2])
    existing_count = connection.execute(
        "SELECT COUNT(*) FROM captured_article WHERE keyword=?",
        (keyword,),
    ).fetchone()[0]
    now = datetime.now().isoformat(timespec="seconds")
    connection.execute(
        """
        INSERT INTO crawl_task(keyword, project, next_result_index, old_streak, status, updated_at)
        VALUES (?, ?, ?, 0, 'pending', ?)
        """,
        (keyword, project, int(existing_count), now),
    )
    connection.commit()
    checkpoint_state_backup(connection)
    return int(existing_count), 0, "pending"


def update_task_state(
    connection: sqlite3.Connection,
    keyword: str,
    project: str,
    next_result_index: int,
    old_streak: int,
    status: str,
    last_error: str | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO crawl_task(
            keyword, project, next_result_index, old_streak, status, last_error, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(keyword) DO UPDATE SET
            project=excluded.project,
            next_result_index=excluded.next_result_index,
            old_streak=excluded.old_streak,
            status=excluded.status,
            last_error=excluded.last_error,
            updated_at=excluded.updated_at
        """,
        (
            keyword, project, next_result_index, old_streak, status, last_error,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    connection.commit()
    checkpoint_state_backup(connection)


def save_collection_issue(
    connection: sqlite3.Connection,
    keyword: str,
    result_index: int,
    issue_type: str,
    title: str,
    detail: str,
    fingerprint: str | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO collection_issue(
            keyword, result_index, issue_type, title, fingerprint, detail, observed_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(keyword, result_index, issue_type) DO UPDATE SET
            title=excluded.title,
            fingerprint=excluded.fingerprint,
            detail=excluded.detail,
            observed_at=excluded.observed_at
        """,
        (
            keyword, result_index, issue_type, title, fingerprint, detail,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    connection.commit()
    checkpoint_state_backup(connection)


def save_row(connection: sqlite3.Connection, row: ResultRow) -> None:
    connection.execute(
        """
        INSERT INTO captured_article (
            article_key, project, keyword, title, url, author_name, published_at,
            location, content, image_count, video_count, like_count, comment_count,
            favorite_count, share_count, captured_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(article_key, keyword) DO UPDATE SET
            project=excluded.project,
            title=COALESCE(NULLIF(excluded.title, ''), captured_article.title),
            url=excluded.url,
            author_name=COALESCE(excluded.author_name, captured_article.author_name),
            published_at=COALESCE(excluded.published_at, captured_article.published_at),
            location=COALESCE(excluded.location, captured_article.location),
            content=COALESCE(NULLIF(excluded.content, ''), captured_article.content),
            image_count=MAX(COALESCE(excluded.image_count, 0), COALESCE(captured_article.image_count, 0)),
            video_count=MAX(COALESCE(excluded.video_count, 0), COALESCE(captured_article.video_count, 0)),
            like_count=COALESCE(excluded.like_count, captured_article.like_count),
            comment_count=COALESCE(excluded.comment_count, captured_article.comment_count),
            favorite_count=COALESCE(excluded.favorite_count, captured_article.favorite_count),
            share_count=COALESCE(excluded.share_count, captured_article.share_count),
            captured_at=excluded.captured_at
        """,
        (
            normalize_article_key(row.url), row.project, row.keyword, row.title, row.url,
            row.author_name, row.published_at, row.location, row.content,
            row.image_count, row.video_count,
            row.like_count, row.comment_count, row.favorite_count,
            row.share_count, row.captured_at,
        ),
    )
    connection.commit()


def find_controls(root, *, control_type: str | None = None, name: str | None = None, class_contains: str | None = None):
    controls = root.descendants(control_type=control_type) if control_type else root.descendants()
    found = []
    for control in controls:
        try:
            info = control.element_info
            if name is not None and (info.name or "") != name:
                continue
            if class_contains is not None and class_contains not in (info.class_name or ""):
                continue
            found.append(control)
        except Exception:
            continue
    return found


def wait_for(predicate, timeout: float, interval: float = 0.25):
    deadline = time.time() + timeout
    last_error = None
    while time.time() < deadline:
        try:
            value = predicate()
            if value:
                return value
        except Exception as exc:
            last_error = exc
        time.sleep(interval)
    if last_error:
        raise TimeoutError(str(last_error)) from last_error
    raise TimeoutError("等待页面状态超时")


def visible_controls(root, *, name: str | None = None, control_type: str | None = None):
    candidates = find_controls(root, control_type=control_type, name=name)
    visible = []
    for control in candidates:
        try:
            rect = control.rectangle()
            if control.is_visible() and rect.width() > 2 and rect.height() > 2:
                visible.append(control)
        except Exception:
            continue
    return visible


def find_search_edit(surface):
    edits = visible_controls(surface, control_type="Edit")
    preferred = [
        control for control in edits
        if (control.element_info.automation_id or "") == "weixin-search-input"
    ]
    if len(preferred) == 1:
        return preferred[0]
    named = [control for control in edits if (control.window_text() or "").strip() in {"搜索", ""}]
    pool = named or edits
    if not pool:
        raise RuntimeError("搜一搜页面未找到可见搜索输入框")
    # The full-page search input is normally the widest visible Edit.
    return max(pool, key=lambda control: control.rectangle().width())


def click_named_surface_control(surface, name: str) -> bool:
    for control_type in ("Button", "TabItem", "Hyperlink", "Text"):
        matches = visible_controls(surface, name=name, control_type=control_type)
        if matches:
            try:
                matches[0].click_input()
            except Exception:
                rect = matches[0].rectangle()
                mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            return True
    # Some embedded Chromium builds expose category labels as Group/ListItem
    # instead of Text or Button. Accept any visible exact-name control last.
    for control in surface.descendants():
        try:
            rect = control.rectangle()
            if (
                (control.window_text() or "").strip() == name
                and control.is_visible() and rect.width() > 2 and rect.height() > 2
            ):
                try:
                    control.click_input()
                except Exception:
                    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
                return True
        except Exception:
            continue
    return False


def search_page_titles(surface) -> list[str]:
    """Return visible Chromium document titles for the current 搜一搜 page."""
    titles = []
    try:
        documents = surface.descendants(control_type="Document")
    except Exception:
        documents = []
    for document in documents:
        try:
            value = (document.window_text() or "").strip()
            if value and "搜一搜" in value:
                titles.append(value)
        except Exception:
            continue
    return titles


def is_article_results_page(surface) -> bool:
    """Verify the selected 搜一搜 category, instead of trusting a click.

    The detached Chromium window title itself often stays as ``微信``.  The
    inner RootWebArea document, however, is named ``关键词 - 文章 - 搜一搜``.
    This check prevents a coordinate drift from silently collecting 全部 or
    视频 results.
    """
    titles = search_page_titles(surface)
    return any(re.search(r"(?:^|\s-\s)文章\s-\s搜一搜$", title) for title in titles)


def ensure_article_results_page(surface, config: dict) -> None:
    """Select 文章 and refuse to continue until the page proves it is active."""
    wait_seconds = max(0.5, float(config.get("wait_short", 0.8)))

    if is_article_results_page(surface):
        print(json.dumps({
            "event": "article_tab_verified",
            "method": "already_active",
            "page_titles": search_page_titles(surface),
        }, ensure_ascii=False), flush=True)
        return

    # Prefer semantic UIA whenever this account/build exposes the category.
    semantic_clicked = click_named_surface_control(surface, "文章")
    if semantic_clicked:
        time.sleep(wait_seconds)
        surface = get_search_window()
        if is_article_results_page(surface):
            print(json.dumps({
                "event": "article_tab_verified",
                "method": "uia_name",
                "page_titles": search_page_titles(surface),
            }, ensure_ascii=False), flush=True)
            return

    # Current Chromium builds expose category labels as unnamed ListItems.
    # Their stable order is AI搜索、全部、视频、文章、贴图... . Selecting the
    # fourth item is substantially safer than a display-dependent coordinate,
    # and the RootWebArea title below still proves that 文章 really opened.
    try:
        edit = find_search_edit(surface)
        edit_rect = edit.rectangle()
        strict_items = []
        band_items = []
        for control in visible_controls(surface, control_type="ListItem"):
            rect = control.rectangle()
            if not (edit_rect.bottom <= rect.top <= edit_rect.bottom + 60):
                continue
            class_name = control.element_info.class_name or ""
            if "unit__item" in class_name:
                strict_items.append(control)
            # Newer WeChat builds renamed the CSS class, so also keep every
            # tab-sized item in the category band as a fallback candidate.
            # Each click below is still verified through the RootWebArea
            # title before the collector is allowed to proceed.
            if 20 <= rect.width() <= 220 and rect.height() <= 80:
                band_items.append(control)
        strict_items.sort(key=lambda control: control.rectangle().left)
        if len(strict_items) >= 4:
            strict_items[3].click_input()
            time.sleep(wait_seconds)
            surface = get_search_window()
            if is_article_results_page(surface):
                print(json.dumps({
                    "event": "article_tab_verified",
                    "method": "uia_category_ordinal",
                    "category_count": len(strict_items),
                    "page_titles": search_page_titles(surface),
                }, ensure_ascii=False), flush=True)
                return
        band_items.sort(key=lambda control: control.rectangle().left)
        for candidate in band_items:
            candidate.click_input()
            time.sleep(wait_seconds)
            if not has_search_window():
                break
            surface = get_search_window()
            if is_search_results_page(surface) and is_article_results_page(surface):
                print(json.dumps({
                    "event": "article_tab_verified",
                    "method": "uia_band_sweep",
                    "category_count": len(band_items),
                    "page_titles": search_page_titles(surface),
                }, ensure_ascii=False), flush=True)
                return
    except Exception:
        pass

    # Chromium sometimes hides category labels from UIA.  Coordinates are
    # anchored to the live search input, then *verified* from RootWebArea.
    edit = find_search_edit(surface)
    edit_rect = edit.rectangle()
    embedded = config.get("embedded_web_points", {})
    version = os.environ.get("WECHAT_CLIENT_VERSION", "4.1.13")
    profiles = embedded.get("article_tab_by_version", {})
    key = next((item for item in profiles if version.startswith(item)), None)
    configured_ratio, y_offset = profiles.get(
        key, embedded.get("article_tab_default", [0.29, 28])
    )
    # 0.41 is the centre of 文章 relative to the live search input in both
    # validated 2560x1440 and 1920x1080 layouts. WeChat builds occasionally
    # add/reorder category tabs (e.g. AI搜索), which shifts 文章 further
    # right; clicking the old ratio then re-selects 全部/视频. Sweep a wide
    # verified range so a layout change can no longer strand the collector,
    # while every candidate click is still proven by the page title first.
    ratios = []
    for ratio in (
        float(configured_ratio), 0.41, 0.44, 0.47, 0.50, 0.53, 0.56, 0.59,
        0.40, 0.42, 0.39, 0.43, 0.38, 0.45, 0.48, 0.35, 0.32, 0.29,
    ):
        if ratio not in ratios:
            ratios.append(ratio)

    attempts = []
    for ratio in ratios:
        x = edit_rect.left + int(edit_rect.width() * ratio)
        y = edit_rect.bottom + scale_pixels(y_offset)
        mouse.click(coords=(x, y))
        time.sleep(wait_seconds)
        surface = get_search_window()
        titles = search_page_titles(surface)
        attempts.append({"ratio": ratio, "point": [x, y], "page_titles": titles})
        if is_article_results_page(surface):
            print(json.dumps({
                "event": "article_tab_verified",
                "method": "anchored_fallback",
                "layout": SEARCH_LAYOUT,
                "wechat_version": version,
                "ratio": ratio,
                "point": [x, y],
                "input_rect": [edit_rect.left, edit_rect.top, edit_rect.right, edit_rect.bottom],
                "page_titles": titles,
            }, ensure_ascii=False), flush=True)
            return

    raise RuntimeError(
        "未能确认已切换到‘文章’栏目；为避免误采‘全部/视频’已停止。"
        + json.dumps({"attempts": attempts}, ensure_ascii=False)
    )


def _search_in_web_once(keyword: str, config: dict) -> None:
    surface = try_get_search_window()
    submitted_from_global_search = False
    if surface is None:
        main = get_wechat_window()
        # New-account UI exposes 搜一搜 directly; legacy-account UI nests it
        # under 发现. Prefer semantic UIA navigation before using coordinates.
        opened = click_named_surface_control(main, "搜一搜")
        if not opened and click_named_surface_control(main, "发现"):
            time.sleep(config.get("wait_short", 0.8))
            opened = click_named_surface_control(get_wechat_window(), "搜一搜")
        if opened:
            try:
                wait_for(has_search_window, timeout=6.0)
                surface = get_search_window()
            except Exception:
                surface = None
        if surface is None:
            # Last-resort route retained for the validated 4.1.13 account UI.
            main = get_wechat_window()
            rect = main.rectangle()
            search_box = config.get("points", {}).get("main_search_box", [150, 55])
            # ``points.main_search_box`` is calibrated at 100% DPI, just like
            # every other legacy-account coordinate in this collector.  This
            # fallback is especially important on Qt builds that expose only
            # ``MMUIRenderSubWindowHW`` (no semantic UIA controls).  Applying
            # the configured 100%/150% scale here prevents a 150%-scaled host
            # from clicking the chat list instead of the global search box.
            mouse.click(coords=(
                rect.left + scale_pixels(search_box[0]),
                rect.top + scale_pixels(search_box[1]),
            ))
            time.sleep(config.get("wait_short", 0.8))
            paste_text(keyword)
            time.sleep(config.get("wait_search_suggest", 1.5))
            keyboard.send_keys("{DOWN}{ENTER}", pause=0.05)
            wait_for(has_search_window, timeout=10.0)
            time.sleep(config.get("wait_page", 3.0))
            surface = get_search_window()
            submitted_from_global_search = True

    if not submitted_from_global_search:
        if not is_search_results_page(surface):
            return_to_search_results()
        surface = get_search_window()
        edit = find_search_edit(surface)
        edit.click_input()
        set_clipboard(keyword)
        keyboard.send_keys("^a{BACKSPACE}^v", pause=0.02)
        if not click_named_surface_control(surface, "搜索"):
            keyboard.send_keys("{ENTER}", pause=0.02)
        time.sleep(config.get("wait_page", 3.0))

    surface = get_search_window()
    ensure_article_results_page(surface, config)
    time.sleep(config.get("wait_page", 3.0))
    surface = get_search_window()
    if not is_article_results_page(surface):
        raise RuntimeError("‘文章’栏目校验在加载后失效，已停止采集")
    # Date-cutoff termination is safe only after explicitly selecting 最新.
    # A deeply scrolled/recovered Chromium page can omit the filter label from
    # UIA, so use the same live-search-input anchoring strategy as 文章.
    latest_method = "uia_name"
    if not click_named_surface_control(surface, "最新"):
        edit = find_search_edit(surface)
        edit_rect = edit.rectangle()
        latest_profile = config.get("embedded_web_points", {}).get(
            "latest_sort_default", [0.155, 75]
        )
        ratio, y_offset = latest_profile
        x = edit_rect.left + int(edit_rect.width() * float(ratio))
        y = edit_rect.bottom + scale_pixels(y_offset)
        mouse.click(coords=(x, y))
        latest_method = "anchored_fallback"
    time.sleep(config.get("wait_page", 3.0))
    print(json.dumps({
        "event": "latest_sort_selected",
        "method": latest_method,
        "page_titles": search_page_titles(get_search_window()),
    }, ensure_ascii=False), flush=True)
    # Selecting a new keyword/filter does not reliably reset Chromium's scroll
    # offset.  Establish the absolute newest-result origin before any global
    # result index is interpreted, including on checkpoint recovery.
    reset_search_results_to_top(config)


def search_in_web(keyword: str, config: dict) -> None:
    """Open and validate the current keyword, rebuilding lost search UI.

    This recovery is bounded so an invalid/login UI cannot become an infinite
    loop.  Every attempt starts from semantic new/legacy-account navigation,
    reselects 文章 and 最新, and returns to the absolute top.
    """
    attempts = max(1, int(config.get("search_window_recovery_attempts", 5)))
    wait_seconds = max(0.0, float(config.get("search_window_recovery_wait_seconds", 2.0)))
    errors: list[str] = []
    for attempt in range(1, attempts + 1):
        if STOP_REQUESTED:
            raise SearchWindowUnavailableError("已请求安全停止，不再重建搜一搜页面")
        try:
            _search_in_web_once(keyword, config)
            surface = get_search_window(wait_seconds=max(2.0, wait_seconds))
            if not is_article_results_page(surface):
                raise RuntimeError("恢复后未确认处于‘文章’栏目")
            print(json.dumps({
                "event": "search_window_ready",
                "keyword": keyword,
                "attempt": attempt,
                "layout": SEARCH_LAYOUT,
            }, ensure_ascii=False), flush=True)
            return
        except Exception as exc:
            errors.append(str(exc))
            print(json.dumps({
                "event": "search_window_recovery_retry",
                "keyword": keyword,
                "attempt": attempt,
                "max_attempts": attempts,
                "error": str(exc),
            }, ensure_ascii=False), flush=True)
            if attempt < attempts:
                time.sleep(wait_seconds)
    raise SearchWindowUnavailableError(
        f"搜一搜页面连续 {attempts} 次恢复失败：{errors[-1] if errors else 'unknown'}"
    )


def search_result_buttons(web):
    return [
        control for control in web.descendants(control_type="Button")
        if "search_list_item" in (control.element_info.class_name or "")
        and "article-thumb" in (control.element_info.class_name or "")
    ]


def reset_search_results_to_top(config: dict) -> list:
    """Return the result list after proving the page is at its absolute top.

    WeChat's embedded Chromium surface preserves the scroll offset when a new
    keyword is submitted or the ``最新`` filter is selected.  On builds that
    virtualize the result list, UIA then exposes the first *currently mounted*
    card as index 0, even though it is a middle article.  Always drive the live
    document to the top before taking the index-0 snapshot.
    """
    wait_seconds = max(0.15, float(config.get("wait_short", 0.8)) / 3)
    max_rounds = max(4, int(config.get("reset_top_rounds", 12)))
    stable_rounds = 0
    previous_first = None
    last_payload = None

    for attempt in range(1, max_rounds + 1):
        surface = get_search_window()
        documents = []
        for document in surface.descendants(control_type="Document"):
            try:
                rect = document.rectangle()
                if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                    documents.append(document)
            except Exception:
                continue
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        point = (
            viewport.left + int(viewport.width() * 0.55),
            viewport.top + int(viewport.height() * 0.45),
        )
        surface.set_focus()
        # Never click the document to obtain focus here: its center is often an
        # article card, which would open a middle result before index 0 is
        # established.  Wheel input is dispatched directly at the document.
        mouse.scroll(coords=point, wheel_dist=60)
        time.sleep(wait_seconds)

        surface = get_search_window()
        results = search_result_buttons(surface)
        vertical_scroll_percent = None
        try:
            if target.iface_scroll.CurrentVerticallyScrollable:
                vertical_scroll_percent = float(target.iface_scroll.CurrentVerticalScrollPercent)
            else:
                vertical_scroll_percent = 0.0
        except Exception:
            pass
        first_signature = None
        if results:
            first = results[0]
            first_rect = first.rectangle()
            first_signature = (
                (first.window_text() or "").strip(),
                first_rect.left, first_rect.top, first_rect.right, first_rect.bottom,
            )
        if first_signature == previous_first:
            stable_rounds += 1
        else:
            stable_rounds = 0
        previous_first = first_signature
        last_payload = {
            "event": "search_results_reset_top",
            "attempt": attempt,
            "result_count": len(results),
            "first_title": first_signature[0] if first_signature else None,
            "first_rect": list(first_signature[1:]) if first_signature else None,
            "vertical_scroll_percent": vertical_scroll_percent,
            "stable_rounds": stable_rounds,
        }
        # The Scroll pattern is authoritative when exposed by RootWebArea.
        # Re-read controls only after it reports the absolute 0% position so a
        # virtualized former middle-card can never masquerade as index 0.
        scroll_pattern_proves_top = (
            vertical_scroll_percent is not None
            and vertical_scroll_percent <= 0.01
        )
        fallback_proves_top = vertical_scroll_percent is None and stable_rounds >= 2
        if results and (scroll_pattern_proves_top or fallback_proves_top):
            print(json.dumps(last_payload, ensure_ascii=False), flush=True)
            return results

    raise RuntimeError(
        "无法确认文章结果列表已经回到顶部："
        + json.dumps(last_payload or {}, ensure_ascii=False)
    )


def load_more_search_results(web, previous_count: int, config: dict):
    """Drive Chromium lazy loading and return once the UIA result list grows.

    scroll_into_view alone often succeeds without emitting a wheel event, so
    WeChat keeps only the first one or two days of high-frequency results in
    the accessibility tree. Always wheel inside RootWebArea and require
    several stable bottom attempts before declaring the result set exhausted.
    """
    rounds = max(3, int(config.get("load_more_rounds", 8)))
    wait_seconds = max(0.5, float(config.get("wait_load_more", 1.5)))
    results = search_result_buttons(web)
    for attempt in range(1, rounds + 1):
        if results:
            try:
                results[-1].scroll_into_view()
            except Exception:
                pass
        surface = get_search_window()
        documents = []
        for document in surface.descendants(control_type="Document"):
            try:
                rect = document.rectangle()
                if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                    documents.append(document)
            except Exception:
                continue
        target = max(documents, key=lambda item: item.rectangle().width() * item.rectangle().height()) if documents else surface
        rect = target.rectangle()
        scroll_x = rect.left + max(40, min(rect.width() - 40, int(rect.width() * 0.55)))
        scroll_y = max(rect.top + 40, rect.bottom - 120)
        surface.set_focus()
        # Multiple wheel pulses are required by the embedded Chromium view;
        # the coordinate is guaranteed to remain inside the live document.
        mouse.scroll(coords=(scroll_x, scroll_y), wheel_dist=-15)
        time.sleep(wait_seconds)
        results = search_result_buttons(get_search_window())
        print(json.dumps({
            "event": "search_load_more",
            "attempt": attempt,
            "previous_count": previous_count,
            "current_count": len(results),
            "point": [scroll_x, scroll_y],
        }, ensure_ascii=False), flush=True)
        if len(results) > previous_count:
            return results, True
    return results, False


def bring_search_result_into_view(result_index: int, config: dict):
    """Use real wheel input until a virtualized result is physically visible.

    Chromium's UIA ScrollItem pattern can report success without changing the
    page. This happened at result 31: its rectangle remained around y=4465
    while the window ended at y=1392. A physical click then silently did
    nothing. Re-read rectangles after every wheel pulse and only return once
    the target is inside the document viewport.
    """
    wait_seconds = max(0.2, float(config.get("wait_short", 0.8)) / 2)
    # Deep yearly result sets can put a card tens of thousands of pixels away
    # from the current viewport. Thirty wheel rounds was insufficient on a
    # 1920x1080 worker at index 76, especially after browser-back reflow.
    max_rounds = max(80, int(config.get("result_scroll_rounds", 120)))
    last_rect = None
    last_position = None
    stagnant_rounds = 0
    direction_changes = 0
    previous_direction = None
    for attempt in range(1, max_rounds + 1):
        surface = get_search_window()
        results = search_result_buttons(surface)
        if result_index >= len(results):
            return None
        result = results[result_index]
        result_rect = result.rectangle()
        documents = []
        for document in surface.descendants(control_type="Document"):
            try:
                rect = document.rectangle()
                if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                    documents.append(document)
            except Exception:
                continue
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        visible_top = viewport.top + min(180, max(60, int(viewport.height() * 0.12)))
        visible_bottom = viewport.bottom - 80
        last_rect = [result_rect.left, result_rect.top, result_rect.right, result_rect.bottom]
        intersection = max(
            0,
            min(result_rect.bottom, visible_bottom) - max(result_rect.top, visible_top),
        )
        required_visible = min(40, max(16, int(result_rect.height() * 0.30)))
        if intersection >= required_visible:
            print(json.dumps({
                "event": "search_result_visible",
                "index": result_index,
                "attempt": attempt,
                "rect": last_rect,
                "scroll_direction": previous_direction or "none",
            }, ensure_ascii=False), flush=True)
            return result

        position = (result_rect.top, result_rect.bottom)
        if last_position is not None and all(
            abs(current - previous) <= 2
            for current, previous in zip(position, last_position)
        ):
            stagnant_rounds += 1
        else:
            stagnant_rounds = 0
        last_position = position

        # ScrollItem is effective on some Chromium builds but a no-op on
        # others. Try it first and again only when physical wheel input has
        # demonstrably stopped moving the card. The next loop obtains a fresh
        # UIA object, so an element cached before browser-back is never clicked.
        if attempt == 1 or stagnant_rounds >= 4:
            try:
                result.scroll_into_view()
                time.sleep(wait_seconds)
                stagnant_rounds = 0
                print(json.dumps({
                    "event": "search_result_scrollitem",
                    "index": result_index,
                    "attempt": attempt,
                    "rect": last_rect,
                }, ensure_ascii=False), flush=True)
                continue
            except Exception:
                stagnant_rounds = 0

        scroll_x = viewport.left + int(viewport.width() * 0.55)
        scroll_y = max(viewport.top + 80, viewport.bottom - 160)
        # A card can be partially visible but still below the safe click area
        # (for example only 32px of a 127px card at the bottom).  Its *bottom*
        # edge determines that we must continue downward; using top > bottom
        # incorrectly selected upward and caused an endless reset/retry loop.
        direction_name = "down" if result_rect.bottom > visible_bottom else "up"
        if previous_direction and direction_name != previous_direction:
            direction_changes += 1
        previous_direction = direction_name
        if direction_name == "down":
            distance = max(0, result_rect.top - visible_bottom)
            direction = -30 if distance > 2500 else (-12 if distance > 900 else -4)
        else:
            distance = max(0, visible_top - result_rect.bottom)
            direction = 30 if distance > 2500 else (12 if distance > 900 else 4)
        surface.set_focus()
        mouse.scroll(coords=(scroll_x, scroll_y), wheel_dist=direction)
        time.sleep(wait_seconds)

    print(json.dumps({
        "event": "search_result_not_visible",
        "index": result_index,
        "last_rect": last_rect,
        "rounds": max_rounds,
        "scroll_direction": previous_direction,
        "direction_changes": direction_changes,
    }, ensure_ascii=False), flush=True)
    # If both ScrollItem and adaptive wheel input fail, restore a deterministic
    # top position before the caller rebuilds/retries this same keyword.
    try:
        surface = get_search_window()
        documents = [
            document for document in surface.descendants(control_type="Document")
            if document.is_visible()
        ]
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        point = (
            viewport.left + int(viewport.width() * 0.55),
            viewport.top + int(viewport.height() * 0.45),
        )
        surface.set_focus()
        for _ in range(4):
            mouse.scroll(coords=point, wheel_dist=60)
            time.sleep(0.15)
        print(json.dumps({
            "event": "search_scroll_reset_top",
            "index": result_index,
            "point": list(point),
        }, ensure_ascii=False), flush=True)
    except Exception:
        pass
    return None


def control_element_records(control, *, include_self: bool = False) -> list[tuple[str, str, str]]:
    records = []
    controls = ([control] if include_self else []) + list(control.descendants())
    for item in controls:
        try:
            info = item.element_info
            value = (info.name or "").strip()
            if value:
                records.append((value, info.control_type or "", info.class_name or ""))
        except Exception:
            continue
    return records


def page_element_records(web) -> list[tuple[str, str, str]]:
    return control_element_records(web)


def parse_publish_time(records: list[tuple[str, str, str]]) -> datetime | None:
    full_patterns = (
        re.compile(r"(20\d{2})年(\d{1,2})月(\d{1,2})日(?:\s+(\d{1,2}):(\d{2}))?"),
        re.compile(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})(?:\s+(\d{1,2}):(\d{2}))?"),
    )
    now = datetime.now()
    for value, _control_type, _class_name in records:
        for pattern in full_patterns:
            match = pattern.search(value)
            if match:
                return datetime(
                    int(match.group(1)), int(match.group(2)), int(match.group(3)),
                    int(match.group(4) or 0), int(match.group(5) or 0),
                )
        # Search result cards sometimes omit the year.  Because results are
        # newest-first, a month/day later than today belongs to the prior year.
        month_day = re.search(r"(?<!\d)(\d{1,2})月(\d{1,2})日(?:\s+(\d{1,2}):(\d{2}))?", value)
        if month_day:
            month, day = int(month_day.group(1)), int(month_day.group(2))
            year = now.year - (1 if (month, day) > (now.month, now.day) else 0)
            return datetime(
                year, month, day,
                int(month_day.group(3) or 0), int(month_day.group(4) or 0),
            )
        if "昨天" in value:
            from datetime import timedelta
            yesterday = now - timedelta(days=1)
            clock = re.search(r"(\d{1,2}):(\d{2})", value)
            return yesterday.replace(
                hour=int(clock.group(1)) if clock else 0,
                minute=int(clock.group(2)) if clock else 0,
                second=0,
                microsecond=0,
            )
        if "今天" in value:
            clock = re.search(r"(\d{1,2}):(\d{2})", value)
            return now.replace(
                hour=int(clock.group(1)) if clock else 0,
                minute=int(clock.group(2)) if clock else 0,
                second=0,
                microsecond=0,
            )
    return None


def normalize_identity_text(value: str | None) -> str:
    return re.sub(r"\s+", "", value or "").casefold()


def parse_search_card_author(
    records: list[tuple[str, str, str]], title: str
) -> str | None:
    """Return an author only when the card exposes a semantic author node.

    Positional guessing is deliberately forbidden: snippets and authors are
    both short text in some WeChat builds, and a false match could skip a
    different account's same-title article.
    """
    semantic_tokens = ("author", "account", "nickname", "profile", "source", "wx_tap_link")
    normalized_title = normalize_identity_text(title)
    for value, control_type, class_name in records:
        candidate = value.strip()
        lowered_class = class_name.lower()
        if not candidate or normalize_identity_text(candidate) == normalized_title:
            continue
        if len(candidate) > 100 or parse_publish_time([(candidate, control_type, class_name)]):
            continue
        if any(token in lowered_class for token in semantic_tokens):
            return candidate
    return None


def article_identity_key(
    title: str | None,
    author_name: str | None,
    published_at: datetime | str | None,
) -> tuple[str, str, str] | None:
    if not title or not author_name or not published_at:
        return None
    if isinstance(published_at, str):
        try:
            published_at = datetime.fromisoformat(published_at)
        except ValueError:
            return None
    return (
        normalize_identity_text(title),
        normalize_identity_text(author_name),
        published_at.date().isoformat(),
    )


def search_card_fingerprint(
    title: str,
    records: list[tuple[str, str, str]],
) -> str | None:
    """Build a stable skip key only from a sufficiently identified card."""
    published_at = parse_publish_time(records)
    author = parse_search_card_author(records, title)
    identity = article_identity_key(title, author, published_at)
    if identity is None:
        return None
    extras = []
    normalized_title = normalize_identity_text(title)
    normalized_author = normalize_identity_text(author)
    for value, control_type, class_name in records:
        normalized = normalize_identity_text(value)
        if not normalized or normalized in {normalized_title, normalized_author}:
            continue
        if parse_publish_time([(value, control_type, class_name)]):
            continue
        extras.append(normalized[:160])
    payload = json.dumps(
        {"identity": identity, "card_text": extras[:8]},
        ensure_ascii=False,
        sort_keys=True,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def count_from_button(records: list[tuple[str, str, str]], class_token: str, zero_when_blank: bool) -> int | None:
    matched_button = False
    for value, control_type, class_name in records:
        if control_type == "Button" and class_token in class_name:
            matched_button = True
            parsed = parse_count(value)
            if parsed is not None:
                return parsed
            # A WeChat article may expose both a hidden empty button and the
            # visible metric button. Keep scanning instead of treating the
            # first empty accessibility node as a real zero.
    if matched_button and zero_when_blank:
        return 0
    return None


def parse_article_metrics(records: list[tuple[str, str, str]]) -> tuple[int | None, int | None, int | None, int | None]:
    like_count = count_from_button(records, "sns_praise_btn", zero_when_blank=True)
    favorite_count = count_from_button(records, "sns_love_btn", zero_when_blank=True)
    share_count = count_from_button(records, "sns_share_btn", zero_when_blank=False)
    comment_count = None
    if any("还没有留言" in value for value, _type, _class in records):
        comment_count = 0
    else:
        for value, control_type, class_name in records:
            if "comment" not in class_name and "留言" not in class_name:
                continue
            match = re.search(r"(?:留言|评论)\s*(\d+)", value)
            if match:
                comment_count = int(match.group(1))
                break
    return like_count, comment_count, favorite_count, share_count


def parse_article_details(records: list[tuple[str, str, str]]) -> dict:
    """Extract ODS-ready fields exposed by the PC WeChat accessibility tree."""
    published = parse_publish_time(records)
    date_index = next((i for i, (value, _type, _class) in enumerate(records)
                       if re.search(r"20\d{2}年\d{1,2}月\d{1,2}日", value)), None)
    author_name = None
    if date_index is not None:
        for value, control_type, class_name in reversed(records[:date_index]):
            if control_type in {"Button", "Hyperlink"} and "wx_tap_link" in class_name:
                author_name = value.strip()
                break
    location = None
    if date_index is not None:
        for value, control_type, class_name in records[date_index + 1:date_index + 5]:
            if control_type == "Text" and "rich_media_meta" in class_name:
                location = value.strip() or None
                break

    start = (date_index + 1) if date_index is not None else 0
    while start < len(records) and (
        "rich_media_meta" in records[start][2] or records[start][1] == "Hyperlink"
    ):
        start += 1
    end = next((i for i, (_value, _type, class_name) in enumerate(records[start:], start)
                if "media_tool_meta" in class_name), len(records))
    content_lines = []
    for value, control_type, class_name in records[start:end]:
        if control_type not in {"Text", "Hyperlink"}:
            continue
        if any(token in class_name for token in ("wx_profile", "appmsg_card_context", "rich_media_meta")):
            continue
        cleaned = re.sub(r"\s+", " ", value).strip()
        if cleaned and cleaned not in {"听全文", "图片"}:
            content_lines.append(cleaned)
    content = "\n".join(content_lines).strip() or None
    image_count = sum(1 for _value, control_type, class_name in records[start:end]
                      if control_type == "Image" and "rich_pages" in class_name)
    video_count = sum(1 for _value, control_type, class_name in records[start:end]
                      if control_type in {"Video", "Media"} or "video" in class_name.lower())
    return {
        "author_name": author_name,
        "published_at": published.isoformat(sep=" ", timespec="seconds") if published else None,
        "location": location,
        "content": content,
        "image_count": image_count,
        "video_count": video_count,
    }


def copy_article_link_from_menu(web, config: dict) -> str:
    old_clipboard = get_clipboard()
    more = [
        control for control in web.descendants(control_type="Button")
        if control.window_text() == "更多" and control.element_info.class_name == "AppMenuButton"
    ]
    if len(more) != 1:
        raise RuntimeError(f"未找到唯一的文章菜单，当前找到 {len(more)} 个")
    more[0].click_input()
    time.sleep(config.get("wait_short", 0.8))

    def menu_items():
        items = []
        for window in Desktop(backend="uia").windows():
            items.extend([
                control for control in window.descendants(control_type="MenuItem")
                if control.window_text() == "复制链接"
            ])
        return items if len(items) == 1 else None

    items = wait_for(menu_items, timeout=3.0)
    items[0].click_input()
    time.sleep(config.get("wait_short", 0.8))
    link = get_clipboard().strip()
    # Retrying the same article legitimately produces the same clipboard URL.
    # Freshness is not a validity signal; only the URL shape is authoritative.
    if not link.startswith("https://mp.weixin.qq.com/"):
        raise RuntimeError("复制链接失败，剪贴板不是有效的微信公众号链接")
    return link


def is_search_results_page(web) -> bool:
    return any("搜一搜" in (control.window_text() or "") for control in web.descendants(control_type="Document"))


def close_stacked_detail_tabs_until_search_results(
    surface, *, max_tabs: int = 32, settle_seconds: float = 0.35,
) -> int:
    """Close accumulated article/video tabs until the 搜一搜 tab is active."""
    current = surface
    closed = 0
    for _attempt in range(max(1, int(max_tabs))):
        if is_search_results_page(current):
            break
        current.set_focus()
        keyboard.send_keys("^w", pause=0.02)
        closed += 1
        time.sleep(max(0.1, float(settle_seconds)))
        windows = web_windows()
        if not windows:
            current = None
            break
        if len(windows) != 1:
            raise SearchWindowUnavailableError(
                f"关闭历史详情标签时检测到 {len(windows)} 个微信网页窗口"
            )
        current = windows[0]
        if is_search_results_page(current):
            break
    else:
        raise SearchWindowUnavailableError(
            f"连续关闭 {max_tabs} 个详情标签后仍未恢复搜索结果页"
        )
    print(json.dumps({
        "event": "stacked_detail_tabs_closed",
        "closed": closed,
        "search_results_ready": bool(current is not None and is_search_results_page(current)),
    }, ensure_ascii=False), flush=True)
    return closed


def return_to_search_results() -> None:
    global SEARCH_LAYOUT
    surface = try_get_search_window()
    if surface is None:
        return
    if is_search_results_page(surface):
        return
    if SEARCH_LAYOUT == "detached":
        closed_tabs = close_stacked_detail_tabs_until_search_results(surface)
        if closed_tabs == 0:
            backs = [
                control for control in surface.descendants(control_type="Button")
                if (control.window_text() or "").startswith("点按该按钮可后退")
            ]
            if backs:
                backs[0].click_input()
            else:
                surface.set_focus()
                keyboard.send_keys("%{LEFT}", pause=0.02)
    else:
        backs = [
            control for control in surface.descendants(control_type="Button")
            if (control.window_text() or "").startswith("点按该按钮可后退")
            or (control.window_text() or "").strip() in {"返回", "后退"}
        ]
        if backs:
            backs[0].click_input()
        else:
            # The embedded Chromium view keeps normal navigation history even
            # when its back button is omitted from the UIA tree.
            surface.set_focus()
            keyboard.send_keys("%{LEFT}", pause=0.02)
    wait_for(
        lambda: (
            (SEARCH_LAYOUT == "detached" and not web_windows())
            or (has_search_window() and is_search_results_page(get_search_window()))
        ),
        timeout=8.0,
    )


def collect_live_keyword(
    keyword: str,
    project: str,
    cutoff: datetime,
    config: dict,
    connection: sqlite3.Connection,
    screenshot_dir: Path,
    max_results: int,
    stop_file: Path | None = None,
) -> dict:
    search_in_web(keyword, config)
    screenshot_dir.mkdir(parents=True, exist_ok=True)
    save_window_screenshot(get_search_window(), screenshot_dir / f"{keyword}_results.png")
    result_index, old_streak, previous_status = load_task_state(connection, keyword, project)
    # Search results are sorted newest-first and new articles can be inserted
    # above a persisted numeric index between runs.  Therefore a fresh process
    # must rebuild the scan from the true top after any partial/interrupted
    # status.  URL-key upserts keep this idempotent while preventing both the
    # visual "start in the middle" problem and missed newly published posts.
    restart_partial_from_top = bool(config.get("restart_partial_from_top", True))
    if restart_partial_from_top and previous_status and not previous_status.startswith("complete_"):
        print(json.dumps({
            "event": "partial_checkpoint_rescan_from_top",
            "keyword": keyword,
            "previous_status": previous_status,
            "previous_next_result_index": result_index,
        }, ensure_ascii=False), flush=True)
        result_index = 0
        old_streak = 0
    # Pagination v2 rechecks legacy complete_exhausted checkpoints because
    # older builds could mistake a non-growing first lazy-load batch for EOF.
    recheck_exhausted = bool(config.get("recheck_legacy_exhausted", True))
    if previous_status.startswith("complete_") and not (
        previous_status == "complete_exhausted" and recheck_exhausted
    ):
        return {
            "keyword": keyword,
            "saved": 0,
            "failed": 0,
            "examined": 0,
            "status": previous_status,
            "complete": True,
        }
    existing_links = {
        row[0] for row in connection.execute(
            "SELECT url FROM captured_article WHERE keyword=?", (keyword,)
        ).fetchall()
    }
    # A title alone is never a safe identity key: different accounts often
    # republish the same headline. Two conservative accelerators are used:
    # (1) skip without opening only when title+semantic card author+date maps to
    # exactly one stored row; (2) otherwise open briefly and verify the stable
    # article URL before skipping expensive detail/metric extraction.
    saved_titles = {
        row[0] for row in connection.execute(
            "SELECT title FROM captured_article WHERE keyword=? AND title != ''",
            (keyword,),
        ).fetchall()
    }
    known_identity_counts: dict[tuple[str, str, str], int] = {}
    for known_title, known_author, known_published_at in connection.execute(
        """
        SELECT title, author_name, published_at
        FROM captured_article
        WHERE keyword=? AND title != '' AND author_name IS NOT NULL
          AND published_at IS NOT NULL
        """,
        (keyword,),
    ).fetchall():
        identity = article_identity_key(known_title, known_author, known_published_at)
        if identity:
            known_identity_counts[identity] = known_identity_counts.get(identity, 0) + 1
    skipped_position_fingerprints = {
        row[0] for row in connection.execute(
            """
            SELECT fingerprint FROM collection_issue
            WHERE keyword=? AND issue_type='position_unreachable_skipped'
              AND fingerprint IS NOT NULL AND fingerprint != ''
            """,
            (keyword,),
        ).fetchall()
    }
    skipped_known_articles = 0
    consecutive_position_skips = 0
    saved = 0
    failed = 0
    processed_this_run = 0
    first_failed_index = None
    status = "running"
    update_task_state(connection, keyword, project, result_index, old_streak, status)

    while max_results <= 0 or processed_this_run < max_results:
        if STOP_REQUESTED or (stop_file and stop_file.exists()):
            status = "partial_stopped"
            update_task_state(
                connection, keyword, project, result_index, old_streak, status,
                "graceful stop requested",
            )
            break
        if not has_search_window():
            search_in_web(keyword, config)
        web = get_search_window()
        results = search_result_buttons(web)
        # A resumed checkpoint can be several lazy-load batches beyond the
        # initial DOM (for example index 120 while only 36 controls exist).
        # Keep loading until the requested index actually exists. The former
        # single load attempt could grow 36 -> 56 and then still access [120].
        exhausted_before_index = False
        while result_index >= len(results):
            previous_count = len(results)
            results, grew = load_more_search_results(web, previous_count, config)
            if result_index >= len(results) and not grew:
                status = "complete_exhausted"
                update_task_state(connection, keyword, project, result_index, old_streak, status)
                exhausted_before_index = True
                break
            web = get_search_window()
        if exhausted_before_index:
            break

        result = results[result_index]
        title = result.window_text().strip()
        initial_card_records = control_element_records(result, include_self=True)
        initial_card_fingerprint = search_card_fingerprint(title, initial_card_records)
        if (
            initial_card_fingerprint
            and initial_card_fingerprint in skipped_position_fingerprints
        ):
            failed_index = result_index
            result_index += 1
            processed_this_run += 1
            update_task_state(
                connection, keyword, project, result_index, old_streak, "running",
                f"结果索引 {failed_index} 命中历史定位失败指纹，快速跳过",
            )
            print(json.dumps({
                "event": "rescan_skip_position_fingerprint",
                "keyword": keyword,
                "index": failed_index,
                "title": title,
                "fingerprint": initial_card_fingerprint,
            }, ensure_ascii=False), flush=True)
            continue
        item_success = False
        last_item_exception = None
        for item_attempt in range(1, 4):
            try:
                # Refresh virtualized UIA elements before every retry. Some
                # WeChat builds expose Invoke but silently ignore it; use a
                # physical click for subsequent attempts.
                refreshed = search_result_buttons(get_search_window())
                if result_index < len(refreshed):
                    result = refreshed[result_index]
                    title = result.window_text().strip()
                visible_result = bring_search_result_into_view(result_index, config)
                if visible_result is None:
                    raise ResultPositionError(f"结果索引 {result_index} 无法滚动到可视区域")
                result = visible_result
                title = result.window_text().strip()
                card_records = control_element_records(result, include_self=True)
                initial_card_records = card_records
                initial_card_fingerprint = search_card_fingerprint(title, card_records)
                card_published_at = parse_publish_time(card_records)
                card_author = parse_search_card_author(card_records, title)
                card_identity = article_identity_key(title, card_author, card_published_at)
                if (
                    bool(config.get("skip_known_fingerprint_on_rescan", True))
                    and card_identity
                    and known_identity_counts.get(card_identity) == 1
                    and result_index != first_failed_index
                ):
                    skipped_known_articles += 1
                    old_streak = 0
                    result_index += 1
                    processed_this_run += 1
                    update_task_state(
                        connection, keyword, project,
                        first_failed_index if first_failed_index is not None else result_index,
                        old_streak, "running",
                    )
                    print(json.dumps({
                        "event": "rescan_skip_known_fingerprint",
                        "keyword": keyword,
                        "index": result_index - 1,
                        "title": title,
                        "author": card_author,
                        "published_date": card_published_at.date().isoformat(),
                    }, ensure_ascii=False), flush=True)
                    item_success = True
                    break
                result.click_input()
                full_article_wait = float(config.get("wait_article", 3.0))
                known_title_candidate = title in saved_titles and result_index != first_failed_index
                initial_wait = (
                    min(full_article_wait, float(config.get("wait_known_article", 1.2)))
                    if known_title_candidate
                    else full_article_wait
                )
                time.sleep(initial_wait)
                wait_for(
                    lambda: any(
                        "搜一搜" not in (document.window_text() or "")
                        for document in get_search_window().descendants(control_type="Document")
                    ),
                    timeout=5.0,
                )
                web = get_search_window()
                link = copy_article_link_from_menu(web, config)
                if (
                    bool(config.get("fast_verify_known_url_on_rescan", True))
                    and known_title_candidate
                    and link in existing_links
                ):
                    return_to_search_results()
                    skipped_known_articles += 1
                    old_streak = 0
                    result_index += 1
                    processed_this_run += 1
                    update_task_state(
                        connection, keyword, project,
                        first_failed_index if first_failed_index is not None else result_index,
                        old_streak, "running",
                    )
                    print(json.dumps({
                        "event": "rescan_skip_verified_url",
                        "keyword": keyword,
                        "index": result_index - 1,
                        "title": title,
                        "url": link,
                    }, ensure_ascii=False), flush=True)
                    item_success = True
                    break
                if initial_wait < full_article_wait:
                    time.sleep(full_article_wait - initial_wait)
                    web = get_search_window()
                records = page_element_records(web)
                published_at = parse_publish_time(records) or card_published_at
                details = parse_article_details(records)
                if published_at is not None and details["published_at"] is None:
                    details["published_at"] = published_at.isoformat(
                        sep=" ", timespec="seconds"
                    )
                like_count, comment_count, favorite_count, share_count = parse_article_metrics(records)
                return_to_search_results()

                # A Chromium navigation can expose the menu/link before the
                # article metadata has hydrated. Treating that intermediate
                # state as success silently drops the row because its date is
                # unknown. Re-open it through the normal retry loop instead.
                if published_at is None:
                    raise ArticleMetadataUnavailableError(
                        "详情页和搜索结果卡片均未提供发布时间"
                    )

                in_period = published_at is not None and published_at >= cutoff
                if published_at is None:
                    # Keep an unknown date out of the requested date window;
                    # it also must not reset an existing old-article streak.
                    pass
                elif published_at < cutoff:
                    old_streak += 1
                else:
                    old_streak = 0
                    save_row(connection, ResultRow(
                        project=project,
                        keyword=keyword,
                        title=title,
                        url=link,
                        author_name=details["author_name"],
                        published_at=details["published_at"],
                        location=details["location"],
                        content=details["content"],
                        image_count=details["image_count"],
                        video_count=details["video_count"],
                        like_count=like_count,
                        comment_count=comment_count,
                        favorite_count=favorite_count,
                        share_count=share_count,
                    ))
                    if link not in existing_links:
                        existing_links.add(link)
                        saved += 1
                result_index += 1
                processed_this_run += 1
                update_task_state(
                    connection, keyword, project,
                    first_failed_index if first_failed_index is not None else result_index,
                    old_streak, "running"
                )
                print(json.dumps({
                    "keyword": keyword,
                    "index": result_index - 1,
                    "title": title,
                    "published_at": published_at.isoformat() if published_at else None,
                    "url": link,
                    "like": like_count,
                    "comment": comment_count,
                    "favorite_or_recommend": favorite_count,
                    "share": share_count,
                    "saved": in_period,
                }, ensure_ascii=False))
                item_success = True
                break
            except Exception as exc:
                last_item_exception = exc
                failed += 1
                print(json.dumps({
                    "keyword": keyword,
                    "index": result_index,
                    "title": title,
                    "attempt": item_attempt,
                    "error": str(exc),
                }, ensure_ascii=False))
                try:
                    if has_search_window() and not is_search_results_page(get_search_window()):
                        return_to_search_results()
                except Exception:
                    pass
                if not has_search_window():
                    try:
                        search_in_web(keyword, config)
                    except Exception:
                        pass
                # This is a page-state failure, not a malformed article.
                # Rebuild the same keyword from 最新/top, then lazy-load back
                # to the unchanged checkpoint before the next retry.
                if isinstance(exc, ResultPositionError) and item_attempt < 3:
                    try:
                        print(json.dumps({
                            "event": "search_result_position_recovery",
                            "keyword": keyword,
                            "index": result_index,
                            "attempt": item_attempt,
                        }, ensure_ascii=False), flush=True)
                        search_in_web(keyword, config)
                        recovered = search_result_buttons(get_search_window())
                        while result_index >= len(recovered):
                            previous_count = len(recovered)
                            recovered, grew = load_more_search_results(
                                get_search_window(), previous_count, config
                            )
                            if result_index >= len(recovered) and not grew:
                                break
                    except Exception as recovery_exc:
                        print(json.dumps({
                            "event": "search_result_position_recovery_failed",
                            "keyword": keyword,
                            "index": result_index,
                            "error": str(recovery_exc),
                        }, ensure_ascii=False), flush=True)
                if has_search_window() and is_search_results_page(get_search_window()):
                    refreshed = search_result_buttons(get_search_window())
                    if result_index < len(refreshed):
                        result = refreshed[result_index]
        if not item_success:
            failed_index = result_index
            if isinstance(last_item_exception, ArticleMetadataUnavailableError):
                consecutive_position_skips = 0
                # Three independent opens have already failed.  Persist an
                # explicit audit record and continue toward the cutoff; this
                # single unverifiable/deleted article must not hold an entire
                # keyword (or worker) hostage forever.
                save_collection_issue(
                    connection, keyword, failed_index, "missing_publish_time",
                    title, str(last_item_exception),
                )
                result_index += 1
                processed_this_run += 1
                update_task_state(
                    connection, keyword, project, result_index, old_streak,
                    "running", f"结果索引 {failed_index} 无可信发布时间，已审计后跳过",
                )
                print(json.dumps({
                    "event": "article_skipped_missing_publish_time",
                    "keyword": keyword,
                    "index": failed_index,
                    "title": title,
                    "checkpoint_advanced": True,
                }, ensure_ascii=False), flush=True)
                continue
            if isinstance(last_item_exception, ResultPositionError):
                card_fingerprint = initial_card_fingerprint or search_card_fingerprint(
                    title, initial_card_records
                )
                card_published_at = parse_publish_time(initial_card_records)
                card_author = parse_search_card_author(initial_card_records, title)
                detail = json.dumps({
                    "error": str(last_item_exception),
                    "title": title,
                    "author": card_author,
                    "published_at": (
                        card_published_at.isoformat() if card_published_at else None
                    ),
                    "card_text": [record[0] for record in initial_card_records[:12]],
                }, ensure_ascii=False)
                save_collection_issue(
                    connection, keyword, failed_index,
                    "position_unreachable_skipped", title, detail,
                    fingerprint=card_fingerprint,
                )
                if card_fingerprint:
                    skipped_position_fingerprints.add(card_fingerprint)
                consecutive_position_skips += 1
                result_index += 1
                processed_this_run += 1
                status = "running"
                update_task_state(
                    connection, keyword, project, result_index, old_streak,
                    status,
                    f"结果索引 {failed_index} 定位失败，已记录指纹并跳过",
                )
                print(json.dumps({
                    "event": "search_result_position_skipped",
                    "keyword": keyword,
                    "index": failed_index,
                    "title": title,
                    "fingerprint": card_fingerprint,
                    "checkpoint_advanced": True,
                    "consecutive_position_skips": consecutive_position_skips,
                }, ensure_ascii=False), flush=True)
                max_position_skips = max(
                    1, int(config.get("max_consecutive_position_skips", 3))
                )
                if consecutive_position_skips >= max_position_skips:
                    status = "partial_position_error"
                    update_task_state(
                        connection, keyword, project, result_index, old_streak,
                        status,
                        f"连续 {consecutive_position_skips} 篇无法定位，疑似整个结果页失效",
                    )
                    print(json.dumps({
                        "event": "search_result_position_circuit_breaker",
                        "keyword": keyword,
                        "next_index": result_index,
                        "consecutive_position_skips": consecutive_position_skips,
                    }, ensure_ascii=False), flush=True)
                    break
                continue
            consecutive_position_skips = 0
            if first_failed_index is None:
                first_failed_index = failed_index
            save_collection_issue(
                connection, keyword, failed_index, "article_retry_exhausted",
                title, str(last_item_exception or "unknown article error"),
            )
            # One malformed/unresponsive result must not prevent the keyword
            # from reaching later dates. Keep the earliest failed index as the
            # durable retry checkpoint while advancing locally in this run.
            result_index += 1
            processed_this_run += 1
            status = "running"
            update_task_state(
                connection, keyword, project, first_failed_index, old_streak,
                status, f"结果索引 {failed_index} 连续 3 次采集失败；本轮继续向后采集",
            )
            print(json.dumps({
                "event": "article_skipped_after_retries",
                "keyword": keyword,
                "index": failed_index,
                "title": title,
                "retry_checkpoint": first_failed_index,
            }, ensure_ascii=False), flush=True)
            continue
        consecutive_position_skips = 0
        if STOP_REQUESTED or (stop_file and stop_file.exists()):
            status = "partial_stopped"
            update_task_state(
                connection, keyword, project, result_index, old_streak, status,
                "graceful stop requested after current article",
            )
            break
        # The search page is required to select 最新 before this loop starts.
        # With cutoff_old_streak=1, the first dated article older than cutoff is
        # the time boundary and completes the current keyword immediately.
        cutoff_old_streak = int(config.get("cutoff_old_streak", 0))
        if cutoff_old_streak > 0 and old_streak >= cutoff_old_streak:
            status = "complete_cutoff"
            update_task_state(connection, keyword, project, result_index, old_streak, status)
            break

    if status == "running" and max_results > 0 and processed_this_run >= max_results:
        status = "partial_limit"
        update_task_state(
            connection, keyword, project, result_index, old_streak, status,
            f"本次达到安全上限 {max_results} 条，需继续当前关键词",
        )
    if (
        first_failed_index is not None
        and not status.startswith("complete_")
        and status != "partial_stopped"
    ):
        status = "partial_error"
        update_task_state(
            connection, keyword, project, first_failed_index, old_streak, status,
            "已继续采集后续结果，但仍有失败文章待下次从最早失败索引复查",
        )
    complete = status.startswith("complete_")
    return {
        "keyword": keyword,
        "saved": saved,
        "failed": failed,
        "examined": processed_this_run,
        "skipped_known_articles": skipped_known_articles,
        "next_result_index": result_index,
        "status": status,
        "complete": complete,
    }


def export_template(template: Path, output: Path, connection: sqlite3.Connection, keyword: str | None = None) -> None:
    workbook = openpyxl.load_workbook(template)
    sheet = workbook["采集结果"]
    # Keep the supplied layout while adding the requested title field after keyword.
    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    if "文章标题" not in headers:
        sheet.insert_cols(3, 1)
        source = sheet.cell(1, 4)
        target = sheet.cell(1, 3)
        if source.has_style:
            from copy import copy
            target._style = copy(source._style)
        target.value = "文章标题"
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    if keyword:
        rows = connection.execute(
            """
            SELECT project, keyword, title, url, like_count, comment_count, favorite_count, share_count
            FROM captured_article WHERE keyword=?
            ORDER BY project, keyword, captured_at, url
            """,
            (keyword,),
        ).fetchall()
    else:
        rows = connection.execute(
            """
            SELECT project, keyword, title, url, like_count, comment_count, favorite_count, share_count
            FROM captured_article
            ORDER BY project, keyword, captured_at, url
            """
        ).fetchall()
    for values in rows:
        sheet.append(values)
    # Make the newly added title and URL columns readable without changing metric semantics.
    sheet.column_dimensions["C"].width = max(sheet.column_dimensions["C"].width or 0, 55)
    sheet.column_dimensions["D"].width = max(sheet.column_dimensions["D"].width or 0, 60)
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=3, max_col=4):
        for cell in row_cells:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
        sheet.row_dimensions[row_cells[0].row].height = 32
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def export_ods_local(output_dir: Path, connection: sqlite3.Connection, keywords: list[str] | None = None) -> dict:
    params: list[str] = []
    where = ""
    if keywords:
        where = f"WHERE keyword IN ({','.join('?' for _ in keywords)})"
        params.extend(keywords)
    rows = connection.execute(
        f"""
        SELECT article_key, project, keyword, title, url, author_name, published_at,
               location, content, image_count, video_count, like_count, comment_count,
               favorite_count, share_count, captured_at
        FROM captured_article {where}
        ORDER BY captured_at, keyword, url
        """,
        params,
    ).fetchall()
    records = []
    seen_urls = set()
    for row in rows:
        (article_key, project, keyword, title, url, author_name, published_at,
         location, content, image_count, video_count, like_count, comment_count,
         favorite_count, share_count, captured_at) = row
        if url in seen_urls:
            continue
        seen_urls.add(url)
        content_type = "视频" if (video_count or 0) else "图文" if (image_count or 0) else "纯文本"
        record = {field: None for field in ODS_FIELDS}
        record.update({
            "作者名字": author_name,
            "文章标题": title,
            "文章内容": content,
            "文章时间地点": published_at,
            "文章点赞数": like_count or 0,
            "文章收藏数": favorite_count or 0,
            "文章评论数": comment_count or 0,
            "文章链接": url,
            "文章id": article_key,
            "关键词": keyword,
            "作者粉丝数": 0,
            "作者点赞数": 0,
            "文章分享数": share_count or 0,
            "平台": "微信公众号",
            "文章获取时间": captured_at,
            "项目": project,
            "地点": location,
            "内容类型": content_type,
        })
        records.append(record)
    output_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = output_dir / "ods_social_media_post_wechat.jsonl"
    csv_path = output_dir / "ods_social_media_post_wechat.csv"
    jsonl_path.write_text("".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records), encoding="utf-8")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=ODS_FIELDS)
        writer.writeheader()
        writer.writerows(records)
    return {"records": len(records), "jsonl": str(jsonl_path), "csv": str(csv_path)}


def backfill_titles_live(keyword: str, config: dict, connection: sqlite3.Connection, screenshot_dir: Path) -> dict:
    """Recover titles for already captured rows by matching each live result URL."""
    targets = {row[0]: row[1] for row in connection.execute(
        "SELECT url, article_key FROM captured_article WHERE keyword=?", (keyword,)
    ).fetchall()}
    search_in_web(keyword, config)
    result_index = 0
    matched = 0
    seen = set()
    while targets:
        web = get_search_window()
        results = search_result_buttons(web)
        if result_index >= len(results):
            if results:
                results[-1].scroll_into_view()
            time.sleep(config.get("wait_page", 3.0))
            results = search_result_buttons(get_search_window())
            if result_index >= len(results):
                break
        result = results[result_index]
        title = result.window_text().strip()
        for attempt in range(3):
            try:
                result.invoke()
                time.sleep(config.get("wait_article", 3.0))
                wait_for(lambda: any("搜一搜" not in (d.window_text() or "")
                                    for d in get_search_window().descendants(control_type="Document")), 5.0)
                link = copy_article_link_from_menu(get_search_window(), config)
                if link in targets:
                    connection.execute("UPDATE captured_article SET title=? WHERE keyword=? AND url=?",
                                       (title, keyword, link))
                    connection.commit()
                    matched += 1
                    targets.pop(link, None)
                    print(json.dumps({"keyword": keyword, "index": result_index, "title": title, "url": link, "matched": matched}, ensure_ascii=False))
                return_to_search_results()
                break
            except Exception as exc:
                if attempt == 2:
                    print(json.dumps({"keyword": keyword, "index": result_index, "error": str(exc)}, ensure_ascii=False))
                try:
                    if has_search_window() and not is_search_results_page(get_search_window()):
                        return_to_search_results()
                except Exception:
                    pass
        result_index += 1
    return {"keyword": keyword, "matched": matched, "remaining": len(targets), "examined": result_index}


def search_keyword(keyword: str, config: dict, screenshot_dir: Path) -> None:
    """进入 PC 微信搜一搜文章页。结果遍历和链接复制由配置坐标控制。"""
    win = get_wechat_window()
    keyboard.send_keys("^f", pause=0.02)
    time.sleep(config.get("wait_short", 0.8))
    paste_text(keyword)
    time.sleep(config.get("wait_search_suggest", 1.5))
    save_window_screenshot(win, screenshot_dir / f"{keyword}_01_suggest.png")
    if config.get("submit_search_with_enter", False):
        keyboard.send_keys("{ENTER}", pause=0.02)
    else:
        click_relative(win, config["points"]["soso_entry"])
    time.sleep(config.get("wait_page", 3.0))
    if "article_tab" in config["points"]:
        click_relative(win, config["points"]["article_tab"])
        time.sleep(config.get("wait_page", 3.0))
    save_window_screenshot(win, screenshot_dir / f"{keyword}_02_articles.png")


def copy_current_article_link(config: dict) -> str:
    win = get_wechat_window()
    before = get_clipboard()
    click_relative(win, config["points"]["article_more"])
    time.sleep(config.get("wait_short", 0.8))
    click_relative(win, config["points"]["copy_link"])
    time.sleep(config.get("wait_short", 0.8))
    value = get_clipboard().strip()
    if value == before or not value.startswith(("http://", "https://")):
        raise RuntimeError("未从剪贴板取得新的文章链接；请重新校准 article_more/copy_link")
    return value


def command_ingest(args) -> None:
    connection = init_state(Path(args.state_db))
    row = ResultRow(
        project=args.project,
        keyword=args.keyword,
        title=args.title or "",
        url=args.url,
        like_count=parse_count(args.like),
        comment_count=parse_count(args.comment),
        favorite_count=parse_count(args.favorite),
        share_count=parse_count(args.share),
    )
    save_row(connection, row)
    print(json.dumps(asdict(row), ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description="PC 微信公众号搜一搜采集器")
    sub = parser.add_subparsers(dest="command", required=True)

    probe = sub.add_parser("probe", help="检查微信窗口")

    search = sub.add_parser("search", help="打开一个关键词的文章搜索页")
    search.add_argument("keyword")
    search.add_argument("--config", default="collector_config.json")
    search.add_argument("--screenshots", default="work/screenshots")

    ingest = sub.add_parser("ingest", help="将已核验的链接和计数写入断点库")
    ingest.add_argument("--project", required=True)
    ingest.add_argument("--keyword", required=True)
    ingest.add_argument("--url", required=True)
    ingest.add_argument("--title", default="")
    ingest.add_argument("--like")
    ingest.add_argument("--comment")
    ingest.add_argument("--favorite")
    ingest.add_argument("--share")
    ingest.add_argument("--state-db", default="work/wechat_collect.sqlite")

    export = sub.add_parser("export", help="将断点库导出到用户结果表")
    export.add_argument("--template", required=True)
    export.add_argument("--output", required=True)
    export.add_argument("--state-db", default="work/wechat_collect.sqlite")
    export.add_argument("--keyword")

    export_ods = sub.add_parser("export-ods", help="按32字段ODS结构导出本地JSONL和CSV")
    export_ods.add_argument("--output-dir", required=True)
    export_ods.add_argument("--state-db", default="work/wechat_collect.sqlite")
    export_ods.add_argument("--keyword", action="append")

    backfill = sub.add_parser("backfill-titles", help="从微信搜一搜结果回填已采集文章标题")
    backfill.add_argument("--keyword", required=True)
    backfill.add_argument("--config", default="collector_config.json")
    backfill.add_argument("--state-db", default="work/wechat_collect.sqlite")
    backfill.add_argument("--screenshots", default="work/screenshots")

    keywords = sub.add_parser("keywords", help="预览待采集关键词")
    keywords.add_argument("--input", required=True)
    keywords.add_argument("--limit", type=int)

    live = sub.add_parser("collect-live", help="在已经打开的微信搜一搜网页窗口中批量实采")
    live.add_argument("--input", required=True, help="关键词工作簿")
    live.add_argument("--config", default="collector_config.json")
    live.add_argument("--keyword", action="append", help="只采指定检索词，可重复传入")
    live.add_argument("--keyword-limit", type=int, default=4)
    live.add_argument("--cutoff", default=DEFAULT_START_DATE)
    # Zero means unlimited. Production workers must stop on the requested date
    # boundary or true list exhaustion, not after an arbitrary 100 articles.
    live.add_argument("--max-results", type=int, default=0)
    live.add_argument("--state-db", default="work/wechat_collect.sqlite")
    live.add_argument("--stop-file", help="存在时在当前文章完成后安全停止")
    live.add_argument("--screenshots", default="work/screenshots")

    args = parser.parse_args()
    if args.command == "probe":
        win = get_wechat_window()
        rect = win.rectangle()
        print(json.dumps({"title": win.window_text(), "rect": [rect.left, rect.top, rect.right, rect.bottom]}, ensure_ascii=False))
    elif args.command == "search":
        search_keyword(args.keyword, load_config(Path(args.config)), Path(args.screenshots))
    elif args.command == "ingest":
        command_ingest(args)
    elif args.command == "export":
        connection = init_state(Path(args.state_db))
        export_template(Path(args.template), Path(args.output), connection, args.keyword)
    elif args.command == "export-ods":
        connection = init_state(Path(args.state_db))
        print(json.dumps(export_ods_local(Path(args.output_dir), connection, args.keyword), ensure_ascii=False, indent=2))
    elif args.command == "backfill-titles":
        connection = init_state(Path(args.state_db))
        result = backfill_titles_live(args.keyword, load_config(Path(args.config)), connection, Path(args.screenshots))
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif args.command == "keywords":
        print(json.dumps(load_keywords(Path(args.input), args.limit), ensure_ascii=False, indent=2))
    elif args.command == "collect-live":
        signal.signal(signal.SIGINT, request_graceful_stop)
        if hasattr(signal, "SIGBREAK"):
            signal.signal(signal.SIGBREAK, request_graceful_stop)
        all_keywords = load_keywords(Path(args.input))
        selected = []
        requested = set(args.keyword or [])
        for keyword, project in all_keywords:
            if requested and keyword not in requested:
                continue
            selected.append((keyword, project))
            if not requested and len(selected) >= args.keyword_limit:
                break
        if requested and {keyword for keyword, _ in selected} != requested:
            missing = requested - {keyword for keyword, _ in selected}
            raise ValueError(f"关键词表中不存在：{sorted(missing)}")
        config = load_config(Path(args.config))
        connection = init_state(Path(args.state_db))
        cutoff = datetime.fromisoformat(args.cutoff)
        summaries = []
        for keyword, project in selected:
            summary = collect_live_keyword(
                keyword=keyword,
                project=project,
                cutoff=cutoff,
                config=config,
                connection=connection,
                screenshot_dir=Path(args.screenshots),
                max_results=args.max_results,
                stop_file=Path(args.stop_file) if args.stop_file else None,
            )
            summaries.append(summary)
            # Never advance to another keyword while the current keyword is
            # only partially collected. A later invocation resumes from its
            # persisted next_result_index.
            if not summary["complete"]:
                break
        print(json.dumps({"summaries": summaries}, ensure_ascii=False, indent=2))
        # A partial keyword is not considered complete. The worker wrapper
        # will still export/import the rows already committed to SQLite, then
        # retry this same keyword from crawl_task.next_result_index.
        if summaries and any(not summary.get("complete") for summary in summaries):
            raise SystemExit(4)


if __name__ == "__main__":
    main()
