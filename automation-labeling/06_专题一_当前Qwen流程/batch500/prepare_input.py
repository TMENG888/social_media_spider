"""Read the actual user workbook; freeze the verified matching raw input only."""
import collections,hashlib,json,sys,decimal
from pathlib import Path
import openpyxl
sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(__file__).parent
SOURCE=Path('E:/Codex/outputs/20261022_两专题待打标原始数据/专题一_原始文章500篇_待Qwen打标.xlsx')
PREP=Path('E:/Codex/work/20261022_两专题原始数据准备/T1_documents.json')
wb=openpyxl.load_workbook(SOURCE,read_only=True,data_only=False)
ws=wb['原始文章'];hs=list(next(ws.iter_rows(min_row=5,max_row=5,values_only=True)))
rows=[dict(zip(hs,r)) for r in ws.iter_rows(min_row=6,values_only=True) if r[0] is not None]
docs=json.loads(PREP.read_text(encoding='utf8'))
assert len(rows)==len(docs)==500
for row,doc in zip(rows,docs):
    assert str(row['id'])==doc['source_id']
    assert row['项目']==doc['topic']=='全国地产设计产品力监测'
    for h,v in doc['row'].items():
        actual='' if row.get(h) is None else str(row[h]);expected='' if v is None else str(v)
        if isinstance(row.get(h),(int,float)) and expected:
            assert decimal.Decimal(actual)==decimal.Decimal(expected),(doc['record_id'],h)
        else:assert actual==expected,(doc['record_id'],h)
    assert row[doc['field']]==doc['text'] and doc['kind']=='文章'
assert len({d['source_id'] for d in docs})==500
assert len({(d['row']['平台'],d['text']) for d in docs})==500
wb.close()
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
(ROOT/'documents.json').write_text(json.dumps(docs,ensure_ascii=False,indent=2),encoding='utf8')
manifest={'source_workbook':str(SOURCE),'source_workbook_sha256':sha(SOURCE),
          'input_sha256':sha(ROOT/'documents.json'),'article_count':len(docs),
          'platforms':dict(collections.Counter(d['row']['平台'] for d in docs)),
          'characters':sum(len(d['text']) for d in docs),'empty_texts':sum(not d['text'].strip() for d in docs),
          'policy':'用户授权处理该500篇专题一文章，不筛广告；沿用r49，非生产切换。',
          'output_dir':'E:/Codex/outputs/20261022_专题一500篇Qwen切片标注',
          'work_dir':str(ROOT),'temp_dir':'E:/Codex/temp/专题一500篇Qwen切片标注'}
(ROOT/'input_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(manifest,ensure_ascii=False))
