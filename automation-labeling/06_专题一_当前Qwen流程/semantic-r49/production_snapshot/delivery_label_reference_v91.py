"""Presentation-only label references; never select or modify assigned tags."""
from copy import deepcopy
from pathlib import Path
import openpyxl
from multilabel_v2 import T1, T2, load_taxonomy

DISPLAY_FIELD = '标签强调方向'
TAG_KEYS = ['标签路径ID','标签编号','一级标签','二级标签','三级标签','四级标签','五级标签','标签强调方向']
REFERENCE_FIELDS = TAG_KEYS + ['打标说明','专题标签词','属性','备注','原表行']


def display_emphases(tags):
    if not isinstance(tags, list):
        raise ValueError('标签明细必须是JSON数组，不得以空字符串掩盖格式错误')
    values = []
    for tag in tags:
        if not isinstance(tag, dict):
            raise ValueError('标签明细必须包含JSON对象')
        value = tag.get(DISPLAY_FIELD)
        if value is not None and not isinstance(value, str):
            raise ValueError('标签强调方向必须是字典文本或null')
        if value and value not in values:
            values.append(value)
    return ';'.join(values)


def display_schema(fields):
    fields = list(fields)
    if DISPLAY_FIELD in fields:
        fields.remove(DISPLAY_FIELD)
    fields.insert(fields.index('标签明细') + 1, DISPLAY_FIELD)
    return fields


def reference_tables(path):
    """Use the existing topic eligibility and row-order ID contract unchanged."""
    path = Path(path)
    taxonomy = load_taxonomy(path)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        headers = [str(v or '').strip() for v in next(sheet.iter_rows(min_row=2,max_row=2,values_only=True))]
        source = {n: dict(zip(headers, values)) for n, values in enumerate(sheet.iter_rows(min_row=3,values_only=True),3)}
        tables = {T1: [], T2: []}
        for tag in taxonomy:
            row = source[tag['source_row']]
            number = tag['source_number']
            if number and str(number).isdigit():
                number = int(number)
            canonical = dict(zip(TAG_KEYS,[tag['tag_id'],number or None,
                *[tag.get('l'+str(n)) or None for n in range(1,6)],tag['emphasis'] or None]))
            for topic in tag['topics']:
                topic_word = row.get('专题一标签' if topic == T1 else '专题二标签')
                if topic == T2 and tag['l1'] == '客户' and not str(topic_word or '').strip():
                    topic_word = canonical[DISPLAY_FIELD]
                tables[topic].append({**canonical,'打标说明':tag['instruction'],
                    '专题标签词':topic_word,'属性':row.get('属性'),'备注':row.get('备注'),'原表行':tag['source_row']})
        return tables, {'file':path.name,'sheet':sheet.title,
            'id_rule':'标签路径ID沿用全部有效标签的行序生成，不按专题重新编号',
            'topic_rule':'专题一按原表专题一标签列；专题二按专题二标签列并包含全部客户标签'}
    finally:
        book.close()


def enrich_rows(rows, reference):
    registry = {r['标签路径ID']:r for r in reference}
    output = []
    for row in rows:
        for tag in row['标签明细']:
            expected = registry.get(tag.get('标签路径ID'))
            if expected is None:
                raise ValueError('当前专题字典没有标签：'+str(tag.get('标签路径ID')))
            if set(tag) != set(TAG_KEYS) or any(tag[k] != expected[k] or type(tag[k]) is not type(expected[k]) for k in TAG_KEYS):
                raise ValueError('标签JSON与字典不一致，不得在展示导出时静默改写：'+str(tag.get('标签路径ID')))
        output.append({**deepcopy(row),DISPLAY_FIELD:display_emphases(row['标签明细'])})
    return output
