"""Independent, read-only checks of the label-reference delivery edit."""
import argparse
import copy
import json
from pathlib import Path

import openpyxl

TAG_KEYS = ['标签路径ID','标签编号','一级标签','二级标签','三级标签','四级标签','五级标签','标签强调方向']
REFERENCE_FIELDS = TAG_KEYS + ['打标说明','专题标签词','属性','备注','原表行']


def source_references(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        headers = [str(v or '').strip() for v in next(sheet.iter_rows(min_row=2,max_row=2,values_only=True))]
        result = [[], []]
        index = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
            raw = dict(zip(headers,values))
            text = {h: str(v).strip() if v is not None else '' for h,v in raw.items()}
            levels = [text.get(h,'') or None for h in TAG_KEYS[2:7]]
            if not any(levels):
                continue
            index += 1
            number = text.get('编号','')
            if number.isdigit():
                number = int(number)
            canonical = dict(zip(TAG_KEYS,[f'VOC-{index:04d}',number or None,*levels,text.get('标签强调方向','') or None]))
            for topic,field in enumerate(['专题一标签','专题二标签']):
                if not text.get(field) and not (topic == 1 and levels[0] == '客户'):
                    continue
                word = raw.get(field)
                if topic == 1 and levels[0] == '客户' and not str(word or '').strip():
                    word = canonical['标签强调方向']
                result[topic].append({**canonical,'打标说明':text['打标说明'],'专题标签词':word,
                    '属性':raw.get('属性'),'备注':raw.get('备注'),'原表行':row_number})
        return result
    finally:
        book.close()


def signature(cell):
    # Object-level formatting comparison avoids workbook-specific style IDs.
    return (cell.number_format,str(cell.font),str(cell.fill),str(cell.border),
        str(cell.alignment),str(cell.protection))


def same_cell(before,after,where):
    assert before.value == after.value, (where,'value changed',before.value,after.value)
    assert before.data_type == after.data_type, (where,'cell type changed')
    assert signature(before) == signature(after), (where,'format changed')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--prior-dir',type=Path,required=True)
    ap.add_argument('--output-dir',type=Path,required=True)
    ap.add_argument('--taxonomy',type=Path,required=True)
    a = ap.parse_args()
    references = source_references(a.taxonomy)
    checks = []
    for topic,kind in enumerate(['产品信息单元','观点信息单元']):
        dictionary = '专题一标签体系' if topic == 0 else '专题二标签体系'
        prior_payload = json.loads((a.prior_dir/(kind+'_workbook_payload.json')).read_text(encoding='utf-8'))
        payload = json.loads((a.output_dir/(kind+'_workbook_payload.json')).read_text(encoding='utf-8'))
        output = a.output_dir/payload['filename']
        before = openpyxl.load_workbook(a.prior_dir/prior_payload['filename'], data_only=False)
        after = openpyxl.load_workbook(output, data_only=False)
        before_cache = openpyxl.load_workbook(a.prior_dir/prior_payload['filename'],data_only=True)
        after_cache = openpyxl.load_workbook(output,data_only=True)
        try:
            assert after.sheetnames == [kind,dictionary,*before.sheetnames[1:]]
            assert list(before.defined_names) == list(after.defined_names)
            old = before[kind]
            new = after[kind]
            old_headers = [c.value for c in old[1]]
            new_headers = [c.value for c in new[1]]
            expected_headers = old_headers.copy()
            expected_headers.insert(old_headers.index('标签明细')+1,'标签强调方向')
            assert new_headers == expected_headers
            assert old.max_row == new.max_row
            assert old.freeze_panes == new.freeze_panes
            tag_count = 0
            for row_number in range(1,old.max_row+1):
                for old_column,header in enumerate(old_headers,1):
                    same_cell(old.cell(row_number,old_column),new.cell(row_number,new_headers.index(header)+1),
                        (kind,row_number,header))
                if row_number == 1:
                    continue
                tags = json.loads(new.cell(row_number,new_headers.index('标签明细')+1).value)
                tag_count += len(tags)
                emphases = list(dict.fromkeys(t['标签强调方向'] for t in tags if t['标签强调方向']))
                assert new.cell(row_number,new_headers.index('标签强调方向')+1).value == ';'.join(emphases)
            assert list(old.tables) == list(new.tables)
            for name in before.sheetnames[1:]:
                prior_sheet = before[name]
                new_sheet = after[name]
                assert prior_sheet.freeze_panes == new_sheet.freeze_panes, name
                assert list(prior_sheet.tables) == list(new_sheet.tables), name
                assert str(prior_sheet.data_validations) == str(new_sheet.data_validations), name
                assert list(prior_sheet.merged_cells.ranges) == list(new_sheet.merged_cells.ranges), name
                for row in prior_sheet:
                    for cell in row:
                        if name == '测试渠道统计' and cell.column == 13 and cell.row >= 3:
                            continue  # Only the delivery notes were extended.
                        same_cell(cell,new_sheet.cell(cell.row,cell.column),(name,cell.coordinate))
                        assert before_cache[name].cell(cell.row,cell.column).value == after_cache[name].cell(cell.row,cell.column).value
                for table_name in prior_sheet.tables:
                    old_table = prior_sheet.tables[table_name]
                    new_table = copy.deepcopy(new_sheet.tables[table_name])
                    # Adding a sheet may renumber package-local table IDs.
                    new_table.id = old_table.id
                    assert old_table == new_table, (name,table_name,'native table changed')
            reference = after[dictionary]
            assert [c.value for c in reference[5]][:len(REFERENCE_FIELDS)] == REFERENCE_FIELDS
            assert reference.max_row == 5+len(references[topic])
            assert reference.freeze_panes == 'A6'
            for n,target in enumerate(references[topic],6):
                for column,header in enumerate(REFERENCE_FIELDS,1):
                    value = reference.cell(n,column).value
                    expected = target[header]
                    assert value == expected or value is None and expected == '', (dictionary,n,header)
            registry = {t['标签路径ID']:t for t in references[topic]}
            for row in payload['tables'][kind]:
                for tag in row['标签明细']:
                    assert set(tag) == set(TAG_KEYS)
                    assert tag == {key:registry[tag['标签路径ID']][key] for key in TAG_KEYS}
            for row in after_cache['测试渠道统计']:
                assert not any(c.data_type == 'e' for c in row)
            check = {'file':str(output),'passed':True,'main_rows':new.max_row-1,
                'dictionary_rows':len(references[topic]),'assigned_tags_unchanged':tag_count,
                'all_original_main_values_and_formats_unchanged':True,'all_other_sheet_data_and_formulas_unchanged':True,
                'new_column_matches_json':True,'dictionary_matches_user_source':True,
                'native_tables_validations_and_freeze_panes_preserved':True}
            checks.append(check)
        finally:
            for book in (before,after,before_cache,after_cache):
                book.close()
    report = {'passed':True,'checks':checks,'note':'This checks the display edit, not semantic tagging accuracy.'}
    (a.output_dir/'independent_xlsx_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))


if __name__ == '__main__':
    main()
