#!/usr/bin/env python3
"""地产专题产品/观点单元自动切片与打标。

主流程：唯一文本源 -> 清洗/去噪 -> 规则原子切片 -> 项目锚定 ->
最新VOC五级标签召回 -> TF-IDF补召回 -> 情绪/表达 -> 置信度门控。

脚本只读源数据和模板，输出 CSV/JSON，不更新数据库。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# The two rule modules are kept one directory above this distributable script.
# Add that directory explicitly so the CLI works regardless of the current
# working directory used to launch it.
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

import process_topic1_samples as t1
import process_topic_samples as t2


TOPIC1 = "全国地产设计产品力监测"
TOPIC2 = "北京臻澐项目口碑"
RULE_VERSION = "voc-five-level-v2.1-20261001"

PRODUCT_FIELDS = [
    "unit_id", "content_id", "证据序号", "原文摘录", "证据定位", "品牌组", "经营主体", "项目", "城市",
    "产品体系", "产品系标签", "产品系匹配状态", "产品档", "标签路径ID", "标签编号",
    "一级标签", "二级标签", "三级标签", "四级标签", "五级标签", "标签强调方向", "标签属性", "标签备注", "打标说明",
    "具体产品做法", "创新类型", "创新等级", "创新依据", "备注",
]

VIEWPOINT_FIELDS = [
    "unit_id", "record_id", "project_id", "项目名称", "项目角色", "发布日期", "平台", "发布者角色",
    "原始摘录", "观点摘要", "标签路径ID", "标签编号", "一级标签", "二级标签", "三级标签", "四级标签", "五级标签",
    "标签强调方向", "标签属性", "标签备注", "打标说明", "表达类型",
    "情绪态度", "阅读/播放", "点赞", "收藏", "评论", "转发", "互动量", "明确评价标记",
    "正负评价标记", "需求提问标记", "备注",
]

STATUS_FIELDS = [
    "源表id", "专题", "平台", "文章标题", "切片来源", "来源字符数", "候选句段数", "产品单元数", "观点单元数",
    "切片数量", "自动通过单元数", "需复核单元数", "处理状态", "处理说明", "识别项目", "规则版本",
]

SPACE = re.compile(r"[\t\u3000 ]+")
NOISE = re.compile(
    r"(?:400[-—\s]?\d{3,4}[-—\s]?\d{4}|1[3-9]\d{9}|微信号|扫码|二维码|关注公众号|点击蓝字|"
    r"点赞收藏分享|欢迎关注|私信预约|售楼处.{0,12}(?:电话|热线)|免责声明|版权声明|END点赞|"
    r"相关文章推荐|火\s*热\s*申\s*报|长按.{0,8}(?:识别|添加)|商务合作)", re.I
)
POSITIVE = re.compile(
    r"优势|优越|完善|成熟|突出|不错|认可|友好|稳定|纯粹|稀缺|舒适|实用|通透|便捷|高标准|"
    r"高品质|性价比|值得|优选|首选|领先|丰富|充足|安全|放心|开阔|方正|采光充足|很稳|拉满|出众"
)
NEGATIVE = re.compile(
    r"劣势|硬伤|不足|问题|风险|偏弱|过高|太高|昂贵|压抑|噪音|临路|缩水|去化慢|踩坑|延期|"
    r"投诉|整改|不建议|不适合|极差|焦虑|虚高|缺点|门槛高|成本高|堵车|减配|维权|瘫痪|缺陷|折旧"
)
QUESTION = re.compile(r"[?？]|是否|能买吗|值得买吗|怎么样|怎么选|为什么|会不会|有没有")
FUTURE = re.compile(r"预计|计划|将会|未来|可能|大概率|有望|或将")
COMPARE = re.compile(r"相比|对比|更[高低大好强弱]|不如|优于|低于|高于|vs|VS|之间|怎么选")
FACT = re.compile(r"\d+(?:\.\d+)?\s*(?:万|亿|元|㎡|m²|平米|平方米|套|户|层|%|米)|开发商|容积率|交房|网签|均价|户型")
PRODUCT_GUARD = re.compile(
    r"户型|玄关|客厅|餐厅|厨房|卧室|卫浴|卫生间|阳台|露台|会所|架空层|园林|景观|大堂|地库|"
    r"车库|收纳|衣柜|材料|石材|铝板|玻璃|门窗|空调|地暖|新风|智能|防水|隔音|层高|面宽|进深|"
    r"采光|动线|人车分流|抬板|立面|精装|设备|楼间距|绿化率|得房率|赠送|实得"
)
VIEW_GUARD = re.compile(
    r"价格|总价|单价|预算|折扣|交付|收房|验房|施工|质量|物业|维修|交通|通勤|地铁|学校|学区|"
    r"商业|配套|噪音|风险|问题|优势|劣势|值得|户型|园林|会所|精装|得房率|层高|面宽|采光|"
    r"低密|洋房|别墅|豪宅|销售|房源|拥堵|污染|圈层"
)
SUMMARY_MIN_LENGTH = 20
SUMMARY_MAX_LENGTH = 80
SUMMARY_LONG_SOURCE = 40
SUMMARY_LEADING = re.compile(
    r"^(?:值得一提的是|更值得关注的是|需要指出的是|可以看到|可以发现|不难发现|不得不说|"
    r"众所周知|毋庸置疑|毫无疑问|总体来看|总的来说|简单来说|事实上|据悉|据了解|此外|同时|而且|尤其是)[，,:：\s]*"
)
SUMMARY_PROMO = re.compile(
    r"(?:重磅|震撼|惊艳|极致|封神|王炸|天花板|逆天|拉满|顶流|爆款|匠心打造|精心打造|"
    r"全面赋能|实力演绎|焕新升级|荣耀呈现|诚意巨献|不容错过|强势来袭)"
)
SUMMARY_PLATFORM_NOISE = re.compile(
    r"(?:本文由[^，。；]{0,30}(?:发布|整理)|来源[：:]?[^，。；]{0,30}|编辑[：:]?[^，。；]{0,20}|"
    r"点击(?:蓝字|链接)[^，。；]{0,30}|欢迎(?:关注|点赞|转发|收藏)[^，。；]{0,30})"
)
SUMMARY_FILLER_CLAUSE = re.compile(
    r"^(?:说点.{0,24}(?:不爱听的|题外话)|再说个细节|再说一点|先说结论|先说重点|"
    r"去各个.{0,20}看看|我们(?:先|再)?来看看|话说回来|回到正题|首先|其次|最后)$"
)


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ").replace("\u3000", " ")).strip()


def clip(value: Any, limit: int) -> str:
    value = norm(value)
    return value if len(value) <= limit else value[: limit - 1] + "…"


def _summary_key(value: Any) -> str:
    return re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", norm(value)).lower()


def _clean_summary_text(value: Any) -> str:
    text = norm(value).strip("，,。；;：:-—|/（）()[]【】‘’“”\"'")
    previous = None
    while text and text != previous:
        previous = text
        text = SUMMARY_LEADING.sub("", text)
    text = SUMMARY_PLATFORM_NOISE.sub("", text)
    text = SUMMARY_PROMO.sub("", text)
    text = text.replace("“", "").replace("”", "").replace("‘", "").replace("’", "")
    text = re.sub(r"([，。；！？])\1+", r"\1", text)
    text = re.sub(r"[，,；;：:]\s*[，,；;：:]", "，", text)
    return norm(text).strip("，,。；;：:-—|/（）()[]【】‘’“”\"'")


def _limit_summary(value: Any, limit: int = SUMMARY_MAX_LENGTH) -> str:
    text = _clean_summary_text(value)
    if len(text) <= limit:
        return text
    window = text[: limit + 1]
    cuts = [window.rfind(mark) for mark in "。！？；，"]
    cut = max(cuts)
    if cut >= SUMMARY_MIN_LENGTH:
        return window[:cut].rstrip("，,。；;：:!?！？")
    return text[:limit].rstrip("，,。；;：:!?！？")


def viewpoint_summary_issues(evidence: Any, summary: Any) -> list[str]:
    """Return review reasons for a viewpoint summary without rewriting evidence."""
    source = norm(evidence)
    value = norm(summary)
    issues: list[str] = []
    if not value:
        return ["摘要待复核：未生成有效观点摘要"]
    if len(value) > SUMMARY_MAX_LENGTH:
        issues.append("摘要待复核：观点摘要超过80字")
    if len(source) > SUMMARY_LONG_SOURCE and len(value) < SUMMARY_MIN_LENGTH:
        issues.append("摘要待复核：长原文的观点摘要少于20字")
    if len(source) > SUMMARY_LONG_SOURCE and _summary_key(source) == _summary_key(value):
        issues.append("摘要待复核：原文超过40字且观点摘要未压缩")
    return issues


def summarize_viewpoint(evidence: Any, object_terms: list[str] | None = None) -> tuple[str, list[str]]:
    """Create a conservative extractive fallback summary.

    The fallback only removes boilerplate/promotional wording and selects
    clauses already present in the evidence. It never invents a fact. Short,
    already-complete evidence may remain unchanged.
    """
    source = norm(evidence)
    if not source:
        return "", ["摘要待复核：原始摘录为空"]
    if source.count("|") >= 3 or len(re.findall(r"(?:攻略|答疑|合集|流程)", source)) >= 3:
        return "", ["摘要待复核：原始摘录疑似导航或列表文本"]
    cleaned_source = _clean_summary_text(source) or source
    if len(source) <= SUMMARY_LONG_SOURCE:
        summary = _limit_summary(cleaned_source)
        return summary, viewpoint_summary_issues(source, summary)

    raw_clauses = re.split(r"(?<=[。！？!?；;])|[，,]", source)
    clauses: list[tuple[int, int, str]] = []
    seen: set[str] = set()
    terms = [norm(x) for x in (object_terms or []) if len(norm(x)) >= 2]
    for index, raw in enumerate(raw_clauses):
        clause = _clean_summary_text(raw)
        key = _summary_key(clause)
        if len(clause) < 4 or not key or key in seen:
            continue
        if SUMMARY_FILLER_CLAUSE.fullmatch(clause):
            continue
        if is_noise(clause):
            continue
        seen.add(key)
        score = 0
        score += 5 * sum(term in clause for term in terms)
        score += 3 if VIEW_GUARD.search(clause) else 0
        score += 2 if POSITIVE.search(clause) or NEGATIVE.search(clause) or QUESTION.search(clause) else 0
        score += 1 if FACT.search(clause) else 0
        clauses.append((score, index, clause))

    if not clauses:
        summary = _limit_summary(cleaned_source)
    else:
        ranked = sorted(clauses, key=lambda x: (-x[0], x[1]))
        selected: list[str] = []
        for score, _, clause in ranked:
            if selected and score <= 0:
                break
            proposal = "；".join(selected + [clause])
            if len(proposal) <= SUMMARY_MAX_LENGTH:
                selected.append(clause)
            if len("；".join(selected)) >= SUMMARY_MIN_LENGTH:
                break
        if not selected:
            selected = [ranked[0][2]]
        summary = _limit_summary("；".join(selected))

    issues = viewpoint_summary_issues(source, summary)
    return summary, issues


def validate_api_viewpoint_summary(
    evidence: Any,
    candidate: Any,
    fallback: Any = "",
    object_terms: list[str] | None = None,
) -> tuple[str, list[str]]:
    """Validate the model summary and return a safe fallback when necessary."""
    source = norm(evidence)
    raw_candidate = norm(candidate)
    fallback_value, fallback_issues = summarize_viewpoint(source, object_terms)
    if norm(fallback):
        fallback_value = norm(fallback)
        fallback_issues = viewpoint_summary_issues(source, fallback_value)
    if not raw_candidate:
        return fallback_value, ["摘要待复核：API未返回有效观点摘要", *fallback_issues]

    summary = _clean_summary_text(raw_candidate)
    issues: list[str] = []
    if len(summary) > SUMMARY_MAX_LENGTH:
        summary = _limit_summary(summary)
        issues.append("摘要待复核：API摘要超过80字，已保守截断")

    source_numbers = set(re.findall(r"\d+(?:\.\d+)?%?", source))
    summary_numbers = set(re.findall(r"\d+(?:\.\d+)?%?", summary))
    if summary_numbers - source_numbers:
        summary = fallback_value
        issues.append("摘要待复核：API摘要含原始摘录外数字，已回退本地摘要")
        issues.extend(fallback_issues)
    issues.extend(viewpoint_summary_issues(source, summary))
    return summary, list(dict.fromkeys(issues))


def number(value: Any) -> float:
    text = norm(value).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else 0.0


def choose_source(row: dict[str, str]) -> tuple[str, str]:
    body = norm(row.get("文章内容"))
    if body:
        return "文章内容", body
    ocr = norm(row.get("文章图片信息"))
    if ocr:
        return "图片OCR", ocr
    return "无可用文本", ""


def is_noise(text: str) -> bool:
    hits = len(NOISE.findall(text))
    meaningful = len(re.findall(r"[\u4e00-\u9fff]", text))
    return hits >= 2 or (hits >= 1 and meaningful < 90)


def atomic_split(text: str) -> list[tuple[int, str]]:
    text = norm(text)
    text = re.sub(r"(?<!\d)(0?[1-9]|1\d)[、.](?!\d)", r"\n\1、", text)
    text = re.sub(r"([一二三四五六七八九十]+)[、](?=[^\n])", r"\n\1、", text)
    sentences = re.split(r"(?<=[。！？!?；;])|[\n\r]+", text)
    output: list[tuple[int, str]] = []
    for pos, sentence in enumerate(sentences, 1):
        sentence = norm(sentence).strip("，,。；;：:-—|/（）()[]【】")
        if len(sentence) < 12:
            continue
        pieces = [sentence]
        if len(sentence) > 150:
            pieces = re.split(r"(?<=[，,：:])|(?=【[^】]{2,20}】)|(?=(?:优势|劣势|亮点|痛点|总结|改造重点)[：:]?)", sentence)
        current = ""
        for piece in pieces:
            piece = norm(piece).strip("，,。；;：:-—|/（）()[]【】")
            if not piece:
                continue
            if current and len(current) + len(piece) <= 115:
                current = norm(current + "，" + piece)
            else:
                if current and len(current) >= 12:
                    output.append((pos, current))
                current = piece
        if current and len(current) >= 12:
            output.append((pos, current))
    seen, deduped = set(), []
    for pos, text in output:
        key = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", text.lower())
        if key and key not in seen:
            seen.add(key)
            deduped.append((pos, text))
    return deduped


def load_voc_taxonomy(path: Path) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    """Load the latest five-level VOC taxonomy.

    Topic 1 follows the workbook's 专题一标签 marker. Topic 2 follows the
    专题二标签 marker and, per business rule, always includes every 客户 row.
    The source 编号 is not unique, so 标签路径ID is generated from row order.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [norm(cell.value) for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    required = {"编号", "一级标签", "二级标签", "三级标签", "四级标签", "五级标签", "打标说明", "专题一标签", "专题二标签"}
    missing = sorted(required - set(headers))
    if missing:
        wb.close()
        raise ValueError("VOC标签体系缺少字段：" + "、".join(missing))
    records: list[dict[str, Any]] = []
    for excel_row, values in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        source = {headers[i]: norm(values[i]) if i < len(values) else "" for i in range(len(headers))}
        levels = [source.get(f"{name}标签", "") for name in ["一级", "二级", "三级", "四级", "五级"]]
        if not any(levels):
            continue
        sequence = len(records) + 1
        topics = set()
        if source.get("专题一标签"):
            topics.add(TOPIC1)
        if source.get("专题二标签") or source.get("一级标签") == "客户":
            topics.add(TOPIC2)
        record = {
            "tag_id": f"VOC-{sequence:04d}",
            "source_row": excel_row,
            "source_number": source.get("编号"),
            "l1": source.get("一级标签"), "l2": source.get("二级标签"),
            "l3": source.get("三级标签"), "l4": source.get("四级标签"),
            "l5": source.get("五级标签"), "note": source.get("备注"),
            "attribute": source.get("属性"), "emphasis": source.get("标签强调方向"),
            "instruction": source.get("打标说明"),
            "topic1_label": source.get("专题一标签"),
            "topic2_label": source.get("专题二标签") or (source.get("标签强调方向") if source.get("一级标签") == "客户" else ""),
            "topics": topics,
        }
        records.append(record)
    wb.close()
    labels_by_topic = {
        TOPIC1: [x for x in records if TOPIC1 in x["topics"]],
        TOPIC2: [x for x in records if TOPIC2 in x["topics"]],
    }
    if not labels_by_topic[TOPIC1] or not labels_by_topic[TOPIC2]:
        raise ValueError("VOC标签体系未形成两个专题的有效标签集合")
    return records, labels_by_topic


def load_brand_dictionary(path: Path) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["03_品牌产品系字典"]
    headers = [norm(x.value) for x in next(ws.iter_rows(min_row=1, max_row=1))]
    result = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[i]: norm(row[i]) if i < len(row) else "" for i in range(len(headers))}
        if item.get("品牌组") and item.get("经营主体"):
            result.append(item)
    wb.close()
    return result


def load_projects(path: Path, selected_names: set[str], extra_aliases: dict[str, list[str]]) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["04_项目竞品字典"]
    projects = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        pid, role, name, _, aliases = row[:5]
        if not pid or not name or (selected_names and norm(name) not in selected_names):
            continue
        name = norm(name)
        alias_list = [name]
        for alias in re.split(r"[；;]", norm(aliases)):
            alias = re.sub(r"（[^）]*）", "", alias).strip()
            if alias and alias not in alias_list:
                alias_list.append(alias)
        for alias in extra_aliases.get(name, []):
            if alias and alias not in alias_list:
                alias_list.append(alias)
        role_text = norm(role)
        project_role = "本品" if "本品" in role_text else ("次核心竞品" if "次核心" in role_text else "核心竞品")
        projects.append({"project_id": norm(pid), "name": name, "role": project_role, "aliases": alias_list})
    wb.close()
    return projects


class SemanticIndex:
    def __init__(self, labels_by_topic: dict[str, list[dict[str, Any]]]):
        self.labels_by_topic = labels_by_topic
        self.models = {}
        for topic, labels in labels_by_topic.items():
            docs = [" ".join([x["l1"], x["l2"], x["l3"], x["l4"], x["l5"], x["emphasis"], x["attribute"], x["note"], x["instruction"]]) for x in labels]
            vectorizer = TfidfVectorizer(analyzer="char", ngram_range=(2, 4), sublinear_tf=True, max_features=30000)
            matrix = vectorizer.fit_transform(docs)
            self.models[topic] = (vectorizer, matrix)

    def best_n(self, topic: str, text: str, limit: int = 8) -> list[tuple[dict[str, Any], float]]:
        vectorizer, matrix = self.models[topic]
        scores = cosine_similarity(vectorizer.transform([text]), matrix)[0]
        indexes = np.argsort(scores)[::-1][:limit]
        return [(self.labels_by_topic[topic][int(idx)], float(scores[int(idx)])) for idx in indexes]


def _match_terms(label: dict[str, Any], topic: str) -> list[str]:
    topic_label = label["topic1_label"] if topic == TOPIC1 else label["topic2_label"]
    deepest = next((label[key] for key in ["l5", "l4", "l3", "l2", "l1"] if label[key]), "")
    terms = [deepest, label["emphasis"], topic_label]
    output = []
    generic = {"整体", "其他", "通用", "布局", "面积", "功能", "位置", "属性", "产品", "客户", "土地", "服务", "市场", "品牌"}
    for value in terms:
        for term in re.split(r"[、，,；;/（）()]+", norm(value)):
            term = term.strip()
            if len(term) >= 2 and term not in generic and term not in output:
                output.append(term)
    return output


def taxonomy_candidates(topic: str, text: str, semantic: SemanticIndex, limit: int) -> list[tuple[dict[str, Any], float, str]]:
    candidates: dict[str, tuple[dict[str, Any], float, str]] = {}
    for label in semantic.labels_by_topic[topic]:
        hits = [term for term in _match_terms(label, topic) if term in text]
        deepest = next((label[key] for key in ["l5", "l4", "l3", "l2", "l1"] if label[key]), "")
        context_sensitive = {"地面", "墙体", "天花板", "面积", "数量", "布局", "尺度", "位置", "风格", "材质与工艺", "功能"}
        if hits and deepest in context_sensitive:
            contexts = [x for x in [label["l3"], label["l2"]] if len(x) >= 2 and x not in {"产品", "客户", "土地", "服务", "市场", "品牌"}]
            if contexts and not any(context in text for context in contexts):
                hits = []
        if hits:
            score = min(1.0, 0.90 + max(len(x) for x in hits) / 100)
            candidates[label["tag_id"]] = (label, score, "VOC标签词精确命中")
    if candidates:
        return sorted(candidates.values(), key=lambda x: x[1], reverse=True)[:limit]
    threshold = 0.22 if topic == TOPIC1 else 0.20
    guard = PRODUCT_GUARD if topic == TOPIC1 else VIEW_GUARD
    if guard.search(text):
        for label, score in semantic.best_n(topic, text, max(limit * 2, 8)):
            if score >= threshold:
                candidates[label["tag_id"]] = (label, score, "VOC五级TF-IDF召回")
    return sorted(candidates.values(), key=lambda x: x[1], reverse=True)[:min(limit, 2)]


def infer_sentiment(text: str) -> str:
    pos, neg = len(POSITIVE.findall(text)), len(NEGATIVE.findall(text))
    if pos > neg and pos:
        return "正向"
    if neg > pos and neg:
        return "负向"
    if pos and neg:
        return "中性"
    return "无态度"


def expression_type(text: str, emo: str, project_count: int) -> str:
    if QUESTION.search(text):
        return "需求提问"
    if FUTURE.search(text):
        return "未来推测"
    if project_count > 1 or COMPARE.search(text):
        return "对比判断"
    if emo == "正向":
        return "正向评价"
    if emo == "负向":
        return "负向评价"
    return "事实介绍" if FACT.search(text) else "中性评价"


def project_hits(text: str, projects: list[dict[str, Any]]) -> list[dict[str, Any]]:
    hits = []
    for project in projects:
        alias = next((x for x in project["aliases"] if x and x in text), None)
        if alias:
            item = dict(project)
            item["matched_alias"] = alias
            hits.append(item)
    return hits


def write_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def product_practice(evidence: list[str]) -> str:
    cleaned = []
    for text in evidence[:3]:
        value = t1.clean_practice(text)
        value = re.sub(r"(?:封神|王炸|天花板|逆天|巨巨巨大|拉满)", "", value)
        if value:
            cleaned.append(value)
    return clip("；".join(cleaned), 420)


def product_units_for_row(row, text, semantic, brand_dict, project_overrides):
    brand = t1.infer_brand(row)
    entity = t1.infer_entity(row, brand)
    project = project_overrides.get(str(row["id"])) or t1.infer_project(row)
    city = t1.infer_city(row)
    full_text = " ".join([norm(row.get("文章标题")), norm(row.get("关键词")), text[:2500]])
    system, series, match_status = t1.product_series(full_text, brand, entity, brand_dict)
    grade = t1.infer_grade(full_text)
    groups = defaultdict(lambda: {"label": None, "evidence": [], "positions": [], "sources": [], "scores": []})

    segments = atomic_split(text)
    for pos, segment in segments:
        if is_noise(segment):
            continue
        matches = taxonomy_candidates(TOPIC1, segment, semantic, 3)
        for label, score, source in matches:
            group = groups[label["tag_id"]]
            group["label"] = label
            if segment not in group["evidence"] and len(group["evidence"]) < 4:
                group["evidence"].append(segment)
                group["positions"].append(pos)
            group["sources"].append(source)
            group["scores"].append(score)

    ordered = sorted(groups.items(), key=lambda x: (max(x[1]["scores"]), len(x[1]["evidence"])), reverse=True)[:24]
    units = []
    for seq, (_, group) in enumerate(ordered, 1):
        label = group["label"]
        evidence = group["evidence"][:3]
        joined = " ".join(evidence)
        if t1.UPGRADE.search(joined) and t1.COMPARE.search(joined):
            innovation_level = "I2 有升级迭代证据"
            innovation_basis = "原文同时出现升级变化和比较基线，仍需复核比较对象与版本。"
        elif t1.CLAIM.search(joined):
            innovation_level = "I1 来源宣称创新"
            innovation_basis = "来源使用首创、首次或创新等表述，但缺少独立同类验证。"
        else:
            innovation_level = "I0 无创新判断"
            innovation_basis = "原文支持具体产品做法，但没有明确升级基线或同类差异证据。"
        innovation_type = (
            "工艺创新" if re.search(r"材料|工艺|工法", " ".join([label["l1"], label["l2"], label["l3"], label["l4"], label["l5"]])) else
            "技术集成创新" if re.search(r"智慧|智能|科技", " ".join([label["l1"], label["l2"], label["l3"], label["l4"], label["l5"]])) else
            "设计创新" if innovation_level != "I0 无创新判断" else "无明确创新类型"
        )
        confidence = min(group["scores"]) if group["scores"] else 0.0
        review = []
        if project.startswith("待判断") or "多项目" in project:
            review.append("项目归属待核")
        if "VOC五级TF-IDF召回" in group["sources"]:
            review.append("含VOC五级语义召回标签")
        if innovation_level == "I2 有升级迭代证据":
            review.append("I2需人工核对比较基线")
        if confidence < 0.85:
            review.append("置信度低于自动通过阈值")
        status = "需人工复核" if review else "自动通过"
        units.append({
            "unit_id": f"T1-{row['id']}-M{seq:02d}", "content_id": f"T1-{row['id']}", "证据序号": seq,
            "原文摘录": clip("；".join(evidence), 700), "证据定位": "正文句段" + "、".join(map(str, group["positions"][:3])),
            "品牌组": brand, "经营主体": entity, "项目": project, "城市": city, "产品体系": system or "",
            "产品系标签": series or "", "产品系匹配状态": match_status, "产品档": grade,
            "标签路径ID": label["tag_id"], "标签编号": label["source_number"],
            "一级标签": label["l1"], "二级标签": label["l2"], "三级标签": label["l3"],
            "四级标签": label["l4"], "五级标签": label["l5"], "标签强调方向": label["emphasis"],
            "标签属性": label["attribute"], "标签备注": label["note"], "打标说明": label["instruction"],
            "具体产品做法": product_practice(evidence), "创新类型": innovation_type, "创新等级": innovation_level,
            "创新依据": innovation_basis,
            "备注": f"{status}；置信度={confidence:.2f}；标签来源={'/'.join(sorted(set(group['sources'])))}；规则版本={RULE_VERSION}" + ("；" + "；".join(review) if review else ""),
            "_review": status == "需人工复核", "_confidence": round(confidence, 3),
        })
    return units, len(segments), project


def viewpoint_units_for_row(row, text, semantic, projects):
    segments = atomic_split(text)
    title_hits = project_hits(" ".join([norm(row.get("关键词")), norm(row.get("文章标题"))]), projects)
    current = title_hits if title_hits else []
    carry = 0
    publisher_role = t2.infer_publisher_role(row)
    interaction = sum(number(row.get(x)) for x in ["文章点赞数", "文章收藏数", "文章评论数", "文章分享数"])
    candidates = []
    seen = set()

    for pos, segment in segments:
        if is_noise(segment):
            continue
        explicit = project_hits(segment, projects)
        if explicit:
            current = explicit
            carry = 0
        elif current and carry < 7:
            explicit = current
            carry += 1
        else:
            continue
        emo = infer_sentiment(segment)
        matches = taxonomy_candidates(TOPIC2, segment, semantic, 5)
        for project in explicit[:3]:
            for label, score, source in matches:
                key = (project["project_id"], label["tag_id"], re.sub(r"\W+", "", segment))
                if key not in seen:
                    seen.add(key)
                    candidates.append((pos, segment, project, label, score, source, emo, len(explicit)))

    units = []
    for seq, (pos, segment, project, label, score, source, emo, project_count) in enumerate(candidates[:100], 1):
        expression = expression_type(segment, emo, project_count)
        confidence = 0.92 if source == "VOC标签词精确命中" else min(0.84, 0.62 + score)
        review = []
        object_terms = [
            project.get("name", ""), project.get("matched_alias", ""),
            label.get("l2", ""), label.get("l3", ""), label.get("l4", ""),
            label.get("l5", ""), label.get("emphasis", ""),
        ]
        viewpoint_summary, summary_review = summarize_viewpoint(segment, object_terms)
        if project_count > 1 or expression == "对比判断":
            confidence -= 0.10
            review.append("多项目比较需核对观点归属")
        if source == "VOC五级TF-IDF召回":
            review.append("VOC五级语义召回标签需复核")
        if summary_review:
            confidence = min(confidence, 0.84)
            review.extend(summary_review)
        if confidence < 0.85:
            review.append("置信度低于自动通过阈值")
        review = list(dict.fromkeys(review))
        status = "需人工复核" if review else "自动通过"
        units.append({
            "unit_id": f"T2-{row['id']}-V{seq:03d}", "record_id": f"T2-{row['id']}",
            "project_id": project["project_id"], "项目名称": project["name"], "项目角色": project["role"],
            "发布日期": row.get("文章时间地点"), "平台": row.get("平台"), "发布者角色": publisher_role,
            "原始摘录": segment, "观点摘要": viewpoint_summary,
            "标签路径ID": label["tag_id"], "标签编号": label["source_number"],
            "一级标签": label["l1"], "二级标签": label["l2"], "三级标签": label["l3"],
            "四级标签": label["l4"], "五级标签": label["l5"], "标签强调方向": label["emphasis"],
            "标签属性": label["attribute"], "标签备注": label["note"], "打标说明": label["instruction"],
            "表达类型": expression, "情绪态度": emo,
            "阅读/播放": number(row.get("文章阅读数") or row.get("文章播放数")), "点赞": number(row.get("文章点赞数")),
            "收藏": number(row.get("文章收藏数")), "评论": number(row.get("文章评论数")), "转发": number(row.get("文章分享数")),
            "互动量": interaction, "明确评价标记": int(expression in {"正向评价", "中性评价", "负向评价", "对比判断"}),
            "正负评价标记": int(emo in {"正向", "负向"}), "需求提问标记": int(expression == "需求提问"),
            "备注": f"正文句段{pos}；{status}；置信度={confidence:.2f}；标签来源={source}；命中项目别名={project.get('matched_alias', project['name'])}；规则版本={RULE_VERSION}" + ("；" + "；".join(review) if review else ""),
            "_review": status == "需人工复核", "_confidence": round(max(confidence, 0.0), 3),
        })
    return units, len(segments), sorted({x["项目名称"] for x in units})


def parse_args():
    parser = argparse.ArgumentParser(description="地产专题混合切片和打标")
    parser.add_argument("--input", required=True, type=Path, help="ODS文章CSV")
    parser.add_argument("--topic1-template", required=True, type=Path)
    parser.add_argument("--topic2-template", required=True, type=Path)
    parser.add_argument("--voc-taxonomy", required=True, type=Path, help="最新五级VOC标签体系xlsx")
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    return parser.parse_args()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    args = parse_args()
    config = json.loads(args.config.read_text(encoding="utf-8"))
    selected_ids = set(map(str, config.get("sample_ids", [])))
    project_overrides = {str(k): v for k, v in config.get("topic1_project_overrides", {}).items()}
    selected_projects = set(config.get("topic2_selected_projects", []))
    extra_aliases = config.get("topic2_extra_aliases", {})

    taxonomy, labels_by_topic = load_voc_taxonomy(args.voc_taxonomy)
    semantic = SemanticIndex(labels_by_topic)
    brand_dict = load_brand_dictionary(args.topic1_template)
    projects = load_projects(args.topic2_template, selected_projects, extra_aliases)
    t2.PROJECTS = projects

    with args.input.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if selected_ids:
        rows = [x for x in rows if str(x.get("id")) in selected_ids]

    product_units, viewpoint_units, status_rows = [], [], []
    for row in rows:
        source_type, text = choose_source(row)
        p_units, v_units, segments, identified = [], [], 0, []
        if text and row.get("项目") == TOPIC1:
            p_units, segments, project = product_units_for_row(row, text, semantic, brand_dict, project_overrides)
            identified = [project]
        elif text and row.get("项目") == TOPIC2:
            v_units, segments, identified = viewpoint_units_for_row(row, text, semantic, projects)
        product_units.extend(p_units)
        viewpoint_units.extend(v_units)
        units = p_units + v_units
        reviews = sum(bool(x.get("_review")) for x in units)
        if not text:
            state, note = "无可用文本", "文章内容与图片OCR均为空；视频文本未参与本轮切片"
        elif not units:
            state, note = "需人工复核", "未生成满足项目和标签证据的单元"
        elif reviews:
            state, note = "部分完成需复核", f"共{len(units)}个单元，其中{reviews}个进入复核队列"
        else:
            state, note = "处理完成", f"共{len(units)}个单元，均达到自动通过阈值"
        status_rows.append({
            "源表id": str(row.get("id")), "专题": row.get("项目"), "平台": row.get("平台"), "文章标题": row.get("文章标题"),
            "切片来源": source_type, "来源字符数": len(text), "候选句段数": segments, "产品单元数": len(p_units),
            "观点单元数": len(v_units), "切片数量": len(units), "自动通过单元数": len(units) - reviews,
            "需复核单元数": reviews, "处理状态": state, "处理说明": note, "识别项目": "；".join(identified), "规则版本": RULE_VERSION,
        })

    clean_product = [{k: v for k, v in x.items() if not k.startswith("_")} for x in product_units]
    clean_view = [{k: v for k, v in x.items() if not k.startswith("_")} for x in viewpoint_units]
    taxonomy_index = {x["tag_id"]: x for x in taxonomy}
    invalid_units = []
    for topic, units in [(TOPIC1, clean_product), (TOPIC2, clean_view)]:
        for unit in units:
            label = taxonomy_index.get(str(unit.get("标签路径ID")))
            fields_match = label and all(
                norm(unit.get(field)) == norm(label[key])
                for field, key in [("一级标签", "l1"), ("二级标签", "l2"), ("三级标签", "l3"), ("四级标签", "l4"), ("五级标签", "l5")]
            )
            if not label or topic not in label["topics"] or not fields_match:
                invalid_units.append(str(unit.get("unit_id")))
    if invalid_units:
        raise ValueError("发现不属于最新VOC标签体系的输出单元：" + "、".join(invalid_units[:10]))
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "产品信息单元.csv", PRODUCT_FIELDS, clean_product)
    write_csv(args.output_dir / "观点信息单元.csv", VIEWPOINT_FIELDS, clean_view)
    write_csv(args.output_dir / "文章处理状态.csv", STATUS_FIELDS, status_rows)

    summary_same = sum(
        _summary_key(x.get("原始摘录")) == _summary_key(x.get("观点摘要"))
        for x in clean_view
    )
    summary_long_same = sum(
        len(norm(x.get("原始摘录"))) > SUMMARY_LONG_SOURCE
        and _summary_key(x.get("原始摘录")) == _summary_key(x.get("观点摘要"))
        for x in clean_view
    )
    summary = {
        "rule_version": RULE_VERSION,
        "input_articles": len(rows),
        "topic1_articles": sum(x.get("项目") == TOPIC1 for x in rows),
        "topic2_articles": sum(x.get("项目") == TOPIC2 for x in rows),
        "product_units": len(clean_product),
        "viewpoint_units": len(clean_view),
        "review_units": sum(x.get("_review") for x in product_units + viewpoint_units),
        "auto_pass_units": sum(not x.get("_review") for x in product_units + viewpoint_units),
        "unique_output_tag_paths": len({x["标签路径ID"] for x in clean_product + clean_view}),
        "fourth_level_units": sum(bool(x.get("四级标签")) for x in clean_product + clean_view),
        "fifth_level_units": sum(bool(x.get("五级标签")) for x in clean_product + clean_view),
        "invalid_taxonomy_units": len(invalid_units),
        "viewpoint_summary_same_as_evidence": summary_same,
        "viewpoint_summary_long_same_as_evidence": summary_long_same,
        "viewpoint_summary_below_20": sum(len(norm(x.get("观点摘要"))) < SUMMARY_MIN_LENGTH for x in clean_view),
        "viewpoint_summary_long_below_20": sum(
            len(norm(x.get("原始摘录"))) > SUMMARY_LONG_SOURCE
            and len(norm(x.get("观点摘要"))) < SUMMARY_MIN_LENGTH
            for x in clean_view
        ),
        "viewpoint_summary_above_80": sum(len(norm(x.get("观点摘要"))) > SUMMARY_MAX_LENGTH for x in clean_view),
        "viewpoint_summary_review_units": sum("摘要待复核" in norm(x.get("备注")) for x in clean_view),
        "topic2_selected_projects": [x["name"] for x in projects],
        "voc_taxonomy_rows": len(taxonomy),
        "topic1_voc_paths": len(labels_by_topic[TOPIC1]),
        "topic2_voc_paths": len(labels_by_topic[TOPIC2]),
        "topic2_customer_paths": sum(x["l1"] == "客户" for x in labels_by_topic[TOPIC2]),
        "label_rule": "仅使用最新VOC标签体系；旧三级标签和旧补充标签规则均不参与召回、回退或输出。专题二强制包含全部客户标签行。",
        "text_source_rule": "文章内容非空时使用文章内容；为空时回退文章图片信息；视频文本不参与本轮切片。",
        "viewpoint_summary_rule": "原始摘录保持证据文本；观点摘要控制在20至80字，去除宣传和平台套话，不新增原文事实。原文超过40字且摘要仍相同，或API未返回有效摘要时自动进入复核。短且完整的观点允许与原文相同。",
        "templates": {"topic1_metadata": str(args.topic1_template), "topic2_metadata": str(args.topic2_template), "voc_taxonomy": str(args.voc_taxonomy)},
    }
    payload = {"summary": summary, "product_units": clean_product, "viewpoint_units": clean_view, "article_status": status_rows}
    (args.output_dir / "sample_delivery_data.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
