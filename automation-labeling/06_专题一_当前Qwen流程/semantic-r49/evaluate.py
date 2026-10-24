"""Offline only. Never imported by the model runner; reference answers are not API input."""
import sys,json,re,hashlib
from pathlib import Path
from collections import defaultdict
import openpyxl
sys.stdout.reconfigure(encoding='utf8')
R=Path(__file__).parent
OUT=Path('E:/Codex/outputs/20261021_专题一参考集ABC')
def read(p):return json.loads(Path(p).read_text(encoding='utf8'))
def tags(v):
 if isinstance(v,str):
  try:v=json.loads(v)
  except ValueError:return set(re.findall(r'VOC-\d{4}',v))
 return {x.get('tag_id') or x.get('标签路径ID') for x in (v or []) if isinstance(x,dict)}-{None}
def refs(batch):
 w=openpyxl.load_workbook(OUT/f'专题一_Codex参考集{batch}_10篇文章.xlsx',read_only=True,data_only=True)
 s=w['产品信息单元'];hs=list(next(s.iter_rows(min_row=5,max_row=5,values_only=True)))
 rs=[dict(zip(hs,row)) for row in s.iter_rows(min_row=6,values_only=True) if row[0]];w.close()
 return rs
def evaluate(batch,path):
 ref=refs(batch);new=read(path);g=defaultdict(list);p=defaultdict(list)
 for u in ref:g[str(u['源表id'])].append({'a':u['起点'],'b':u['终点'],'tags':tags(u['标签明细']),'summary':u['具体产品做法'] or '', 'brand':u['品牌组'],'entity':u['经营主体'] or ''})
 invalid=[];blank=0;duplicates=0
 for result in new:
  d=result['document'];sid=d['source_id'];seen=set()
  for i,u in enumerate(result.get('rows',[])):
   tr=result.get('traces',[]);t=tr[i] if i<len(tr) else {}
   a=u.get('起点',t.get('起点'));b=u.get('终点',t.get('终点'))
   quote=u.get('原文摘录','')
   if not isinstance(a,int) or not isinstance(b,int) or d['text'][a:b]!=quote:invalid.append([sid,i,'offset'])
   if (a,b) in seen:duplicates+=1
   seen.add((a,b));blank+=not bool(u.get('具体产品做法'))
   p[sid].append({'a':a,'b':b,'tags':tags(u.get('标签明细')),'summary':u.get('具体产品做法') or '', 'brand':u.get('品牌组'),'entity':u.get('经营主体') or ''})
 tp=fp=fn=0;ledger=[];span_tp=span_fp=span_fn=0
 for sid in sorted(g.keys()|p.keys()):
  gt=set().union(*(x['tags'] for x in g[sid]));pr=set().union(*(x['tags'] for x in p[sid]))
  tp+=len(gt&pr);fp+=len(pr-gt);fn+=len(gt-pr)
  ledger.append({'source_id':sid,'reference_units':len(g[sid]),'model_units':len(p[sid]),'missing':sorted(gt-pr),'extra':sorted(pr-gt)})
  used=set()
  for j,x in enumerate(p[sid]):
   for tag in x['tags']:
    matches=[]
    for k,y in enumerate(g[sid]):
     if (k,tag) in used or tag not in y['tags'] or None in [x['a'],x['b'],y['a'],y['b']]:continue
     overlap=max(0,min(x['b'],y['b'])-max(x['a'],y['a']))
     if overlap/max(1,min(x['b']-x['a'],y['b']-y['a']))>=.5:matches.append((overlap,k))
    if matches:span_tp+=1;used.add((max(matches)[1],tag))
    else:span_fp+=1
  span_fn+=sum(len(x['tags']) for x in g[sid])-len(used)
 result={'batch':batch,'model_sources':len(new),'reference_sources':len(g),'source_tag':{'tp':tp,'fp':fp,'fn':fn,'precision':tp/(tp+fp) if tp+fp else None,'recall':tp/(tp+fn) if tp+fn else None},'span_tag':{'tp':span_tp,'fp':span_fp,'fn':span_fn,'precision':span_tp/(span_tp+span_fp) if span_tp+span_fp else None,'recall':span_tp/(span_tp+span_fn) if span_tp+span_fn else None},'structural_errors':invalid,'duplicate_slices':duplicates,'empty_practice':blank,'ledger':ledger,'limitations':'字符重叠只作切片标签定位初筛；差异须逐证据复核，不代表Codex全部正确。摘要忠实性、主体、品牌和错误严重度另行审核。'}
 dest=Path(path).with_name('evaluation_'+batch+'.json');dest.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':evaluate(sys.argv[1],sys.argv[2])
