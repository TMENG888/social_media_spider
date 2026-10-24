"""Read-only final workbook checks, with independent aggregate calculations."""
import sys,json,math,statistics,re
from pathlib import Path
from collections import defaultdict
import openpyxl

R=Path(__file__).parent
OUT=Path('E:/Codex/outputs/20261021_专题一Qwen回归验收')
FORBIDDEN={'项目','验收状态','互动深度指数','涉及文章数','参与文章数','命中切片数','组内全部切片数','有热度文章数','缺热度文章数','字典ID'}

def read(p):return json.loads(p.read_text(encoding='utf8'))
def eq(a,b):
    if a in ('',None) and b in ('',None):return True
    if isinstance(a,(float,int)) and isinstance(b,(float,int)):return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12)
    return a==b
def brands(v):
    if v in ('',None,'待判断','未明确','不明'):return set()
    return {x.strip() for x in re.split('[;；、,，]',v) if x.strip()}

def validate(batch,version,file_override=None,run_dir=None,source_path=None):
    run_dir=Path(run_dir) if run_dir else R/f'run_{version}'
    plan=read(run_dir/f'{batch}_delivery_plan.json')
    results=read(source_path if source_path else run_dir/f'{batch}_results.json')
    metrics=read(run_dir/f'{batch}_delivery_metrics.json')['sources']
    tax={t['tag_id']:t for t in read(R/'inputs/taxonomy.json')}
    bypath={tuple([t['l1'],t['l2'],t['l3'],t['l4'] or '不设四级（原字典）',t['l5'] or '不设五级（原字典）']):tid for tid,t in tax.items()}
    file=Path(file_override) if file_override else OUT/f'Qwen回归集{batch}_10篇文章.xlsx'
    wb=openpyxl.load_workbook(file,read_only=True,data_only=False)
    assert wb.sheetnames==[s['name'] for s in plan['sheets']]
    cells=0
    for sheet in plan['sheets']:
        ws=wb[sheet['name']];hs=list(next(ws.iter_rows(min_row=5,max_row=5,values_only=True)))
        assert ws.sheet_state=='visible' and not FORBIDDEN.intersection(hs)
        n=sheet['headers'].index(None) if sheet['notes'] else len(sheet['headers'])
        assert hs[:n]==sheet['headers'][:n]
        actual=list(ws.iter_rows(min_row=6,max_row=len(sheet['rows'])+5,max_col=n,values_only=True))
        assert len(actual)==len(sheet['rows'])
        if sheet['name']=='产品信息单元':
            label_col=sheet['headers'].index('标签明细')
            for row in actual:
                value=row[label_col]
                labels=json.loads(value) if isinstance(value,str) else value
                assert isinstance(labels,list) and labels,'最终交付不能包含空标签行'
        for i,(got,expected) in enumerate(zip(actual,sheet['rows']),6):
            for j,(a,b) in enumerate(zip(got,expected[:n]),1):
                assert eq(a,b),(sheet['name'],i,j,a,b)
                assert not (isinstance(a,str) and a in {'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}), (sheet['name'],i,j)
                cells+=1
    # Independently assemble source/partition groups from frozen Qwen output.
    parts=[];units=[]
    excluded=set(plan.get('empty_label_filter',{}).get('excluded',[]))
    for result in results:
        sid=result['document']['source_id'];role=result['publisher_role']
        unique=set(b for row in result['rows'] for b in brands(row['品牌组']))
        row_by_span={(u['起点'],u['终点']):u for u in result['rows']}
        for p in result['semantic_partition']:
            row=row_by_span.get((p['start'],p['end']))
            bs=brands(row['品牌组'] if row else p.get('brand'))
            if not bs and len(unique)==1:bs=unique.copy()
            bs={b if b in ['中海','保利','招商','金茂'] else '其他品牌' for b in bs} or {'未明确'}
            v={'sid':sid,'brands':bs,'role':role,'tags':{l['tag_id'] for l in p['labels']},'span':(p['start'],p['end'])}
            parts.append(v)
            if row and row['unit_id'] not in excluded:
                # Partition context inheritance is only a denominator rule.
                # A delivered tag keeps its own confirmed row brand; never
                # attribute an unidentified second project to the sole named
                # brand elsewhere in the same article.
                row_brands={b if b in ['中海','保利','招商','金茂'] else '其他品牌' for b in brands(row['品牌组'])} or {'未明确'}
                units.append({**v,'brands':row_brands})
    assert len(units)==plan['unit_count'],'交付过滤后的产品单元计数不一致'
    checks=0
    cache={}
    def expected(group,role,tag):
        key=(group,role)
        if key not in cache:
            match=lambda p:(group=='全部' or group in p['brands']) and (role=='全部' or role==p['role'])
            ps=[p for p in parts if match(p)];us=[p for p in units if match(p)]
            den_sources=len({p['sid'] for p in ps});den_parts=len(ps)
            ts=defaultdict(set);tu=defaultdict(int)
            for p in us:
                for t in p['tags']:ts[t].add(p['sid']);tu[t]+=1
            cache[key]=(den_sources,den_parts,ts,tu)
        ns,np,ts,tu=cache[key];heats=[metrics[s]['heat'] for s in ts[tag] if metrics[s]['heat'] is not None]
        return [len(ts[tag])/ns if ns else None,tu[tag]/np if np else None,statistics.median(heats) if heats else None]
    for sheet in plan['sheets']:
        name=sheet['name']
        if name=='议题指标' or name.endswith('_议题指标') or name=='品牌发布者角色指标':
            for row in sheet['rows']:
                if name=='品牌发布者角色指标':g,role=row[:2];offset=2
                else:g='全部' if name=='议题指标' else name.split('_')[0];role='全部';offset=0
                tag=bypath[tuple(row[offset:offset+5])]
                expected_metrics=expected(g,role,tag)
                assert all(eq(a,b) for a,b in zip(row[offset+6:offset+9],expected_metrics)),(name,g,role,tag,{'workbook':row[offset+6:offset+9],'independent':expected_metrics})
                checks+=3
    wb.close()
    receipt={'batch':batch,'file':str(file),'sheets':len(plan['sheets']),'source_count':len(results),'units':len(units),'full_partition_units':len(parts),'readback_cells':cells,'independent_metric_checks':checks,'prohibited_fields':0,'unexpected_formula_errors':0}
    (run_dir/f'{batch}_export_validation.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(receipt,ensure_ascii=False))

if __name__=='__main__':validate(sys.argv[1],sys.argv[2])
