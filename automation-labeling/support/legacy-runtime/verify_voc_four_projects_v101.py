"""Independent read-only final XLSX verification against frozen CSV and Qwen evidence."""
import csv
import copy
import json
import math
import statistics
import sys
from collections import Counter,defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
import voc_four_projects_v101 as p
from voc_four_delivery_contract_v101 import targets as delivery_targets, preserved_cohort
from voc_four_single_truth_v101 import derive_unit
from voc_four_identity_citation_v101 import unit_for_export
from voc_four_targeted_tag_review_v101 import unit_for_export as apply_targeted_tag_review
from voc_four_new_cohort_review_v101 import unit_for_export as apply_new_cohort_review
from voc_four_evidence_packing_v101 import unpack as unpack_evidence

BOOK='观点标签指标_四项目1000有效来源.xlsx'
FIELDS=['一级标签','二级标签','三级标签','四级标签','五级标签','标签名','占比（%）','议题热度','情绪评分','数据量']

def assert_own_date(actual,expected):
 # Export rounding may differ by a millisecond; timezone shifts may not.
 assert isinstance(actual,datetime), 'Saved publication date is not typed'
 assert abs((actual-datetime.fromisoformat(expected)).total_seconds())<0.01, 'Saved publication date differs from its own CSV timestamp'

def close(a,b,key):
 if b is None:
  if a not in (None,''):raise AssertionError(f'Expected missing {key}: {a}')
 elif isinstance(a,bool) or not isinstance(a,(float,int)) or not math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-8):
  raise AssertionError(f'Metric mismatch {key}: {a}/{b}')

def csv_rows(name):
 with (p.OUT/'inputs'/name).open(encoding='utf-8-sig',newline='') as f:return {str(r['id']):r for r in csv.DictReader(f)}

def sheet_rows(wb,name):
 rows=list(wb[name].iter_rows(min_row=5,values_only=True));header=list(rows[0]);return header,rows[1:]


def label_evidence(view,catalogue):
 evidence={item['source_key']:item['text'] for item in catalogue}
 return [{'tag_id':tag['tag_id'],'project_id':view['project_id'],'project_identity':tag['project_identity'],
  'attribute_quotes':[evidence[key] for key in tag['attribute_keys']],
  'context_quotes':[evidence[key] for key in tag['context_keys']]} for tag in view['labels']]


def project_platform_rows(real_sources,expected_units,projects):
 platforms=sorted({record['document']['row']['平台'] for record,_ in real_sources.values()});rows={}
 for project in projects:
  pid=project['id']
  for platform in platforms:
   sources=[record for record,pids in real_sources.values() if pid in pids and record['document']['row']['平台']==platform]
   kinds=Counter(record['kind'] for record in sources)
   views=sum(entry['view']['project_id']==pid and entry['doc']['row']['平台']==platform for entry in expected_units.values())
   rows[(pid,platform)]=[pid,project['name'],project['role'],platform,kinds['文章'],kinds['评论'],len(sources),
    len({str(record['document']['parent']['id']) for record in sources}),views]
 return rows

def normalized_branch_response(raw):
 result=copy.deepcopy(raw)
 units=result.get('units')
 if not isinstance(units,list):raise ValueError('Branch units must be an array')
 for unit in units:
  if not isinstance(unit,dict) or not isinstance(unit.get('views'),(list,dict)):raise ValueError('Branch unit views invalid')
  views=unit['views'].values() if isinstance(unit['views'],dict) else unit['views']
  if any(not isinstance(view,dict) or not isinstance(view.get('labels'),list) for view in views):raise ValueError('Branch project view invalid')
 result['units']=[derive_unit(unit) for unit in units]
 return result

def assert_raw_source_text(doc,own):
 field=doc.get('field');text=doc.get('text')
 assert field in own and isinstance(text,str) and text and text==own[field], 'Source text differs from its own frozen CSV field'


def model_response_index(selected,branch_responses=None):
 from voc_four_literal_recovery_v101 import normalize_role as role_aliases
 wanted={r['record_id']:r['document'] for r in selected};model_units=defaultdict(lambda:defaultdict(set));model_publishers=defaultdict(set)
 for folder in ['calls','call_attempt_history']:
  for file in (p.OUT/folder).glob('*.json'):
   call=p.read(file);rid=call['metrics']['record_id'];raw=call.get('response')
   if rid not in wanted or call['metrics']['status']!='success' or not isinstance(raw,dict):continue
   if branch_responses is not None and isinstance(raw.get('checked_tag_ids'),list):
    try:
     branch_digest=p.digest(normalized_branch_response(raw))
     branch_responses[rid].add(branch_digest)
    except (ValueError,KeyError,TypeError):pass
   if isinstance(raw.get('publisher'),dict):
    pub,_=role_aliases(raw['publisher'],p.own_fields(wanted[rid]));model_publishers[rid].add(p.digest(pub))
   units=raw.get('units')
   for unit in units if isinstance(units,list) else []:
    if isinstance(unit,dict) and 'unit_id' in unit and isinstance(unit.get('views'),(list,dict)):
     views=unit['views'].values() if isinstance(unit['views'],dict) else unit['views']
     if all(isinstance(view,dict) and isinstance(view.get('labels'),list) for view in views):
      try:model_units[rid][unit['unit_id']].add(p.digest(derive_unit(unit)))
      except (ValueError,KeyError,TypeError):pass
 # An unfinished unit may be completed by an identity-fields-only formal
 # response. Reproduce its exact two-response chain, never trust a free-standing
 # normalized unit or diagnostic HTTP record as model provenance.
 from voc_four_missing_identity_stage_v101 import validated_chain
 folder=p.OUT/'unit_field_corrections'
 for file in folder.glob('*.json'):
  proof=p.read(file);rid=proof['record_id']
  if rid not in wanted:continue
  unit,_=validated_chain(proof,wanted[rid],p.read(p.OUT/'taxonomy.json'))
  model_units[rid][unit['unit_id']].add(p.digest(unit))
 return model_units,model_publishers

def verify():
 if p.LOCK.exists():raise RuntimeError('Live generation lock: not final')
 completion=p.read(p.OUT/'completion.json')
 assert completion['state']=='generation_complete_pending_export' and not completion['active_sources']
 book=p.OUT/BOOK;assert book.is_file() and book.stat().st_size>10000
 selected=p.read(p.OUT/'selected_results.json');tax=p.read(p.OUT/'taxonomy.json');projects=p.read(p.OUT/'projects.json')
 assert len(selected)==1000 and Counter(r['kind'] for r in selected)==delivery_targets()
 preserved_cohort(selected)
 assert len({r['record_id'] for r in selected})==1000
 article_keys=[(r['document']['row']['平台'],r['document']['row']['文章id']) for r in selected if r['kind']=='文章']
 assert len(article_keys)==len(set(article_keys)), 'Duplicate original article under different source ids'
 comment_keys=[(r['document']['row']['平台'],r['document']['row']['文章id'],r['document']['row']['评论id']) for r in selected if r['kind']=='评论']
 assert len(comment_keys)==len(set(comment_keys)), 'Duplicate original comment under different source ids'
 assert len(tax)==580 and {t['l1'] for t in tax}=={'客户','市场','品牌','土地','产品','服务'}
 assert all(t['topics']==[p.v.b.T2] for t in tax)
 posts=csv_rows('ods_social_media_post_02.csv');comments=csv_rows('ods_social_media_comment_02.csv')
 frozen=p.read(p.OUT/'manifest.json')
 for f in frozen['inputs']:assert p.v.sha(f['snapshot'])==f['sha256']
 by_project={r['id']:r for r in projects};idx={t['tag_id']:t for t in tax}
 # A literal quote alone is not proof that Qwen produced the delivered label.
 # Bind every selected baseline unit and publisher to a saved formal response.
 branch_responses=defaultdict(set)
 model_units,model_publishers=model_response_index(selected,branch_responses)
 source_groups=defaultdict(set);links=defaultdict(dict);heat=defaultdict(dict);expected_units={};real_sources={};platform_articles={}
 for r in selected:
  d=r['document'];rid=r['record_id'];kind=d['kind'];own=(posts if kind=='文章' else comments)[str(d['source_id'])]
  assert own==d['row'],f'Frozen source row differs: {rid}'
  assert_raw_source_text(d,own)
  parent=posts[str(d['parent']['id'])];assert parent==d['parent']
  assert str(own['文章id'])==str(parent['文章id']) and own['平台']==parent['平台']
  assert r['signature']==p.signature(d,tax) and r['state']=='accepted' and not r['unfinished_units']
  if r.get('six_branch_recall_complete'):
   import complete_voc_four_branches_v101 as recall
   items={u['unit_id']:u for u in r['units']};parts=defaultdict(list);seen=set()
   for proof in r['branch_recall_proof']:
    pair=(proof['unit_id'],proof['level']);assert pair not in seen;seen.add(pair)
    branch=[t for t in tax if t['l1']==proof['level']];stored=p.read(proof['branch_file'])
    assert p.v.sha(proof['branch_file'])==proof['sha256'] and stored['original_source_signature']==r['signature']
    assert p.digest(normalized_branch_response(stored['response'])) in branch_responses[rid], f'Branch response not bound to a formal Qwen call: {rid}/{pair}'
    assert stored['signature']==p.digest([r['signature'],items[proof['unit_id']],branch,recall.SYSTEM])
    assert not recall.checked_keys_errors(stored['response'],branch)
    parts[proof['unit_id']].append((proof['level'],stored['response']['units'][0]))
   assert seen=={(uid,level) for uid in items for level in recall.LEVELS}
   for uid in items:assert recall.merge_outcomes(uid,parts[uid])==r['outcomes'][uid]
  if r.get('literal_recovery_proof'):
   literal_base=p.read(p.OUT/'branch_original_exclusions'/(p.digest(rid)+'.json')) if r.get('six_branch_recall_complete') else r
   for proof in r['literal_recovery_proof']:
    raw=p.read(proof['call_file']);assert p.v.sha(proof['call_file'])==proof['call_sha256']
    raw_unit=next(u for u in raw['response']['units'] if u['unit_id']==proof['unit_id'])
    assert derive_unit(raw_unit)==literal_base['outcomes'][proof['unit_id']]
  assert not p.product.role_errors(r['publisher'],p.own_fields(d))
  assert p.digest(r['publisher']) in model_publishers[rid], f'Publisher not bound to a formal Qwen response: {rid}'
  own_score=p.number(own['情绪评分']);assert own_score is not None and 0<=own_score<=100 and own_score==r['score']
  own_date=own['文章时间地点'] if kind=='文章' else own['评论时间'];assert own_date and own['平台']
  parent_id=str(parent['id']);parent_heat=p.number(parent['热度指数'])
  platform_articles[parent_id]=str(parent['文章id'])
  cursor=0;cat=p.v.supplement.evidence_catalog(r['units']);seen_projects=set()
  assert {u['unit_id'] for u in r['units']}==set(r['outcomes'])
  for unit in r['units']:
   assert unit['start']==cursor and d['text'][unit['start']:unit['end']]==unit['text']
   cursor=unit['end'];result=r['outcomes'][unit['unit_id']]
   if not r.get('six_branch_recall_complete'):
    assert p.digest(result) in model_units[rid][unit['unit_id']], f'Unit not bound to a formal Qwen response: {unit["unit_id"]}'
   assert not p.validate_unit(result,unit,cat,d,idx)
   result=unit_for_export(r,result)
   result=apply_targeted_tag_review(r,result)
   result=apply_new_cohort_review(r,result)
   assert not p.validate_unit(result,unit,cat,d,idx)
   for view in result['views']:
    pid=view['project_id'];assert pid in p.PIDS;seen_projects.add(pid)
    key=unit['unit_id']+'@'+pid;assert key not in expected_units
    expected_units[key]={'r':r,'doc':d,'unit':unit,'view':view,'date':own_date}
    for label in view['labels']:
     tid=label['tag_id']
     for group in ['全部',pid]:
      links[(group,tid)][rid]=own_score
      heat[(group,tid)][parent_id]=parent_heat
  assert cursor==len(d['text']) and seen_projects
  for group in ['全部',*seen_projects]:source_groups[group].add(rid)
  real_sources[rid]=(r,seen_projects)
 wb=load_workbook(book,read_only=True,data_only=True)
 group_names={'全部':'标签指标',**{q['id']:q['name']+'_标签指标' for q in projects}}
 expected_sheets=set(group_names.values())|{'打标明细','来源范围','历史撤标与补标','标签来源明细','热度文章来源','专题二标签体系','测试渠道统计','项目平台文章评论统计','源记录与文章关联'}
 assert set(wb.sheetnames)==expected_sheets
 summaries={}
 for group,name in group_names.items():
  ws=wb[name];header=[ws.cell(5,c).value for c in range(1,11)];assert header==FIELDS
  assert ws['F3'].value==len(source_groups[group])
  rows=[r[:10] for r in ws.iter_rows(min_row=6,max_row=585,values_only=True)];assert len(rows)==580
  source_metrics=p.read(p.OUT/'metric_results_v101.json')['metrics']
  expected_order=[g['tag_id'] for g in source_metrics if g['group']==group]
  for row,tid in zip(rows,expected_order):
   t=idx[tid];path=[t.get('l'+str(i)) or '不设'+['一','二','三','四','五'][i-1]+'级（原字典）' for i in range(1,6)]
   assert list(row[:6])==path+[t['emphasis']]
   scores=list(links[(group,tid)].values());hs=[v for v in heat[(group,tid)].values() if v is not None]
   expected=[len(scores)/len(source_groups[group]) if source_groups[group] else None,
    statistics.median(hs) if hs else None,statistics.mean(scores) if scores else None,len(scores)]
   for c,(a,b) in enumerate(zip(row[6:],expected)):close(a,b,f'{group}/{tid}/{c}')
  summaries[group]={'sources':len(source_groups[group]),'nonzero_labels':sum(bool(links[(group,t['tag_id'])]) for t in tax)}
 unit_head,unit_rows=sheet_rows(wb,'打标明细');unit_rows=[dict(zip(unit_head,row)) for row in unit_rows if row[0]]
 for key in ['project_id','项目名称','项目角色','发布日期','平台','发布者角色']:assert unit_head.count(key)==1
 assert {u['切片ID'] for u in unit_rows}==set(expected_units) and len(unit_rows)==len(expected_units)
 for u in unit_rows:
  e=expected_units[u['切片ID']];r=e['r'];d=e['doc'];v=e['view'];pid=v['project_id']
  assert u['record_id']==r['record_id'] and u['project_id']==pid and u['项目名称']==by_project[pid]['name'] and u['项目角色']==by_project[pid]['role']
  assert str(u['源表id'])==d['source_id'] and str(u['父文章ODS主键'])==str(d['parent']['id'])
  assert u['平台']==d['row']['平台'] and u['发布者角色']==r['publisher']['role']
  assert_own_date(u['发布日期'],e['date'])
  assert u['原始摘录']==e['unit']['text'] and u['观点摘要']==v['summary']
  assert json.loads(u['标签明细'])==[p.v.b.label_json(idx[t['tag_id']]) for t in v['labels']]
  assert unpack_evidence(json.loads(u['项目与标签证据']))==label_evidence(v,p.v.supplement.evidence_catalog(r['units'])), 'Saved evidence differs from literal/formal Qwen evidence'
  assert u['情绪评分']==r['score']
 h,rows=sheet_rows(wb,'来源范围');assert sum(bool(r[0]) for r in rows)==1000
 assert {r[0] for r in rows if r[0]}==set(real_sources)
 for row in rows:
  if row[0]:
   record=real_sources[row[0]][0];own=record['document']['row']
   assert_own_date(row[7],own['文章时间地点'] if record['kind']=='文章' else own['评论时间'])
 h,rows=sheet_rows(wb,'源记录与文章关联');observed={(r[0],r[4]) for r in rows if r[0]}
 assert observed=={(rid,pid) for rid,(_,pids) in real_sources.items() for pid in pids}
 for row in rows:
  if row[0]:
   assert isinstance(row[17],str) and row[17]==platform_articles[str(row[3])]
   record=real_sources[row[0]][0];own=record['document']['row']
   assert_own_date(row[7],own['文章时间地点'] if record['kind']=='文章' else own['评论时间'])
 h,rows=sheet_rows(wb,'专题二标签体系');assert {r[0] for r in rows if r[0]}==set(idx)
 taxonomy_mismatches=[]
 for row in rows:
  if not row[0]:continue
  tag=idx[row[0]]
  expected_taxonomy=[tag['tag_id'],p.number(tag['source_number']),*[tag.get('l'+str(i)) or None for i in range(1,6)],tag['emphasis'],tag['instruction'],'专题二']
  if list(row[:10])!=expected_taxonomy:taxonomy_mismatches.append(row[0])
 # Verify statistical source-pair rows, parent dedup, and channel/article/comment counts.
 h,rows=sheet_rows(wb,'标签来源明细');pairs=[r for r in rows if r[0]]
 pair_keys=[(r[0],r[1],r[3]) for r in pairs];assert len(pair_keys)==len(set(pair_keys))
 assert set(pair_keys)=={(group,tid,rid) for (group,tid),ss in links.items() for rid in ss}
 for row in pairs:close(row[7],links[(row[0],row[1])][row[3]],'pair score')
 h,rows=sheet_rows(wb,'热度文章来源');hs=[r for r in rows if r[0]]
 heat_keys=[(r[0],r[1],str(r[2])) for r in hs];assert len(heat_keys)==len(set(heat_keys))
 assert set(heat_keys)=={(group,tid,parent) for (group,tid),ss in heat.items() for parent in ss}
 for row in hs:
  close(row[5],heat[(row[0],row[1])][str(row[2])],'parent heat')
  assert isinstance(row[7],str) and row[7]==platform_articles[str(row[2])]
 h,rows=sheet_rows(wb,'测试渠道统计')
 channel=Counter((r['document']['row']['平台'],r['kind']) for r in selected)
 assert {(r[0],r[1]):r[2] for r in rows if r[0]}=={(plat,kind):channel[(plat,kind)] for plat in sorted({a for a,b in channel}) for kind in ['文章','评论']}
 for row in rows:
  if not row[0]:continue
  cohort=[record for record in selected if record['document']['row']['平台']==row[0] and record['kind']==row[1]]
  assert list(row[:6])==[row[0],row[1],len(cohort),len(cohort),
   sum(p.number(record['document']['parent']['热度指数']) is not None for record in cohort),
   sum(entry['doc']['row']['平台']==row[0] and entry['doc']['kind']==row[1] for entry in expected_units.values())]
 h,rows=sheet_rows(wb,'项目平台文章评论统计');actual_rows=[list(row[:9]) for row in rows if row[0]]
 expected_platform=project_platform_rows(real_sources,expected_units,projects)
 assert len(actual_rows)==len(expected_platform) and len({(row[0],row[3]) for row in actual_rows})==len(actual_rows)
 assert {(row[0],row[3]):row for row in actual_rows}==expected_platform, 'Project/platform article/comment or parent counts differ'
 for ws in wb:
  for row in ws.iter_rows():
   for cell in row:
    assert cell.data_type!='e',f'Excel error: {ws.title}!{cell.coordinate}'
    assert '互动深度指数' not in str(cell.value or ''),f'Removed field still present: {ws.title}!{cell.coordinate}'
 wb.close()
 formula_wb=load_workbook(book,read_only=True,data_only=False)
 for name in group_names.values():
  for row in formula_wb[name].iter_rows(min_row=6,max_row=585,min_col=7,max_col=10):assert all(c.data_type=='f' for c in row)
 formula_wb.close()
 proof=p.read(p.OUT/'workbook_verification.json');assert proof['recalculation_verified'] and proof['score_input_change_verified'] and proof['heat_input_change_verified'] and not proof['partial_not_delivery']
 assert proof['metrics_sha256']==p.v.sha(p.OUT/'metric_results_v101.json')
 assert proof['builder_sha256']==p.v.sha(Path(__file__).with_name('build_voc_four_projects_v101.mjs'))
 previews=p.OUT/'previews';assert all(any(previews.glob(name+'_*.png')) for name in expected_sheets)
 assert not taxonomy_mismatches, 'Saved taxonomy values differ: '+','.join(taxonomy_mismatches)
 final={'passed':True,'book':str(book),'book_sha256':p.v.sha(book),'sources':1000,'source_counts':delivery_targets(),
  'slice_project_views':len(unit_rows),'topic2_dictionary':580,'project_groups':summaries,'independent_frozen_csv_verification':True,
  'parent_heat_dedup_verified':True,'own_sentiment_verified':True,'full_partition_and_literal_evidence_verified':True,
  'six_metadata_fields_verified':True,'interaction_index_removed':True,'saved_formula_values_verified':True,
  'saved_literal_evidence_exactly_matches_verified_qwen_and_overlays':True,'project_platform_all_counts_independently_verified':True,
  'full_saved_taxonomy_paths_and_definitions_verified':True,
  'all_delivered_tags_and_publisher_verdicts_bound_to_formal_qwen_responses':True,
  'codex_semantic_gold_gate':False,'semantic_boundary':'Grounded Qwen reasons and full-directory/fact selfcheck; not human gold or absolute zero-miss',
  'visual_review_required_before_delivery':True}
 p.save('saved_workbook_verification.json',final);return final

if __name__=='__main__':
 sys.stdout.reconfigure(encoding='utf-8');print(json.dumps(verify(),ensure_ascii=False))
