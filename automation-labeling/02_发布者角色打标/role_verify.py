import json,hashlib,os
from pathlib import Path
import openpyxl
out=Path(os.environ['ROLE_TEST_OUTPUT'])
data=json.loads((out/'comparison_data.json').read_text(encoding='utf-8'))
p=out/'发布者角色_Codex与Qwen_两专题各50条全文复测.xlsx'
values=openpyxl.load_workbook(p,data_only=True);formulas=openpyxl.load_workbook(p,data_only=False)
assert values.sheetnames==['评估汇总','逐条对照','共同输入']
def norm(v):return '' if v is None else v
for i,row in enumerate(data['rows'],6):
 for j,v in enumerate(row,1):
  got=values['逐条对照'].cell(i,j)
  assert norm(got.value)==norm(v),(i,j,'saved value differs')
  if j in [4,5]:assert got.data_type=='s'
assert sum(values['逐条对照'].cell(i,1).value is not None for i in range(6,values['逐条对照'].max_row+1))==100
for i,stat in enumerate(data['stats'],11):
 assert values['评估汇总'].cell(i,2).value==stat['n']
 assert values['评估汇总'].cell(i,4).value==stat['agree']
 assert abs(values['评估汇总'].cell(i,5).value-stat['agreement'])<1e-9
matrix=[[values['评估汇总'].cell(i,j).value for j in range(2,6)] for i in range(34,38)]
assert sum(map(sum,matrix))==100 and sum(matrix[i][i] for i in range(4))==data['stats'][2]['agree']
raw={r['sample_id']:r for r in data['raw']}
audits=[json.loads(l) for l in (out/'request_input_audit.jsonl').read_text(encoding='utf-8').splitlines()]
assert len(audits)==100 and len({a['sample_id'] for a in audits})==100
amap={a['sample_id']:a for a in audits}
for i,packet in enumerate(data['packets'],7):
 sid=packet['sample_id'];a=amap[sid];record=raw[sid]
 assert not packet['truncated_fields'] and a['full_text'] and a['serialized_input_exact']
 for k,v in packet['fields'].items():
  assert v==(record['row'].get(k) or '')
  assert hashlib.sha256(v.encode('utf-8')).hexdigest()==a['field_sha256'][k]
 expected=[sid,packet['dataset'],record['source_line'],record['row']['id'],record['row']['文章id'],*[packet['fields'][k] for k in ['平台','作者名字','作者主页链接','文章标题','文章标签','文章内容','文章视频内容','文章图片信息']],'']
 for j,v in enumerate(expected,1):assert norm(values['共同输入'].cell(i,j).value)==norm(v),(sid,j,'full saved input differs')
 assert values['共同输入'].cell(i,4).data_type=='s' and values['共同输入'].cell(i,5).data_type=='s'
errors=[]
for sh in values:
 for row in sh:
  for c in row:
   if c.data_type=='e':errors.append((sh.title,c.coordinate,c.value))
assert not errors
assert len(formulas['逐条对照'].data_validations.dataValidation)==1
assert hashlib.sha256(Path(data['manifest'][0]['source']).read_bytes()).hexdigest()==data['manifest'][0]['sha256']
assert hashlib.sha256(Path(data['manifest'][1]['source']).read_bytes()).hexdigest()==data['manifest'][1]['sha256']
result={'passed':True,'sheets':values.sheetnames,'detail_rows':100,'full_input_saved_and_http_hash_verified':100,'zero_truncated_inputs':True,'saved_ids_exact_text':True,'saved_labels_reasons_evidence_exact':True,'summary_and_matrix_reconciled':True,'formula_errors':errors,'source_csv_unchanged':True,'xlsx_bytes':p.stat().st_size}
(out/'saved_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
