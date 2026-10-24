"""Independent read-only XLSX verification and contact sheets for visual QA."""
import argparse,json,datetime
from pathlib import Path
import openpyxl
from PIL import Image,ImageDraw,ImageFont

def main():
    ap=argparse.ArgumentParser();ap.add_argument('output_dir',type=Path);a=ap.parse_args();checks=[]
    for kind in ['产品信息单元','观点信息单元']:
        p=json.loads((a.output_dir/(kind+'_workbook_payload.json')).read_text(encoding='utf-8'))
        file=a.output_dir/p['filename'];wb=openpyxl.load_workbook(file,data_only=True,read_only=True);ws=wb[kind]
        headers=p['schemas'][kind];assert list(next(ws.iter_rows(values_only=True)))==headers
        rows=list(ws.iter_rows(min_row=2,values_only=True));assert len(rows)==len(p['tables'][kind])
        for actual,expected in zip(rows,p['tables'][kind]):
            for h,v in zip(headers,actual):
                target=expected[h]
                if isinstance(v,datetime.datetime):assert str(target)==v.strftime('%Y-%m-%d %H:%M:%S') or str(target)==v.strftime('%Y-%m-%d')
                elif isinstance(target,(dict,list)):assert v==json.dumps(target,ensure_ascii=False,separators=(',',':'))
                else:assert v==target or v is None and target=='', (h,v,target)
        stats=list(wb['测试渠道统计'].iter_rows(min_row=6,max_row=5+len(p['tables']['测试渠道统计']),max_col=11,values_only=True))
        for row,expected in zip(stats,p['tables']['测试渠道统计']):assert list(row)==list(expected.values())
        formulas=openpyxl.load_workbook(file,data_only=False,read_only=False)
        assert formulas[kind].freeze_panes=='C2'
        counts=sum(c.data_type=='f' for row in formulas['测试渠道统计'] for c in row)
        assert counts==6*len(stats)
        assert formulas['人工审核'].data_validations.count>0
        verify=json.loads((a.output_dir/(kind+'_核验')/'export_verification.json').read_text(encoding='utf-8'))
        assert verify['verifiedSheets']==len(p['tables']) and verify['mainRows']==len(rows)
        checks.append({'file':str(file),'main_rows':len(rows),'exact_all_main_values':True,'statistics_formula_count':counts,'statistics_cache_matches_source':True,'sheets':wb.sheetnames})
        qa=a.output_dir/(kind+'_核验');images=sorted(qa.glob('*.png'));canvas=Image.new('RGB',(1600,len(images)*450),'white');draw=ImageDraw.Draw(canvas)
        font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',20)
        for i,imfile in enumerate(images):
            im=Image.open(imfile).convert('RGB');im.thumbnail((1560,405))
            draw.text((20,i*450+5),imfile.stem,font=font,fill='#17365D');canvas.paste(im,(20,i*450+38))
        canvas.save(qa/'全部工作表视觉核验.png')
    (a.output_dir/'independent_xlsx_verification.json').write_text(json.dumps({'passed':True,'checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':True,'checks':checks},ensure_ascii=False))

if __name__=='__main__':main()
