"""Independent read-only XLSX, dictionary, source preservation and delta checks."""
import argparse
import json
from collections import Counter
from pathlib import Path
import openpyxl
from verify_label_reference_v91 import source_references,TAG_KEYS,REFERENCE_FIELDS


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--prior',type=Path,required=True)
    ap.add_argument('--output-dir',type=Path,required=True);ap.add_argument('--taxonomy',type=Path,required=True);a=ap.parse_args()
    prior=json.loads(a.prior.read_text(encoding='utf-8'));data=json.loads((a.output_dir/'delivery_results.json').read_text(encoding='utf-8'))
    refs=source_references(a.taxonomy);old={r['document']['record_id']:r for r in prior['results']}
    assert len(data['results'])==len(old)
    delta=Counter();ids=set();new_total=0;removed_total=0
    for r in data['results']:
        d=r['document'];p=old[d['record_id']]
        assert d==p['document'], 'Source metadata/text changed'
        assert len(r['rows'])==len(r['traces'])
        before={u['unit_id']:u for u in p['rows']};after={u['unit_id']:u for u in r['rows']}
        after_tags={uid:{t['标签路径ID'] for t in u['标签明细']} for uid,u in after.items()}
        before_tags={uid:{t['标签路径ID'] for t in u['标签明细']} for uid,u in before.items()}
        for uid,tags in before_tags.items():
            removed=tags-after_tags.get(uid,set());removed_total+=len(removed)
            for tid in removed:
                assert any(e['处理']=='两项核验一致撤回错误标签' and e['unit_id']==uid and 'VOC-'+e['key'][1:]==tid for e in r['label_supplement_events'])
                delta['withdrawn_tags']+=1
            if uid not in after:
                assert uid in {q['unit_id'] for q in r['pending_units']},'Removed unit not traceable'
                delta['moved_to_pending']+=1
        for uid,u in after.items():
            assert uid not in ids;ids.add(uid)
            added=after_tags[uid]-before_tags.get(uid,set());new_total+=len(added);delta['new_tags']+=len(added)
            if uid not in before:delta['restored_units']+=1
            else:
                for field,value in before[uid].items():
                    if field in ['标签明细','标签强调方向','打标说明','备注','验收状态']:continue
                    if field in ['project_id','项目名称','项目角色'] and not value:continue
                    assert u[field]==value,(uid,'Unrelated field changed',field)
        for u,t in zip(r['rows'],r['traces']):
            assert d['text'][t['起点']:t['终点']]==u.get('原文摘录',u.get('原始摘录'))
            assert u['标签明细'] and u['打标说明'] and u.get('具体产品做法',u.get('观点摘要'))
            if d['kind']=='评论':
                assert len(r['rows'])==1 and u['原始摘录']==d['text']
                assert u['映射文章id']==(d['row'].get('文章id') or None)
                assert '无态度' not in u['情绪态度'].split(';') or u['情绪态度']=='无态度'
            if '观点摘要' in u:assert len(u['观点摘要'])<=80
            for proof in t.get('全字典补选证据',[]):
                assert proof.get('quotes') and not proof.get('program_issues')
                for span in proof.get('attribute_spans',[]):
                    assert t['起点']<=span['start']<span['end']<=t['终点']
                    assert d['text'][span['start']:span['end']]==span['quote']
    assert new_total==data['supplement_status']['new_tags']
    assert removed_total==data['supplement_status']['withdrawn_wrong_tags']
    checks=[]
    for topic,kind in enumerate(['产品信息单元','观点信息单元']):
        pack=json.loads((a.output_dir/(kind+'_workbook_payload.json')).read_text(encoding='utf-8'))
        book=openpyxl.load_workbook(a.output_dir/pack['filename'],read_only=False,data_only=False)
        cache=openpyxl.load_workbook(a.output_dir/pack['filename'],read_only=True,data_only=True)
        try:
            dictionary='专题一标签体系' if topic==0 else '专题二标签体系'
            assert book.sheetnames[:2]==[kind,dictionary]
            assert book[kind].freeze_panes=='C2';assert dictionary in book
            headers=[c.value for c in book[kind][1]];assert headers==pack['schemas'][kind]
            assert book[kind].max_row==len(pack['tables'][kind])+1
            registry={t['标签路径ID']:t for t in refs[topic]}
            assigned=0
            for n,u in enumerate(pack['tables'][kind],2):
                actual={h:book[kind].cell(n,i+1).value for i,h in enumerate(headers)}
                tags=json.loads(actual['标签明细']);assigned+=len(tags)
                assert len(tags)==len({t['标签路径ID'] for t in tags})
                for tag in tags:
                    assert list(tag)==TAG_KEYS and tag=={k:registry[tag['标签路径ID']][k] for k in TAG_KEYS}
                assert actual['标签强调方向']==';'.join(dict.fromkeys(t['标签强调方向'] for t in tags if t['标签强调方向']))
                for h in ['unit_id','打标说明','原文摘录' if topic==0 else '原始摘录','具体产品做法' if topic==0 else '观点摘要','热度指数','互动深度指数']:
                    assert actual[h]==u[h],(kind,n,h)
            reference=book[dictionary]
            assert [c.value for c in reference[5]][:len(REFERENCE_FIELDS)]==REFERENCE_FIELDS
            assert reference.max_row==len(refs[topic])+5
            for n,item in enumerate(refs[topic],6):
                for i,h in enumerate(REFERENCE_FIELDS,1):
                    value=reference.cell(n,i).value;expected=item[h]
                    assert value==expected or value is None and expected=='',(dictionary,n,h)
            for name in book.sheetnames:
                for row in cache[name].iter_rows():assert all(c.data_type!='e' for c in row),('Formula error',name)
            source_sheet=book['源记录与文章关联']
            for column in source_sheet.iter_cols():
                if column[0].value not in ['源表id','父文章ODS主键','评论映射父文章id']:continue
                for c in column[1:]:
                    if c.value is not None and str(c.value).isdigit() and len(str(c.value))>=16:
                        assert c.data_type=='s' and c.number_format=='@',('Long source ID not stored as text',kind,c.coordinate)
            for n,row in enumerate(pack['tables']['测试渠道统计'],6):
                stat_headers=[c.value for c in book['测试渠道统计'][5]][:len(row)]
                for i,h in enumerate(stat_headers,1):assert cache['测试渠道统计'].cell(n,i).value==row[h],('Statistic mismatch',h)
                assert book['测试渠道统计'].cell(n,stat_headers.index('范围内行数')+1).data_type=='f'
            for name in ['人工审核','未完成待核','标签补选与纠错','未解决事实']:
                if pack['tables'].get(name):assert book[name].data_validations.dataValidation
            if data.get('summary',{}).get('version','').startswith(('v93-','v94-','v95-')):
                usage_report=json.loads((a.output_dir/'token_cost_summary.json').read_text(encoding='utf-8'))
                scoped=[x for x in usage_report['by_topic_source'] if x['topic']==pack['topic']]
                costs=list(cache['本轮补选Token费用'].iter_rows(values_only=True))
                fee_headers=list(costs[0]);fee_rows=costs[1:]
                for heading,key in [('实际请求记录数','request_records'),('HTTP尝试数','http_or_network_attempts'),
                    ('已知输入Token','prompt_tokens'),('已知缓存Token','cached_input_tokens'),
                    ('已知输出Token','completion_tokens'),('未知usage的HTTP尝试数','attempts_without_usage')]:
                    actual=sum(row[fee_headers.index(heading)] or 0 for row in fee_rows)
                    expected=sum(item[key] for item in scoped)
                    assert actual==expected,('Cost detail and actual usage mismatch',kind,heading,actual,expected)
                actual=sum(row[fee_headers.index('已知估算费用元')] or 0 for row in fee_rows)
                expected=sum(item['known_estimated_yuan'] for item in scoped)
                assert abs(actual-expected)<1e-8,('Cost detail price mismatch',kind)
            qa=json.loads((a.output_dir/(kind+'_核验')/'export_verification.json').read_text(encoding='utf-8'))
            assert Path(qa['output']).resolve()==(a.output_dir/pack['filename']).resolve()
            assert qa['mainRows']==len(pack['tables'][kind]) and qa['verifiedSheets']==len(pack['tables'])
            checks.append({'file':str(a.output_dir/pack['filename']),'main_rows':len(pack['tables'][kind]),
                'dictionary_rows':len(refs[topic]),'assigned_tags':assigned,'passed':True})
        finally:book.close();cache.close()
    report={'passed':True,'sources':len(old),'delta':dict(delta),'checks':checks,
        'note':'Checks export, literal evidence and dictionary contracts; not a claim of zero semantic errors.'}
    (a.output_dir/'independent_xlsx_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':main()
