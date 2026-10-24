from __future__ import annotations

import argparse
import calendar
import ctypes
import csv
import hashlib
import io
import json
import math
import os
import queue
import random
import re
import signal
import socket
import sqlite3
import subprocess
import sys
import time
import tempfile
import threading
import wave
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import openpyxl
import win32clipboard
from PIL import ImageGrab
from pywinauto import Desktop, keyboard, mouse

SRC_DIR = Path(__file__).resolve().parent
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))
from wechat_media_cache import discover_cache_dirs, snapshot_cache, wait_for_unique_media
from channels_resilience import RedisVideoCheckpointStore, RedisCheckpointUnavailable, atomic_write_text
from asr_model_loader import load_asr_model
from subtitle_review import init_review_tables, stage_candidate, has_internal_prompt, has_untrusted_transcript, force_retranscribe, checkpoint_repair_reason
from transcript_quality import transcript_artifact
from search_scan_proof import POLICY as SEARCH_SCAN_POLICY, TERMINAL_TEXTS, SearchEndTracker, valid_origin
from operational_gate import report_from_environment, classify_restriction
from audio_process_guard import stop_owned_recorder, cleanup_audio_file

# Windows consoles often default to GBK; force UTF-8 so real WeChat titles
# containing Japanese punctuation/emoji cannot terminate a long collection.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


WINDOW_TITLE = "微信"
DEFAULT_START_DATE = "2026-01-01"
STOP_REQUESTED = False
SEARCH_LAYOUT: str | None = None
LAST_MULTI_WINDOW_SIGNATURE: tuple | None = None
_ASR_MODEL = None
_AUDIO_SESSION_IMPORT_WARNING_EMITTED = False


class SearchWindowUnavailableError(RuntimeError):
    """The WeChat 搜一搜 surface disappeared while Chromium was navigating."""


class ResultPositionError(RuntimeError):
    """A virtualized search card could not be brought into the safe viewport."""


class VideoDetailIdentityMismatchError(RuntimeError):
    """The open feed item is not the video card that was clicked."""


class ClientAudioCaptureError(RuntimeError):
    """The client playback could not yield a trustworthy audio recording."""


class VideoPlaybackUnavailableError(RuntimeError):
    """Bounded network-player reloads failed; retain the current video cursor."""


class VideoDetailRecoveryBlockedError(ClientAudioCaptureError):
    """A stop/login/security gate forbids further detail refreshes."""


class AsrQualityRejectedError(RuntimeError):
    """Inference worked; retrying the identical service locally is not recovery."""


def request_graceful_stop(_signum=None, _frame=None) -> None:
    global STOP_REQUESTED
    STOP_REQUESTED = True
    print("Graceful stop requested. Finishing the current article, then exporting and importing committed rows.", flush=True)


def configured_range(
    config: dict, key: str, default: tuple[float, float],
) -> tuple[float, float]:
    """Return a normalized non-negative [min,max] configuration range."""
    value = config.get(key, list(default))
    if isinstance(value, (int, float)):
        low = high = float(value)
    elif isinstance(value, (list, tuple)) and len(value) >= 2:
        low, high = float(value[0]), float(value[1])
    else:
        low, high = map(float, default)
    low, high = max(0.0, low), max(0.0, high)
    return (low, high) if low <= high else (high, low)


def random_delay_seconds(
    config: dict, key: str, default: tuple[float, float],
) -> float:
    if not bool(config.get("jitter_enabled", False)):
        return 0.0
    low, high = configured_range(config, key, default)
    return round(random.uniform(low, high), 3)


def interruptible_wait(seconds: float) -> bool:
    """Wait in short slices so a safe-stop is observed during long rests."""
    deadline = time.monotonic() + max(0.0, float(seconds))
    while True:
        if STOP_REQUESTED:
            return False
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return True
        time.sleep(min(0.5, remaining))


def jitter_pause(
    config: dict, key: str, default: tuple[float, float], reason: str,
    **context,
) -> bool:
    seconds = random_delay_seconds(config, key, default)
    if seconds <= 0:
        return not STOP_REQUESTED
    print(json.dumps({
        "event": "rpa_jitter_pause",
        "reason": reason,
        "seconds": seconds,
        **context,
    }, ensure_ascii=False), flush=True)
    return interruptible_wait(seconds)


REDIS_SEARCH_TOKEN_LUA = """
local clock = redis.call('TIME')
local now = tonumber(clock[1]) * 1000 + math.floor(tonumber(clock[2]) / 1000)
local current = tonumber(redis.call('GET', KEYS[1]) or '0')
if current <= now then
  local interval = tonumber(ARGV[1])
  local next_at = now + interval
  redis.call('SET', KEYS[1], next_at, 'PX', math.max(60000, interval * 4))
  return {1, 0, next_at}
end
return {0, current - now, current}
""".strip()


def wait_for_distributed_search_token(config: dict, keyword: str) -> bool:
    """Serialize keyword submissions across Redis workers without blocking ASR.

    Only the instant that submits/rebuilds a search is rate limited. Detail
    parsing, playback, offline ASR, export and MySQL work remain independent.
    """
    if not bool(config.get("redis_search_token_enabled", False)):
        return True
    url = os.environ.get("REDIS_URL", "").strip()
    task_id = os.environ.get("WECHAT_REDIS_TASK_ID", "").strip()
    if not url or not task_id:
        return True
    from redis_task_queue import redis_command

    password = os.environ.get("REDIS_PASSWORD") or None
    run_id = re.sub(
        r"[^0-9A-Za-z_-]+", "_", os.environ.get("REDIS_RUN_ID", "default")
    ).strip("_") or "default"
    token_key = f"wechat:channels:{run_id}:search-token"
    low, high = configured_range(
        config, "redis_search_token_seconds", (12.0, 25.0),
    )
    while not STOP_REQUESTED:
        interval_seconds = random.uniform(low, high)
        try:
            reply = redis_command(
                url, password, "EVAL", REDIS_SEARCH_TOKEN_LUA, 1,
                token_key, max(1, int(interval_seconds * 1000)),
            )
        except Exception as error:
            print(json.dumps({
                "event": "redis_search_token_error",
                "keyword": keyword,
                "error": str(error),
                "fail_open": bool(config.get("redis_search_token_fail_open", False)),
            }, ensure_ascii=False), flush=True)
            if bool(config.get("redis_search_token_fail_open", False)):
                return True
            raise SearchWindowUnavailableError(
                f"Redis搜索节奏令牌不可用，已停止提交关键词：{error}"
            ) from error
        granted = bool(reply and int(reply[0]) == 1)
        if granted:
            print(json.dumps({
                "event": "redis_search_token_acquired",
                "keyword": keyword,
                "interval_seconds": round(interval_seconds, 3),
                "token_key": token_key,
            }, ensure_ascii=False), flush=True)
            return True
        wait_seconds = max(0.2, float(reply[1]) / 1000.0)
        print(json.dumps({
            "event": "redis_search_token_wait",
            "keyword": keyword,
            "seconds": round(wait_seconds, 3),
        }, ensure_ascii=False), flush=True)
        if not interruptible_wait(wait_seconds):
            return False
    return False
ODS_FIELDS = [
    "id", "文章图片链接", "文章图片信息", "作者名字", "文章标题", "文章内容", "文章标签",
    "文章时间地点", "文章点赞数", "文章收藏数", "文章评论数", "文章链接", "文章id", "作者主页链接",
    "文章视频链接", "文章视频内容", "关键词", "作者id", "作者粉丝数", "作者点赞数", "文章分享数",
    "文章视频时长", "平台", "文章获取时间", "文章标注", "分类", "清洗标注", "项目", "地点",
    "内容类型", "是否广告", "广告关键词",
]


@dataclass
class ResultRow:
    project: str
    keyword: str
    title: str
    url: str
    author_name: str | None = None
    tags: str | None = None
    published_at: str | None = None
    location: str | None = None
    content: str | None = None
    image_count: int = 0
    video_count: int = 0
    video_content: str | None = None
    video_duration: int | None = None
    like_count: int | None = None
    comment_count: int | None = None
    favorite_count: int | None = None
    share_count: int | None = None
    captured_at: str = ""

    def __post_init__(self) -> None:
        if not self.captured_at:
            self.captured_at = datetime.now().isoformat(timespec="seconds")


@dataclass
class PendingClientAsr:
    wav_path: Path
    title: str
    prompt_context: str | None
    duration: int
    capture_seconds: int
    recorder_seconds: int
    recorder_audit: dict
    seeked: bool
    reloaded_from_start: bool
    source: str = "windows_wasapi_loopback"


def load_config(path: Path) -> dict:
    if not path.exists():
        raise FileNotFoundError(f"配置文件不存在：{path}")
    return json.loads(path.read_text(encoding="utf-8"))


def parse_video_duration_seconds(value: str) -> int | None:
    text = str(value or "")
    # Detail pages expose ``elapsed / total``; the right-hand value is the
    # actual duration, while search-card autoplay can show a decreasing value.
    total = re.search(r"/\s*((?:\d+:)?\d{1,2}:\d{2})(?:\s|$)", text)
    if total:
        text = total.group(1)
    match = re.match(r"^\s*(?:(\d+):)?(\d{1,2}):(\d{2})(?:\s|$)", text)
    if not match:
        return None
    hours = int(match.group(1) or 0)
    minutes = int(match.group(2))
    seconds = int(match.group(3))
    return hours * 3600 + minutes * 60 + seconds


def clean_asr_text(value: str | None) -> str | None:
    text = str(value or "").replace("\r", "\n")
    for boilerplate in (
        "请不吝点赞订阅转发打赏支持明镜与点点栏目",
        "请点赞订阅转发打赏支持明镜与点点栏目",
        "请不吝点赞订阅转发打赏支持明镜与点点",
    ):
        text = text.replace(boilerplate, "")
    # Faster-Whisper can hallucinate a generic creator outro over long music
    # or near-silence.  Remove it only when the same short suffix contains a
    # viewing phrase plus engagement prompts; ordinary spoken “谢谢” remains.
    text = re.sub(
        r"(?:[，。；！？])?(?:多谢|感谢)(?:大家|各位|这边)?(?:的)?"
        r"(?:观看|收看).{0,16}(?:订阅|点赞|转发).*$",
        "",
        text,
    )
    text = re.sub(r"[\U0001F000-\U0001FAFF\u2600-\u27BF]", "", text)
    text = re.sub(r"\s*\n+\s*", "，", text)
    text = re.sub(r"[\t ]+", "", text)
    text = re.sub(r"，{2,}", "，", text).strip("， ")
    # Whisper can repeat a complete phrase block, especially during a no-VAD
    # recovery pass. Collapse exact adjacent blocks without touching ordinary
    # repeated words or emphatic single characters.
    for _ in range(3):
        collapsed = re.sub(r"(.{6,}?[，。！？；])(?:\1)+", r"\1", text)
        if collapsed == text:
            break
        text = collapsed
    clauses = [part.strip() for part in re.split(r"[，。！？；]+", text) if part.strip()]
    original_clauses = list(clauses)
    changed = True
    while changed and len(clauses) >= 2:
        changed = False
        for width in range(len(clauses) // 2, 0, -1):
            for start in range(0, len(clauses) - width * 2 + 1):
                if clauses[start:start + width] == clauses[start + width:start + width * 2]:
                    del clauses[start + width:start + width * 2]
                    changed = True
                    break
            if changed:
                break
    if clauses != original_clauses:
        text = "，".join(clauses)
    return text or None


def asr_speech_compare_key(value: str) -> str:
    """Normalize punctuation and Chinese script variants for ASR overlap checks."""
    comparable = value
    # Independent head/full passes may spell a spoken number differently.
    # Normalize the common ``10分钟``/``十分钟`` form before overlap matching.
    comparable = re.sub(r"10(?=分钟|秒|米|平|套|个|层|年|月|日)", "十", comparable)
    try:
        converter = ctypes.windll.kernel32.LCMapStringEx
        needed = converter("zh-CN", 0x02000000, comparable, -1, None, 0, None, None, 0)
        if needed > 0:
            buffer = ctypes.create_unicode_buffer(needed)
            if converter("zh-CN", 0x02000000, comparable, -1, buffer, needed, None, None, 0) > 0:
                comparable = buffer.value
    except Exception:
        # This collector runs on Windows, but keep a deterministic fallback
        # for unit tests and stripped-down runtimes.
        comparable = comparable.translate(str.maketrans({"說": "说", "萬": "万", "發": "发"}))
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", comparable).lower()


def merge_asr_head_prefix(head_text: str | None, full_text: str | None) -> str | None:
    """Restore a missing opening phrase when head/full transcripts overlap."""
    head = clean_asr_text(head_text)
    full = clean_asr_text(full_text)
    if not head:
        return full
    if not full:
        return head
    if head in full[: max(len(head) * 2, 24)]:
        return full
    if full in head:
        return head
    # The independent head pass and the full pass often disagree only on
    # punctuation (for example ``臻澐,三个`` versus ``臻澐三个``). Compare a
    # punctuation-free speech key before prepending the head; otherwise the
    # recovery path would duplicate an opening sentence that is already in
    # the full transcript.
    head_key = asr_speech_compare_key(head)
    full_key = asr_speech_compare_key(full)
    full_opening_key = full_key[: max(len(head_key) * 2, 24)]
    if head_key and head_key in full_opening_key:
        return full
    if full_key and full_key in head_key:
        return head
    # Recover a demonstrably missing phrase before fuzzy containment checks.
    maximum = min(len(head), len(full), 80)
    for overlap in range(maximum, 3, -1):
        if head[-overlap:] == full[:overlap]:
            return clean_asr_text(head + full[overlap:])
    # Whisper may emit the same opening once in simplified Chinese and once
    # in traditional Chinese. When the shorter transcript closely matches
    # the beginning of the longer one, keep the longer version instead of
    # concatenating both scripts as duplicate speech.
    shorter_key, longer_key = sorted((head_key, full_key), key=len)
    if len(shorter_key) >= 8:
        longer_prefix = longer_key[: len(shorter_key)]
        prefix_similarity = SequenceMatcher(
            None,
            shorter_key,
            longer_prefix,
        ).ratio()
        # Compare the entire anchored prefix, not the start of its longest
        # exact substring: homophone errors near the beginning otherwise
        # turn two versions of the same speech into a duplicated opening.
        if prefix_similarity >= 0.72:
            return head if len(head_key) >= len(full_key) else full
    if head_key and full_opening_key:
        comparable = full_opening_key[: min(len(full_opening_key), len(head_key) + 8)]
        match = SequenceMatcher(None, head_key, comparable).find_longest_match()
        if (
            match.a <= 2
            and match.b <= 2
            and match.size >= min(8, len(head_key))
            and match.size / max(1, len(head_key)) >= 0.82
        ):
            return full
    # Without overlap or timestamp evidence we cannot establish that speech
    # is missing. Never concatenate two alternative decodes speculatively.
    return full


def extract_wav_head(source: Path, seconds: float) -> Path | None:
    """Create a short WAV containing the opening seconds for prefix recovery."""
    if source.suffix.lower() != ".wav" or seconds <= 0:
        return None
    handle, name = tempfile.mkstemp(prefix="channels-head-", suffix=".wav", dir=source.parent)
    os.close(handle)
    target_path = Path(name)
    try:
        with wave.open(str(source), "rb") as input_wav:
            params = input_wav.getparams()
            frame_count = min(input_wav.getnframes(), int(input_wav.getframerate() * seconds))
            frames = input_wav.readframes(frame_count)
        with wave.open(str(target_path), "wb") as output_wav:
            output_wav.setparams(params)
            output_wav.writeframes(frames)
        return target_path
    except Exception:
        target_path.unlink(missing_ok=True)
        return None


def assess_asr_text_quality(
    value: str | None,
    prompt_context: str | None = None,
    capture_seconds: float | None = None,
    speech_seconds: float | None = None,
    fallback_without_vad: bool = False,
) -> tuple[str | None, dict]:
    """Reject obvious Whisper hallucinations before they reach business data."""
    if os.environ.get('SOCIAL_VOLUME_ENABLED') == '1':
        return (str(value).strip() or None) if value is not None else None, {
            'accepted':None,'qualityChecked':False,
            'reason':'quality_check_disabled_by_volume_policy'}
    artifact = transcript_artifact(value)
    if artifact:
        return None, {"accepted": False, "reason": "known_whisper_hallucination",
                      "artifact_marker": artifact, "known_hallucination": True,
                      "repeat_ngram_ratio": 0.0}
    text = clean_asr_text(value)
    prompt_text = str(prompt_context or "")
    # Project names are often rare characters and Whisper returns a common
    # homophone.  Apply narrowly-scoped correction only when the current task
    # explicitly contains the canonical term.
    if text and "臻澐" in prompt_text:
        text = re.sub(r"(?:浸云|进云|近云|真云|臻云|臻雲)", "臻澐", text)
    if text and "华润臻澐" in prompt_text:
        text = re.sub(r"华人臻澐", "华润臻澐", text)
    if not text:
        return None, {"accepted": False, "reason": "empty", "repeat_ngram_ratio": 0.0}
    normalized = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]", "", text)
    prompt = re.sub(
        r"[^0-9A-Za-z\u4e00-\u9fff]", "",
        prompt_text[:64],
    )
    grams = [normalized[index:index + 4] for index in range(max(0, len(normalized) - 3))]
    repeat_ratio = 1.0 - (len(set(grams)) / len(grams)) if grams else 0.0
    capture = float(capture_seconds or 0)
    speech = float(speech_seconds or 0)
    speech_ratio = speech / capture if capture > 0 else None
    # A spoken sentence can legitimately repeat the caption/title.  Only
    # classify it as prompt echo when VAD found virtually no actual speech.
    prompt_echo = bool(
        len(prompt) >= 8
        and prompt in normalized
        and len(normalized) <= int(len(prompt) * 2.6)
        and (speech_seconds is None or speech < 0.8)
    )
    common_hallucination = normalized in {
        "字幕by索兰娅",
        "请不吝点赞订阅转发打赏支持明镜与点点栏目",
        "请点赞订阅转发打赏支持明镜与点点栏目",
        "请不吝点赞订阅转发打赏支持明镜与点点",
    } or any(marker in normalized.casefold() for marker in (
        "amaraorg",
        "mingpaocanada",
        "mingpaotoronto",
        "优优独播剧场",
        "yoyotelevisionseriesexclusive",
    ))
    sparse_long_audio = bool(
        capture >= 30
        and speech_ratio is not None
        and speech_ratio < 0.05
        and len(normalized) < max(12, int(capture * 0.15))
    )
    reason = None
    letters = [ch for ch in text if ch.isalpha()]
    unsupported_letters = [ch for ch in letters if not re.match(r"[A-Za-z\u3400-\u9fff]", ch)]
    if len(unsupported_letters) >= 4 and len(unsupported_letters) / max(1, len(letters)) > 0.25:
        reason = "unexpected_script_for_zh_decoder"
    elif prompt_echo:
        reason = "prompt_echo"
    elif len(normalized) >= 40 and repeat_ratio >= 0.45:
        reason = "looping_repetition"
    elif common_hallucination:
        reason = "known_whisper_hallucination"
    audit = {
        "accepted": reason is None,
        "reason": reason,
        "repeat_ngram_ratio": round(repeat_ratio, 4),
        "prompt_echo": prompt_echo,
        "known_hallucination": common_hallucination,
        "capture_seconds": capture_seconds,
        "speech_seconds": speech_seconds,
        "speech_ratio": round(speech_ratio, 4) if speech_ratio is not None else None,
        "fallback_without_vad": fallback_without_vad,
        "low_speech_density_warning": sparse_long_audio,
    }
    return (text if reason is None else None), audit


def require_trustworthy_asr_result(text, audits):
    """Rejected nonempty decoding is not evidence that the source is silent."""
    if os.environ.get('SOCIAL_VOLUME_ENABLED') == '1':
        return
    def evidence(values):
        for value in values:
            if not isinstance(value,dict):
                continue
            yield value
            yield from evidence(value.get("chunks") or [])
    rejected = any((a.get("quality") or {}).get("accepted") is False
                   and (a.get("quality") or {}).get("reason") not in (None, "empty")
                   for a in evidence(audits))
    if not text and rejected:
        raise AsrQualityRejectedError("ASR_QUALITY_REVIEW_REQUIRED：解码候选被质量门槛拒绝，不得标无内容")


def asr_output_needs_review(
    value: str | None,
    capture_seconds: float | None,
    speech_seconds: float | None = None,
    config: dict | None = None,
) -> bool:
    """Flag implausibly short transcripts without discarding partial text."""
    if os.environ.get('SOCIAL_VOLUME_ENABLED') == '1':
        return False
    text = asr_speech_compare_key(clean_asr_text(value) or "")
    capture = float(capture_seconds or 0)
    review_from = float((config or {}).get("client_asr_review_min_duration_seconds", 10.0))
    if capture < review_from:
        return False
    minimum_density = float((config or {}).get("client_asr_min_chars_per_second", 0.65))
    # Short videos previously bypassed completeness review entirely. A
    # 21-second capture containing only nine characters could therefore be
    # marked complete even when playback/recording stopped near the opening.
    minimum_chars = max(8, int(capture * minimum_density + 0.999))
    if len(text) < minimum_chars:
        return True
    if speech_seconds is not None:
        speech_ratio = float(speech_seconds or 0) / max(1.0, capture)
        if speech_ratio < 0.05 and len(text) < max(30, int(capture * 0.8)):
            return True
    return False


def extract_wav_slice(source: Path, start_seconds: float, length_seconds: float) -> Path | None:
    """Create a temporary WAV slice for sparse/quiet speech recovery."""
    if source.suffix.lower() != ".wav" or length_seconds <= 0:
        return None
    handle, name = tempfile.mkstemp(prefix="channels-chunk-", suffix=".wav", dir=source.parent)
    os.close(handle)
    target_path = Path(name)
    try:
        with wave.open(str(source), "rb") as input_wav:
            params = input_wav.getparams()
            rate = input_wav.getframerate()
            start_frame = min(input_wav.getnframes(), max(0, int(rate * start_seconds)))
            input_wav.setpos(start_frame)
            frames = input_wav.readframes(max(1, int(rate * length_seconds)))
        if not frames:
            target_path.unlink(missing_ok=True)
            return None
        with wave.open(str(target_path), "wb") as output_wav:
            output_wav.setparams(params)
            output_wav.writeframes(frames)
        return target_path
    except Exception:
        target_path.unlink(missing_ok=True)
        return None


def merge_asr_chunks(values: list[str]) -> str | None:
    """Join ordered ASR chunks while removing overlap at chunk boundaries."""
    merged = clean_asr_text(values[0]) if values else None
    for value in values[1:]:
        current = clean_asr_text(value)
        if not current:
            continue
        if not merged:
            merged = current
            continue
        if current in merged[-max(80, len(current) * 2):]:
            continue
        maximum = min(len(merged), len(current), 80)
        overlap = next((size for size in range(maximum, 5, -1) if merged[-size:] == current[:size]), 0)
        merged = clean_asr_text(merged + (current[overlap:] if overlap else "，" + current))
    return merged


def transcribe_sparse_audio_in_chunks(
    model, media_path: Path, base_options: dict, config: dict,
    prompt_context: str | None, capture_seconds: float,
) -> tuple[str | None, dict, object | None, list]:
    """Recover speech that whole-file VAD/no-VAD decoding truncated."""
    if media_path.suffix.lower() != ".wav":
        return None, {"pass": "chunked_recovery", "skipped": "not_wav"}, None, []
    chunk_seconds = max(10.0, float(config.get("client_asr_chunk_seconds", 30.0)))
    overlap_seconds = max(0.0, min(3.0, float(config.get("client_asr_chunk_overlap_seconds", 1.0))))
    step = max(5.0, chunk_seconds - overlap_seconds)
    texts: list[str] = []
    segment_list: list = []
    chunk_audits: list[dict] = []
    last_info = None
    offset = 0.0
    while offset < max(0.1, float(capture_seconds or 0)):
        chunk_path = extract_wav_slice(media_path, offset, min(chunk_seconds, capture_seconds - offset))
        if chunk_path is None:
            break
        try:
            options = dict(base_options)
            # Short/music-heavy chunks have unreliable language detection.
            # Mandarin and Cantonese both use Whisper's zh language token.
            options["language"] = "zh"
            options["vad_filter"] = True
            options["vad_parameters"] = {
                "threshold": float(config.get("client_asr_chunk_vad_threshold", 0.20)),
                "min_speech_duration_ms": 80,
                "min_silence_duration_ms": 500,
                "speech_pad_ms": int(config.get("client_asr_speech_pad_ms", 800)),
            }
            segments_iter, info = model.transcribe(str(chunk_path), **options)
            segments = list(segments_iter)
            raw = clean_asr_text("\n".join(
                (segment.text or "").strip() for segment in segments
                if (segment.text or "").strip()
            ))
            accepted, quality = assess_asr_text_quality(
                raw, prompt_context,
                capture_seconds=min(chunk_seconds, capture_seconds - offset),
                speech_seconds=getattr(info, "duration_after_vad", None),
            )
            fallback_used = False
            if not accepted:
                options["vad_filter"] = False
                options.pop("vad_parameters", None)
                segments_iter, info = model.transcribe(str(chunk_path), **options)
                segments = list(segments_iter)
                raw = clean_asr_text("\n".join(
                    (segment.text or "").strip() for segment in segments
                    if (segment.text or "").strip()
                ))
                accepted, quality = assess_asr_text_quality(
                    raw, prompt_context,
                    capture_seconds=min(chunk_seconds, capture_seconds - offset),
                    speech_seconds=getattr(info, "duration_after_vad", None),
                    fallback_without_vad=True,
                )
                fallback_used = True
            if accepted:
                texts.append(accepted)
            segment_list.extend(segments)
            last_info = info
            chunk_audits.append({
                "offset_seconds": round(offset, 2),
                "text_length": len(accepted or ""),
                "segments": len(segments),
                "fallback_without_vad": fallback_used,
                "quality": quality,
            })
        finally:
            chunk_path.unlink(missing_ok=True)
        offset += step
    merged = merge_asr_chunks(texts)
    accepted, quality = assess_asr_text_quality(
        merged, prompt_context, capture_seconds=capture_seconds,
        speech_seconds=None, fallback_without_vad=True,
    )
    audit = {
        "pass": "chunked_recovery",
        "chunk_seconds": chunk_seconds,
        "overlap_seconds": overlap_seconds,
        "chunks": chunk_audits,
        "raw_text_length": len(merged or ""),
        "quality": quality,
    }
    return accepted, audit, last_info, segment_list


def transcribe_audio_with_recovery(
    model, media_path: Path, config: dict,
    prompt_context: str | None, capture_seconds: float,
) -> tuple[str | None, object, list, bool, list[dict]]:
    """Decode full audio, then recover sparse results by no-VAD and chunks."""
    # Old persisted jobs may still contain business labels. Never feed them
    # into decoding, including recovery chunks and the head-prefix pass.
    prompt_context = None
    domain_prompt = None
    base_options = {
        "language": "zh",
        "beam_size": int(config.get("client_asr_beam_size", 8)),
        "temperature": 0.0,
        "vad_filter": True,
        "vad_parameters": {
            "threshold": float(config.get("client_asr_vad_threshold", 0.35)),
            "min_speech_duration_ms": 100,
            "min_silence_duration_ms": 300,
            "speech_pad_ms": int(config.get("client_asr_speech_pad_ms", 800)),
        },
        "condition_on_previous_text": False,
        "repetition_penalty": 1.05,
    }
    if domain_prompt:
        base_options["hotwords"] = domain_prompt
    pass_audits: list[dict] = []
    candidates: list[tuple[str, object, list, bool, dict]] = []
    for pass_index, use_vad in enumerate((True, False), start=1):
        options = dict(base_options)
        options["vad_filter"] = use_vad
        if not use_vad:
            options.pop("vad_parameters", None)
        segments, info = model.transcribe(str(media_path), **options)
        segment_list = list(segments)
        raw_text = clean_asr_text("\n".join(
            (segment.text or "").strip() for segment in segment_list
            if (segment.text or "").strip()
        ))
        accepted_text, quality = assess_asr_text_quality(
            raw_text, domain_prompt,
            capture_seconds=capture_seconds,
            speech_seconds=getattr(info, "duration_after_vad", None),
            fallback_without_vad=not use_vad,
        )
        pass_audits.append({
            "pass": pass_index,
            "vad_filter": use_vad,
            "raw_text_length": len(raw_text or ""),
            "segments": len(segment_list),
            "quality": quality,
        })
        if accepted_text:
            candidates.append((accepted_text, info, segment_list, not use_vad, quality))
            if not asr_output_needs_review(
                accepted_text, capture_seconds,
                getattr(info, "duration_after_vad", None), config,
            ):
                break
    if candidates:
        accepted_text, final_info, final_segments, accepted_fallback, _quality = max(
            candidates, key=lambda item: len(asr_speech_compare_key(item[0]))
        )
    else:
        accepted_text, final_info, final_segments, accepted_fallback = None, None, [], True

    if asr_output_needs_review(
        accepted_text, capture_seconds,
        getattr(final_info, "duration_after_vad", None) if final_info else None,
        config,
    ):
        chunk_text, chunk_audit, chunk_info, chunk_segments = transcribe_sparse_audio_in_chunks(
            model, media_path, base_options, config, prompt_context, capture_seconds,
        )
        pass_audits.append(chunk_audit)
        if chunk_text and len(asr_speech_compare_key(chunk_text)) > len(asr_speech_compare_key(accepted_text or "")):
            accepted_text = chunk_text
            final_info = chunk_info or final_info
            final_segments = chunk_segments or final_segments
            accepted_fallback = True
    if not accepted_text:
        return None, final_info, final_segments, True, pass_audits

    # Decode the opening separately. Full-length decoding can miss the first
    # one or two words even when the audio is present. The overlap merger only
    # prepends text when a reliable >=4-character suffix/prefix match exists.
    head_seconds = min(
        float(capture_seconds or 0),
        max(0.0, float(config.get("client_asr_head_seconds", 8.0))),
    )
    head_path = extract_wav_head(media_path, head_seconds)
    if head_path is not None:
        try:
            head_options = dict(base_options)
            head_options["vad_filter"] = False
            head_options.pop("vad_parameters", None)
            head_segments_iter, head_info = model.transcribe(str(head_path), **head_options)
            head_segments = list(head_segments_iter)
            head_raw = clean_asr_text("\n".join(
                (segment.text or "").strip() for segment in head_segments
                if (segment.text or "").strip()
            ))
            head_text, head_quality = assess_asr_text_quality(
                head_raw, domain_prompt,
                capture_seconds=head_seconds,
                speech_seconds=getattr(head_info, "duration_after_vad", None),
                fallback_without_vad=True,
            )
            merged = merge_asr_head_prefix(head_text, accepted_text)
            pass_audits.append({
                "pass": "head_prefix",
                "vad_filter": False,
                "head_seconds": head_seconds,
                "raw_text_length": len(head_raw or ""),
                "merged_prefix_chars": max(0, len(merged or "") - len(accepted_text or "")),
                "head_text_preview": (head_text or "")[:80],
                "full_text_preview_before_merge": (accepted_text or "")[:80],
                "merged_text_preview": (merged or "")[:120],
                "quality": head_quality,
            })
            accepted_text = merged
        finally:
            head_path.unlink(missing_ok=True)
    return accepted_text, final_info, final_segments, accepted_fallback, pass_audits


def wait_for_loopback_recorder_ready(
    process: subprocess.Popen, timeout_seconds: float,
) -> dict:
    """Wait for the recorder's JSON readiness line without blocking forever."""
    line_queue: queue.Queue = queue.Queue(maxsize=1)

    def read_line() -> None:
        try:
            line_queue.put(process.stdout.readline() if process.stdout else "")
        except Exception as exc:
            line_queue.put(json.dumps({"event": "reader_error", "error": str(exc)}))

    threading.Thread(target=read_line, daemon=True).start()
    try:
        line = line_queue.get(timeout=max(1.0, timeout_seconds))
    except queue.Empty as exc:
        raise RuntimeError("等待系统音频回采器就绪超时") from exc
    normalized_line = (line or "").strip()
    if not normalized_line:
        exit_code = process.poll()
        stderr = ""
        if exit_code is not None and process.stderr:
            try:
                stderr = (process.stderr.read() or "").strip()
            except Exception:
                stderr = ""
        detail = f"exitCode={exit_code}"
        if stderr:
            detail += f"; stderr={stderr[-2000:]}"
        raise RuntimeError(f"系统音频回采器未返回就绪信息：{detail}")
    try:
        payload = json.loads(normalized_line)
    except Exception as exc:
        raise RuntimeError(f"系统音频回采器未返回有效就绪信息：{line!r}") from exc
    if payload.get("event") != "loopback_recorder_ready" or not payload.get("ready"):
        raise RuntimeError(f"系统音频回采器初始化失败：{payload}")
    return payload


def truncate_wav_duration(path: Path, keep_seconds: float) -> float:
    """Trim recorder buffer so autoplay audio from the next video is excluded."""
    with wave.open(str(path), "rb") as source:
        params = source.getparams()
        total_frames = source.getnframes()
        keep_frames = min(total_frames, max(1, int(source.getframerate() * keep_seconds)))
        frames = source.readframes(keep_frames)
    with wave.open(str(path), "wb") as target:
        target.setparams(params)
        target.writeframes(frames)
    return keep_frames / max(1, params.framerate)


def ensure_wechat_audio_session_unmuted() -> dict:
    """Unmute Windows audio sessions owned by the WeChat web runtime."""
    audit = {"matched": 0, "sessions": [], "error": None}
    global _AUDIO_SESSION_IMPORT_WARNING_EMITTED
    try:
        from pycaw.pycaw import AudioUtilities

        for session in AudioUtilities.GetAllSessions():
            process = getattr(session, "Process", None)
            if process is None:
                continue
            try:
                process_name = str(process.name() or "")
            except Exception:
                continue
            if process_name.lower() not in {
                "weixin.exe", "wechat.exe", "wechatappex.exe",
            }:
                continue
            volume = session.SimpleAudioVolume
            before_mute = int(volume.GetMute())
            before_volume = float(volume.GetMasterVolume())
            volume.SetMute(0, None)
            if before_volume < 0.8:
                volume.SetMasterVolume(1.0, None)
            audit["matched"] += 1
            audit["sessions"].append({
                "pid": int(process.pid),
                "process": process_name,
                "beforeMute": before_mute,
                "beforeVolume": round(before_volume, 4),
                "afterMute": int(volume.GetMute()),
                "afterVolume": round(float(volume.GetMasterVolume()), 4),
            })
    except ImportError as exc:
        audit['error'] = str(exc)
        if not _AUDIO_SESSION_IMPORT_WARNING_EMITTED:
            print(json.dumps({'event': 'audio_session_control_unavailable',
                'errorCode': 'DEPENDENCY_MISSING', 'module': getattr(exc, 'name', 'pycaw'),
                'pythonExecutable': sys.executable, 'waveformVerificationRequired': True},
                ensure_ascii=False), flush=True)
            _AUDIO_SESSION_IMPORT_WARNING_EMITTED = True
        return audit
    except Exception as exc:
        audit["error"] = str(exc)
    print(json.dumps({
        "event": "wechat_audio_session_unmuted",
        **audit,
    }, ensure_ascii=False), flush=True)
    return audit


def select_wechat_audio_process_id(session_audit: dict | None) -> int | None:
    """Prefer the process that owns WeChat's live Windows audio session."""
    sessions = list((session_audit or {}).get("sessions") or [])
    preferred_names = ("wechatappex.exe", "weixin.exe", "wechat.exe")
    for preferred_name in preferred_names:
        for session in sessions:
            if str(session.get("process") or "").lower() == preferred_name:
                try:
                    pid = int(session.get("pid") or 0)
                except (TypeError, ValueError):
                    continue
                if pid > 0:
                    return pid
    return None


def video_surface_process_id(surface) -> int | None:
    """Return the WeChat renderer PID that owns the current video surface.

    A paused WeChat video may temporarily disappear from the Windows audio
    session list.  The Chromium surface still belongs to WeChatAppEx, and the
    process-loopback helper can capture that process tree once playback starts.
    """
    candidates = [surface]
    try:
        root = active_video_root(surface)
        if root is not surface:
            candidates.append(root)
    except Exception:
        pass
    for candidate in candidates:
        info = getattr(candidate, "element_info", None)
        value = getattr(info, "process_id", None)
        if value is None:
            try:
                value = candidate.process_id()
            except Exception:
                value = None
        try:
            process_id = int(value or 0)
        except (TypeError, ValueError):
            continue
        if process_id > 0:
            return process_id
    return None


def loopback_recorder_command(
    recorder: Path, duration: int, sample_rate: int, output: Path,
    session_audit: dict | None, preferred_process_id: int | None = None,
    prefer_explicit_process_id: bool = False,
) -> tuple[list[str], str, int | None]:
    """Use endpoint-independent process loopback whenever WeChat exposes a session."""
    process_id = None
    if prefer_explicit_process_id:
        try:
            process_id = int(preferred_process_id or 0) or None
        except (TypeError, ValueError):
            process_id = None
    if not process_id:
        process_id = select_wechat_audio_process_id(session_audit)
    if not process_id:
        try:
            process_id = int(preferred_process_id or 0) or None
        except (TypeError, ValueError):
            process_id = None
    command = [
        sys.executable, str(recorder), "--duration", str(duration),
        "--sample-rate", str(sample_rate), "--output", str(output),
    ]
    if process_id:
        command.extend(["--process-id", str(process_id)])
        return command, "windows_process_loopback", process_id
    return command, "windows_wasapi_loopback", None


def verify_client_audio_before_full_capture(
    surface, config: dict, temp_root: Path, sample_rate: int,
    probe_seconds: int, *, assume_audible_from_search_card: bool,
) -> dict:
    """Prove a live waveform before spending a full video duration recording."""
    recorder = Path(__file__).with_name("loopback-recorder.py")
    attempts = []
    for probe_attempt in (1, 2):
        ensure_capture_not_stopped(config)
        try:
            click_named_surface_control(active_video_root(surface), "暂停")
            seek_active_video_to_start(surface)
        except Exception:
            pass
        handle, probe_name = tempfile.mkstemp(
            prefix="channels-audio-probe-", suffix=".wav", dir=temp_root,
        )
        os.close(handle)
        probe_path = Path(probe_name)
        process = None
        try:
            recorder_env = os.environ.copy()
            recorder_env["PYTHONIOENCODING"] = "utf-8"
            recorder_env["PYTHONUTF8"] = "1"
            session_audit = ensure_wechat_audio_session_unmuted()
            recorder_command, recorder_source, recorder_process_id = loopback_recorder_command(
                recorder, probe_seconds, sample_rate, probe_path, session_audit,
                video_surface_process_id(surface),
            )
            process = subprocess.Popen(
                recorder_command,
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                encoding="utf-8", errors="replace", env=recorder_env,
            )
            recorder_ready = wait_for_loopback_recorder_ready(
                process,
                float(config.get("client_asr_recorder_ready_timeout_seconds", 20.0)),
            )
            unmuted, unmute_method = ensure_video_audio_unmuted(
                surface,
                assume_audible_from_search_card=(
                    assume_audible_from_search_card and probe_attempt == 1
                ),
                toggle_generic_volume=probe_attempt > 1,
            )
            playing, play_method = ensure_video_playing(surface)
            stdout, stderr = process.communicate(timeout=probe_seconds + 30)
            final_lines = (stdout or "").strip().splitlines()
            recorder_audit = {}
            if final_lines:
                try:
                    recorder_audit = json.loads(final_lines[-1])
                except Exception:
                    recorder_audit = {}
            attempt = {
                "attempt": probe_attempt,
                "ok": process.returncode == 0,
                "unmuted": unmuted,
                "unmuteMethod": unmute_method,
                "playing": playing,
                "playMethod": play_method,
                "session": session_audit,
                "source": recorder_source,
                "processId": recorder_process_id,
                "recorderReady": recorder_ready,
                "recorder": recorder_audit,
            }
            attempts.append(attempt)
            print(json.dumps({
                "event": "client_audio_preflight",
                **attempt,
            }, ensure_ascii=False), flush=True)
            if process.returncode == 0:
                return attempt
            error = (stderr or stdout or "音频预检未捕获到有效波形").strip()
            if probe_attempt == 2:
                raise ClientAudioCaptureError(error)
        finally:
            stop_owned_recorder(process)
            cleanup_audio_file(probe_path)
    raise ClientAudioCaptureError(
        "微信音频预检失败：" + json.dumps(attempts, ensure_ascii=False)
    )


VIDEO_PROGRESS_PATTERN = re.compile(
    r"(?:(\d+):)?(\d+):(\d{2})\s*/\s*"
    r"(?:(\d+):)?(\d+):(\d{2})"
)


def parse_video_progress_records(
    records: list[tuple[str, str, str]], expected_seconds: float | None = None,
) -> tuple[int, int] | None:
    """Return the active player's elapsed/total seconds from UIA text."""
    candidates: list[tuple[int, int]] = []
    for value, _control_type, _class_name in records:
        match = VIDEO_PROGRESS_PATTERN.fullmatch(str(value or "").strip())
        if not match:
            continue
        elapsed = (
            int(match.group(1) or 0) * 3600
            + int(match.group(2)) * 60
            + int(match.group(3))
        )
        total = (
            int(match.group(4) or 0) * 3600
            + int(match.group(5)) * 60
            + int(match.group(6))
        )
        if total > 0 and 0 <= elapsed <= total + 2:
            candidates.append((elapsed, total))
    if not candidates:
        return None
    if expected_seconds:
        return min(candidates, key=lambda item: abs(item[1] - float(expected_seconds)))
    return max(candidates, key=lambda item: item[1])


def read_active_video_progress(
    surface, expected_seconds: float | None = None,
) -> tuple[int, int] | None:
    """Best-effort playback clock read, including controls revealed on hover."""
    root = active_video_root(surface)
    progress = parse_video_progress_records(page_element_records(root), expected_seconds)
    if progress:
        return progress
    try:
        rect = root.rectangle()
        mouse.move(coords=(int((rect.left + rect.right) / 2), rect.bottom - 28))
        time.sleep(0.15)
        return parse_video_progress_records(
            page_element_records(active_video_root(surface)), expected_seconds,
        )
    except Exception:
        return None


def resolve_video_detail_duration(
    surface, records: list[tuple[str, str, str]],
    card_duration_seconds: int | None = None,
) -> int | None:
    """Prefer the active detail player's total over a card countdown."""
    progress = read_active_video_progress(surface)
    if progress and int(progress[1] or 0) > 0:
        return int(progress[1])
    parsed = parse_video_duration_seconds(
        " ".join(value for value, _control_type, _class_name in records)
    )
    return parsed or card_duration_seconds


def playback_capture_is_complete(
    audit: dict, expected_seconds: float, tolerance_seconds: float = 4.0,
) -> bool | None:
    """Return True/False when the UI clock proves coverage, else None."""
    max_elapsed = audit.get("maxElapsed")
    if max_elapsed is None:
        return None
    # Reaching the requested video's own end is conclusive. The feed may
    # auto-advance during the recorder's safety buffer, making the *last*
    # observed clock belong to a shorter next video (for example max=81/81,
    # then observedTotal=8). Do not discard an already complete capture or
    # replay the full video solely because of that later clock.
    if float(max_elapsed) >= max(
        0.0, float(expected_seconds or 0) - float(tolerance_seconds)
    ):
        return True
    observed_total = audit.get("observedTotal")
    if (
        observed_total is not None
        and float(observed_total) + float(tolerance_seconds)
        < float(expected_seconds or 0)
    ):
        # The active feed auto-advanced to a shorter neighbouring video, or
        # UIA attached to the wrong virtual player. Never let ``min()`` turn
        # that mismatch into an apparently complete recording.
        return False
    target = min(
        float(expected_seconds or 0),
        float(observed_total or expected_seconds or 0),
    )
    return float(max_elapsed) >= max(0.0, target - float(tolerance_seconds))


def recorder_audio_coverage_is_suspicious(audit: dict, expected_seconds: float) -> bool:
    """Detect a live player whose render endpoint went silent near the start."""
    expected = float(expected_seconds or 0)
    if expected < 30:
        return False
    active = audit.get("activeSecondCount")
    last_active = audit.get("lastActiveSecond")
    silent_tail = audit.get("silentTailSeconds")
    if active is None or last_active is None or silent_tail is None:
        return False
    return (
        int(active) < max(5, int(expected * 0.10))
        and float(last_active) < expected * 0.25
        and float(silent_tail) > expected * 0.50
    )


def ensure_capture_not_stopped(config):
    stop_path = config.get('_capture_stop_file')
    if STOP_REQUESTED or (stop_path and Path(stop_path).exists()):
        raise ClientAudioCaptureError('STOP_REQUESTED: audio capture/retry cancelled')


def monitor_video_playback_during_capture(
    process: subprocess.Popen, expected_seconds: float, config: dict,
) -> dict:
    """Keep the video moving while WASAPI records and audit timeline coverage."""
    interval = max(
        0.5, float(config.get("client_asr_playback_watchdog_interval_seconds", 2.0))
    )
    stall_seconds = max(
        interval * 2,
        float(config.get("client_asr_playback_stall_seconds", 6.0)),
    )
    tolerance = max(
        1.0,
        float(config.get("client_asr_playback_coverage_tolerance_seconds", 4.0)),
    )
    maximum_recoveries = max(
        0, int(config.get("client_asr_playback_max_recoveries", 8))
    )
    samples: list[dict] = []
    max_elapsed: int | None = None
    observed_total: int | None = None
    first_elapsed: int | None = None
    last_advance_at = time.monotonic()
    recoveries = 0
    identity = config.get('_capture_identity') or {}
    switched_after_end = False
    recovery_pending = None
    while process.poll() is None:
        ensure_capture_not_stopped(config)
        sample_time = time.monotonic()
        progress = None
        try:
            current_surface = get_video_detail_window()
            if identity and not switched_after_end:
                matched, _, _ = video_detail_identity_match(
                    identity.get('title'), identity.get('author_name'),
                    active_video_records(current_surface))
                if not matched:
                    if max_elapsed is None or max_elapsed < expected_seconds - tolerance:
                        raise VideoDetailIdentityMismatchError('Video identity changed before full playback coverage')
                    # Once the target ended, pause the recommendation while
                    # the native writer finalizes its safety buffer.
                    pause_active_video_if_running(current_surface)
                    switched_after_end = True
            if not switched_after_end:
                progress = read_active_video_progress(current_surface, expected_seconds=expected_seconds)
        except (ClientAudioCaptureError, VideoDetailIdentityMismatchError):
            raise
        except Exception:
            progress = None
        if progress:
            elapsed, total = progress
            if first_elapsed is None:
                first_elapsed = elapsed
            observed_total = total
            if max_elapsed is None or elapsed > max_elapsed:
                max_elapsed = elapsed
                last_advance_at = sample_time
            if recovery_pending is not None:
                advanced = elapsed > recovery_pending['elapsed']
                if advanced or sample_time - recovery_pending['at'] >= stall_seconds:
                    print(json.dumps({
                        'event': 'client_audio_playback_recovery_result',
                        'attempt': recovery_pending['attempt'], 'recovered': advanced,
                        'beforeElapsed': recovery_pending['elapsed'], 'afterElapsed': elapsed,
                        'verification': 'timeline_advanced' if advanced else 'timeline_stalled',
                    }, ensure_ascii=False), flush=True)
                    recovery_pending = None
                    if not advanced and recoveries >= maximum_recoveries:
                        raise ClientAudioCaptureError('Playback stalled after bounded recovery attempts')
            samples.append({"elapsed": elapsed, "total": total})
            target = min(float(expected_seconds), float(total))
            if (
                elapsed < max(0.0, target - tolerance)
                and sample_time - last_advance_at >= stall_seconds
                and recoveries < maximum_recoveries
            ):
                control_ready, method = ensure_video_playing(get_video_detail_window())
                recoveries += 1
                last_advance_at = time.monotonic()
                recovery_pending = {'elapsed': elapsed, 'at': last_advance_at, 'attempt': recoveries}
                print(json.dumps({
                    "event": "client_audio_playback_stall_recovery",
                    "elapsed": elapsed,
                    "total": total,
                    "attempt": recoveries,
                    "controlReady": control_ready,
                    "recovered": None,
                    "verification": "awaiting_timeline_advance",
                    "method": method,
                }, ensure_ascii=False), flush=True)
        time.sleep(interval)
    audit = {
        "firstElapsed": first_elapsed,
        "maxElapsed": max_elapsed,
        "observedTotal": observed_total,
        "progressSamples": samples,
        "stallRecoveries": recoveries,
        "switchedAfterTargetEnd": switched_after_end,
    }
    audit["coverageComplete"] = playback_capture_is_complete(
        audit, expected_seconds, tolerance,
    )
    return audit


def seek_active_video_to_start(surface) -> bool:
    """Verify the clock, including reloads already at zero without a slider."""
    for attempt in range(1, 4):
        current = surface if attempt == 1 else get_video_detail_window()
        progress = read_active_video_progress(current)
        if progress is not None and progress[0] <= 2:
            return True
        _seek_active_video_to_start_once(current)
        time.sleep(0.25)
        progress = read_active_video_progress(get_video_detail_window())
        verified = progress is not None and progress[0] <= 2
        print(json.dumps({'event': 'video_seek_start_verified', 'attempt': attempt,
            'verified': verified, 'elapsed': progress[0] if progress else None},
            ensure_ascii=False), flush=True)
        if verified:
            return True
    return False


def _seek_active_video_to_start_once(surface) -> bool:
    """Best-effort seek to 0 seconds before loopback recording starts."""
    root = active_video_root(surface)
    candidates = []
    for control_type in ("Slider", "ProgressBar"):
        try:
            candidates.extend(root.descendants(control_type=control_type))
        except Exception:
            continue
    for control in candidates:
        try:
            rect = control.rectangle()
            if not control.is_visible() or rect.width() < 30:
                continue
            try:
                range_value = control.iface_range_value
                range_value.SetValue(range_value.CurrentMinimum)
            except Exception:
                mouse.click(coords=(rect.left + 2, int((rect.top + rect.bottom) / 2)))
            return True
        except Exception:
            continue
    # Current WeChat Chromium exposes the player time and buttons through
    # UIA, but not the HTML progress range itself. Hover the control bar, then
    # click the verified left edge between the time label and 倍速 button.
    try:
        rect = root.rectangle()
        mouse.move(coords=(int((rect.left + rect.right) / 2), rect.bottom - 120))
        time.sleep(0.25)
        root = active_video_root(surface)
        time_labels = []
        for control in root.descendants(control_type="Text"):
            value = (control.window_text() or "").strip()
            match = re.fullmatch(r"(\d+):(\d{2})\s*/\s*(?:(\d+):)?(\d+):(\d{2})", value)
            if match:
                time_labels.append((control, match))
        speed_buttons = [
            control for control in root.descendants(control_type="Button")
            if (control.window_text() or "").strip() == "倍速"
        ]
        if time_labels and speed_buttons:
            label, _match = time_labels[0]
            label_rect = label.rectangle()
            speed_rect = speed_buttons[0].rectangle()
            # The gap grows with display scaling (100% / 150%). A fixed
            # 14 physical pixels landed BEFORE the track on 150% displays.
            progress_left = label_rect.right + max(14, round(label_rect.height() * 1.4))
            progress_y = int((label_rect.top + label_rect.bottom) / 2)
            if progress_left + 30 < speed_rect.left:
                mouse.click(coords=(progress_left + 2, progress_y))
                time.sleep(0.4)
                refreshed = page_element_records(active_video_root(surface))
                elapsed = next((
                    int(match.group(1)) * 60 + int(match.group(2))
                    for value, _type, _class in refreshed
                    if (match := re.fullmatch(
                        r"(\d+):(\d{2})\s*/\s*(?:(\d+):)?(\d+):(\d{2})",
                        value.strip(),
                    ))
                ), None)
                return elapsed is not None and elapsed <= 2
    except Exception:
        pass
    return False


CARD_TIME_PATTERN = re.compile(
    r"(?:刚刚|昨天|前天|\d+\s*(?:分钟|小时|天|周|个月|年)前|"
    r"20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}日?)"
)
CARD_TAG_PATTERN = re.compile(r"#([^#\s,，;；.!！?？、]{1,50})")


def parse_video_card_text(raw: str) -> dict:
    """Split a flattened 视频 result card into business fields."""
    value = re.sub(r"\s+", " ", str(raw or "")).strip()
    duration = parse_video_duration_seconds(value)
    value = re.sub(r"^\s*(?:\d+:)?\d{1,2}:\d{2}(?:\s+|$)", "", value, count=1)
    # Current WeChat builds do not always expose the “打开声音” label.  The
    # isolated numeric token immediately after duration is the card metric,
    # not the title; remove it in both UI variants. A title such as “143平” is
    # unaffected because the number is not an isolated whitespace token.
    value = re.sub(
        r"^\s*(?!20\d{2}\s)\d{1,6}(?:\.\d+)?\s*[万wW]?\s+(?:(?:打开声音|播放)\s*)?",
        "", value, count=1,
    )
    # Once sound was enabled on a previous visit the same card is flattened as
    # “关闭声音 ...”.  That state label is UI chrome, never article content.
    # Keeping it made tag-only cards acquire the false title “关闭声音” and
    # subsequently fail the detail identity check on every resume.
    value = re.sub(
        r"^\s*(?:打开声音|开启声音|关闭声音|静音|播放)\s*",
        "", value, count=1,
    )

    relative_time = None
    author_name = None
    tail = re.search(rf"\s+({CARD_TIME_PATTERN.pattern})\s*$", value)
    if tail:
        relative_time = tail.group(1)
        before_time = value[:tail.start()].rstrip()
        # The flattened card tail is always ``... caption author time``.  Do
        # not treat everything after the final hashtag as the author: current
        # WeChat cards can append normal caption text after the last hashtag,
        # for example ``#东港买房，必选❤️ 安鑫美玲 3分钟前``.  In that layout
        # only the final whitespace-delimited unit is the nickname; the
        # preceding ``，必选❤️`` still belongs to the title.
        match = re.search(r"(?:^|\s)(\S+)$", before_time)
        if match:
            author_name = match.group(1).strip()
            value = before_time[:match.start(1)].rstrip()
        else:
            value = before_time

    tags = [match.group(1).strip() for match in CARD_TAG_PATTERN.finditer(value)]
    tags = list(dict.fromkeys(tag for tag in tags if tag))
    title = CARD_TAG_PATTERN.sub("", value).replace("#", "")
    title = re.sub(r"\s+", " ", title)
    # Removing adjacent hashtags can leave a synthetic space immediately
    # before punctuation (``买听澜 ，必选``).  That space never belonged to
    # the caption and would make both export and detail identity checks wrong.
    title = re.sub(r"\s+([，,。.!！?？；;：:])", r"\1", title)
    title = title.strip(" ，,;；") or None
    return {
        "title": title,
        "tags": ",".join(tags) or None,
        "author_name": author_name,
        "relative_time": relative_time,
        "duration_seconds": duration,
    }


def capture_client_audio_asr(
    title: str, surface, config: dict, duration_seconds: int | None = None,
    prompt_context: str | None = None,
    media_cache_before: dict | None = None,
    during_capture_callback=None,
    stop_file=None,
) -> tuple[str | PendingClientAsr | None, int | None]:
    """Prefer a verified cache MP4; then capture the WeChat render process."""
    duration = duration_seconds or parse_video_duration_seconds(title)
    asr_enabled = config.get("enable_client_asr", True) and os.environ.get("WECHAT_CLIENT_ASR", "1") != "0"
    if not asr_enabled or not duration:
        return None, duration
    config = {**config, '_capture_stop_file': stop_file or config.get('_capture_stop_file')}
    ensure_capture_not_stopped(config)
    surface = prepare_video_detail(surface, config)
    # UIA is single-threaded: finish comments on a PAUSED target first. No
    # comment scrolling may block the recording watchdog or change its surface.
    if callable(during_capture_callback):
        if not pause_active_video_if_running(surface):
            raise ClientAudioCaptureError('Cannot confirm paused target before detail/comment collection')
        started = time.monotonic()
        during_capture_callback()
        ensure_capture_not_stopped(config)
        surface = get_video_detail_window()
        identity = config.get('_capture_identity') or {}
        if identity:
            matched, _, _ = video_detail_identity_match(
                identity.get('title'), identity.get('author_name'), active_video_records(surface))
            if not matched:
                raise VideoDetailIdentityMismatchError('Video identity changed during paused detail collection')
        print(json.dumps({'event': 'detail_fields_collected_before_audio_capture',
            'elapsed_seconds': round(time.monotonic() - started, 3),
            'recordingActive': False}, ensure_ascii=False), flush=True)
    max_seconds = max(1, int(config.get("client_asr_max_seconds", 720)))
    capture_seconds = min(duration, max_seconds)
    capture_buffer = max(0, int(config.get("client_asr_capture_buffer_seconds", 8)))
    recorder_seconds = capture_seconds + capture_buffer
    temp_root = Path(config.get("client_asr_temp_dir") or (Path(__file__).resolve().parents[1] / ".state" / "asr-temp"))
    temp_root.mkdir(parents=True, exist_ok=True)

    if config.get("enable_wechat_media_cache", True) and media_cache_before is not None:
        cache_dirs = discover_cache_dirs(config.get("wechat_media_cache_dirs"))
        try:
            cache_match = wait_for_unique_media(
                media_cache_before,
                duration,
                temp_root,
                cache_dirs=cache_dirs,
                timeout_seconds=float(config.get("wechat_media_cache_timeout_seconds", 6.0)),
                poll_seconds=float(config.get("wechat_media_cache_poll_seconds", 0.35)),
                stable_polls=int(config.get("wechat_media_cache_stable_polls", 2)),
                min_bytes=int(config.get("wechat_media_cache_min_bytes", 65536)),
                tolerance_seconds=float(config.get("wechat_media_cache_duration_tolerance_seconds", 2.0)),
                tolerance_ratio=float(config.get("wechat_media_cache_duration_tolerance_ratio", 0.06)),
            )
            cache_audit = asdict(cache_match)
            print(json.dumps({
                "event": "wechat_media_cache_probe",
                **cache_audit,
            }, ensure_ascii=False), flush=True)
            if cache_match.status == "matched" and cache_match.copied_media:
                media_path = Path(cache_match.copied_media)
                pending = PendingClientAsr(
                    wav_path=media_path,
                    title=title,
                    prompt_context=prompt_context,
                    duration=duration,
                    capture_seconds=capture_seconds,
                    recorder_seconds=0,
                    recorder_audit={
                        "cacheSource": cache_match.selected_source,
                        "cacheProbeSeconds": cache_match.elapsed_seconds,
                        "mediaDurationSeconds": cache_match.candidates[0].duration_seconds,
                        "mediaBytes": cache_match.candidates[0].size,
                    },
                    seeked=False,
                    reloaded_from_start=False,
                    source="wechat_media_cache",
                )
                if config.get("client_asr_background", False):
                    print(json.dumps({
                        "event": "client_audio_asr_queued",
                        "source": pending.source,
                        "duration_seconds": duration,
                        "cache_probe_seconds": cache_match.elapsed_seconds,
                        "media_path": str(media_path),
                    }, ensure_ascii=False), flush=True)
                    return pending, duration
                try:
                    return transcribe_pending_client_audio(pending, config), duration
                finally:
                    pending.wav_path.unlink(missing_ok=True)
        except Exception as exc:
            print(json.dumps({
                "event": "wechat_media_cache_failed",
                "duration_seconds": duration,
                "error": str(exc),
                "fallback": "windows_wasapi_loopback",
            }, ensure_ascii=False), flush=True)

    try:
        # Prevent losing the opening sentence: pause and seek first, start the
        # loopback recorder, then unmute/play while recording is already live.
        root = active_video_root(surface)
        paused = click_named_surface_control(root, "暂停")
        seeked = seek_active_video_to_start(surface)
        sample_rate = int(config.get("client_asr_sample_rate", 16000))
        preflight_seconds = max(
            0, int(config.get("client_asr_audio_preflight_seconds", 3))
        )
        preflight_audit = None
        if preflight_seconds:
            preflight_audit = verify_client_audio_before_full_capture(
                surface, config, temp_root, sample_rate, preflight_seconds,
                assume_audible_from_search_card=bool(
                    config.get("_search_card_audio_enabled", False)
                ),
            )
            # The short probe has consumed a few seconds. Always pause and seek
            # again so the full recording includes the opening sentence.
            click_named_surface_control(active_video_root(surface), "暂停")
            seeked = seek_active_video_to_start(surface)
        # The 3-second preflight has already proved which Windows process
        # actually carries this video's audio. Reusing that PID is materially
        # safer than asking UIA again: the detail surface may belong to an
        # AppEx renderer while the audible process-loopback tree is rooted at
        # Weixin.exe (or vice versa). Re-selecting here produced a valid
        # timeline but only 4 seconds of audio, then replayed the whole video.
        verified_audio_process_id = None
        if (
            preflight_audit
            and preflight_audit.get("ok")
            and preflight_audit.get("source") == "windows_process_loopback"
            and preflight_audit.get("processId")
        ):
            verified_audio_process_id = int(preflight_audit["processId"])
        handle, wav_name = tempfile.mkstemp(prefix="channels-", suffix=".wav", dir=temp_root)
        os.close(handle)
        wav_path = Path(wav_name)
        process = None
        keep_wav = False
        try:
            recorder = Path(__file__).with_name("loopback-recorder.py")
            silence_retries = max(0, int(config.get("client_asr_silence_retries", 1)))
            completed = None
            recorder_audit = {}
            reloaded_from_start = False
            unmuted, unmute_method = False, "not_attempted"
            playing, play_method = False, "not_attempted"
            # This state belongs to the current video, not to only the first
            # recording attempt.  A silence retry reloads the detail page but
            # must not blindly click its state-less 音量 button, which can mute
            # an already-audible player.  Pop once, then carry the evidence
            # through every bounded retry of this same video.
            audible_from_search_card = bool(
                config.pop("_search_card_audio_enabled", False)
            )
            recorder_source = "windows_wasapi_loopback"
            recorder_process_id = None
            force_endpoint_loopback = False
            for audio_attempt in range(1, silence_retries + 2):
                ensure_capture_not_stopped(config)
                # Reload BEFORE launching the full recorder. A reload destroys
                # the previous audibility proof; never blindly toggle 音量.
                reloaded_from_start = audio_attempt > 1 or not seeked
                if reloaded_from_start:
                    surface.set_focus()
                    keyboard.send_keys("^r", pause=0.02)
                    time.sleep(max(0.5, float(config.get("client_asr_reload_wait_seconds", 1.0))))
                    surface = get_video_detail_window()
                    surface = prepare_video_detail(surface, config)
                    preflight_audit = verify_client_audio_before_full_capture(
                        surface, config, temp_root, sample_rate,
                        max(3, preflight_seconds),
                        assume_audible_from_search_card=False,
                    )
                    verified_audio_process_id = (
                        int(preflight_audit['processId'])
                        if preflight_audit.get('source') == 'windows_process_loopback'
                        and preflight_audit.get('processId') else None
                    )
                    click_named_surface_control(active_video_root(surface), "暂停")
                    seeked = seek_active_video_to_start(surface)
                    if not seeked:
                        raise ClientAudioCaptureError('重载预检后无法回到视频开头，保留精准补采断点')
                recorder_env = os.environ.copy()
                recorder_env["PYTHONIOENCODING"] = "utf-8"
                recorder_env["PYTHONUTF8"] = "1"
                session_audit = ensure_wechat_audio_session_unmuted()
                recorder_command, recorder_source, recorder_process_id = loopback_recorder_command(
                    recorder, recorder_seconds, sample_rate, wav_path,
                    None if force_endpoint_loopback else session_audit,
                    None if force_endpoint_loopback else (
                        verified_audio_process_id or video_surface_process_id(surface)
                    ),
                    prefer_explicit_process_id=bool(
                        verified_audio_process_id and not force_endpoint_loopback
                    ),
                )
                recorder_command.extend([
                    '--initial-silence-seconds',
                    str(max(3, float(config.get('client_asr_initial_silence_seconds', 12)))),
                ])
                process = subprocess.Popen(
                    recorder_command,
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
                    encoding="utf-8", errors="replace",
                    env=recorder_env,
                )
                recorder_ready = wait_for_loopback_recorder_ready(
                    process,
                    float(config.get("client_asr_recorder_ready_timeout_seconds", 20.0)),
                )
                # A retry always reloads the detail page. Some WeChat builds
                # report the player as active after a previous replay while no
                # render endpoint receives audio; reload reliably recreates the
                # media session and returns it to zero.
                force_reload = reloaded_from_start
                ensure_wechat_audio_session_unmuted()
                unmuted, unmute_method = ensure_video_audio_unmuted(
                    surface,
                    assume_audible_from_search_card=(
                        audible_from_search_card or preflight_audit is not None
                    ),
                )
                # A short real waveform is stronger evidence than a UI action.
                if preflight_audit and preflight_audit.get('ok') and not unmuted:
                    unmuted, unmute_method = True, 'verified_preflight_waveform'
                playing, play_method = ensure_video_playing(surface)
                print(json.dumps({
                    "event": "client_audio_playback_ready",
                    "attempt": audio_attempt,
                    "unmuted": unmuted,
                    "unmute_method": unmute_method,
                    "playing": playing,
                    "play_method": play_method,
                    "forced_reload": force_reload,
                    "source": recorder_source,
                    "process_id": recorder_process_id,
                    "process_id_source": (
                        "all_output_endpoints"
                        if force_endpoint_loopback else
                        "verified_preflight"
                        if verified_audio_process_id else "video_surface"
                    ),
                    "recorder_ready": recorder_ready,
                    "audio_preflight": preflight_audit,
                }, ensure_ascii=False), flush=True)
                playback_audit = monitor_video_playback_during_capture(
                    process, capture_seconds, config,
                )
                stdout, stderr = process.communicate(timeout=30)
                completed = subprocess.CompletedProcess(
                    process.args, process.returncode, stdout, stderr
                )
                attempt_recorder_audit = {}
                try:
                    attempt_recorder_audit = json.loads(
                        (completed.stdout or "").strip().splitlines()[-1]
                    )
                except Exception:
                    pass
                attempt_recorder_audit.update(playback_audit)
                recorder_audit = attempt_recorder_audit
                if completed.returncode == 0:
                    suspicious_audio_coverage = recorder_audio_coverage_is_suspicious(
                        attempt_recorder_audit, capture_seconds,
                    )
                    if (
                        playback_audit.get("coverageComplete") is not False
                        and not suspicious_audio_coverage
                    ):
                        break
                    if suspicious_audio_coverage:
                        recorder_error = (
                            "微信视频音频时间覆盖不足："
                            f"active={attempt_recorder_audit.get('activeSecondCount')}s, "
                            f"lastActive={attempt_recorder_audit.get('lastActiveSecond')}s, "
                            f"silentTail={attempt_recorder_audit.get('silentTailSeconds')}s, "
                            f"expected={capture_seconds}s"
                        )
                        # A process-loopback tree can lose the render stream
                        # when WeChat migrates playback to another AppEx
                        # process. Reload first and retry the verified process
                        # tree by default; deployments may opt into all-output
                        # endpoint capture as the bounded fallback.
                        retry_process_loopback = bool(config.get(
                            "client_asr_retry_process_loopback_after_reload",
                                True,
                        ))
                        force_endpoint_loopback = (
                            recorder_source == "windows_process_loopback"
                            and not retry_process_loopback
                        )
                    else:
                        recorder_error = (
                            "微信视频播放时间轴覆盖不足："
                            f"max={playback_audit.get('maxElapsed')}s, "
                            f"expected={capture_seconds}s, "
                            f"recoveries={playback_audit.get('stallRecoveries')}"
                        )
                    if audio_attempt <= silence_retries:
                        print(json.dumps({
                            "event": "client_audio_incomplete_playback_retry",
                            "attempt": audio_attempt,
                            "next_attempt": audio_attempt + 1,
                            "playback": playback_audit,
                            "error": recorder_error,
                        }, ensure_ascii=False), flush=True)
                        continue
                    raise RuntimeError(recorder_error)
                recorder_error = (
                    completed.stderr or completed.stdout or "系统音频回采失败"
                ).strip()
                if audio_attempt <= silence_retries:
                    print(json.dumps({
                        "event": "client_audio_silence_retry",
                        "attempt": audio_attempt,
                        "next_attempt": audio_attempt + 1,
                        "error": recorder_error[-1000:],
                    }, ensure_ascii=False), flush=True)
                    continue
                raise RuntimeError(recorder_error)
            if completed is None or completed.returncode != 0:
                raise RuntimeError("系统音频回采失败")
            # The recorder starts before playback and intentionally has a
            # safety buffer.  Never pass the whole buffer to Whisper: short
            # videos auto-play the next feed item and contaminate the subtitle
            # field.  Retain only the expected video window plus startup slack.
            reload_wait = max(0.5, float(config.get("client_asr_reload_wait_seconds", 1.0)))
            startup_slack = reload_wait + 1.0 if reloaded_from_start else 1.0
            trimmed_seconds = truncate_wav_duration(
                wav_path, min(recorder_seconds, capture_seconds + startup_slack),
            )
            recorder_audit["trimmedSeconds"] = trimmed_seconds
            if config.get("client_asr_background", False):
                # Recording must happen while this exact WeChat video is open,
                # but Whisper inference does not. Keep the WAV and return the
                # job immediately so RPA can advance to the next video while a
                # background worker runs the configured ASR model.
                keep_wav = True
                pending = PendingClientAsr(
                    wav_path=wav_path,
                    title=title,
                    prompt_context=prompt_context,
                    duration=duration,
                    capture_seconds=capture_seconds,
                    recorder_seconds=recorder_seconds,
                    recorder_audit=recorder_audit,
                    seeked=seeked,
                    reloaded_from_start=reloaded_from_start,
                    source=recorder_source,
                )
                print(json.dumps({
                    "event": "client_audio_asr_queued",
                    "duration_seconds": duration,
                    "captured_seconds": capture_seconds,
                    "recorded_seconds": recorder_seconds,
                    "trimmed_seconds": trimmed_seconds,
                    "unmuted": unmuted,
                    "unmute_method": unmute_method,
                    "playing": playing,
                    "play_method": play_method,
                    "wav_path": str(wav_path),
                    "source": recorder_source,
                    "process_id": recorder_process_id,
                }, ensure_ascii=False), flush=True)
                return pending, duration
            global _ASR_MODEL
            if _ASR_MODEL is None:
                _ASR_MODEL = load_asr_model(config)
            text, info, segment_list, fallback_used, pass_audits = transcribe_audio_with_recovery(
                _ASR_MODEL, wav_path, config, prompt_context, capture_seconds,
            )
            try:
                require_trustworthy_asr_result(text, pass_audits)
            except Exception:
                keep_wav = True
                raise
            quality = next((
                item["quality"] for item in pass_audits
                if item.get("pass") in (1, 2) and item.get("quality", {}).get("accepted")
            ), pass_audits[-1]["quality"] if pass_audits else {})
            print(json.dumps({
                "event": "client_audio_asr",
                "duration_seconds": duration,
                "captured_seconds": capture_seconds,
                "recorded_seconds": recorder_seconds,
                "text_length": len(text or ""),
                "source": recorder_source,
                "model": config.get("client_asr_model", "small"),
                "segments": len(segment_list),
                "duration_after_vad": getattr(info, "duration_after_vad", None),
                "rewound_to_start": seeked,
                "reloaded_from_start": reloaded_from_start,
                "audio_peak": recorder_audit.get("peak"),
                "audio_rms": recorder_audit.get("rms"),
                "active_second_count": recorder_audit.get("activeSecondCount"),
                "last_active_second": recorder_audit.get("lastActiveSecond"),
                "silent_tail_seconds": recorder_audit.get("silentTailSeconds"),
                "playback_coverage_complete": recorder_audit.get("coverageComplete"),
                "playback_max_elapsed": recorder_audit.get("maxElapsed"),
                "playback_stall_recoveries": recorder_audit.get("stallRecoveries"),
                "unmuted": unmuted,
                "unmute_method": unmute_method,
                "playing": playing,
                "play_method": play_method,
                "fallback_without_vad": fallback_used,
                "quality": quality,
                "passes": pass_audits,
            }, ensure_ascii=False), flush=True)
            return text, duration
        finally:
            stop_owned_recorder(process)
            if not keep_wav:
                cleanup_audio_file(wav_path)
    except (VideoPlaybackUnavailableError, VideoDetailRecoveryBlockedError):
        raise
    except Exception as exc:
        print(json.dumps({
            "event": "client_audio_asr_failed",
            "duration_seconds": duration,
            "error": str(exc),
        }, ensure_ascii=False), flush=True)
        # A recorder/device/cache failure is not evidence that the video is
        # silent.  Returning None made the ASR backfill path persist
        # ``no_content`` and permanently XACK a video whose audio had never
        # been captured.  Raise a typed error so the item becomes retry_pending
        # (and eventually a visible dead letter) instead.
        raise ClientAudioCaptureError(str(exc)) from exc


def volume_export_local_target(local_posts, sql_posts, threshold):
    """An export trigger, never acceptance: add the SQL deficit to local rows.

    Hydrated/local counts can differ from MySQL. Reusing the absolute SQL P80
    as a local cap either exports without progress or needlessly overcollects.
    The parent must still read MySQL after import to confirm completion.
    """
    if (any(type(n) is not int or n < 0 for n in (local_posts, sql_posts, threshold))
            or threshold < 1):
        raise ValueError('VOLUME_EXPORT_COUNTS_INVALID')
    return local_posts + max(0, threshold-sql_posts)


def capture_primary_audio_optional(*args, stop_file=None, **kwargs):
    """Quantity mode must not discard verified metadata on a recorder failure.

    This waives audio only, not identity, access controls, lease/stop failures,
    or the caller's subsequent fresh detail-field validation.
    """
    try:
        text, duration = capture_client_audio_asr(*args, stop_file=stop_file, **kwargs)
        return text, duration, None
    except ClientAudioCaptureError as error:
        if (os.environ.get('SOCIAL_VOLUME_ENABLED') != '1' or STOP_REQUESTED
                or (stop_file is not None and stop_file.exists())):
            raise
        cause = error
        seen = set()
        while cause is not None and id(cause) not in seen:
            seen.add(id(cause))
            if (isinstance(cause, (VideoDetailIdentityMismatchError, SearchWindowUnavailableError, VideoPlaybackUnavailableError))
                    or re.search(r'LEASE_LOST|STOP_REQUESTED|STOP_SECURITY|STOP_LOGIN|验证码|安全验证|扫码|登录|访问频繁', str(cause), re.I)):
                raise
            cause = cause.__cause__ or cause.__context__
        print(json.dumps({'event':'volume_optional_audio_failed_metadata_retained',
            'error':str(error), 'qualityVerified':False,
            'audioCaptured':False, 'notEvidenceOfSilence':True},ensure_ascii=False),flush=True)
        return None, kwargs.get('duration_seconds'), str(error)


def transcribe_pending_client_audio(
    job: PendingClientAsr, config: dict, *, raise_errors: bool = False
) -> str | None:
    """Run Whisper after the WeChat RPA has already moved to another video."""
    try:
        persistent_host = os.environ.get("WECHAT_PERSISTENT_ASR_HOST", "").strip()
        persistent_port = os.environ.get("WECHAT_PERSISTENT_ASR_PORT", "").strip()
        persistent_token = os.environ.get("WECHAT_PERSISTENT_ASR_TOKEN", "")
        if persistent_host and persistent_port and persistent_token:
            request_id = uuid.uuid4().hex
            timeout_seconds = max(
                60.0,
                float(config.get("persistent_asr_request_timeout_seconds", 1800.0)),
            )
            request = {
                "type": "transcribe",
                "id": request_id,
                "token": persistent_token,
                "path": str(job.wav_path.resolve()),
                "prompt": job.prompt_context,
                "capture_seconds": job.capture_seconds,
            }
            print(json.dumps({
                "event": "persistent_asr_request_started",
                "id": request_id,
                "host": persistent_host,
                "port": int(persistent_port),
                "duration_seconds": job.duration,
            }, ensure_ascii=False), flush=True)
            try:
                with socket.create_connection(
                    (persistent_host, int(persistent_port)), timeout=15.0,
                ) as client:
                    client.settimeout(timeout_seconds)
                    stream = client.makefile("rwb")
                    stream.write(
                        (json.dumps(request, ensure_ascii=False) + "\n").encode("utf-8")
                    )
                    stream.flush()
                    raw_response = stream.readline(16 * 1024 * 1024)
                if not raw_response:
                    raise RuntimeError("persistent ASR service closed without a response")
                response = json.loads(raw_response.decode("utf-8"))
                if response.get("type") != "result" or response.get("id") != request_id:
                    raise RuntimeError(
                        str(response.get("error") or f"invalid persistent ASR response: {response}")
                    )
                passes = list(response.get("passes") or [])
                quality = next((
                    item.get("quality") for item in passes
                    if item.get("pass") in (1, 2)
                    and (item.get("quality") or {}).get("accepted")
                ), passes[-1].get("quality") if passes else {})
                text = clean_asr_text(response.get("text"))
                require_trustworthy_asr_result(text, passes)
                print(json.dumps({
                    "event": "persistent_asr_request_complete",
                    "id": request_id,
                    "service_pid": response.get("service_pid"),
                    "service_request_count": response.get("service_request_count"),
                    "model_reused": int(response.get("service_request_count") or 0) > 1,
                    "elapsed_seconds": response.get("elapsed_seconds"),
                    "duration_seconds": job.duration,
                    "text_length": len(text or ""),
                    "segments": response.get("segments"),
                    "duration_after_vad": response.get("duration_after_vad"),
                    "fallback_without_vad": response.get("fallback_without_vad"),
                    "quality": quality,
                }, ensure_ascii=False), flush=True)
                return text
            except AsrQualityRejectedError:
                raise
            except Exception as persistent_error:
                # The local in-process implementation remains a correctness
                # fallback. A service crash must slow one task down, never
                # lose its already-recorded audio or poison the Redis item.
                print(json.dumps({
                    "event": "persistent_asr_request_failed_fallback_local",
                    "id": request_id,
                    "error": str(persistent_error),
                }, ensure_ascii=False), flush=True)
        global _ASR_MODEL
        if _ASR_MODEL is None:
            _ASR_MODEL = load_asr_model(config)
        text, info, segment_list, fallback_used, pass_audits = transcribe_audio_with_recovery(
            _ASR_MODEL, job.wav_path, config, job.prompt_context, job.capture_seconds,
        )
        require_trustworthy_asr_result(text, pass_audits)
        quality = next((
            item["quality"] for item in pass_audits
            if item.get("pass") in (1, 2) and item.get("quality", {}).get("accepted")
        ), pass_audits[-1]["quality"] if pass_audits else {})
        print(json.dumps({
            "event": "client_audio_asr_background_complete",
            "source": job.source,
            "duration_seconds": job.duration,
            "text_length": len(text or ""),
            "segments": len(segment_list),
            "duration_after_vad": getattr(info, "duration_after_vad", None),
            "audio_peak": job.recorder_audit.get("peak"),
            "audio_rms": job.recorder_audit.get("rms"),
            "active_second_count": job.recorder_audit.get("activeSecondCount"),
            "last_active_second": job.recorder_audit.get("lastActiveSecond"),
            "silent_tail_seconds": job.recorder_audit.get("silentTailSeconds"),
            "playback_coverage_complete": job.recorder_audit.get("coverageComplete"),
            "playback_max_elapsed": job.recorder_audit.get("maxElapsed"),
            "playback_stall_recoveries": job.recorder_audit.get("stallRecoveries"),
            "trimmed_seconds": job.recorder_audit.get("trimmedSeconds"),
            "cache_probe_seconds": job.recorder_audit.get("cacheProbeSeconds"),
            "media_duration_seconds": job.recorder_audit.get("mediaDurationSeconds"),
            "fallback_without_vad": fallback_used,
            "quality": quality,
            "passes": pass_audits,
        }, ensure_ascii=False), flush=True)
        return text
    except Exception as exc:
        print(json.dumps({
            "event": "client_audio_asr_background_failed",
            "duration_seconds": job.duration,
            "error": str(exc),
        }, ensure_ascii=False), flush=True)
        if raise_errors:
            raise
        return None


class BackgroundAsrQueue:
    """Single-model ASR worker; SQLite updates are isolated from the UI thread."""

    def __init__(
        self,
        state_db: Path,
        config: dict,
        *,
        recovery_keyword: str | None = None,
        recovery_article_keys: set[str] | None = None,
    ):
        self.state_db = state_db
        self.config = dict(config)
        self.jobs: queue.Queue = queue.Queue()
        self.pending_keys: set[tuple[str, str]] = set()
        self.infrastructure_error: RedisCheckpointUnavailable | None = None
        self.thread = threading.Thread(
            target=self._run, name="channels-asr-background", daemon=True
        )
        self.thread.start()

        # Recover audio captured before a process/power interruption. Jobs are
        # durable in SQLite; only files that still exist are enqueued.
        recovery = sqlite3.connect(self.state_db, timeout=60)
        try:
            recovery_sql = (
                "SELECT article_key, keyword, payload_json FROM asr_job "
                "WHERE status IN ('pending','running','retry_pending')"
            )
            recovery_parameters: list[object] = []
            if recovery_keyword:
                recovery_sql += " AND keyword=?"
                recovery_parameters.append(recovery_keyword)
            # ``None`` means normal keyword recovery.  An empty set means that
            # an exact Redis retry has no durable article key yet, so it must
            # locate/capture that one card instead of draining unrelated WAVs.
            if recovery_article_keys is not None:
                normalized_keys = sorted(
                    str(value).strip() for value in recovery_article_keys
                    if str(value).strip()
                )
                if not normalized_keys:
                    recovery_sql += " AND 1=0"
                else:
                    recovery_sql += " AND article_key IN (" + ",".join(
                        "?" for _ in normalized_keys
                    ) + ")"
                    recovery_parameters.extend(normalized_keys)
            recovery_sql += " ORDER BY created_at"
            rows = recovery.execute(
                recovery_sql, tuple(recovery_parameters)
            ).fetchall()
            print(json.dumps({
                "event": "client_audio_asr_recovery_scope",
                "keyword": recovery_keyword,
                "article_keys": (
                    sorted(recovery_article_keys)
                    if recovery_article_keys is not None else None
                ),
                "recovered": len(rows),
            }, ensure_ascii=False), flush=True)
            for article_key, keyword, payload_json in rows:
                payload = json.loads(payload_json)
                wav_path = Path(payload["wav_path"])
                if not wav_path.exists():
                    recovery.execute(
                        "UPDATE asr_job SET status='missing_audio', last_error=?, updated_at=? "
                        "WHERE article_key=? AND keyword=?",
                        ("ASR audio file missing after recovery", datetime.now().isoformat(timespec="seconds"), article_key, keyword),
                    )
                    continue
                job = PendingClientAsr(
                    wav_path=wav_path,
                    title=payload.get("title") or "",
                    prompt_context=payload.get("prompt_context"),
                    duration=int(payload.get("duration") or 0),
                    capture_seconds=int(payload.get("capture_seconds") or 0),
                    recorder_seconds=int(payload.get("recorder_seconds") or 0),
                    recorder_audit=dict(payload.get("recorder_audit") or {}),
                    seeked=bool(payload.get("seeked")),
                    reloaded_from_start=bool(payload.get("reloaded_from_start")),
                    source=payload.get("source") or "windows_wasapi_loopback",
                )
                key = (str(article_key), str(keyword))
                if key not in self.pending_keys:
                    self.pending_keys.add(key)
                    self.jobs.put((str(article_key), str(keyword), job, payload))
            recovery.commit()
        finally:
            recovery.close()

    def submit(
        self,
        article_key: str,
        keyword: str,
        job: PendingClientAsr,
        *,
        fingerprint: str,
        project: str,
        remote_store: RedisVideoCheckpointStore | None = None,
    ) -> None:
        self.raise_if_infrastructure_failed()
        key = (article_key, keyword)
        if key in self.pending_keys:
            job.wav_path.unlink(missing_ok=True)
            return
        payload = asdict(job)
        payload["wav_path"] = str(job.wav_path)
        payload["fingerprint"] = fingerprint
        payload["project"] = project
        persistence = sqlite3.connect(self.state_db, timeout=60)
        try:
            now = datetime.now().isoformat(timespec="seconds")
            persistence.execute(
                """
                INSERT INTO asr_job(
                    article_key, keyword, wav_path, payload_json, status,
                    attempts, created_at, updated_at
                ) VALUES (?, ?, ?, ?, 'pending', 0, ?, ?)
                ON CONFLICT(article_key, keyword) DO UPDATE SET
                    wav_path=excluded.wav_path,
                    payload_json=excluded.payload_json,
                    status='pending',
                    last_error=NULL,
                    updated_at=excluded.updated_at
                """,
                (article_key, keyword, str(job.wav_path), json.dumps(payload, ensure_ascii=False), now, now),
            )
            persistence.commit()
        finally:
            persistence.close()
        self.pending_keys.add(key)
        self.jobs.put((article_key, keyword, job, payload))
        print(json.dumps({
            "event": "client_audio_asr_background_started",
            "article_key": article_key,
            "keyword": keyword,
            "queue_size": self.jobs.qsize(),
        }, ensure_ascii=False), flush=True)

    def _run(self) -> None:
        connection = sqlite3.connect(self.state_db, timeout=60)
        connection.execute("PRAGMA busy_timeout=60000")
        try:
            while True:
                item = self.jobs.get()
                try:
                    if item is None:
                        return
                    if self.infrastructure_error is not None:
                        # Keep durable pending jobs untouched after the fuse.
                        continue
                    article_key, keyword, job, payload = item
                    now = datetime.now().isoformat(timespec="seconds")
                    disposition = connection.execute(
                        "SELECT status FROM video_checkpoint WHERE keyword=? AND fingerprint=?",
                        (keyword, payload.get("fingerprint")),
                    ).fetchone()
                    if disposition and disposition[0] == "skipped_unresolved":
                        connection.execute(
                            "UPDATE asr_job SET status='skipped_unresolved',updated_at=? WHERE article_key=? AND keyword=?",
                            (now, article_key, keyword),
                        )
                        connection.commit()
                        continue
                    attempt_row = connection.execute(
                        "SELECT attempts FROM asr_job WHERE article_key=? AND keyword=?",
                        (article_key, keyword),
                    ).fetchone()
                    attempts = int(attempt_row[0] or 0) + 1 if attempt_row else 1
                    connection.execute(
                        "UPDATE asr_job SET status='running', attempts=?, updated_at=? "
                        "WHERE article_key=? AND keyword=?",
                        (attempts, now, article_key, keyword),
                    )
                    connection.commit()
                    text = transcribe_pending_client_audio(job, self.config, raise_errors=True)
                    needs_review = has_untrusted_transcript(text) or asr_output_needs_review(
                        text, job.capture_seconds, config=self.config,
                    )
                    asr_status = (
                        "partial_review" if text and needs_review
                        else "completed" if text
                        else "no_content"
                    )
                    asr_error = (
                        "ASR transcript is implausibly short for captured duration"
                        if needs_review else
                        None if text else
                        "No trustworthy speech recognized after VAD, no-VAD and chunked passes"
                    )
                    stage_candidate(connection, article_key, keyword, text,
                                    approved=bool(text) and not needs_review,
                                    reason=asr_error or 'quality gates passed; replacement audited')
                    if needs_review or not text:
                        # A short/empty decode is not proof of absent speech.
                        # Use the same bounded failure path and real counter
                        # as decoder-quality exceptions, retaining the audit.
                        raise AsrQualityRejectedError(
                            'ASR_QUALITY_REVIEW_REQUIRED：录音识别为空或不完整，不能作为已验收口播或无内容'
                        )
                    if text and not needs_review:
                        connection.execute(
                            "UPDATE captured_article SET video_content=?, "
                            "video_duration=COALESCE(video_duration, ?) "
                            "WHERE article_key=? AND keyword=?",
                            (text, job.duration, article_key, keyword),
                        )
                    connection.execute(
                        "UPDATE asr_job SET status=?, last_error=?, updated_at=? "
                        "WHERE article_key=? AND keyword=?",
                        (
                            asr_status,
                            asr_error,
                            datetime.now().isoformat(timespec="seconds"), article_key, keyword,
                        ),
                    )
                    checkpoint = connection.execute(
                        "SELECT project, retry_count, row_json, comments_json, fingerprint "
                        "FROM video_checkpoint WHERE keyword=? AND fingerprint=?",
                        (keyword, payload.get("fingerprint")),
                    ).fetchone()
                    connection.commit()
                    if checkpoint and checkpoint[2]:
                        row_data = json.loads(checkpoint[2])
                        if text and not needs_review:
                            row_data["video_content"] = text
                        row_data["video_duration"] = row_data.get("video_duration") or job.duration
                        allowed = ResultRow.__dataclass_fields__.keys()
                        row = ResultRow(**{key: value for key, value in row_data.items() if key in allowed})
                        comments = json.loads(checkpoint[3]) if checkpoint[3] else []
                        remote_store = RedisVideoCheckpointStore()
                        checkpoint_payload = upsert_video_checkpoint(
                            connection,
                            keyword=keyword,
                            project=checkpoint[0] or payload.get("project") or "",
                            fingerprint=checkpoint[4],
                            status="completed" if text and not needs_review else "retry_pending",
                            retry_count=max(attempts, int(checkpoint[1] or 0)),
                            row=row,
                            comments=comments,
                            asr_status=asr_status,
                            last_error=asr_error,
                            remote_store=remote_store,
                        )
                    if text and not needs_review:
                        job.wav_path.unlink(missing_ok=True)
                except RedisCheckpointUnavailable as exc:
                    # Decoding/local commit may already have succeeded. A
                    # transport outage is not another ASR/content attempt.
                    self.pause_for_infrastructure(exc)
                except Exception as exc:
                    if item is not None:
                        article_key, keyword, job, payload = item
                        maximum = max(1, int(self.config.get("asr_job_max_attempts", 3)))
                        attempt_row = connection.execute(
                            "SELECT attempts FROM asr_job WHERE article_key=? AND keyword=?",
                            (article_key, keyword),
                        ).fetchone()
                        attempts = int(attempt_row[0] or 0) if attempt_row else 1
                        asr_job_status = "dead_letter" if attempts >= maximum else "retry_pending"
                        connection.execute(
                            "UPDATE asr_job SET status=?, last_error=?, updated_at=? "
                            "WHERE article_key=? AND keyword=?",
                            (asr_job_status, str(exc), datetime.now().isoformat(timespec="seconds"), article_key, keyword),
                        )
                        connection.execute(
                            "UPDATE video_checkpoint SET status=?, asr_status=?, last_error=?, updated_at=?, "
                            "retry_count=MAX(COALESCE(retry_count,0),?) "
                            "WHERE keyword=? AND fingerprint=? AND status!='skipped_unresolved'",
                            (
                                "dead_letter" if asr_job_status == "dead_letter" else "retry_pending",
                                asr_job_status, str(exc), datetime.now().isoformat(timespec="seconds"),
                                attempts, keyword, payload.get("fingerprint"),
                            ),
                        )
                        connection.commit()
                        checkpoint = connection.execute(
                            "SELECT row_json, comments_json, retry_count FROM video_checkpoint "
                            "WHERE keyword=? AND fingerprint=?",
                            (keyword, payload.get("fingerprint")),
                        ).fetchone()
                        row = None
                        comments = []
                        if checkpoint and checkpoint[0]:
                            row_data = json.loads(checkpoint[0])
                            allowed = ResultRow.__dataclass_fields__.keys()
                            row = ResultRow(**{key: value for key, value in row_data.items() if key in allowed})
                            comments = json.loads(checkpoint[1]) if checkpoint[1] else []
                        try:
                            checkpoint_payload = upsert_video_checkpoint(
                                connection,
                                keyword=keyword,
                                project=payload.get("project") or "",
                                fingerprint=payload.get("fingerprint") or article_key,
                                status="dead_letter" if asr_job_status == "dead_letter" else "retry_pending",
                                retry_count=max(attempts, int(checkpoint[2] or 0) if checkpoint else 0),
                                row=row,
                                comments=comments,
                                asr_status=asr_job_status,
                                last_error=str(exc),
                                remote_store=RedisVideoCheckpointStore(),
                            )
                        except RedisCheckpointUnavailable as network_error:
                            # A failure while saving failure evidence must not
                            # kill the queue thread and hang jobs.join().
                            self.pause_for_infrastructure(network_error)
                            continue
                        if asr_job_status == "dead_letter":
                            dead_letter_time = datetime.now().isoformat(timespec="seconds")
                            connection.execute(
                                """
                                INSERT INTO video_dead_letter(
                                    keyword, fingerprint, project, result_index, title,
                                    author_name, raw_card_text, retry_count, last_error,
                                    screenshot_path, created_at, updated_at
                                ) VALUES (?, ?, ?, NULL, ?, ?, ?, ?, ?, NULL, ?, ?)
                                ON CONFLICT(keyword, fingerprint) DO UPDATE SET
                                    title=excluded.title,
                                    author_name=excluded.author_name,
                                    raw_card_text=excluded.raw_card_text,
                                    retry_count=excluded.retry_count,
                                    last_error=excluded.last_error,
                                    updated_at=excluded.updated_at
                                """,
                                (
                                    keyword, payload.get("fingerprint") or article_key,
                                    payload.get("project") or "",
                                    row.title if row else payload.get("title"),
                                    row.author_name if row else None,
                                    "ASR persistent job",
                                    attempts, str(exc), dead_letter_time, dead_letter_time,
                                ),
                            )
                            connection.commit()
                            checkpoint_payload.update({
                                "title": row.title if row else payload.get("title"),
                                "author_name": row.author_name if row else None,
                                "raw_card_text": "ASR persistent job",
                                "result_index": None,
                            })
                            try:
                                RedisVideoCheckpointStore().put(
                                    payload.get("fingerprint") or article_key,
                                    checkpoint_payload,
                                )
                            except RedisCheckpointUnavailable as network_error:
                                self.pause_for_infrastructure(network_error)
                                continue
                        print(json.dumps({
                            "event": "client_audio_asr_job_persisted_failure",
                            "article_key": article_key,
                            "keyword": keyword,
                            "status": asr_job_status,
                            "attempts": attempts,
                            "error": str(exc),
                        }, ensure_ascii=False), flush=True)
                finally:
                    if item is not None:
                        self.pending_keys.discard((item[0], item[1]))
                    self.jobs.task_done()
        finally:
            connection.close()

    def pause_for_infrastructure(self, error: RedisCheckpointUnavailable) -> None:
        self.infrastructure_error = error
        print(json.dumps({'event': 'client_audio_asr_remote_sync_paused',
            'localCheckpointRetained': True, 'remoteSyncConfirmed': False,
            'asrRetryCharged': False, 'exitCode': 75}), flush=True)

    def raise_if_infrastructure_failed(self) -> None:
        if self.infrastructure_error is not None:
            raise self.infrastructure_error

    def close_and_wait(self) -> None:
        pending = self.jobs.qsize()
        print(json.dumps({
            "event": "client_audio_asr_background_drain",
            "pending": pending,
        }, ensure_ascii=False), flush=True)
        self.jobs.join()
        self.jobs.put(None)
        self.thread.join()
        self.raise_if_infrastructure_failed()


def _wechat_main_window_candidates():
    # The PC WeChat main window is Qt-based. The separate 搜一搜 article
    # window has the same title but is Chromium-based, so do not count it as
    # a second main window.
    candidates = []
    for candidate in Desktop(backend="uia").windows():
        try:
            class_name = candidate.element_info.class_name or ""
            rect = candidate.rectangle()
            # Legacy accounts use a Qt51514... top-level window; newer account
            # rollouts use mmui::MainWindow.  Do not depend on the title: some
            # Windows locale combinations expose 微信 as mojibake through UIA.
            ui_family = class_name.startswith("Qt51514") or class_name.startswith("mmui::MainWindow")
            main_sized = rect.width() >= 600 and rect.height() >= 500
            if ui_family and candidate.is_visible() and main_sized:
                candidates.append(candidate)
        except Exception:
            continue
    return candidates


def _wechat_hidden_main_candidates():
    candidates = []
    for candidate in Desktop(backend="uia").windows(visible_only=False):
        try:
            class_name = candidate.element_info.class_name or ""
            if class_name.startswith("Qt51514") or class_name.startswith("mmui::MainWindow"):
                if not class_name.startswith("mmui::LoginWindow"):
                    candidates.append(candidate)
        except Exception:
            continue
    return candidates


def _wechat_login_windows():
    matches = []
    for candidate in Desktop(backend="uia").windows():
        try:
            if (candidate.element_info.class_name or "").startswith("mmui::LoginWindow"):
                matches.append(candidate)
        except Exception:
            continue
    return matches


def wake_wechat_main_window() -> bool:
    """Wake WeChat when all background processes remain but its UI is hidden.

    Launching the already-running executable is WeChat's supported single-
    instance activation path and restores a tray-hidden main window.  The
    executable is supplied by the environment probe, with standard install
    locations only as portable fallbacks.
    """
    # A visible login window is an authentication boundary, not a tray-hidden
    # main window. Launching Weixin.exe again here creates another login
    # window on current 4.1 clients and eventually leaves several overlapping
    # prompts. Stop before any launch and let the operator finish login.
    login_windows = _wechat_login_windows()
    if login_windows:
        report_from_environment('STOP_LOGIN_REQUIRED：请人工登录微信')
        print(json.dumps({
            "event": "wechat_window_login_required",
            "login_windows": len(login_windows),
        }, ensure_ascii=False), flush=True)
        return False

    # First restore an existing tray-hidden instance.  Some WeChat builds do
    # not react to starting Weixin.exe again, but pywinauto restore does.
    hidden = _wechat_hidden_main_candidates()
    for candidate in hidden:
        try:
            candidate.restore()
            candidate.set_focus()
            time.sleep(1.0)
            if _wechat_main_window_candidates():
                print(json.dumps({
                    "event": "wechat_window_recovered",
                    "method": "restore_existing_window",
                }, ensure_ascii=False), flush=True)
                return True
        except Exception:
            continue

    candidates = [
        os.environ.get("WECHAT_CLIENT_PATH", ""),
        os.path.join(os.environ.get("ProgramFiles", ""), "Tencent", "Weixin", "Weixin.exe"),
        os.path.join(os.environ.get("ProgramFiles(x86)", ""), "Tencent", "WeChat", "WeChat.exe"),
    ]
    executable = next((path for path in candidates if path and Path(path).is_file()), None)
    if not executable:
        print(json.dumps({
            "event": "wechat_window_recovery_failed",
            "reason": "client executable path not found",
        }, ensure_ascii=False), flush=True)
        return False
    print(json.dumps({
        "event": "wechat_window_recovery_started",
        "executable": executable,
    }, ensure_ascii=False), flush=True)
    try:
        subprocess.Popen(
            [executable],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            close_fds=True,
        )
    except Exception as exc:
        print(json.dumps({
            "event": "wechat_window_recovery_failed",
            "reason": str(exc),
        }, ensure_ascii=False), flush=True)
        return False
    deadline = time.time() + 10.0
    while time.time() < deadline:
        time.sleep(0.5)
        if _wechat_main_window_candidates():
            print(json.dumps({
                "event": "wechat_window_recovered",
                "method": "single_instance_launch",
            }, ensure_ascii=False), flush=True)
            return True
    print(json.dumps({
        "event": "wechat_window_recovery_failed",
        "reason": "main window did not appear within 10 seconds",
    }, ensure_ascii=False), flush=True)
    return False


def get_wechat_window():
    candidates = _wechat_main_window_candidates()
    if not candidates:
        wake_wechat_main_window()
        candidates = _wechat_main_window_candidates()
    if not candidates and _wechat_login_windows():
        raise RuntimeError("微信客户端已打开，但仍停留在‘需在手机上完成登录’界面；请先完成登录")
    if len(candidates) == 1:
        try:
            rect = candidates[0].rectangle()
            if rect.width() <= 10 or rect.height() <= 10:
                candidates[0].restore()
                time.sleep(0.5)
        except Exception:
            pass
    matches = []
    for candidate in candidates:
        try:
            rect = candidate.rectangle()
            if rect.width() > 10 and rect.height() > 10:
                matches.append(candidate)
        except Exception:
            continue
    if len(matches) != 1:
        raise RuntimeError(f"需要且只能有一个可见标题为‘{WINDOW_TITLE}’的窗口，当前找到 {len(matches)} 个")
    win = matches[0]
    rect = win.rectangle()
    if rect.width() <= 10 or rect.height() <= 10:
        win.restore()
        time.sleep(0.5)
    try:
        win.maximize()
        time.sleep(0.4)
    except Exception:
        pass
    win.set_focus()
    return win


def web_windows():
    matches = []
    for win in Desktop(backend="uia").windows():
        try:
            title = (win.window_text() or "").strip()
            class_name = win.element_info.class_name or ""
            if class_name == "Chrome_WidgetWin_0" and (
                title == WINDOW_TITLE or "搜一搜" in title or "视频号" in title
            ):
                rect = win.rectangle()
                if rect.width() > 100 and rect.height() > 100:
                    matches.append(win)
        except Exception:
            continue
    return matches


def detached_web_window_role(win) -> str:
    """Classify a detached WeChat Chromium window without changing focus.

    Current WeChat builds keep the 视频 search results and the opened video
    player in two independent ``Chrome_WidgetWin_0`` top-level windows.  Both
    windows are titled ``微信``, so counting/title matching alone cannot tell
    them apart.  The RootWebArea document, however, remains semantic: search
    pages contain ``搜一搜`` while the player document is named ``视频号``.
    """
    document_names: list[str] = []
    try:
        for document in win.descendants(control_type="Document"):
            try:
                name = (document.window_text() or "").strip()
                if name:
                    document_names.append(name)
            except Exception:
                continue
    except Exception:
        pass
    if any("搜一搜" in name for name in document_names):
        return "search_results"
    if any(name == "视频号" for name in document_names):
        return "video_detail"
    try:
        title = (win.window_text() or "").strip()
    except Exception:
        title = ""
    if "搜一搜" in title:
        return "search_results"
    if "视频号" in title:
        return "video_detail"
    return "unknown"


def detached_window_handle(win) -> int | None:
    try:
        return int(win.handle)
    except Exception:
        try:
            return int(win.element_info.handle)
        except Exception:
            return None


def select_detached_window_by_role(windows, expected_role: str):
    """Select one detached surface by semantic role, never by foreground.

    The search result and video player are independent top-level windows in
    current WeChat builds.  Foreground selection made callers that needed the
    result list repeatedly receive the player, producing an unbounded wait.
    """
    if expected_role not in {"search_results", "video_detail"}:
        raise ValueError(f"unsupported detached window role: {expected_role}")
    roles = [(detached_web_window_role(win), win) for win in windows]
    matches = [win for role, win in roles if role == expected_role]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise RuntimeError(
            f"检测到多个{expected_role}微信网页窗口："
            + json.dumps([detached_window_handle(win) for win in matches], ensure_ascii=False)
        )
    # UIA can briefly report an empty RootWebArea during Chromium navigation.
    # A lone unknown window is a safe bounded fallback; with multiple windows
    # we must not guess because that recreates the search/detail role inversion.
    unknown = [win for role, win in roles if role == "unknown"]
    if len(windows) == 1 and len(unknown) == 1:
        return unknown[0]
    opposite_role = "video_detail" if expected_role == "search_results" else "search_results"
    opposite = [win for role, win in roles if role == opposite_role]
    if len(windows) == 2 and len(unknown) == 1 and len(opposite) == 1:
        return unknown[0]
    return None


def _log_multi_window_resolution(windows, selected, expected_role: str) -> None:
    """Log a changed window topology once instead of on every polling tick."""
    global LAST_MULTI_WINDOW_SIGNATURE
    if len(windows) <= 1:
        LAST_MULTI_WINDOW_SIGNATURE = None
        return
    role_handles = tuple(sorted([
        (
            detached_window_handle(win),
            detached_web_window_role(win),
        )
        for win in windows
    ], key=lambda item: (item[0] is None, item[0] or 0, item[1])))
    signature = (
        expected_role,
        role_handles,
        detached_window_handle(selected) if selected is not None else None,
    )
    if signature == LAST_MULTI_WINDOW_SIGNATURE:
        return
    LAST_MULTI_WINDOW_SIGNATURE = signature
    print(json.dumps({
        "event": "multiple_wechat_web_windows_resolved",
        "count": len(windows),
        "roles": [role for _handle, role in role_handles],
        "expected_role": expected_role,
        "selected_role": detached_web_window_role(selected) if selected is not None else None,
        "selected_handle": detached_window_handle(selected) if selected is not None else None,
    }, ensure_ascii=False), flush=True)


def close_detached_window(win) -> bool:
    """Close one detached Chromium window, never the WeChat Qt main window."""
    try:
        win.close()
        return True
    except Exception:
        pass
    try:
        win.set_focus()
        keyboard.send_keys("%{F4}", pause=0.02)
        return True
    except Exception:
        return False


def cleanup_stale_video_windows(*, reason: str, require_search_peer: bool = False) -> int:
    """Close detached video-player windows left behind by completed items.

    ``require_search_peer`` protects same-window navigation: during normal item
    cleanup we only close a player when a separate result window is visible.
    At the beginning of a new keyword, a standalone stale player is also safe to
    close because the search surface will be rebuilt from the WeChat main UI.
    """
    windows = web_windows()
    roles = [(detached_web_window_role(win), win) for win in windows]
    has_search = any(role == "search_results" for role, _win in roles)
    # During Chromium hydration the search RootWebArea can briefly be unknown
    # while the sibling video-detail role is still unambiguous.
    if not has_search and len(roles) == 2:
        role_names = [role for role, _win in roles]
        has_search = role_names.count("video_detail") == 1 and role_names.count("unknown") == 1
    if require_search_peer and not has_search:
        return 0
    details = [win for role, win in roles if role == "video_detail"]
    target_handles = {
        detached_window_handle(win) for win in details
        if detached_window_handle(win) is not None
    }
    remaining_handles = set(target_handles)
    for _attempt in range(1, 3):
        if not remaining_handles:
            break
        for win in details:
            if detached_window_handle(win) in remaining_handles:
                close_detached_window(win)
        deadline = time.time() + 3.0
        while time.time() < deadline:
            live_handles = {detached_window_handle(win) for win in web_windows()}
            remaining_handles = target_handles & live_handles
            if not remaining_handles:
                break
            time.sleep(0.1)
    closed = len(target_handles - remaining_handles)
    if target_handles:
        print(json.dumps({
            "event": "stale_video_windows_closed",
            "reason": reason,
            "closed": closed,
            "remaining": len(web_windows()),
        }, ensure_ascii=False), flush=True)
    if remaining_handles:
        print(json.dumps({
            "event": "stale_video_window_close_failed",
            "reason": reason,
            "handles": sorted(remaining_handles),
            "attempts": 2,
        }, ensure_ascii=False), flush=True)
        raise SearchWindowUnavailableError(
            "视频详情窗口连续2次关闭失败："
            + json.dumps(sorted(remaining_handles), ensure_ascii=False)
        )
    return closed


def get_web_window():
    matches = web_windows()
    if len(matches) != 1:
        raise RuntimeError(f"需要且只能有一个微信搜一搜网页窗口，当前找到 {len(matches)} 个")
    win = matches[0]
    try:
        win.maximize()
        time.sleep(0.4)
    except Exception:
        pass
    win.set_focus()
    return win


def _is_embedded_search_surface(win) -> bool:
    """Recognize the account-specific 搜一搜 page embedded in the Qt main window."""
    names = set()
    has_search_action = False
    has_search_edit = False
    has_search_document = False
    try:
        for control in win.descendants():
            try:
                name = (control.window_text() or "").strip()
                if name:
                    names.add(name)
                if (control.element_info.automation_id or "") == "weixin-search-input":
                    has_search_edit = True
                if (
                    control.element_info.control_type == "Document"
                    and (control.element_info.automation_id or "") == "RootWebArea"
                    and "搜一搜" in name
                ):
                    has_search_document = True
                if name == "搜索" and control.element_info.control_type in {"Button", "Text", "Hyperlink"}:
                    has_search_action = True
            except Exception:
                continue
    except Exception:
        return False
    # Requiring category labels avoids confusing the global chat-list search
    # box with the full embedded 搜一搜 page.
    return (
        has_search_document and has_search_edit and has_search_action
    ) or (
        has_search_action and "视频" in names and "公众号" in names
        and ("搜一搜" in names or "全部" in names)
    )


def _prepare_selected_window(win, *, focus: bool) -> None:
    if not focus:
        return
    try:
        win.maximize()
        time.sleep(0.2)
    except Exception:
        pass
    win.set_focus()


def _try_get_main_surface_without_focus():
    """Inspect the visible WeChat main surface without restoring or focusing it."""
    visible = []
    for candidate in _wechat_main_window_candidates():
        try:
            rect = candidate.rectangle()
            if rect.width() > 10 and rect.height() > 10:
                visible.append(candidate)
        except Exception:
            continue
    return visible[0] if len(visible) == 1 else None


def try_get_search_results_window(*, focus: bool = True):
    """Return only the search-results surface, never a separate video player."""
    global SEARCH_LAYOUT
    detached = web_windows()
    if detached:
        win = select_detached_window_by_role(detached, "search_results")
        _log_multi_window_resolution(detached, win, "search_results")
        if win is not None:
            # A detached video player is not proof that the search surface is
            # detached.  Updating the global layout before role selection made
            # every polling tick oscillate detached -> embedded while a player
            # was open, destabilising both cleanup and focus handling.
            if SEARCH_LAYOUT != "detached":
                SEARCH_LAYOUT = "detached"
                print(json.dumps({"event": "wechat_search_layout", "layout": SEARCH_LAYOUT}, ensure_ascii=False), flush=True)
            _prepare_selected_window(win, focus=focus)
            return win
    main = get_wechat_window() if focus else _try_get_main_surface_without_focus()
    if main is None:
        return None
    # SEARCH_LAYOUT is diagnostic state, not evidence.  Always prove that the
    # current main window still contains the embedded search surface.
    if _is_embedded_search_surface(main):
        if SEARCH_LAYOUT != "embedded":
            SEARCH_LAYOUT = "embedded"
            print(json.dumps({"event": "wechat_search_layout", "layout": SEARCH_LAYOUT}, ensure_ascii=False), flush=True)
        return main
    return None


def try_get_video_detail_window(*, focus: bool = True):
    """Return only the active video-detail surface."""
    detached = web_windows()
    if detached:
        win = select_detached_window_by_role(detached, "video_detail")
        _log_multi_window_resolution(detached, win, "video_detail")
        if win is not None:
            _prepare_selected_window(win, focus=focus)
            return win
    try:
        main = get_wechat_window() if focus else _try_get_main_surface_without_focus()
        if main is not None and is_video_detail_page(main):
            return main
    except Exception:
        pass
    return None


def try_get_search_window():
    """Backward-compatible alias whose semantics are now search-results only."""
    return try_get_search_results_window()


def get_search_results_window(wait_seconds: float = 2.0):
    """Return a stable search surface across short Chromium rebuilds."""
    deadline = time.time() + max(0.0, float(wait_seconds))
    last_error = None
    while True:
        try:
            surface = try_get_search_results_window(focus=False)
            if surface is not None:
                _prepare_selected_window(surface, focus=True)
                return surface
        except Exception as exc:
            last_error = exc
        if time.time() >= deadline:
            detail = f"：{last_error}" if last_error else ""
            raise SearchWindowUnavailableError(
                f"未检测到独立或内嵌的微信搜一搜页面{detail}"
            ) from last_error
        time.sleep(0.25)


def get_search_window(wait_seconds: float = 2.0):
    """Backward-compatible alias for the explicit search-results getter."""
    return get_search_results_window(wait_seconds=wait_seconds)


def get_video_detail_window(wait_seconds: float = 2.0):
    deadline = time.time() + max(0.0, float(wait_seconds))
    last_error = None
    while True:
        try:
            surface = try_get_video_detail_window(focus=False)
            if surface is not None and is_video_detail_page(surface):
                _prepare_selected_window(surface, focus=True)
                return surface
        except Exception as exc:
            last_error = exc
        if time.time() >= deadline:
            detail = f"：{last_error}" if last_error else ""
            raise SearchWindowUnavailableError(f"未检测到微信视频详情页面{detail}") from last_error
        time.sleep(0.25)


def has_search_window() -> bool:
    try:
        return try_get_search_results_window(focus=False) is not None
    except Exception:
        return False


def has_video_detail_window() -> bool:
    # ``select_detached_window_by_role`` intentionally rejects ambiguous
    # multiple-player layouts.  A boolean probe must still report that stale
    # players exist, otherwise the exception recovery path skips cleanup and
    # every subsequent card adds another window.
    try:
        if any(detached_web_window_role(win) == "video_detail" for win in web_windows()):
            return True
    except Exception:
        pass
    try:
        surface = try_get_video_detail_window(focus=False)
        return surface is not None and is_video_detail_page(surface)
    except Exception:
        return False


def click_relative(win, point: list[int]) -> None:
    rect = win.rectangle()
    scale = window_dpi_scale(win)
    mouse.click(coords=(
        rect.left + int(float(point[0]) * scale),
        rect.top + int(float(point[1]) * scale),
    ))


def window_dpi_scale(win) -> float:
    """Return physical-pixel scale for 100%/150% Windows desktops."""
    configured = os.environ.get("WECHAT_DISPLAY_SCALE")
    if configured:
        try:
            value = float(configured) / 100.0
            if 0.75 <= value <= 3.0:
                return value
        except ValueError:
            pass
    try:
        dpi = ctypes.windll.user32.GetDpiForWindow(int(win.handle))
        if dpi:
            return max(0.75, min(3.0, float(dpi) / 96.0))
    except Exception:
        pass
    return 1.0


def set_clipboard(text: str) -> None:
    win32clipboard.OpenClipboard()
    try:
        win32clipboard.EmptyClipboard()
        win32clipboard.SetClipboardText(text, win32clipboard.CF_UNICODETEXT)
    finally:
        win32clipboard.CloseClipboard()


def get_clipboard() -> str:
    win32clipboard.OpenClipboard()
    try:
        if win32clipboard.IsClipboardFormatAvailable(win32clipboard.CF_UNICODETEXT):
            return win32clipboard.GetClipboardData(win32clipboard.CF_UNICODETEXT)
        return ""
    finally:
        win32clipboard.CloseClipboard()


def paste_text(text: str) -> None:
    set_clipboard(text)
    keyboard.send_keys("^a{BACKSPACE}^v", pause=0.02)


def save_window_screenshot(win, path: Path) -> None:
    rect = win.rectangle()
    path.parent.mkdir(parents=True, exist_ok=True)
    ImageGrab.grab(
        bbox=(rect.left, rect.top, rect.right, rect.bottom),
        all_screens=True,
    ).save(path)


def load_keywords(path: Path, limit: int | None = None) -> list[tuple[str, str]]:
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        entries = payload.get("keywords", payload) if isinstance(payload, dict) else payload
        rows = []
        for item in entries:
            if isinstance(item, str):
                keyword, project = item, ""
            else:
                keyword = item.get("keyword")
                topic = str(item.get("topic") or "").strip().lower()
                # ODS `项目` is the business topic name, not the source
                # workbook's internal P01/C01 keyword grouping code.
                project = {
                    "topic1": "全国地产设计产品力监测",
                    "topic2": "北京臻澐项目口碑",
                }.get(topic, item.get("project", ""))
            if keyword:
                rows.append((str(keyword).strip(), str(project or "").strip()))
            if limit and len(rows) >= limit:
                break
        return rows
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["关键词"]
    rows: list[tuple[str, str]] = []
    for keyword, project, *_ in sheet.iter_rows(min_row=2, values_only=True):
        if keyword is None:
            continue
        rows.append((str(keyword).strip(), str(project or "").strip()))
        if limit and len(rows) >= limit:
            break
    return rows


def normalize_article_key(url: str) -> str:
    """Return a stable key for a Channels preview URL.

    Finder links use either ``/finder-preview/pages/sph?id=...`` or a share
    path containing the export id.  Tracking parameters are deliberately not
    part of the key so the same video found by several searches stays one ODS
    post while each keyword hit remains recoverable in SQLite.
    """
    parsed = urlparse(url.strip())
    query = parse_qs(parsed.query)
    finder_id = query.get("id", [""])[0] or query.get("exportkey", [""])[0]
    stable = finder_id or f"{parsed.netloc.lower()}|{parsed.path.rstrip('/')}"
    return hashlib.sha256(stable.encode("utf-8")).hexdigest()


def extract_finder_id(url: str | None) -> str | None:
    """Extract a stable finder/export id when WeChat exposes one."""
    if not url:
        return None
    parsed = urlparse(str(url).strip())
    query = parse_qs(parsed.query)
    for name in ("id", "finder_id", "object_id", "exportkey"):
        value = (query.get(name) or [""])[0].strip()
        if value:
            return value
    match = re.search(r"/(?:sph|finder|object)/([0-9A-Za-z_-]{8,})", parsed.path)
    return match.group(1) if match else None


def make_video_card_fingerprint(
    title: str | None, author_name: str | None, duration_seconds: int | None
) -> str:
    title_key = normalize_video_identity_text(title)
    author_key = normalize_video_identity_text(author_name)
    material = "|".join([title_key, author_key, str(duration_seconds or "")])
    if not material.replace("|", ""):
        return ""
    return hashlib.sha256(f"channels-card-v2|{material}".encode("utf-8")).hexdigest()


def normalize_video_card_fingerprint(raw_card_text: str) -> str:
    """Build a stable pre-open fingerprint from title, author and duration.

    Relative date text is deliberately excluded because it changes between
    retries.  Unlike the legacy title-only key, this fingerprint does not merge
    different authors' videos just because their captions are identical.
    """
    card = parse_video_card_text(raw_card_text or "")
    value = re.sub(r"\s+", " ", card.get("title") or raw_card_text or "").strip()
    value = re.sub(
        r"^\d{1,2}:\d{2}\s+(?:\d+(?:\.\d+)?\s*[万wW]?\s+)?(?:播放\s+|打开声音\s+)?",
        "",
        value,
    )
    value = re.sub(
        r"\s+(?:刚刚|昨天|前天|\d+\s*(?:分钟|小时|天|周|个月|年)前|"
        r"20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}日?)$",
        "",
        value,
    )
    return make_video_card_fingerprint(
        value, card.get("author_name"), card.get("duration_seconds")
    )


def retry_checkpoint_matches_card(checkpoint: dict | None, card: dict) -> bool:
    """Match the same search card when its displayed duration has drifted.

    WeChat's result card may expose the autoplay *remaining* time rather than
    the immutable total duration.  Consequently the normal title+author+
    duration fingerprint can change from 00:38 to 00:35/00:27 while the card
    is visible.  Redis retry tasks keep the original fingerprint for durable
    state, but may locate that exact card using stable title and author text.
    """
    checkpoint = checkpoint or {}
    row = checkpoint.get("row") or {}
    expected_title = normalize_video_identity_text(
        row.get("title") or checkpoint.get("title")
    )
    actual_title = normalize_video_identity_text(card.get("title"))
    if not expected_title or not actual_title:
        return False
    shorter, longer = sorted((expected_title, actual_title), key=len)
    title_matches = (
        expected_title == actual_title
        or (
            len(shorter) >= 12
            and (longer.startswith(shorter) or shorter in longer)
        )
        or (
            len(shorter) >= 16
            and SequenceMatcher(None, expected_title, actual_title).ratio() >= 0.94
        )
    )
    if not title_matches:
        return False
    expected_author = normalize_video_identity_text(
        row.get("author_name") or checkpoint.get("author_name")
    )
    actual_author = normalize_video_identity_text(card.get("author_name"))
    if expected_author and actual_author:
        return expected_author == actual_author
    return True


def redis_retry_card_scan_limit(config: dict) -> int:
    """Bound one exact-video retry while preserving deep original indexes.

    A child retry must never rebuild and rescan an entire keyword several
    times.  Recent failures search a small newest-first window; a failure that
    originally occurred deeper in the annual list gets its old index plus a
    bounded look-ahead allowance.
    """
    minimum = max(1, int(config.get("asr_retry_min_scan_results", 36)))
    ahead = max(1, int(config.get("asr_retry_scan_ahead_results", 36)))
    hard_cap = max(minimum, int(config.get("asr_retry_hard_cap_results", 240)))
    try:
        original_index = int(os.environ.get("WECHAT_REDIS_VIDEO_RESULT_INDEX", ""))
    except (TypeError, ValueError):
        original_index = -1
    target = hard_cap if original_index < 0 else max(minimum, original_index + ahead + 1)
    return min(hard_cap, target)


def normalize_video_identity_text(value: str | None) -> str:
    """Normalize a card/detail caption for same-video verification."""
    text = CARD_TAG_PATTERN.sub("", str(value or "")).replace("#", "")
    text = re.sub(r"\s+", "", text).casefold()
    # UIA may expose visually identical punctuation in full/half width forms.
    return re.sub(r"[，,。.!！?？；;：:、|｜·“”‘’《》【】()（）\-—_]", "", text)


def video_detail_identity_match(
    expected_title: str | None,
    expected_author: str | None,
    records: list[tuple[str, str, str]],
) -> tuple[bool, float, str | None]:
    """Verify that the active feed item still matches the clicked card.

    Very short videos can auto-advance before metrics/comments are read.  The
    active root then contains a perfectly valid *different* video, so link
    validation alone cannot detect the corruption.  Compare the card caption
    against every visible text candidate in the active item and require a
    strong caption match.  Author is used only as an extra guard for very short
    captions, never as the sole identity.
    """
    expected = normalize_video_identity_text(expected_title)
    expected_author_key = normalize_video_identity_text(expected_author)
    title_was_truncated = bool(re.match(r"^\s*(?:\.{2,}|…+)", str(expected_title or "")))
    if not expected:
        return False, 0.0, None

    best_score = 0.0
    best_value = None
    author_seen = not expected_author_key
    for value, control_type, _class_name in records:
        if control_type not in {"Text", "Hyperlink", "Button"}:
            continue
        candidate = normalize_video_identity_text(value)
        if not candidate:
            continue
        if expected_author_key and candidate == expected_author_key:
            author_seen = True
        score = SequenceMatcher(None, expected, candidate).ratio()
        shorter, longer = sorted((expected, candidate), key=len)
        if len(shorter) >= 6 and shorter in longer:
            score = max(score, len(shorter) / max(1, len(longer)))
            if len(expected) >= 8 and expected in candidate:
                score = max(score, 0.98)
            # The result card can expose the full caption while the detail
            # accessibility tree truncates that same caption after several
            # lines.  A long exact opening is a stronger identity signal than
            # the raw whole-string ratio (for example 0.76 for a 200-character
            # verified prefix) and must not be sent to retry forever.
            if (
                len(shorter) >= 30
                and (expected.startswith(candidate) or candidate.startswith(expected))
            ):
                score = max(score, 0.96)
        if title_was_truncated and len(expected) >= 3 and expected in candidate:
            # Search cards occasionally expose only ``...精装修`` while the
            # detail page has the full caption.  In that special layout the
            # author must also match before this score can pass below.
            score = max(score, 0.96)
        if score > best_score:
            best_score = score
            best_value = value

    minimum = 0.82 if len(expected) >= 8 else 0.90
    matched = best_score >= minimum and author_seen
    return matched, best_score, best_value


def require_video_detail_identity(
    card: dict,
    records: list[tuple[str, str, str]],
    *,
    keyword: str,
    result_index: int,
) -> None:
    matched, score, detail_text = video_detail_identity_match(
        card.get("title"), card.get("author_name"), records,
    )
    if matched:
        return
    payload = {
        "event": "video_detail_identity_mismatch",
        "keyword": keyword,
        "index": result_index,
        "expected_title": card.get("title"),
        "expected_author": card.get("author_name"),
        "detail_candidate": detail_text,
        "similarity": round(score, 4),
    }
    print(json.dumps(payload, ensure_ascii=False), flush=True)
    raise VideoDetailIdentityMismatchError(
        f"视频详情已自动跳转或打开错位：卡片={card.get('title')!r}，"
        f"详情候选={detail_text!r}，相似度={score:.3f}"
    )


def video_player_failure_state(surface) -> tuple[bool, str | None]:
    """Inspect visible player labels, not hidden feed items or caption substrings."""
    root = active_video_root(surface)
    unavailable = False
    # Access dialogs may be siblings of the active feed item. Inspect them at
    # surface scope, but keep the network-error label scoped to the player.
    for container, check_network in ((surface, False), (root, True)):
        viewport = container.rectangle()
        for control in container.descendants():
            try:
                info = control.element_info
                if info.control_type not in {'Text', 'Button', 'Pane'} or not control.is_visible():
                    continue
                rect = control.rectangle()
                if (min(rect.bottom, viewport.bottom) <= max(rect.top, viewport.top)
                        or min(rect.right, viewport.right) <= max(rect.left, viewport.left)):
                    continue
                name = re.sub(r'\s+', '', control.window_text() or '')
                if check_network and re.fullmatch(r'(?:暂无法|暂时无法|无法)观看此视频(?:[，,。].*)?', name):
                    unavailable = True
                # Only standalone UI prompts qualify, not caption substrings.
                if re.fullmatch(r'(?:请先登录|请登录微信|扫码登录(?:微信)?|请在手机上确认登录|'
                                r'请完成安全验证|安全验证|验证码|请完成拼图验证|'
                                r'操作频繁|访问频繁|网络环境存在风险)(?:[。！!，,].*)?', name):
                    return unavailable, name
            except Exception:
                continue
    return unavailable, None


def prepare_video_detail(surface, config: dict, card: dict | None = None,
                         *, keyword: str = '', result_index: int = -1):
    """Reload the SAME detail tab only after an explicit unwatchable label.

    No recorder is active here. Reacquire UIA after every reload and verify the
    target again before returning; disappearance of the error alone is not
    success. Actual audio/playback coverage is still verified by the recorder.
    """
    max_refreshes = max(0, min(10, int(config.get('video_unavailable_refresh_attempts', 4))))
    wait_seconds = max(1.0, min(60.0, float(config.get('wait_video_detail_ready_seconds', 15))))
    backoff = max(0.5, min(30.0, float(config.get('video_unavailable_refresh_backoff_seconds', 2))))
    deadline = time.monotonic() + wait_seconds
    refreshes = 0
    observed = False
    def check_stop():
        try:
            ensure_capture_not_stopped(config)
        except ClientAudioCaptureError as exc:
            raise VideoDetailRecoveryBlockedError(str(exc)) from exc
    while True:
        check_stop()
        unavailable, restriction = video_player_failure_state(surface) if surface is not None else (False, None)
        if restriction:
            if '登录' in restriction:
                restriction = 'STOP_LOGIN_REQUIRED：' + restriction
            else:
                restriction = 'STOP_PLATFORM_RESTRICTION：' + restriction
            report_from_environment(restriction)
            raise VideoDetailRecoveryBlockedError(restriction)
        if unavailable:
            observed = True
        if surface is not None and not unavailable:
            root = active_video_root(surface)
            # Pause as soon as possible, including while metadata hydrates, to
            # prevent short videos auto-advancing during the readiness wait.
            paused = pause_active_video_if_running(surface)
            actions = [b for b in root.descendants(control_type='Button')
                       if 'op-item' in (b.element_info.class_name or '') and b.is_visible()]
            if len(actions) >= 4:
                identity = card or config.get('_capture_identity')
                if refreshes and not identity:
                    raise VideoDetailIdentityMismatchError('Cannot verify target identity after video refresh')
                if identity:
                    require_video_detail_identity(identity, active_video_records(surface),
                                                  keyword=keyword, result_index=result_index)
                if refreshes:
                    print(json.dumps({'event': 'video_unavailable_refresh_recovered',
                        'policy': 'channels-video-refresh-20261023-v1',
                        'keyword': keyword, 'index': result_index, 'refreshes': refreshes,
                        'identityVerified': bool(identity), 'paused': paused,
                        'audioCoverageVerified': False}, ensure_ascii=False), flush=True)
                return surface
        expired = time.monotonic() >= deadline
        if observed and (refreshes == 0 or expired):
            if refreshes >= max_refreshes:
                print(json.dumps({'event': 'video_unavailable_refresh_exhausted',
                    'keyword': keyword, 'index': result_index, 'refreshes': refreshes,
                    'localCheckpointRetained': True, 'contentAttemptCharged': False},
                    ensure_ascii=False), flush=True)
                raise VideoPlaybackUnavailableError('VIDEO_PLAYBACK_UNAVAILABLE：刷新耗尽，当前视频及断点保留，未标记采集成功')
            if surface is None or not is_video_detail_page(surface):
                raise VideoPlaybackUnavailableError('VIDEO_PLAYBACK_UNAVAILABLE：详情窗口不可确认，停止刷新并保留断点')
            delay = min(30.0, backoff * (refreshes + 1))
            print(json.dumps({'event': 'video_unavailable_refresh_retry',
                'keyword': keyword, 'index': result_index, 'attempt': refreshes + 1,
                'maxAttempts': max_refreshes, 'waitSeconds': delay}, ensure_ascii=False), flush=True)
            # Check both process signals and the worker stop file in slices.
            wait_until = time.monotonic() + delay
            while time.monotonic() < wait_until:
                check_stop()
                interruptible_wait(min(0.25, wait_until - time.monotonic()))
            check_stop()
            # A user may have switched tabs during backoff. Never send Ctrl+R
            # to the cached handle without reacquiring its detail role.
            try:
                surface = get_video_detail_window()
            except SearchWindowUnavailableError as exc:
                raise VideoPlaybackUnavailableError('VIDEO_PLAYBACK_UNAVAILABLE：刷新前详情窗口已消失，保留断点') from exc
            still_unavailable, restriction = video_player_failure_state(surface)
            if restriction:
                continue  # report the gate at the top, without refreshing
            if not still_unavailable and len([
                b for b in active_video_root(surface).descendants(control_type='Button')
                if 'op-item' in (b.element_info.class_name or '') and b.is_visible()
            ]) >= 4:
                continue  # it recovered during backoff; avoid an extra reload
            surface.set_focus()
            keyboard.send_keys('^r', pause=0.02)
            refreshes += 1
            # Old UIA controls can survive briefly during navigation.
            if not interruptible_wait(0.8):
                check_stop()
            deadline = time.monotonic() + wait_seconds
        elif expired:
            raise RuntimeError('视频详情页已打开，但互动栏未在时限内完成加载')
        check_stop()
        interruptible_wait(0.25)
        try:
            surface = get_video_detail_window(wait_seconds=0)
        except SearchWindowUnavailableError:
            surface = None  # navigation gap: never refresh an unverified tab


def article_key_for_row(row: ResultRow) -> str:
    # Finder id is the authoritative identity. When the share surface hides it,
    # use the requested composite fallback instead of the legacy title-only key.
    finder_id = extract_finder_id(row.url)
    if finder_id:
        return hashlib.sha256(f"channels-finder|{finder_id}".encode("utf-8")).hexdigest()
    material = "|".join([
        normalize_video_identity_text(row.title),
        normalize_video_identity_text(row.author_name),
        str(row.published_at or ""),
        str(row.video_duration or ""),
    ])
    if material.replace("|", ""):
        return hashlib.sha256(f"channels-fallback-v2|{material}".encode("utf-8")).hexdigest()
    return normalize_article_key(row.url)


def parse_count(text: str | None) -> int | None:
    if text is None:
        return None
    value = text.strip().replace(",", "")
    if not value or value in {"--", "不可见", "未显示"}:
        return None
    match = re.search(r"(\d+(?:\.\d+)?)\s*([万wW]?)", value)
    if not match:
        return None
    number = float(match.group(1))
    if match.group(2):
        number *= 10000
    return int(number)


def init_state(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    init_review_tables(connection)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS captured_article (
            article_key TEXT NOT NULL,
            project TEXT NOT NULL,
            keyword TEXT NOT NULL,
            url TEXT NOT NULL,
            like_count INTEGER,
            comment_count INTEGER,
            favorite_count INTEGER,
            share_count INTEGER,
            captured_at TEXT NOT NULL,
            PRIMARY KEY (article_key, keyword)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS captured_comment (
            comment_key TEXT NOT NULL,
            article_key TEXT NOT NULL,
            project TEXT NOT NULL,
            keyword TEXT NOT NULL,
            article_url TEXT NOT NULL,
            user_url TEXT,
            user_id TEXT,
            user_name TEXT,
            content TEXT,
            published_at TEXT,
            location TEXT,
            like_count INTEGER,
            reply_count INTEGER,
            captured_at TEXT NOT NULL,
            PRIMARY KEY (comment_key, keyword)
        )
        """
    )
    comment_columns = {row[1] for row in connection.execute("PRAGMA table_info(captured_comment)")}
    for column, definition in {"user_url": "TEXT", "user_id": "TEXT"}.items():
        if column not in comment_columns:
            connection.execute(f"ALTER TABLE captured_comment ADD COLUMN {column} {definition}")
    # Idempotent migration for databases created before title capture was added.
    columns = {row[1] for row in connection.execute("PRAGMA table_info(captured_article)")}
    migrations = {
        "title": "TEXT", "author_name": "TEXT", "published_at": "TEXT",
        "tags": "TEXT",
        "location": "TEXT", "content": "TEXT", "image_count": "INTEGER DEFAULT 0",
        "video_count": "INTEGER DEFAULT 0",
        "video_content": "TEXT", "video_duration": "INTEGER",
    }
    for column, definition in migrations.items():
        if column not in columns:
            connection.execute(f"ALTER TABLE captured_article ADD COLUMN {column} {definition}")
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS crawl_task (
            keyword TEXT PRIMARY KEY,
            project TEXT NOT NULL,
            next_result_index INTEGER NOT NULL DEFAULT 0,
            old_streak INTEGER NOT NULL DEFAULT 0,
            status TEXT NOT NULL DEFAULT 'pending',
            last_error TEXT,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS collection_issue (
            keyword TEXT NOT NULL,
            result_index INTEGER NOT NULL,
            issue_type TEXT NOT NULL,
            fingerprint TEXT,
            title TEXT,
            detail TEXT,
            captured_at TEXT NOT NULL,
            PRIMARY KEY (keyword, result_index, issue_type)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS video_checkpoint (
            keyword TEXT NOT NULL,
            fingerprint TEXT NOT NULL,
            project TEXT NOT NULL,
            result_index INTEGER,
            article_key TEXT,
            title TEXT,
            author_name TEXT,
            published_at TEXT,
            video_duration INTEGER,
            status TEXT NOT NULL,
            retry_count INTEGER NOT NULL DEFAULT 0,
            field_completeness TEXT,
            asr_status TEXT,
            last_error TEXT,
            screenshot_path TEXT,
            row_json TEXT,
            comments_json TEXT,
            updated_at TEXT NOT NULL,
            PRIMARY KEY (keyword, fingerprint)
        )
        """
    )
    checkpoint_columns = {
        row[1] for row in connection.execute("PRAGMA table_info(video_checkpoint)")
    }
    if "result_index" not in checkpoint_columns:
        connection.execute(
            "ALTER TABLE video_checkpoint ADD COLUMN result_index INTEGER"
        )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS video_dead_letter (
            keyword TEXT NOT NULL,
            fingerprint TEXT NOT NULL,
            project TEXT NOT NULL,
            result_index INTEGER,
            title TEXT,
            author_name TEXT,
            raw_card_text TEXT,
            retry_count INTEGER NOT NULL,
            last_error TEXT,
            screenshot_path TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY (keyword, fingerprint)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS asr_job (
            article_key TEXT NOT NULL,
            keyword TEXT NOT NULL,
            wav_path TEXT NOT NULL,
            payload_json TEXT NOT NULL,
            status TEXT NOT NULL,
            attempts INTEGER NOT NULL DEFAULT 0,
            last_error TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            PRIMARY KEY (article_key, keyword)
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS state_meta (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.commit()
    migrate_legacy_video_state(connection)
    return connection


def load_task_state(connection: sqlite3.Connection, keyword: str, project: str) -> tuple[int, int, str]:
    row = connection.execute(
        "SELECT next_result_index, old_streak, status FROM crawl_task WHERE keyword=?",
        (keyword,),
    ).fetchone()
    if row:
        return int(row[0]), int(row[1]), str(row[2])
    existing_count = connection.execute(
        "SELECT COUNT(*) FROM captured_article WHERE keyword=?",
        (keyword,),
    ).fetchone()[0]
    now = datetime.now().isoformat(timespec="seconds")
    connection.execute(
        """
        INSERT INTO crawl_task(keyword, project, next_result_index, old_streak, status, updated_at)
        VALUES (?, ?, ?, 0, 'pending', ?)
        """,
        (keyword, project, int(existing_count), now),
    )
    connection.commit()
    return int(existing_count), 0, "pending"


def update_task_state(
    connection: sqlite3.Connection,
    keyword: str,
    project: str,
    next_result_index: int,
    old_streak: int,
    status: str,
    last_error: str | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO crawl_task(
            keyword, project, next_result_index, old_streak, status, last_error, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(keyword) DO UPDATE SET
            project=excluded.project,
            next_result_index=excluded.next_result_index,
            old_streak=excluded.old_streak,
            status=excluded.status,
            last_error=excluded.last_error,
            updated_at=excluded.updated_at
        """,
        (
            keyword, project, next_result_index, old_streak, status, last_error,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    connection.commit()


def save_collection_issue(
    connection: sqlite3.Connection,
    keyword: str,
    result_index: int,
    issue_type: str,
    title: str,
    detail: str,
    fingerprint: str | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO collection_issue(
            keyword, result_index, issue_type, fingerprint, title, detail, captured_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(keyword, result_index, issue_type) DO UPDATE SET
            fingerprint=excluded.fingerprint,
            title=excluded.title,
            detail=excluded.detail,
            captured_at=excluded.captured_at
        """,
        (
            keyword, result_index, issue_type, fingerprint, title, detail,
            datetime.now().isoformat(timespec="seconds"),
        ),
    )
    connection.commit()


def save_row(connection: sqlite3.Connection, row: ResultRow) -> None:
    connection.execute(
        """
        INSERT INTO captured_article (
            article_key, project, keyword, title, url, author_name, tags, published_at,
            location, content, image_count, video_count, like_count, comment_count,
            video_content, video_duration, favorite_count, share_count, captured_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(article_key, keyword) DO UPDATE SET
            project=excluded.project,
            title=COALESCE(NULLIF(excluded.title, ''), captured_article.title),
            url=excluded.url,
            author_name=COALESCE(excluded.author_name, captured_article.author_name),
            tags=COALESCE(excluded.tags, captured_article.tags),
            published_at=COALESCE(excluded.published_at, captured_article.published_at),
            location=COALESCE(excluded.location, captured_article.location),
            content=COALESCE(NULLIF(excluded.content, ''), captured_article.content),
            image_count=MAX(COALESCE(excluded.image_count, 0), COALESCE(captured_article.image_count, 0)),
            video_count=MAX(COALESCE(excluded.video_count, 0), COALESCE(captured_article.video_count, 0)),
            video_content=COALESCE(NULLIF(excluded.video_content, ''), captured_article.video_content),
            video_duration=COALESCE(excluded.video_duration, captured_article.video_duration),
            like_count=COALESCE(excluded.like_count, captured_article.like_count),
            comment_count=COALESCE(excluded.comment_count, captured_article.comment_count),
            favorite_count=COALESCE(excluded.favorite_count, captured_article.favorite_count),
            share_count=COALESCE(excluded.share_count, captured_article.share_count),
            captured_at=excluded.captured_at
        """,
        (
            article_key_for_row(row), row.project, row.keyword, row.title, row.url,
            row.author_name, row.tags, row.published_at, row.location, None,
            row.image_count, row.video_count,
            row.like_count, row.comment_count, row.video_content,
            row.video_duration, row.favorite_count, row.share_count, row.captured_at,
        ),
    )
    connection.commit()


def save_comments(connection: sqlite3.Connection, row: ResultRow, comments: list[dict]) -> None:
    article_key = article_key_for_row(row)
    for index, item in enumerate(comments):
        material = "|".join([
            article_key,
            str(item.get("user_name") or ""),
            str(item.get("content") or ""),
            str(item.get("published_at") or ""),
            str(index),
        ])
        comment_key = str(item.get("comment_id") or hashlib.sha256(material.encode("utf-8")).hexdigest())
        connection.execute(
            """
            INSERT INTO captured_comment(
                comment_key, article_key, project, keyword, article_url,
                user_url, user_id, user_name, content, published_at, location, like_count,
                reply_count, captured_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(comment_key, keyword) DO UPDATE SET
                article_url=excluded.article_url,
                user_url=COALESCE(excluded.user_url, captured_comment.user_url),
                user_id=COALESCE(excluded.user_id, captured_comment.user_id),
                user_name=COALESCE(excluded.user_name, captured_comment.user_name),
                content=COALESCE(excluded.content, captured_comment.content),
                published_at=COALESCE(excluded.published_at, captured_comment.published_at),
                location=COALESCE(excluded.location, captured_comment.location),
                like_count=COALESCE(excluded.like_count, captured_comment.like_count),
                reply_count=COALESCE(excluded.reply_count, captured_comment.reply_count),
                captured_at=excluded.captured_at
            """,
            (
                comment_key, article_key, row.project, row.keyword, row.url,
                item.get("user_url"), item.get("user_id"), item.get("user_name"),
                item.get("content"), item.get("published_at"),
                item.get("location"), item.get("like_count"), item.get("reply_count"),
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
    connection.commit()


def video_field_completeness(
    row: ResultRow, comments: list[dict] | None = None, *, asr_status: str | None = None
) -> dict[str, bool]:
    comments = comments or []
    return {
        "title": bool((row.title or "").strip()),
        "url": bool((row.url or "").strip()),
        "author": bool((row.author_name or "").strip()),
        "published_at": bool(row.published_at),
        "duration": row.video_duration is not None,
        "like": row.like_count is not None,
        "comment_count": row.comment_count is not None,
        "favorite": row.favorite_count is not None,
        "share": row.share_count is not None,
        "comments": (row.comment_count or 0) == 0 or bool(comments),
        "asr": bool((row.video_content or "").strip()) or asr_status in {
            "completed", "no_content", "not_applicable",
        },
    }


def load_video_checkpoints(connection: sqlite3.Connection, keyword: str) -> dict[str, dict]:
    rows = connection.execute(
        """
        SELECT fingerprint, status, retry_count, published_at, asr_status,
               field_completeness, article_key, row_json, comments_json,
               title, author_name, last_error, result_index
        FROM video_checkpoint WHERE keyword=?
        """,
        (keyword,),
    ).fetchall()
    output = {}
    for item in rows:
        output[item[0]] = {
            "fingerprint": item[0], "status": item[1], "retry_count": int(item[2] or 0),
            "published_at": item[3], "asr_status": item[4],
            "field_completeness": json.loads(item[5]) if item[5] else {},
            "article_key": item[6],
            "row": json.loads(item[7]) if item[7] else None,
            "comments": json.loads(item[8]) if item[8] else [],
            "title": item[9], "author_name": item[10], "last_error": item[11],
            "result_index": item[12],
        }
    return output


def retry_checkpoint_has_identity(checkpoint: dict | None) -> bool:
    """Return whether a failed card can be located again without its index."""
    checkpoint = checkpoint or {}
    row = checkpoint.get("row") or {}
    return bool(
        str(checkpoint.get("article_key") or "").strip()
        or str(row.get("url") or "").strip()
        or str(row.get("title") or checkpoint.get("title") or "").strip()
    )


def reusable_cross_keyword_checkpoint(checkpoint, *, project, article_key, url,
                                      author, duration, cutoff, collect_comments=False):
    """Only an exact, usable source can replace a second audio/detail visit.

    Local row existence, metadata completion and dead letters are NOT success.
    This is a reuse gate, not database acceptance: the final MySQL readback is
    still required for the receiving keyword.
    """
    if not isinstance(checkpoint, dict) or checkpoint.get('status') not in {'completed', 'no_content', 'not_applicable'}:
        return False
    if checkpoint.get('subtitle_review_required'):
        return False
    data = checkpoint.get('row')
    if not isinstance(data, dict) or checkpoint.get('article_key') != article_key or data.get('project') != project:
        return False
    if not data.get('keyword') or (checkpoint.get('keyword') and checkpoint['keyword'] != data['keyword']):
        return False
    if not data.get('title') or not is_valid_channels_share_url(data.get('url')) or not is_valid_channels_share_url(url):
        return False
    if not extract_finder_id(data['url']) or extract_finder_id(data['url']) != extract_finder_id(url):
        return False
    if not normalize_video_identity_text(author) or normalize_video_identity_text(data.get('author_name')) != normalize_video_identity_text(author):
        return False
    try:
        expected_duration, source_duration = float(duration or 0), float(data.get('video_duration') or 0)
        if not all(math.isfinite(value) and value > 0 for value in (expected_duration, source_duration)) or abs(source_duration - expected_duration) > 2:
            return False
        published = datetime.fromisoformat(str(data.get('published_at') or ''))
        if published < cutoff or published.year != cutoff.year:
            return False
        row = ResultRow(**{key: value for key, value in data.items() if key in ResultRow.__dataclass_fields__})
        if article_key_for_row(row) != article_key:
            return False
    except (TypeError, ValueError):
        return False
    text = data.get('video_content')
    asr_status = checkpoint.get('asr_status')
    if has_untrusted_transcript(text) or asr_status != checkpoint['status']:
        return False
    if checkpoint['status'] == 'completed' and (not str(text or '').strip() or asr_output_needs_review(text, duration)):
        return False
    if checkpoint['status'] != 'completed' and str(text or '').strip():
        return False
    if not isinstance(checkpoint.get('comments', []), list) or any(not isinstance(item, dict) for item in checkpoint.get('comments', [])):
        return False
    if not isinstance(checkpoint.get('field_completeness') or {}, dict):
        return False
    completeness = video_field_completeness(row, checkpoint.get('comments'), asr_status=asr_status)
    if any(value is not True for field, value in completeness.items() if field != 'comments' or collect_comments):
        return False
    # A stored explicit failure is not superseded by recalculating a flag.
    if any(value is not True for field, value in (checkpoint.get('field_completeness') or {}).items()
           if field != 'comments' or collect_comments):
        return False
    return True


def persist_cross_keyword_reuse(connection, *, checkpoint, keyword, project,
                                fingerprint, result_index, cutoff, collect_comments=False,
                                remote_store=None):
    """Bind retained source data to this keyword without erasing either history."""
    current = load_video_checkpoints(connection, keyword).get(fingerprint) or {}
    if current.get('status') == 'skipped_unresolved':
        raise RuntimeError('CROSS_KEYWORD_REUSE_EXPLICIT_SKIP')
    source = dict(checkpoint['row'])
    source_key = checkpoint['article_key']
    if not reusable_cross_keyword_checkpoint(checkpoint, project=project, article_key=source_key,
            url=source.get('url'), author=source.get('author_name'), duration=source.get('video_duration'),
            cutoff=cutoff, collect_comments=collect_comments):
        raise RuntimeError('CROSS_KEYWORD_REUSE_SOURCE_NOT_USABLE')
    if current.get('article_key') and current['article_key'] != source_key:
        raise RuntimeError('CROSS_KEYWORD_REUSE_IDENTITY_CONFLICT')
    if (current.get('row') or {}).get('url') and extract_finder_id(current['row']['url']) != extract_finder_id(source['url']):
        raise RuntimeError('CROSS_KEYWORD_REUSE_IDENTITY_CONFLICT')
    row = ResultRow(**{key: value for key, value in source.items() if key in ResultRow.__dataclass_fields__})
    if row.project != project or row.keyword == keyword or article_key_for_row(row) != source_key:
        raise RuntimeError('CROSS_KEYWORD_REUSE_SCOPE_CONFLICT')
    row.keyword = keyword
    comments = checkpoint.get('comments') or []
    previous_cursor = connection.execute('SELECT * FROM video_checkpoint WHERE keyword=? AND fingerprint=?',
                                         (keyword, fingerprint))
    previous_row = previous_cursor.fetchone()
    previous_record = dict(zip((item[0] for item in previous_cursor.description), previous_row)) if previous_row else None
    article_cursor = connection.execute('SELECT * FROM captured_article WHERE keyword=? AND article_key=?',
                                        (keyword, source_key))
    article_row = article_cursor.fetchone()
    previous_article = dict(zip((item[0] for item in article_cursor.description), article_row)) if article_row else None
    comment_cursor = connection.execute('SELECT * FROM captured_comment WHERE keyword=? AND article_key=? ORDER BY comment_key',
                                        (keyword, source_key))
    comment_columns = [item[0] for item in comment_cursor.description]
    previous_comments = [dict(zip(comment_columns, item)) for item in comment_cursor.fetchall()]
    # Append the full source and old receiving checkpoint before any update.
    connection.execute('''CREATE TABLE IF NOT EXISTS cross_keyword_reuse_audit (
        id INTEGER PRIMARY KEY AUTOINCREMENT, keyword TEXT NOT NULL,
        fingerprint TEXT NOT NULL, source_keyword TEXT NOT NULL,
        source_checkpoint TEXT NOT NULL, previous_checkpoint TEXT NOT NULL,
        previous_article TEXT NOT NULL, previous_comments TEXT NOT NULL,
        created_at TEXT NOT NULL)''')
    connection.execute('''INSERT INTO cross_keyword_reuse_audit
        (keyword,fingerprint,source_keyword,source_checkpoint,previous_checkpoint,previous_article,previous_comments,created_at)
        VALUES (?,?,?,?,?,?,?,?)''', (keyword, fingerprint, source['keyword'],
        json.dumps(checkpoint, ensure_ascii=False), json.dumps(previous_record, ensure_ascii=False),
        json.dumps(previous_article, ensure_ascii=False), json.dumps(previous_comments, ensure_ascii=False),
        datetime.now().isoformat(timespec='seconds')))
    connection.commit()
    save_row(connection, row)
    save_comments(connection, row, comments)
    return upsert_video_checkpoint(connection, keyword=keyword, project=project,
        fingerprint=fingerprint, result_index=result_index, status=checkpoint['status'],
        retry_count=int(current.get('retry_count') or 0), row=row, comments=comments,
        asr_status=checkpoint['asr_status'], remote_store=remote_store)


def should_age_missing_retry(
    fingerprint: str,
    checkpoint: dict,
    seen_fingerprints: set[str],
    *,
    asr_backfill_mode: bool,
    retry_fingerprint: str | None,
) -> bool:
    """Limit a missing-card failure to work owned by the current run."""
    if checkpoint.get("status") != "retry_pending":
        return False
    if fingerprint in seen_fingerprints:
        return False
    if asr_backfill_mode and retry_fingerprint and fingerprint != retry_fingerprint:
        return False
    return retry_checkpoint_has_identity(checkpoint)


def unique_checkpoint_for_stable_card_identity(
    checkpoints: dict[str, dict], card: dict,
) -> tuple[str, dict] | None:
    """Return one unambiguous title+author checkpoint despite countdown drift."""
    card_title = normalize_video_identity_text(card.get("title"))
    card_author = normalize_video_identity_text(card.get("author_name"))
    if not card_title or not card_author:
        return None
    matches = []
    for fingerprint, checkpoint in checkpoints.items():
        row = checkpoint.get("row") or {}
        checkpoint_title = normalize_video_identity_text(
            row.get("title") or checkpoint.get("title")
        )
        checkpoint_author = normalize_video_identity_text(
            row.get("author_name") or checkpoint.get("author_name")
        )
        if checkpoint_title == card_title and checkpoint_author == card_author:
            matches.append((fingerprint, checkpoint))
    return matches[0] if len(matches) == 1 else None


def load_captured_result_row(
    connection: sqlite3.Connection, article_key: str, keyword: str
) -> ResultRow | None:
    item = connection.execute(
        """
        SELECT project, keyword, title, url, author_name, tags, published_at,
               location, content, image_count, video_count, video_content,
               video_duration, like_count, comment_count, favorite_count,
               share_count, captured_at
        FROM captured_article WHERE article_key=? AND keyword=?
        """,
        (article_key, keyword),
    ).fetchone()
    if not item:
        return None
    return ResultRow(
        project=item[0], keyword=item[1], title=item[2] or "", url=item[3] or "",
        author_name=item[4], tags=item[5], published_at=item[6], location=item[7],
        content=item[8], image_count=int(item[9] or 0), video_count=int(item[10] or 0),
        video_content=item[11], video_duration=item[12], like_count=item[13],
        comment_count=item[14], favorite_count=item[15], share_count=item[16],
        captured_at=item[17] or "",
    )


def upsert_video_checkpoint(
    connection: sqlite3.Connection,
    *, keyword: str,
    project: str,
    fingerprint: str,
    status: str,
    retry_count: int = 0,
    row: ResultRow | None = None,
    comments: list[dict] | None = None,
    asr_status: str | None = None,
    field_completeness: dict | None = None,
    last_error: str | None = None,
    screenshot_path: str | None = None,
    result_index: int | None = None,
    remote_store: RedisVideoCheckpointStore | None = None,
) -> dict:
    prior_status=connection.execute("SELECT status FROM video_checkpoint WHERE keyword=? AND fingerprint=?",(keyword,fingerprint)).fetchone()
    if prior_status and prior_status[0]=="skipped_unresolved":
        previous=load_video_checkpoints(connection,keyword)[fingerprint]
        if os.environ.get('SOCIAL_VOLUME_ENABLED')!='1':
            return previous
        # An original-keyword reset may revisit a naturally encountered note.
        # Preserve the entire old checkpoint before replacing its metadata;
        # do not create a child retry or erase cumulative failure counts.
        connection.execute('CREATE TABLE IF NOT EXISTS keyword_replay_history (keyword TEXT,fingerprint TEXT,evidence_hash TEXT,payload TEXT,created_at TEXT,PRIMARY KEY(keyword,fingerprint,evidence_hash))')
        encoded=json.dumps(previous,ensure_ascii=False,sort_keys=True)
        connection.execute('INSERT OR IGNORE INTO keyword_replay_history VALUES (?,?,?,?,?)',
            (keyword,fingerprint,hashlib.sha256(encoded.encode('utf-8')).hexdigest(),encoded,datetime.now().isoformat()))
    now = datetime.now().isoformat(timespec="seconds")
    row_value = asdict(row) if row is not None else None
    comments_value = comments if comments is not None else []
    completeness = field_completeness or (
        video_field_completeness(row, comments_value, asr_status=asr_status) if row else {}
    )
    article_key = article_key_for_row(row) if row else None
    connection.execute(
        """
        INSERT INTO video_checkpoint(
            keyword, fingerprint, project, result_index, article_key, title, author_name,
            published_at, video_duration, status, retry_count,
            field_completeness, asr_status, last_error, screenshot_path,
            row_json, comments_json, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(keyword, fingerprint) DO UPDATE SET
            project=excluded.project,
            result_index=COALESCE(excluded.result_index, video_checkpoint.result_index),
            article_key=COALESCE(excluded.article_key, video_checkpoint.article_key),
            title=COALESCE(NULLIF(excluded.title, ''), video_checkpoint.title),
            author_name=COALESCE(NULLIF(excluded.author_name, ''), video_checkpoint.author_name),
            published_at=COALESCE(excluded.published_at, video_checkpoint.published_at),
            video_duration=COALESCE(excluded.video_duration, video_checkpoint.video_duration),
            status=excluded.status,
            retry_count=MAX(video_checkpoint.retry_count,excluded.retry_count),
            field_completeness=CASE WHEN excluded.field_completeness='{}'
                THEN video_checkpoint.field_completeness ELSE excluded.field_completeness END,
            asr_status=COALESCE(excluded.asr_status, video_checkpoint.asr_status),
            last_error=excluded.last_error,
            screenshot_path=COALESCE(excluded.screenshot_path, video_checkpoint.screenshot_path),
            row_json=COALESCE(excluded.row_json, video_checkpoint.row_json),
            comments_json=CASE WHEN excluded.comments_json='[]'
                THEN video_checkpoint.comments_json ELSE excluded.comments_json END,
            updated_at=excluded.updated_at
        """,
        (
            keyword, fingerprint, project, result_index, article_key,
            row.title if row else None, row.author_name if row else None,
            row.published_at if row else None, row.video_duration if row else None,
            status, retry_count, json.dumps(completeness, ensure_ascii=False), asr_status,
            last_error, screenshot_path,
            json.dumps(row_value, ensure_ascii=False) if row_value else None,
            json.dumps(comments_value, ensure_ascii=False), now,
        ),
    )
    connection.commit()
    payload = {
        "keyword": keyword, "project": project, "fingerprint": fingerprint,
        "result_index": result_index,
        "status": status, "retry_count": retry_count, "article_key": article_key,
        "published_at": row.published_at if row else None,
        "asr_status": asr_status, "field_completeness": completeness,
        "last_error": last_error, "screenshot_path": screenshot_path,
        "row": row_value, "comments": comments_value, "updated_at": now,
    }
    # Publish the merged durable record, not the sparse arguments of a status
    # update. Otherwise a failure erases the identity required on another node.
    payload.update(load_video_checkpoints(connection, keyword)[fingerprint])
    if remote_store is not None:
        remote_store.put(fingerprint, payload)
        if row is not None and article_key and status in {
            "completed", "completed_metadata", "dead_letter", "no_content", "not_applicable",
        }:
            remote_store.put_global(project, article_key, payload)
    return payload


def hydrate_remote_video_checkpoints(
    connection: sqlite3.Connection,
    keyword: str,
    project: str,
    remote_store: RedisVideoCheckpointStore,
    only_fingerprints: set[str] | None = None,
) -> int:
    hydrated = 0
    for fingerprint, payload in remote_store.load_all().items():
        if payload.get("keyword") != keyword:
            continue
        if only_fingerprints is not None and fingerprint not in only_fingerprints:
            continue
        checkpoint_project = str(payload.get("project") or project)
        row_data = payload.get("row")
        row = None
        if isinstance(row_data, dict):
            allowed = ResultRow.__dataclass_fields__.keys()
            row = ResultRow(**{key: value for key, value in row_data.items() if key in allowed})
            save_row(connection, row)
            if payload.get("comments"):
                save_comments(connection, row, list(payload.get("comments") or []))
        upsert_video_checkpoint(
            connection,
            keyword=keyword,
            project=checkpoint_project,
            fingerprint=fingerprint,
            status=str(payload.get("status") or "retry_pending"),
            retry_count=int(payload.get("retry_count") or 0),
            row=row,
            comments=list(payload.get("comments") or []),
            asr_status=payload.get("asr_status"),
            field_completeness=dict(payload.get("field_completeness") or {}),
            last_error=payload.get("last_error"),
            screenshot_path=payload.get("screenshot_path"),
            result_index=payload.get("result_index"),
            remote_store=None,
        )
        raw_identity = parse_video_card_text(
            str(payload.get("raw_card_text") or "")
        )
        hydrated_title = (
            payload.get("title")
            or (row.title if row else None)
            or raw_identity.get("title")
        )
        hydrated_author = (
            payload.get("author_name")
            or (row.author_name if row else None)
            or raw_identity.get("author_name")
        )
        connection.execute(
            """
            UPDATE video_checkpoint
            SET article_key=COALESCE(NULLIF(?, ''), article_key),
                title=COALESCE(NULLIF(?, ''), title),
                author_name=COALESCE(NULLIF(?, ''), author_name)
            WHERE keyword=? AND fingerprint=?
            """,
            (
                payload.get("article_key"), hydrated_title, hydrated_author,
                keyword, fingerprint,
            ),
        )
        connection.commit()
        if str(payload.get("status")) == "dead_letter":
            now = datetime.now().isoformat(timespec="seconds")
            connection.execute(
                """
                INSERT INTO video_dead_letter(
                    keyword, fingerprint, project, result_index, title, author_name,
                    raw_card_text, retry_count, last_error, screenshot_path,
                    created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(keyword, fingerprint) DO UPDATE SET
                    result_index=excluded.result_index,
                    title=excluded.title,
                    author_name=excluded.author_name,
                    raw_card_text=excluded.raw_card_text,
                    retry_count=excluded.retry_count,
                    last_error=excluded.last_error,
                    screenshot_path=excluded.screenshot_path,
                    updated_at=excluded.updated_at
                """,
                (
                    keyword, fingerprint, project, payload.get("result_index"),
                    payload.get("title") or (row.title if row else None),
                    payload.get("author_name") or (row.author_name if row else None),
                    payload.get("raw_card_text"), int(payload.get("retry_count") or 0),
                    payload.get("last_error"), payload.get("screenshot_path"), now, now,
                ),
            )
            connection.commit()
        hydrated += 1
    return hydrated


def record_video_failure(
    connection: sqlite3.Connection,
    *, keyword: str,
    project: str,
    fingerprint: str,
    result_index: int,
    title: str,
    author_name: str | None,
    raw_card_text: str,
    error: str,
    screenshot_path: str | None,
    max_attempts: int,
    remote_store: RedisVideoCheckpointStore | None,
) -> tuple[str, int]:
    previous = connection.execute(
        "SELECT retry_count FROM video_checkpoint WHERE keyword=? AND fingerprint=?",
        (keyword, fingerprint),
    ).fetchone()
    retry_count = int(previous[0] or 0) + 1 if previous else 1
    status = "dead_letter" if retry_count >= max(1, max_attempts) else "retry_pending"
    checkpoint_payload = upsert_video_checkpoint(
        connection, keyword=keyword, project=project, fingerprint=fingerprint,
        status=status, retry_count=retry_count, last_error=error,
        screenshot_path=screenshot_path, result_index=result_index,
        remote_store=remote_store,
    )
    connection.execute(
        """
        UPDATE video_checkpoint
        SET title=COALESCE(NULLIF(?, ''), title),
            author_name=COALESCE(NULLIF(?, ''), author_name)
        WHERE keyword=? AND fingerprint=?
        """,
        (title, author_name, keyword, fingerprint),
    )
    connection.commit()
    checkpoint_payload.update({
        "result_index": result_index,
        "title": title,
        "author_name": author_name,
        "raw_card_text": raw_card_text,
    })
    if remote_store is not None:
        # Persist retry identity on the first failure, not only once it has
        # already become dead-letter. A different PC can then locate it.
        remote_store.put(fingerprint, checkpoint_payload)
    if status == "dead_letter":
        now = datetime.now().isoformat(timespec="seconds")
        connection.execute(
            """
            INSERT INTO video_dead_letter(
                keyword, fingerprint, project, result_index, title, author_name,
                raw_card_text, retry_count, last_error, screenshot_path,
                created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(keyword, fingerprint) DO UPDATE SET
                result_index=excluded.result_index,
                title=excluded.title,
                author_name=excluded.author_name,
                raw_card_text=excluded.raw_card_text,
                retry_count=excluded.retry_count,
                last_error=excluded.last_error,
                screenshot_path=excluded.screenshot_path,
                updated_at=excluded.updated_at
            """,
            (
                keyword, fingerprint, project, result_index, title, author_name,
                raw_card_text, retry_count, error, screenshot_path, now, now,
            ),
        )
        connection.commit()
    return status, retry_count


def legacy_annual_rescan_required(status: str, *, asr_backfill=False, recheck_exhausted=True) -> bool:
    """A date streak is not a search endpoint; preserve item history on replay."""
    if asr_backfill:
        return False
    return (status.startswith('complete_cutoff')
            or status.startswith('complete_explicit_empty')
            or (recheck_exhausted and status.startswith('complete_exhausted')))


def observe_old_date_window(config: dict, old_streak: int) -> bool:
    """Diagnostic only. Interleaved target-year cards must still be searched."""
    threshold = max(0, int(config.get('cutoff_old_streak', 0)))
    return threshold > 0 and old_streak >= threshold


def status_with_deferred_video_retries(status: str, retry_pending_count: int) -> str:
    """Keep a finished annual scan terminal while child retries stay durable."""
    if retry_pending_count <= 0 or status == "partial_stopped":
        return status
    if status.startswith("complete_"):
        return status if status.endswith("_with_deferred_retry") else f"{status}_with_deferred_retry"
    return "partial_retry_pending"


def record_precise_search_miss(connection, *, keyword, project, retry_fingerprint,
                              scan_status, seen_fingerprints, max_attempts, remote_store=None):
    """A verified bounded exact search miss is an attempt, not an endless defer.

    This is not proof of deletion or success. Access/stop failures never enter it.
    """
    if scan_status!='partial_retry_target_not_found' or not retry_fingerprint or retry_fingerprint in seen_fingerprints:
        return None
    cp=load_video_checkpoints(connection,keyword).get(retry_fingerprint)
    if not cp or cp.get('status') in {'completed','no_content','not_applicable','skipped_unresolved'}:
        return None
    row=cp.get('row') or {}
    if not str(row.get('title') or cp.get('title') or '').strip() or not (
        cp.get('article_key') or row.get('url') or row.get('author_name') or cp.get('author_name')
    ):
        return None
    return record_video_failure(connection,keyword=keyword,project=project,
        fingerprint=retry_fingerprint,result_index=cp.get('result_index'),
        title=str(row.get('title') or cp.get('title') or ''),
        author_name=row.get('author_name') or cp.get('author_name'),raw_card_text=cp.get('raw_card_text') or '',
        error='DETAIL_UNAVAILABLE: target not observed after verified bounded exact search; not proof of deletion',
        screenshot_path=None,max_attempts=max_attempts,remote_store=remote_store)


def migrate_legacy_video_state(connection: sqlite3.Connection) -> None:
    """One-time migration from title-only ids to finder/composite checkpoints."""
    version = "channels-video-state-v2"
    current = connection.execute(
        "SELECT value FROM state_meta WHERE key='video_state_schema'"
    ).fetchone()
    if current and current[0] == version:
        return
    rows = connection.execute(
        """
        SELECT article_key, project, keyword, title, url, author_name, tags,
               published_at, location, content, image_count, video_count,
               video_content, video_duration, like_count, comment_count,
               favorite_count, share_count, captured_at
        FROM captured_article
        """
    ).fetchall()
    for item in rows:
        old_key = item[0]
        row = ResultRow(
            project=item[1], keyword=item[2], title=item[3] or "", url=item[4] or "",
            author_name=item[5], tags=item[6], published_at=item[7], location=item[8],
            content=item[9], image_count=int(item[10] or 0), video_count=int(item[11] or 0),
            video_content=item[12], video_duration=item[13], like_count=item[14],
            comment_count=item[15], favorite_count=item[16], share_count=item[17],
            captured_at=item[18] or "",
        )
        new_key = article_key_for_row(row)
        if new_key != old_key:
            target = connection.execute(
                "SELECT 1 FROM captured_article WHERE article_key=? AND keyword=?",
                (new_key, row.keyword),
            ).fetchone()
            if target:
                save_row(connection, row)
                connection.execute(
                    "UPDATE captured_comment SET article_key=? WHERE article_key=? AND keyword=?",
                    (new_key, old_key, row.keyword),
                )
                connection.execute(
                    "DELETE FROM captured_article WHERE article_key=? AND keyword=?",
                    (old_key, row.keyword),
                )
            else:
                connection.execute(
                    "UPDATE captured_article SET article_key=? WHERE article_key=? AND keyword=?",
                    (new_key, old_key, row.keyword),
                )
                connection.execute(
                    "UPDATE captured_comment SET article_key=? WHERE article_key=? AND keyword=?",
                    (new_key, old_key, row.keyword),
                )
            connection.execute(
                "UPDATE video_checkpoint SET article_key=? WHERE article_key=? AND keyword=?",
                (new_key, old_key, row.keyword),
            )
            connection.execute(
                "UPDATE OR IGNORE asr_job SET article_key=? WHERE article_key=? AND keyword=?",
                (new_key, old_key, row.keyword),
            )
            connection.execute(
                "DELETE FROM asr_job WHERE article_key=? AND keyword=?",
                (old_key, row.keyword),
            )
            connection.commit()
        fingerprint = make_video_card_fingerprint(
            row.title, row.author_name, row.video_duration
        )
        if fingerprint:
            comments = []
            for comment in connection.execute(
                """
                SELECT comment_key, user_url, user_id, user_name, content,
                       published_at, location, like_count, reply_count
                FROM captured_comment WHERE article_key=? AND keyword=?
                """,
                (new_key, row.keyword),
            ).fetchall():
                comments.append({
                    "comment_id": comment[0], "user_url": comment[1],
                    "user_id": comment[2], "user_name": comment[3],
                    "content": comment[4], "published_at": comment[5],
                    "location": comment[6], "like_count": comment[7],
                    "reply_count": comment[8],
                })
            upsert_video_checkpoint(
                connection,
                keyword=row.keyword,
                project=row.project,
                fingerprint=fingerprint,
                status="completed" if row.video_content else "completed_metadata",
                retry_count=0,
                row=row,
                comments=comments,
                asr_status="completed" if row.video_content else "pending",
                remote_store=None,
            )
    connection.execute(
        """
        INSERT INTO state_meta(key, value, updated_at)
        VALUES ('video_state_schema', ?, ?)
        ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at
        """,
        (version, datetime.now().isoformat(timespec="seconds")),
    )
    connection.commit()


def find_controls(root, *, control_type: str | None = None, name: str | None = None, class_contains: str | None = None):
    controls = root.descendants(control_type=control_type) if control_type else root.descendants()
    found = []
    for control in controls:
        try:
            info = control.element_info
            if name is not None and (info.name or "") != name:
                continue
            if class_contains is not None and class_contains not in (info.class_name or ""):
                continue
            found.append(control)
        except Exception:
            continue
    return found


def wait_for(predicate, timeout: float, interval: float = 0.25):
    deadline = time.time() + timeout
    last_error = None
    while time.time() < deadline:
        try:
            value = predicate()
            if value:
                return value
        except Exception as exc:
            last_error = exc
        time.sleep(interval)
    if last_error:
        raise TimeoutError(str(last_error)) from last_error
    raise TimeoutError("等待页面状态超时")


def visible_controls(root, *, name: str | None = None, control_type: str | None = None):
    candidates = find_controls(root, control_type=control_type, name=name)
    visible = []
    for control in candidates:
        try:
            rect = control.rectangle()
            if control.is_visible() and rect.width() > 2 and rect.height() > 2:
                visible.append(control)
        except Exception:
            continue
    return visible


def find_search_edit(surface):
    edits = visible_controls(surface, control_type="Edit")
    preferred = [
        control for control in edits
        if (control.element_info.automation_id or "") == "weixin-search-input"
    ]
    if len(preferred) == 1:
        return preferred[0]
    named = [control for control in edits if (control.window_text() or "").strip() in {"搜索", ""}]
    pool = named or edits
    if not pool:
        raise RuntimeError("搜一搜页面未找到可见搜索输入框")
    # The full-page search input is normally the widest visible Edit.
    return max(pool, key=lambda control: control.rectangle().width())


def click_named_surface_control(surface, name: str) -> bool:
    for control_type in ("Button", "TabItem", "Hyperlink", "Text"):
        matches = visible_controls(surface, name=name, control_type=control_type)
        if matches:
            try:
                matches[0].click_input()
            except Exception:
                rect = matches[0].rectangle()
                mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            return True
    # Some embedded Chromium builds expose category labels as Group/ListItem
    # instead of Text or Button. Accept any visible exact-name control last.
    for control in surface.descendants():
        try:
            rect = control.rectangle()
            if (
                (control.window_text() or "").strip() == name
                and control.is_visible() and rect.width() > 2 and rect.height() > 2
            ):
                try:
                    control.click_input()
                except Exception:
                    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
                return True
        except Exception:
            continue
    return False


def ensure_video_audio_unmuted(
    surface, *, assume_audible_from_search_card: bool = False,
    toggle_generic_volume: bool = False,
) -> tuple[bool, str]:
    """Unmute current Channels video across old/new Chromium UI families."""
    root = active_video_root(surface)
    # Player controls auto-hide after reload. Hover the lower player edge and
    # re-resolve the virtualized card before looking for the volume button.
    try:
        root_rect = root.rectangle()
        mouse.move(coords=(int((root_rect.left + root_rect.right) / 2), root_rect.bottom - 28))
        time.sleep(0.25)
        root = active_video_root(surface)
    except Exception:
        pass
    # These are action labels, so clicking them always moves toward audible
    # playback. Never click a control labelled 静音/关闭声音: that means the
    # advertised action would mute an already audible video.
    for name in ("打开声音", "开启声音", "取消静音", "恢复声音"):
        if click_named_surface_control(root, name):
            return True, f"uia:{name}"
    records = []
    for control in root.descendants():
        try:
            info = control.element_info
            value = (info.name or "").strip()
            class_name = (info.class_name or "").strip()
            tokens = f"{value} {class_name}".lower()
            if info.control_type == "Button" or any(token in tokens for token in ("mute", "volume", "sound")):
                rect = control.rectangle()
                if control.is_visible() and rect.width() > 2 and rect.height() > 2:
                    records.append((control, value, class_name, rect))
        except Exception:
            continue
    for control, value, class_name, rect in records:
        tokens = f"{value} {class_name}".lower()
        if any(token in tokens for token in ("unmute", "muted", "open-sound", "volume-off")):
            try:
                control.click_input()
            except Exception:
                mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            return True, f"semantic:{value or class_name}"
    # New-account UI exposes only a generic 音量 button regardless of state.
    # Hover it first and set the revealed range above zero. This avoids blindly
    # toggling an already-audible player back to mute.
    for control, value, class_name, rect in records:
        if value == "音量":
            mouse.move(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            time.sleep(0.25)
            refreshed_root = active_video_root(surface)
            for control_type in ("Slider", "ProgressBar"):
                try:
                    sliders = refreshed_root.descendants(control_type=control_type)
                except Exception:
                    sliders = []
                for slider in sliders:
                    try:
                        slider_rect = slider.rectangle()
                        if not slider.is_visible() or slider_rect.width() < 3 or slider_rect.height() < 3:
                            continue
                        try:
                            range_value = slider.iface_range_value
                            maximum = float(range_value.CurrentMaximum)
                            minimum = float(range_value.CurrentMinimum)
                            range_value.SetValue(minimum + (maximum - minimum) * 0.8)
                        except Exception:
                            if slider_rect.height() > slider_rect.width():
                                mouse.click(coords=(int((slider_rect.left + slider_rect.right) / 2), slider_rect.top + max(2, int(slider_rect.height() * 0.2))))
                            else:
                                mouse.click(coords=(slider_rect.left + max(2, int(slider_rect.width() * 0.8)), int((slider_rect.top + slider_rect.bottom) / 2)))
                        return True, f"uia:{value}:range"
                    except Exception:
                        continue
            if assume_audible_from_search_card:
                # The card exposes an unambiguous action-labelled button and
                # WeChat carries that audible preference into the detail tab.
                # The detail's generic “音量” button has no state; blindly
                # clicking it here can toggle audible playback back to mute.
                return False, "search_card:audio_requested_unverified"
            if not toggle_generic_volume:
                # A generic 音量 button is a state-less toggle. Clicking it on
                # the first pass can mute an already-audible video. Capture the
                # current state first; only a proven-silent retry may toggle it.
                return False, "uia:音量:state_unknown_no_click"
            # Older UI has no volume range in UIA; in that family the button
            # itself is the mute toggle. This branch is entered only after the
            # first capture proved silent.
            try:
                control.click_input()
            except Exception:
                mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            return True, "uia:音量"
    print(json.dumps({
        "event": "client_audio_unmute_control_missing",
        "controls": [
            {"name": value, "class": class_name, "rect": [rect.left, rect.top, rect.right, rect.bottom]}
            for _control, value, class_name, rect in records[:30]
        ],
    }, ensure_ascii=False), flush=True)
    return False, "not_found"


def ensure_video_playing(surface) -> tuple[bool, str]:
    """Start the current video and verify the control changes to 暂停."""
    root = active_video_root(surface)
    try:
        rect = root.rectangle()
        mouse.move(coords=(int((rect.left + rect.right) / 2), rect.bottom - 28))
        time.sleep(0.2)
        root = active_video_root(surface)
    except Exception:
        pass
    play_buttons = visible_controls(root, name="播放", control_type="Button")
    if not play_buttons:
        # No 播放 action means the player is normally already running.
        if visible_controls(root, name="暂停", control_type="Button"):
            return True, "already_playing"
        try:
            root.set_focus()
            keyboard.send_keys("{SPACE}", pause=0.02)
            time.sleep(0.25)
            return bool(visible_controls(active_video_root(surface), name="暂停", control_type="Button")), "space"
        except Exception:
            return False, "missing_controls"
    button = play_buttons[0]
    button_rect = button.rectangle()
    try:
        button.click_input()
    except Exception:
        mouse.click(coords=(int((button_rect.left + button_rect.right) / 2), int((button_rect.top + button_rect.bottom) / 2)))
    time.sleep(0.3)
    refreshed = active_video_root(surface)
    if visible_controls(refreshed, name="暂停", control_type="Button"):
        return True, "uia_play"
    # Chromium occasionally accepts focus but drops invoke; retry once using
    # a physical coordinate derived from the freshly visible play button.
    retry = visible_controls(refreshed, name="播放", control_type="Button")
    if retry:
        retry_rect = retry[0].rectangle()
        mouse.click(coords=(int((retry_rect.left + retry_rect.right) / 2), int((retry_rect.top + retry_rect.bottom) / 2)))
        time.sleep(0.3)
    return bool(visible_controls(active_video_root(surface), name="暂停", control_type="Button")), "coordinate_retry"


def click_search_network_result(main, keyword: str, config: dict) -> str:
    """Open the new-account global-search row labelled 搜索网络结果."""
    candidates = []
    for control in main.descendants():
        try:
            name = re.sub(r"\s+", " ", (control.window_text() or "")).strip()
            rect = control.rectangle()
            if (
                "搜索网络结果" in name
                and control.is_visible() and rect.width() > 4 and rect.height() > 4
            ):
                candidates.append(control)
        except Exception:
            continue
    if candidates:
        control = min(candidates, key=lambda item: item.rectangle().top)
        rect = control.rectangle()
        mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
        return "uia_search_network_result"
    # The renderer may expose only one opaque MMUIRenderSubWindowHW pane. The
    # first suggestion row remains anchored below the global search box.
    entry = config.get("points", {}).get("soso_entry", [170, 95])
    rect = main.rectangle()
    scale = window_dpi_scale(main)
    mouse.click(coords=(
        rect.left + int(float(entry[0]) * scale),
        rect.top + int(float(entry[1]) * scale),
    ))
    return "anchored_search_network_result"


def click_search_category_label(surface, name: str) -> bool:
    """Click an exact category label only inside the strip below search."""
    edit_rect = find_search_edit(surface).rectangle()
    matches = []
    for control in surface.descendants():
        try:
            rect = control.rectangle()
            if (
                (control.window_text() or "").strip() == name
                and control.is_visible()
                and rect.width() > 2 and rect.height() > 2
                and edit_rect.bottom <= rect.top <= edit_rect.bottom + 70
            ):
                matches.append(control)
        except Exception:
            continue
    if not matches:
        return False
    control = min(matches, key=lambda item: item.rectangle().top)
    rect = control.rectangle()
    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
    return True


def search_page_titles(surface) -> list[str]:
    """Return visible Chromium document titles for the current 搜一搜 page."""
    titles = []
    try:
        documents = surface.descendants(control_type="Document")
    except Exception:
        documents = []
    for document in documents:
        try:
            value = (document.window_text() or "").strip()
            if value and "搜一搜" in value:
                titles.append(value)
        except Exception:
            continue
    return titles


def search_scan_surface_snapshot(surface, keyword, *, scanned_count=0, surface_generation=0):
    """Read the exact search document; never infer scroll bottom from count."""
    expected = f'{keyword} - 视频 - 搜一搜'
    snapshot = {'keyword': keyword, 'pageTitle': None, 'contextValid': False,
                'loading': False, 'loadFailed': False, 'cardCount': 0,
                'cardDigest': hashlib.sha256(b'[]').hexdigest(),
                'scannedCount': scanned_count, 'surfaceGeneration': surface_generation,
                'scrollPercent': None, 'verticallyScrollable': None,
                'capturedAt': datetime.now().astimezone().isoformat(),
                'terminalText': None, 'terminalOutsideCards': False,
                'terminalBelowCards': False, 'lastCardInViewport': False}
    try:
        documents = [doc for doc in surface.descendants(control_type='Document')
                     if (doc.window_text() or '').strip() == expected and doc.is_visible()]
        if len(documents) != 1:
            return snapshot
        document = documents[0]
        viewport = document.rectangle()
        cards = search_result_buttons(document)
        names = [re.sub(r'\s+', ' ', (card.window_text() or '').strip()) for card in cards]
        snapshot.update(pageTitle=expected, contextValid=bool(cards and all(names)), cardCount=len(cards),
                        cardDigest=hashlib.sha256(json.dumps(names, ensure_ascii=False).encode()).hexdigest())
        try:
            snapshot['verticallyScrollable'] = bool(document.iface_scroll.CurrentVerticallyScrollable)
            snapshot['scrollPercent'] = float(document.iface_scroll.CurrentVerticalScrollPercent) if snapshot['verticallyScrollable'] else 0.
        except Exception:
            pass
        last_rect = cards[-1].rectangle() if cards else None
        snapshot['lastCardInViewport'] = bool(last_rect and last_rect.width() > 0
            and last_rect.height() > 0 and last_rect.top >= viewport.top
            and last_rect.bottom <= viewport.bottom)
        for control in document.descendants():
            try:
                control_kind = control.element_info.control_type
                if control_kind not in {'Text', 'Button', 'Pane', 'ProgressBar'}:
                    continue
                if not control.is_visible():
                    continue
                name = re.sub(r'\s+', '', (control.window_text() or '').strip())
                if not name and control_kind != 'ProgressBar':
                    continue
                # A result title saying "没有更多了" is content, not page EOF.
                parent, outside = control.parent(), False
                for _ in range(12):
                    kind = parent.element_info.control_type
                    class_name = (parent.element_info.class_name or '').lower()
                    if kind == 'ListItem' or 'rich-media' in class_name:
                        break
                    if kind == 'Document':
                        outside = (parent.window_text() or '').strip() == expected
                        break
                    parent = parent.parent()
                if not outside:
                    continue
                rect = control.rectangle()
                if rect.top < viewport.top or rect.bottom > viewport.bottom:
                    continue
                if control_kind == 'ProgressBar' or name in {'加载中', '正在加载', '加载中...'}:
                    snapshot['loading'] = True
                if any(token in name for token in ('加载失败', '网络异常', '访问频繁', '操作频繁', '验证码', '安全验证', '请先登录')):
                    snapshot['loadFailed'] = True
                    if classify_restriction(name):
                        snapshot['restrictionReason'] = name
                if name in TERMINAL_TEXTS and last_rect and rect.top >= last_rect.bottom - 2:
                    snapshot.update(terminalText=name, terminalOutsideCards=True, terminalBelowCards=True)
            except Exception:
                continue
    except Exception:
        snapshot['contextValid'] = False
    if snapshot.get('restrictionReason'):
        report_from_environment(snapshot['restrictionReason'])
        raise RuntimeError('STOP_PLATFORM_RESTRICTION：'+snapshot['restrictionReason'])
    return snapshot


def is_article_results_page(surface) -> bool:
    """Verify the selected 视频 category, instead of trusting a click.

    The detached Chromium window title itself often stays as ``微信``.  The
    inner RootWebArea document is normally named ``关键词 - 视频 - 搜一搜``.
    This check prevents a coordinate drift from silently collecting 全部 or
    文章 results.
    """
    titles = search_page_titles(surface)
    return any(re.search(r"(?:^|\s-\s)视频\s-\s搜一搜$", title) for title in titles)


def is_keyword_video_results_page(surface, keyword: str) -> bool:
    expected = f"{keyword.strip()} - 视频 - 搜一搜"
    return expected in search_page_titles(surface)


def is_all_results_page(surface) -> bool:
    """Return true only for the unqualified ``全部`` search page title."""
    return any(
        title.endswith(" - 搜一搜")
        and not re.search(r"\s-\s[^-]+\s-\s搜一搜$", title)
        for title in search_page_titles(surface)
    )


def click_anchored_search_category(surface, profile: list[float | int]) -> list[int]:
    """Physically click an unnamed top category relative to the search input."""
    edit_rect = find_search_edit(surface).rectangle()
    ratio, y_offset = profile
    point = [
        edit_rect.left + int(edit_rect.width() * float(ratio)),
        edit_rect.bottom + int(float(y_offset) * window_dpi_scale(surface)),
    ]
    surface.set_focus()
    mouse.click(coords=tuple(point))
    return point


def probe_video_category_controls(surface, wait_seconds: float, max_probes: int = 12):
    """Select the 视频 tab from the horizontally virtualized category strip.

    New-account WeChat renders the search categories as unnamed ListItems and
    recent builds horizontally translate the whole strip whenever a category
    is selected.  Keeping rectangles from the first UIA snapshot therefore
    turns later clicks into a different category (朋友圈/表情/听一听 at 150%
    scaling).  Re-acquire both the surface and controls before every click and
    accept a click only when the RootWebArea title proves that 视频 is active.
    """
    probes = []
    for candidate_index in range(max(1, int(max_probes))):
        surface = get_search_window()
        if is_article_results_page(surface):
            return surface, probes
        edit_rect = find_search_edit(surface).rectangle()
        window_rect = surface.rectangle()
        candidates = []
        for control in visible_controls(surface, control_type="ListItem"):
            try:
                rect = control.rectangle()
                class_name = control.element_info.class_name or ""
                center_x = int((rect.left + rect.right) / 2)
                if (
                    "unit__item" in class_name
                    and "unit__item--ai_first" not in class_name
                    and edit_rect.bottom <= rect.top <= edit_rect.bottom + 60
                    and window_rect.left + 4 <= center_x <= window_rect.right - 4
                ):
                    candidates.append(control)
            except Exception:
                continue
        # Preserve DOM order.  Sorting by the translated screen rectangle is
        # precisely what made the old implementation lose category identity.
        if not candidates:
            break
        control = candidates[candidate_index % len(candidates)]
        rect = control.rectangle()
        point = [
            int((rect.left + rect.right) / 2),
            int((rect.top + rect.bottom) / 2),
        ]
        try:
            control.click_input()
        except Exception:
            mouse.click(coords=tuple(point))
        time.sleep(wait_seconds)
        surface = get_search_window()
        titles = search_page_titles(surface)
        probes.append({
            "candidate_index": candidate_index,
            "point": point,
            "page_titles": titles,
        })
        if is_article_results_page(surface):
            return surface, probes
    return surface, probes


def ensure_article_results_page(surface, config: dict) -> None:
    """Select 视频 and refuse to continue until the page proves it is active."""
    wait_seconds = max(0.5, float(config.get("wait_short", 0.8)))

    if is_article_results_page(surface):
        print(json.dumps({
            "event": "video_tab_verified",
            "method": "already_active",
            "page_titles": search_page_titles(surface),
        }, ensure_ascii=False), flush=True)
        return

    # Some builds expose only unnamed, horizontally virtualized ListItems.
    # Re-acquire them after every click because selecting a category translates
    # the strip and invalidates all rectangles from the previous snapshot.
    try:
        surface, probes = probe_video_category_controls(
            surface,
            wait_seconds,
            max_probes=int(config.get("category_probe_attempts", 12)),
        )
        if is_article_results_page(surface):
            print(json.dumps({
                "event": "video_tab_verified",
                "method": "uia_reacquired_category_probe",
                "probe_count": len(probes),
                "probes": probes,
                "page_titles": search_page_titles(surface),
            }, ensure_ascii=False), flush=True)
            return
    except Exception:
        pass

    embedded = config.get("embedded_web_points", {})
    version = os.environ.get("WECHAT_CLIENT_VERSION", "4.1.13")
    profiles = embedded.get("video_tab_by_version", {})
    key = next((item for item in profiles if version.startswith(item)), None)
    video_profile = profiles.get(
        key, embedded.get("video_tab_default", [0.305, 34])
    )
    attempts = []
    # Last-resort coordinate click is still title-verified.  Do not first
    # attempt an anchored 全部 click: on the translated strip that coordinate
    # can be 朋友圈/表情 and used to burn every recovery attempt.
    point = click_anchored_search_category(surface, video_profile)
    time.sleep(wait_seconds)
    surface = get_search_window()
    attempts.append({"category": "视频", "point": point, "page_titles": search_page_titles(surface)})
    if is_article_results_page(surface):
        print(json.dumps({
            "event": "video_tab_verified",
            "method": "all_baseline_then_anchored_video",
            "layout": SEARCH_LAYOUT,
            "wechat_version": version,
            "point": point,
            "page_titles": search_page_titles(surface),
        }, ensure_ascii=False), flush=True)
        return

    raise RuntimeError(
        "未能确认已切换到‘视频’栏目；为避免误采‘全部/文章’已停止。"
        + json.dumps({"attempts": attempts}, ensure_ascii=False)
    )


def _search_in_web_once(keyword: str, config: dict) -> None:
    surface = try_get_search_window()
    submitted_from_global_search = False
    if surface is None:
        main = get_wechat_window()
        # New-account UI exposes 搜一搜 directly; legacy-account UI nests it
        # under 发现. Prefer semantic UIA navigation before using coordinates.
        opened = click_named_surface_control(main, "搜一搜")
        if not opened and click_named_surface_control(main, "发现"):
            time.sleep(config.get("wait_short", 0.8))
            opened = click_named_surface_control(get_wechat_window(), "搜一搜")
        if opened:
            try:
                wait_for(has_search_window, timeout=6.0)
                surface = get_search_window()
            except Exception:
                surface = None
        if surface is None:
            # Last-resort route retained for the validated 4.1.13 account UI.
            main = get_wechat_window()
            rect = main.rectangle()
            # The opaque Qt UI exposes no semantic controls on some accounts.
            # A failed prior attempt can leave WeChat inside 发现/朋友圈, where
            # the old search-box coordinate points at the page title instead of
            # the global search input.  Return to the chat list first so the
            # following coordinate has a stable meaning at both 100% and 150%
            # Windows scaling.
            chat_tab = config.get("points", {}).get("main_chat_tab", [45, 125])
            click_relative(main, chat_tab)
            time.sleep(config.get("wait_short", 0.8))
            main = get_wechat_window()
            rect = main.rectangle()
            search_box = config.get("points", {}).get("main_search_box", [150, 55])
            scale = window_dpi_scale(main)
            mouse.click(coords=(
                rect.left + int(float(search_box[0]) * scale),
                rect.top + int(float(search_box[1]) * scale),
            ))
            time.sleep(config.get("wait_short", 0.8))
            paste_text(keyword)
            time.sleep(config.get("wait_search_suggest", 1.5))
            # WeChat 4.1.15's new-account search dropdown renders the
            # “搜索网络结果” row in an opaque Qt surface.  On 4.1.15 a bare
            # Enter is a no-op. Down selects the first ordinary query
            # suggestion and Enter opens the full 搜一搜 result page (including
            # its category strip). This was verified on the live 150%-scaled
            # client and avoids the fragile popup-row coordinate. Keep the
            # UIA/anchored click only as a bounded legacy fallback for older
            # builds.
            keyboard.send_keys("{DOWN}{ENTER}", pause=0.02)
            entry_method = "keyboard_suggestion"
            try:
                wait_for(has_search_window, timeout=6.0)
            except Exception:
                main = get_wechat_window()
                entry_method = click_search_network_result(main, keyword, config)
            print(json.dumps({
                "event": "global_search_network_result_clicked",
                "keyword": keyword,
                "method": entry_method,
            }, ensure_ascii=False), flush=True)
            wait_for(has_search_window, timeout=10.0)
            jitter_pause(
                config, "jitter_search_page_seconds", (4.0, 8.0),
                "search_page_load", keyword=keyword,
            )
            if not bool(config.get("jitter_enabled", False)):
                time.sleep(config.get("wait_page", 3.0))
            surface = get_search_window()
            submitted_from_global_search = True

    if not submitted_from_global_search:
        if not is_search_results_page(surface):
            return_to_search_results(config)
        surface = get_search_window()
        edit = find_search_edit(surface)
        edit.click_input()
        set_clipboard(keyword)
        keyboard.send_keys("^a{BACKSPACE}^v", pause=0.02)
        if not click_named_surface_control(surface, "搜索"):
            keyboard.send_keys("{ENTER}", pause=0.02)
        jitter_pause(
            config, "jitter_search_page_seconds", (4.0, 8.0),
            "search_page_load", keyword=keyword,
        )
        if not bool(config.get("jitter_enabled", False)):
            time.sleep(config.get("wait_page", 3.0))

    surface = get_search_window()
    # WeChat can transiently return an explicit empty state on the unqualified
    # search page and omit the whole category strip.  There is then no 视频 tab
    # to click.  Let the collector compare that empty claim with its durable
    # local checkpoints instead of burning all recovery attempts here.
    if is_all_results_page(surface) and has_explicit_no_search_results(surface):
        print(json.dumps({
            "event": "search_all_page_explicit_empty_without_categories",
            "keyword": keyword,
            "page_titles": search_page_titles(surface),
        }, ensure_ascii=False), flush=True)
        return
    ensure_article_results_page(surface, config)
    time.sleep(config.get("wait_page", 3.0))
    surface = get_search_window()
    if not is_article_results_page(surface):
        raise RuntimeError("‘视频’栏目校验在加载后失效，已停止采集")
    # Date-cutoff termination is safe only after explicitly selecting 最新.
    # A deeply scrolled/recovered Chromium page can omit the filter label from
    # UIA, so use the same live-search-input anchoring strategy as 文章.
    latest_method = "uia_name"
    if not click_named_surface_control(surface, "最新"):
        edit = find_search_edit(surface)
        edit_rect = edit.rectangle()
        latest_profile = config.get("embedded_web_points", {}).get(
            "latest_sort_default", [0.155, 75]
        )
        ratio, y_offset = latest_profile
        x = edit_rect.left + int(edit_rect.width() * float(ratio))
        y = edit_rect.bottom + int(float(y_offset) * window_dpi_scale(surface))
        mouse.click(coords=(x, y))
        latest_method = "anchored_fallback"
    time.sleep(config.get("wait_page", 3.0))
    print(json.dumps({
        "event": "latest_sort_selected",
        "method": latest_method,
        "page_titles": search_page_titles(get_search_window()),
    }, ensure_ascii=False), flush=True)
    # Selecting a new keyword/filter does not reliably reset Chromium's scroll
    # offset.  Establish the absolute newest-result origin before any global
    # result index is interpreted, including on checkpoint recovery.
    reset_search_results_to_top(config)


def search_in_web(keyword: str, config: dict) -> None:
    """Open the keyword's 视频/最新 page with bounded self-recovery."""
    attempts = max(1, int(config.get("search_window_recovery_attempts", 5)))
    wait_seconds = max(0.0, float(config.get("search_window_recovery_wait_seconds", 2.0)))
    errors: list[str] = []
    for attempt in range(1, attempts + 1):
        if STOP_REQUESTED:
            raise SearchWindowUnavailableError("已请求安全停止，不再重建搜一搜页面")
        try:
            if not jitter_pause(
                config, "jitter_search_submit_seconds", (5.0, 12.0),
                "before_keyword_submit", keyword=keyword, attempt=attempt,
            ):
                raise SearchWindowUnavailableError("安全停止期间不再提交搜一搜关键词")
            if not wait_for_distributed_search_token(config, keyword):
                raise SearchWindowUnavailableError("等待Redis搜索节奏令牌时收到安全停止")
            # A detached player is a per-item child surface, not another search
            # page.  Previous builds left it alive after an exception and every
            # subsequent keyword then failed with “找到多个…窗口：2”.  Keyword
            # recovery starts from a clean result surface instead.
            cleanup_stale_video_windows(
                reason=f"keyword_start:{keyword}:attempt_{attempt}",
                require_search_peer=False,
            )
            _search_in_web_once(keyword, config)
            surface = get_search_window(wait_seconds=max(2.0, wait_seconds))
            explicit_all_empty = (
                is_all_results_page(surface)
                and has_explicit_no_search_results(surface)
                and f"{keyword.strip()} - 搜一搜" in search_page_titles(surface)
            )
            if not is_keyword_video_results_page(surface, keyword) and not explicit_all_empty:
                raise RuntimeError(
                    "恢复后页面关键词或‘视频’栏目不匹配："
                    + json.dumps(search_page_titles(surface), ensure_ascii=False)
                )
            print(json.dumps({
                "event": "search_window_ready",
                "keyword": keyword,
                "attempt": attempt,
                "layout": SEARCH_LAYOUT,
                "explicit_all_empty": explicit_all_empty,
            }, ensure_ascii=False), flush=True)
            return
        except Exception as exc:
            errors.append(str(exc))
            print(json.dumps({
                "event": "search_window_recovery_retry",
                "keyword": keyword,
                "attempt": attempt,
                "max_attempts": attempts,
                "error": str(exc),
            }, ensure_ascii=False), flush=True)
            if attempt < attempts:
                time.sleep(wait_seconds)
    raise SearchWindowUnavailableError(
        f"搜一搜视频页面连续 {attempts} 次恢复失败：{errors[-1] if errors else 'unknown'}"
    )


def search_result_buttons(web):
    """Return 视频号 rich-media cards in newest-to-oldest DOM order."""
    results = []
    for control in web.descendants(control_type="ListItem"):
        try:
            class_name = (control.element_info.class_name or "").lower()
            name = (control.window_text() or "").strip()
            if "rich-media" not in class_name or "rich-media__info" in class_name:
                continue
            if name:
                results.append(control)
        except Exception:
            continue
    return results


def reset_search_results_to_top(config: dict) -> list:
    """Return the result list after proving the page is at its absolute top.

    WeChat's embedded Chromium surface preserves the scroll offset when a new
    keyword is submitted or the ``最新`` filter is selected.  On builds that
    virtualize the result list, UIA then exposes the first *currently mounted*
    card as index 0, even though it is a middle article.  Always drive the live
    document to the top before taking the index-0 snapshot.
    """
    wait_seconds = max(0.15, float(config.get("wait_short", 0.8)) / 3)
    max_rounds = max(4, int(config.get("reset_top_rounds", 12)))
    stable_rounds = 0
    previous_first = None
    last_payload = None

    for attempt in range(1, max_rounds + 1):
        surface = get_search_window()
        documents = []
        for document in surface.descendants(control_type="Document"):
            try:
                rect = document.rectangle()
                if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                    documents.append(document)
            except Exception:
                continue
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        point = (
            viewport.left + int(viewport.width() * 0.55),
            viewport.top + int(viewport.height() * 0.45),
        )
        surface.set_focus()
        # Never click the document to obtain focus here: its center is often an
        # article card, which would open a middle result before index 0 is
        # established.  Wheel input is dispatched directly at the document.
        mouse.scroll(coords=point, wheel_dist=60)
        time.sleep(wait_seconds)

        surface = get_search_window()
        results = search_result_buttons(surface)
        vertical_scroll_percent = None
        try:
            if target.iface_scroll.CurrentVerticallyScrollable:
                vertical_scroll_percent = float(target.iface_scroll.CurrentVerticalScrollPercent)
            else:
                vertical_scroll_percent = 0.0
        except Exception:
            pass
        first_signature = None
        if results:
            first = results[0]
            first_rect = first.rectangle()
            first_signature = (
                (first.window_text() or "").strip(),
                first_rect.left, first_rect.top, first_rect.right, first_rect.bottom,
            )
        if first_signature == previous_first:
            stable_rounds += 1
        else:
            stable_rounds = 0
        previous_first = first_signature
        last_payload = {
            "event": "search_results_reset_top",
            "attempt": attempt,
            "result_count": len(results),
            "first_title": first_signature[0] if first_signature else None,
            "first_rect": list(first_signature[1:]) if first_signature else None,
            "vertical_scroll_percent": vertical_scroll_percent,
            "stable_rounds": stable_rounds,
        }
        # The Scroll pattern is authoritative when exposed by RootWebArea.
        # Re-read controls only after it reports the absolute 0% position so a
        # virtualized former middle-card can never masquerade as index 0.
        scroll_pattern_proves_top = (
            vertical_scroll_percent is not None
            and vertical_scroll_percent <= 0.01
        )
        fallback_proves_top = vertical_scroll_percent is None and stable_rounds >= 2
        # An empty result page is a valid, complete keyword result. Requiring
        # at least one card made zero-hit keywords retry forever instead of
        # completing and advancing to the next keyword.
        empty_proves_top = (
            not results and stable_rounds >= 2
            and (vertical_scroll_percent is None or vertical_scroll_percent <= 0.01)
        )
        if (results and (scroll_pattern_proves_top or fallback_proves_top)) or empty_proves_top:
            print(json.dumps(last_payload, ensure_ascii=False), flush=True)
            return results

    raise RuntimeError(
        "无法确认文章结果列表已经回到顶部："
        + json.dumps(last_payload or {}, ensure_ascii=False)
    )


def precise_retry_load_config(config):
    """Cap one no-growth probe; do not change ordinary annual scan settings."""
    return {**config,
        'load_more_rounds':min(6,max(3,int(config.get('load_more_rounds',8)))),
        'load_more_search_recovery_attempts':min(1,max(0,int(config.get('load_more_search_recovery_attempts',2))))}


def load_more_search_results(
    web, previous_count: int, config: dict, keyword: str | None = None,
):
    """Drive Chromium lazy loading and return once the UIA result list grows.

    scroll_into_view alone often succeeds without emitting a wheel event, so
    WeChat keeps only the first one or two days of high-frequency results in
    the accessibility tree. Always wheel inside RootWebArea and require
    several stable bottom attempts before declaring the result set exhausted.
    """
    rounds = max(3, int(config.get("load_more_rounds", 8)))
    recovery_limit = max(
        0, int(config.get("load_more_search_recovery_attempts", 2))
    )
    recovery_count = 0
    results = search_result_buttons(web)
    attempt = 0
    allowed_rounds = rounds
    while attempt < allowed_rounds:
        attempt += 1
        try:
            if results:
                try:
                    results[-1].scroll_into_view()
                except Exception:
                    pass
            surface = get_search_window()
            documents = []
            for document in surface.descendants(control_type="Document"):
                try:
                    rect = document.rectangle()
                    if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                        documents.append(document)
                except Exception:
                    continue
            target = max(documents, key=lambda item: item.rectangle().width() * item.rectangle().height()) if documents else surface
            rect = target.rectangle()
            scroll_x = rect.left + max(40, min(rect.width() - 40, int(rect.width() * 0.55)))
            scroll_y = max(rect.top + 40, rect.bottom - 120)
            surface.set_focus()
            # Multiple wheel pulses are required by the embedded Chromium view;
            # the coordinate is guaranteed to remain inside the live document.
            mouse.scroll(coords=(scroll_x, scroll_y), wheel_dist=-15)
            jitter_pause(
                config, "jitter_scroll_load_seconds", (1.8, 4.0),
                "scroll_load_more", attempt=attempt,
            )
            if not bool(config.get("jitter_enabled", False)):
                time.sleep(max(0.5, float(config.get("wait_load_more", 1.5))))
            results = search_result_buttons(get_search_window())
        except SearchWindowUnavailableError as exc:
            if not keyword or recovery_count >= recovery_limit:
                raise SearchWindowUnavailableError(
                    f"滚动加载期间搜一搜窗口消失，恢复 {recovery_count} 次后仍失败：{exc}"
                ) from exc
            recovery_count += 1
            print(json.dumps({
                "event": "search_load_more_window_recovery",
                "keyword": keyword,
                "attempt": recovery_count,
                "max_attempts": recovery_limit,
                "previous_count": previous_count,
                "error": str(exc),
            }, ensure_ascii=False), flush=True)
            search_in_web(keyword, config)
            web = get_search_window()
            results = search_result_buttons(web)
            # A rebuilt page starts at the newest batch. Deep scans need enough
            # bounded wheel rounds to rehydrate at least ``previous_count``
            # controls before one additional batch can prove growth/exhaustion.
            restore_rounds = min(
                max(1, int(config.get("load_more_recovery_scroll_cap", 80))),
                max(rounds, previous_count // 6 + 3),
            )
            allowed_rounds = max(allowed_rounds, attempt + restore_rounds)
            continue
        print(json.dumps({
            "event": "search_load_more",
            "attempt": attempt,
            "previous_count": previous_count,
            "current_count": len(results),
            "point": [scroll_x, scroll_y],
        }, ensure_ascii=False), flush=True)
        if len(results) > previous_count:
            return results, True
    return results, False


def bring_search_result_into_view(result_index: int, config: dict):
    """Use real wheel input until a virtualized result is physically visible.

    Chromium's UIA ScrollItem pattern can report success without changing the
    page. This happened at result 31: its rectangle remained around y=4465
    while the window ended at y=1392. A physical click then silently did
    nothing. Re-read rectangles after every wheel pulse and only return once
    the target is inside the document viewport.
    """
    wait_seconds = max(0.2, float(config.get("wait_short", 0.8)) / 2)
    max_rounds = max(80, int(config.get("result_scroll_rounds", 120)))
    last_rect = None
    last_position = None
    stagnant_rounds = 0
    previous_direction = None
    direction_changes = 0
    for attempt in range(1, max_rounds + 1):
        surface = get_search_window()
        results = search_result_buttons(surface)
        if result_index >= len(results):
            return None
        result = results[result_index]
        result_rect = result.rectangle()
        documents = []
        for document in surface.descendants(control_type="Document"):
            try:
                rect = document.rectangle()
                if document.is_visible() and rect.width() > 100 and rect.height() > 100:
                    documents.append(document)
            except Exception:
                continue
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        visible_top = viewport.top + min(180, max(60, int(viewport.height() * 0.12)))
        visible_bottom = viewport.bottom - 80
        last_rect = [result_rect.left, result_rect.top, result_rect.right, result_rect.bottom]
        intersection = max(
            0,
            min(result_rect.bottom, visible_bottom) - max(result_rect.top, visible_top),
        )
        required_visible = min(40, max(16, int(result_rect.height() * 0.30)))
        if intersection >= required_visible:
            print(json.dumps({
                "event": "search_result_visible",
                "index": result_index,
                "attempt": attempt,
                "rect": last_rect,
                "scroll_direction": previous_direction or "none",
            }, ensure_ascii=False), flush=True)
            return result

        position = (result_rect.top, result_rect.bottom)
        if last_position is not None and all(
            abs(current - previous) <= 2
            for current, previous in zip(position, last_position)
        ):
            stagnant_rounds += 1
        else:
            stagnant_rounds = 0
        last_position = position

        if attempt == 1 or stagnant_rounds >= 4:
            try:
                result.scroll_into_view()
                time.sleep(wait_seconds)
                stagnant_rounds = 0
                print(json.dumps({
                    "event": "search_result_scrollitem",
                    "index": result_index,
                    "attempt": attempt,
                    "rect": last_rect,
                }, ensure_ascii=False), flush=True)
                continue
            except Exception:
                stagnant_rounds = 0

        scroll_x = viewport.left + int(viewport.width() * 0.55)
        scroll_y = max(viewport.top + 80, viewport.bottom - 160)
        # A card can be partially visible but still below the safe click area
        # (for example only 32px of a 127px card at the bottom).  Its *bottom*
        # edge determines that we must continue downward; using top > bottom
        # incorrectly selected upward and caused an endless reset/retry loop.
        direction_name = "down" if result_rect.bottom > visible_bottom else "up"
        if previous_direction and direction_name != previous_direction:
            direction_changes += 1
        previous_direction = direction_name
        if direction_name == "down":
            distance = max(0, result_rect.top - visible_bottom)
            direction = -30 if distance > 2500 else (-12 if distance > 900 else -4)
        else:
            distance = max(0, visible_top - result_rect.bottom)
            direction = 30 if distance > 2500 else (12 if distance > 900 else 4)
        surface.set_focus()
        mouse.scroll(coords=(scroll_x, scroll_y), wheel_dist=direction)
        time.sleep(wait_seconds)

    print(json.dumps({
        "event": "search_result_not_visible",
        "index": result_index,
        "last_rect": last_rect,
        "rounds": max_rounds,
        "scroll_direction": previous_direction,
        "direction_changes": direction_changes,
    }, ensure_ascii=False), flush=True)
    # The initial rectangle can be stale immediately after returning from an
    # article. Do not reverse direction in-place (which causes oscillation).
    # Restore a deterministic top position so the normal retry starts above
    # the target and scrolls downward only.
    try:
        surface = get_search_window()
        documents = [
            document for document in surface.descendants(control_type="Document")
            if document.is_visible()
        ]
        target = max(
            documents,
            key=lambda item: item.rectangle().width() * item.rectangle().height(),
        ) if documents else surface
        viewport = target.rectangle()
        point = (
            viewport.left + int(viewport.width() * 0.55),
            viewport.top + int(viewport.height() * 0.45),
        )
        surface.set_focus()
        for _ in range(4):
            mouse.scroll(coords=point, wheel_dist=60)
            time.sleep(0.15)
        print(json.dumps({
            "event": "search_scroll_reset_top",
            "index": result_index,
            "point": list(point),
        }, ensure_ascii=False), flush=True)
    except Exception:
        pass
    return None


def click_visible_search_result(result) -> list[int]:
    """Open a video card from the safe centre of its visible thumbnail.

    The current client first shows inline playback, then asynchronously opens
    a 视频号 detail tab.  Do not issue a second title click during that
    transition: it can land on the blank detail document and interrupt its
    hydration.  The caller separately waits for both navigation and actions.
    """
    surface = get_search_window()
    documents = [
        document for document in surface.descendants(control_type="Document")
        if document.is_visible()
    ]
    viewport = max(
        documents,
        key=lambda item: item.rectangle().width() * item.rectangle().height(),
    ).rectangle() if documents else surface.rectangle()
    rect = result.rectangle()
    visible_top = max(rect.top, viewport.top + min(180, max(60, int(viewport.height() * 0.12))))
    visible_bottom = min(rect.bottom, viewport.bottom - 80)
    if visible_bottom <= visible_top:
        raise RuntimeError(f"结果卡片没有安全可点击区域：{rect}")
    point = [
        int((rect.left + rect.right) / 2),
        int(visible_top + (visible_bottom - visible_top) * 0.38),
    ]
    surface.set_focus()
    mouse.click(coords=tuple(point))
    return point


def enable_search_card_audio(result) -> bool:
    """Enable sound through the card's stateful action before detail opens."""
    labels_seen = []
    for control in result.descendants():
        try:
            label = (control.window_text() or "").strip()
            if label:
                labels_seen.append(label)
            if label not in {"打开声音", "开启声音"}:
                continue
            rect = control.rectangle()
            if not control.is_visible() or rect.width() <= 2 or rect.height() <= 2:
                continue
            # Chromium's nested sound button sometimes rejects Invoke/click_input
            # after virtualization even though its rectangle is visible. A
            # physical click on the freshly read rectangle is deterministic.
            mouse.click(coords=(
                int((rect.left + rect.right) / 2),
                int((rect.top + rect.bottom) / 2),
            ))
            time.sleep(0.25)
            print(json.dumps({
                "event": "search_card_audio_enabled",
                "control": label,
            }, ensure_ascii=False), flush=True)
            return True
        except Exception:
            continue
    print(json.dumps({
        "event": "search_card_audio_control_missing",
        "labels": labels_seen[:20],
    }, ensure_ascii=False), flush=True)
    return False


def control_element_records(root, include_self: bool = False) -> list[tuple[str, str, str]]:
    records = []
    controls = ([root] if include_self else []) + list(root.descendants())
    for control in controls:
        try:
            info = control.element_info
            value = (info.name or "").strip()
            if value:
                records.append((value, info.control_type or "", info.class_name or ""))
        except Exception:
            continue
    return records


def page_element_records(web) -> list[tuple[str, str, str]]:
    return control_element_records(web)


def parse_publish_time(records: list[tuple[str, str, str]]) -> datetime | None:
    pattern = re.compile(r"(20\d{2})[年/-](\d{1,2})[月/-](\d{1,2})日?(?:\s+(\d{1,2}):(\d{2}))?")
    # A listing description can contain contract/validity dates.  Only accept
    # a short, date-like UIA value (normally the dedicated publish-time Text),
    # never the first date embedded in the whole description/document text.
    for value, _control_type, _class_name in records:
        cleaned = re.sub(r"\s+", " ", value).strip()
        if len(cleaned) > 40:
            continue
        match = re.fullmatch(
            rf"(?:发布于?\s*)?{pattern.pattern}(?:\s*发布)?",
            cleaned,
        )
        if match:
            return datetime(
                int(match.group(1)), int(match.group(2)), int(match.group(3)),
                int(match.group(4) or 0), int(match.group(5) or 0),
            )
    return None


def parse_search_card_time(text: str, now: datetime | None = None) -> datetime | None:
    """Convert the date label embedded in a 视频 search card to a timestamp."""
    now = now or datetime.now()
    value = re.sub(r"\s+", " ", str(text or "")).strip()
    tail = re.search(rf"(?:^|\s)({CARD_TIME_PATTERN.pattern})\s*$", value)
    value = tail.group(1) if tail else value
    absolute = re.fullmatch(r"(20\d{2})[年/-](\d{1,2})[月/-](\d{1,2})日?", value)
    if absolute:
        return datetime(int(absolute.group(1)), int(absolute.group(2)), int(absolute.group(3)))
    month_day = re.fullmatch(r"(\d{1,2})月(\d{1,2})日", value)
    if month_day:
        value = datetime(now.year, int(month_day.group(1)), int(month_day.group(2)))
        return value.replace(year=now.year - 1) if value > now + timedelta(days=1) else value
    if value == "前天":
        return now - timedelta(days=2)
    if value == "昨天":
        return now - timedelta(days=1)
    if value == "刚刚":
        return now
    relative = re.fullmatch(r"(\d+)\s*(分钟|小时|天|周|个月|年)前", value)
    if not relative:
        return None
    amount = int(relative.group(1))
    unit = relative.group(2)
    if unit == "分钟":
        return now - timedelta(minutes=amount)
    if unit == "小时":
        return now - timedelta(hours=amount)
    if unit == "天":
        return now - timedelta(days=amount)
    if unit == "周":
        return now - timedelta(weeks=amount)
    if unit == "个月":
        month_index = now.year * 12 + now.month - 1 - amount
        year, month_zero = divmod(month_index, 12)
        month = month_zero + 1
        return now.replace(year=year, month=month, day=min(now.day, calendar.monthrange(year, month)[1]))
    try:
        return now.replace(year=now.year - amount)
    except ValueError:
        return now.replace(year=now.year - amount, day=28)


def count_from_button(records: list[tuple[str, str, str]], class_token: str, zero_when_blank: bool) -> int | None:
    matched_button = False
    for value, control_type, class_name in records:
        if control_type == "Button" and class_token in class_name:
            matched_button = True
            parsed = parse_count(value)
            if parsed is not None:
                return parsed
            # A WeChat article may expose both a hidden empty button and the
            # visible metric button. Keep scanning instead of treating the
            # first empty accessibility node as a real zero.
    if matched_button and zero_when_blank:
        return 0
    return None


def _semantic_metric(records: list[tuple[str, str, str]], labels: tuple[str, ...], class_tokens: tuple[str, ...]) -> int | None:
    for value, _control_type, class_name in records:
        haystack = f"{value} {class_name}".lower()
        if any(token.lower() in haystack for token in class_tokens):
            parsed = parse_count(value)
            if parsed is not None:
                return parsed
        for label in labels:
            match = re.search(rf"(?:{re.escape(label)})\s*[:：]?\s*(\d+(?:\.\d+)?\s*[万wW]?)", value)
            if match:
                return parse_count(match.group(1))
    return None


def parse_article_metrics(records: list[tuple[str, str, str]]) -> tuple[int | None, int | None, int | None, int | None]:
    """Map the four visible Channels actions to the legacy ODS columns.

    视频号详情页从左到右为点赞、转发、推荐、评论。ODS 没有“推荐数”列，
    因此沿用公众号约定把推荐数写入 ``文章收藏数``。
    """
    action_values = [
        value for value, control_type, class_name in records
        if control_type == "Button" and "op-item" in class_name
    ]
    if len(action_values) >= 4:
        # PC 视频号当前顺序固定为：喜欢、转发、推荐（UIA 名称为点赞）、评论。
        # A visible action with no numeric suffix is the UI's real zero state.
        like_count = parse_count(action_values[0]) or 0
        share_count = parse_count(action_values[1]) or 0
        favorite_count = parse_count(action_values[2]) or 0
        comment_count = parse_count(action_values[3]) or 0
        return like_count, comment_count, favorite_count, share_count
    like_count = _semantic_metric(records, ("喜欢",), ("like", "praise"))
    share_count = _semantic_metric(records, ("转发", "分享"), ("share", "forward"))
    favorite_count = _semantic_metric(records, ("推荐", "收藏"), ("recommend", "favorite", "love"))
    comment_count = _semantic_metric(records, ("评论", "留言"), ("comment", "reply"))
    return like_count, comment_count, favorite_count, share_count


def active_video_root(web):
    """Return the feed item occupying most of the current 视频号 viewport."""
    documents = [
        item for item in web.descendants(control_type="Document")
        if (item.window_text() or "").strip() == "视频号"
    ]
    viewport = documents[0].rectangle() if documents else web.rectangle()
    candidates = []
    for control in web.descendants(control_type="Group"):
        try:
            if (control.element_info.class_name or "").strip() != "relative h-full w-full":
                continue
            rect = control.rectangle()
            overlap = max(0, min(rect.bottom, viewport.bottom) - max(rect.top, viewport.top))
            if overlap:
                candidates.append((overlap, control))
        except Exception:
            continue
    return max(candidates, key=lambda item: item[0])[1] if candidates else web


def active_video_records(web) -> list[tuple[str, str, str]]:
    """Read only the active feed item, excluding previous/next virtual cards."""
    root = active_video_root(web)
    if root is web:
        raise VideoDetailIdentityMismatchError('Cannot isolate visible feed item; refusing whole-window residual metadata')
    viewport = root.rectangle()
    records = []
    for control in root.descendants():
        try:
            rect = control.rectangle()
            if not control.is_visible() or min(rect.bottom, viewport.bottom) <= max(rect.top, viewport.top):
                continue
            info = control.element_info
            value = (info.name or '').strip()
            if value:
                records.append((value, info.control_type or '', info.class_name or ''))
        except Exception:
            continue
    return records


def pause_active_video_if_running(web) -> bool:
    """Freeze the opened feed item before metadata/comment extraction.

    Short videos can auto-advance in six seconds.  Without pausing, link and
    comment operations may silently move to the next unrelated feed item.
    Returning True for an already-paused ``播放`` control makes this helper
    safe to call repeatedly during retries.
    """
    root = active_video_root(web)
    buttons = [
        button for button in root.descendants(control_type="Button")
        if button.is_visible()
    ]
    # Virtual feed items retain hidden controls from videos that were already
    # visited.  A hidden ``播放`` button is not evidence that the current card
    # is paused; accepting it caused six-second videos to advance while the
    # collector was scrolling comments.
    play_buttons = [
        button for button in buttons
        if (button.window_text() or "").strip() == "播放"
    ]
    if play_buttons:
        progress = None
        for value, _control_type, _class_name in page_element_records(root):
            match = re.fullmatch(
                r"(?:(\d+):)?(\d+):(\d{2})\s*/\s*"
                r"(?:(\d+):)?(\d+):(\d{2})",
                value.strip(),
            )
            if not match:
                continue
            elapsed = int(match.group(1) or 0) * 3600 + int(match.group(2)) * 60 + int(match.group(3))
            total = int(match.group(4) or 0) * 3600 + int(match.group(5)) * 60 + int(match.group(6))
            progress = (elapsed, total)
            break
        if progress and progress[1] > 0 and progress[0] >= progress[1] - 1:
            # ``播放`` at 0:06/0:06 means ended, not safely paused.  Seeking
            # back to zero cancels the pending auto-next transition while
            # leaving the player paused for link/comment extraction.
            reset = seek_active_video_to_start(web)
            print(json.dumps({
                "event": "video_end_reset_before_metadata",
                "elapsed": progress[0],
                "duration": progress[1],
                "reset": reset,
            }, ensure_ascii=False), flush=True)
            return reset
        return True
    pauses = [
        button for button in buttons
        if (button.window_text() or "").strip() == "暂停"
    ]
    player_fallback = not pauses
    if player_fallback:
        # Native video controls are hidden until hover on some builds.  The
        # full video-player group is still exposed and a centre click toggles
        # play/pause without waiting for the control bar to appear.
        players = []
        for control in root.descendants(control_type="Group"):
            try:
                class_name = (control.element_info.class_name or "").lower()
                rect = control.rectangle()
                if (
                    ("feed-video" in class_name or (control.window_text() or "").strip() == "Video Player")
                    and control.is_visible() and rect.width() > 200 and rect.height() > 200
                ):
                    players.append(control)
            except Exception:
                continue
        if not players:
            return False
        control = max(players, key=lambda item: item.rectangle().width() * item.rectangle().height())
    else:
        control = pauses[0]
    rect = control.rectangle()
    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
    time.sleep(0.15)
    if player_fallback:
        refreshed = active_video_root(get_video_detail_window())
        refreshed_pauses = [
            button for button in refreshed.descendants(control_type="Button")
            if (button.window_text() or "").strip() == "暂停"
        ]
        if refreshed_pauses:
            pause_rect = refreshed_pauses[0].rectangle()
            mouse.click(coords=(
                int((pause_rect.left + pause_rect.right) / 2),
                int((pause_rect.top + pause_rect.bottom) / 2),
            ))
            time.sleep(0.15)
    paused = bool(visible_controls(active_video_root(get_video_detail_window()),
                                   name='播放', control_type='Button'))
    if not paused:
        print(json.dumps({'event': 'video_pause_unverified',
            'reason': 'play_control_not_observed_after_pause_action'}, ensure_ascii=False), flush=True)
    return paused


def parse_article_details(records: list[tuple[str, str, str]]) -> dict:
    """Extract ODS-ready fields exposed by the 视频号 accessibility tree."""
    published = parse_publish_time(records)
    date_index = next((i for i, (value, _type, _class) in enumerate(records)
                       if re.search(r"20\d{2}年\d{1,2}月\d{1,2}日", value)), None)
    author_name = None
    candidates = records[:date_index] if date_index is not None else records
    for value, control_type, class_name in reversed(candidates):
        lowered = class_name.lower()
        if control_type in {"Button", "Hyperlink", "Text"} and any(token in lowered for token in ("author", "nickname", "finder", "profile")):
            if value.strip() not in {"关注", "+关注", "展开"}:
                author_name = value.strip()
                break
    location = None
    for value, _control_type, class_name in records:
        if any(token in class_name.lower() for token in ("location", "ip")) and value.strip():
            location = value.strip()
            break

    excluded = {"关注", "+关注", "展开", "收起", "转发给朋友", "分享到朋友圈", "收藏", "复制链接"}
    content_lines = []
    for value, control_type, class_name in records:
        if control_type not in {"Text", "Hyperlink"}:
            continue
        cleaned = re.sub(r"\s+", " ", value).strip()
        if not cleaned or cleaned in excluded or re.fullmatch(r"\d+(?:\.\d+)?[万wW]?", cleaned):
            continue
        if re.search(r"20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}", cleaned):
            continue
        if any(token in class_name.lower() for token in ("comment", "reply", "toolbar", "action")):
            continue
        content_lines.append(cleaned)
    content = "\n".join(dict.fromkeys(content_lines)).strip() or None
    image_count = sum(1 for _value, control_type, class_name in records
                      if control_type == "Image" and "avatar" not in class_name.lower())
    video_count = max(1, sum(1 for _value, control_type, class_name in records
                      if control_type in {"Video", "Media"} or "video" in class_name.lower()))
    return {
        "author_name": author_name,
        "published_at": published.isoformat(sep=" ", timespec="seconds") if published else None,
        "location": location,
        "content": content,
        "image_count": image_count,
        "video_count": video_count,
    }


def is_valid_channels_share_url(value: str | None) -> bool:
    parsed = urlparse(str(value or "").strip())
    if parsed.scheme != "https":
        return False
    if parsed.hostname == "channels.weixin.qq.com":
        return True
    return parsed.hostname == "weixin.qq.com" and parsed.path.startswith("/sph/")


def wait_for_channels_clipboard(timeout: float = 3.0) -> str:
    """Allow asynchronous copy/brief clipboard contention, never log contents."""
    deadline = time.monotonic() + max(0.0, min(float(timeout), 5.0))
    reads = 0
    while True:
        reads += 1
        try:
            value = get_clipboard().strip()
        except Exception:
            value = ""
        if is_valid_channels_share_url(value):
            if reads > 1:
                print(json.dumps({"event": "video_link_clipboard_delayed_ready", "reads": reads}), flush=True)
            return value
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            # Clipboard can hold passwords or other unrelated user material.
            # Never include it in task logs or remote monitoring.
            raise RuntimeError("复制链接超时或格式无效，将重新打开当前卡片；剪贴板内容不写入日志")
        time.sleep(min(0.1, remaining))


def copy_article_link_from_menu(web, config: dict) -> str:
    # 视频详情隐藏地址栏 URL，但第二个互动按钮会展开四项分享菜单；
    # 第四项固定为“复制链接”。使用物理点击规避该 React 菜单的空 UIA Invoke。
    web.set_focus()
    root = active_video_root(web)
    # Legacy detached players tolerated an unconditional ESC here because it
    # only dismissed a share overlay.  In the new-account embedded player the
    # same key closes/hides the entire detail surface before actions are read,
    # leaving zero op-item buttons and making recovery launch a login window.
    # Dismiss ESC only when an actual share menu is already open.
    open_menu_items = [
        control for control in root.descendants(control_type="Group")
        if (control.element_info.class_name or "") == "menu-item"
        and control.is_visible()
    ]
    if open_menu_items:
        keyboard.send_keys("{ESC}", pause=0.02)
        time.sleep(0.2)
        root = active_video_root(get_video_detail_window())
    actions = [
        control for control in root.descendants(control_type="Button")
        if "op-item" in (control.element_info.class_name or "")
    ]
    if len(actions) < 2:
        raise RuntimeError(f"当前视频未找到转发按钮：互动按钮数={len(actions)}")
    # Never accept a valid-looking URL left by the preceding video when the
    # React menu misses a click. A blank clipboard turns that silent mismatch
    # into a retriable collection error instead of a wrong card/link pairing.
    set_clipboard("")
    rect = actions[1].rectangle()
    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))

    def copy_menu_item():
        refreshed = active_video_root(get_video_detail_window())
        share_buttons = [
            control for control in refreshed.descendants(control_type="Button")
            if "op-item" in (control.element_info.class_name or "")
            and (control.window_text() or "").strip().startswith(("分享", "转发"))
        ]
        for share_button in share_buttons:
            items = [
                control for control in share_button.descendants(control_type="Group")
                if (control.element_info.class_name or "") == "menu-item"
            ]
            # Current and legacy clients both expose the menu in the order:
            # 转发给朋友、分享到朋友圈、收藏、复制链接.  On 4.1.15 the
            # parent button keeps the accessible name “分享” instead of
            # concatenating the four child labels, so inspect its descendants.
            if len(items) >= 4:
                return items[3]
        return None

    try:
        item = wait_for(copy_menu_item, timeout=4.0)
    except Exception as exc:
        raise RuntimeError("视频详情分享菜单未出现或未识别到复制链接") from exc
    rect = item.rectangle()
    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
    return wait_for_channels_clipboard()


def parse_visible_comments(records: list[tuple[str, str, str]]) -> list[dict]:
    """Best-effort extraction of comments currently exposed by UIA.

    The viewer virtualizes the comment pane.  The collector stores every
    visible comment on each pass and SQLite upserts them, so retries and later
    scroll passes are safe.  Unknown fields stay NULL rather than becoming 0.
    """
    comments = []
    for value, control_type, class_name in records:
        lowered = class_name.lower()
        if "comment" not in lowered or control_type not in {"ListItem", "Group", "Text"}:
            continue
        text = re.sub(r"\s+", " ", value).strip()
        if not text or text in {"评论", "回复", "发表"}:
            continue
        comments.append({"content": text})
    return comments


def parse_comment_card(records: list[tuple[str, str, str]], comment_id: str | None = None) -> dict | None:
    """Parse one UIA comment container without inventing unavailable fields."""
    user_name = content = published_at = location = None
    like_count = reply_count = None
    ignored = {"评论", "回复", "发表", "点赞", "展开", "收起"}
    for value, control_type, class_name in records:
        text = re.sub(r"\s+", " ", value).strip()
        lowered = class_name.lower()
        if not text or text in ignored:
            continue
        if user_name is None and any(token in lowered for token in ("nickname", "user-name", "username", "author")):
            user_name = text
            continue
        if published_at is None and (
            any(token in lowered for token in ("time", "date"))
            or re.fullmatch(r"(?:刚刚|昨天|前天|\d+\s*(?:分钟|小时|天|周|个月|年)前|20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}日?)", text)
        ):
            published_at = text
            continue
        if location is None and ("location" in lowered or "ip" in lowered or text.startswith("IP属地")):
            location = text.removeprefix("IP属地：").removeprefix("IP属地:").strip() or None
            continue
        if any(token in lowered for token in ("like", "praise")):
            like_count = parse_count(text)
            continue
        if "reply" in lowered:
            reply_count = parse_count(text)
            continue
        if content is None and any(token in lowered for token in ("content", "comment-text", "comment_text")):
            content = text
    if content is None:
        # Some WeChat versions expose only a flattened comment ListItem. Keep
        # the source text verbatim rather than fabricating user/time fields.
        flattened = next((
            re.sub(r"\s+", " ", value).strip()
            for value, control_type, class_name in records
            if control_type in {"ListItem", "Group"}
            and "comment" in class_name.lower()
            and re.sub(r"\s+", " ", value).strip() not in ignored
        ), None)
        content = flattened
    if not content:
        return None
    return {
        "comment_id": comment_id,
        "user_url": None,
        "user_id": None,
        "user_name": user_name,
        "content": content,
        "published_at": published_at,
        "location": location,
        "like_count": like_count,
        "reply_count": reply_count,
    }


def parse_new_comment_panel(surface) -> tuple[list[dict], int | None]:
    """Parse the 4.1.15 side-panel whose comment fields are flat siblings."""
    panel_titles = []
    for control in surface.descendants(control_type="Text"):
        try:
            if (control.window_text() or "").strip() == "评论" and control.is_visible():
                panel_titles.append(control)
        except Exception:
            continue
    if not panel_titles:
        return [], None
    panel_left = max(item.rectangle().left for item in panel_titles)
    content_controls = []
    user_controls = []
    visible_texts = []
    for control in surface.descendants():
        try:
            rect = control.rectangle()
            if not control.is_visible() or rect.left < panel_left:
                continue
            class_name = (control.element_info.class_name or "").lower()
            control_type = control.element_info.control_type or ""
            if control_type == "Group" and "comment-content" in class_name:
                content_controls.append(control)
            elif control_type == "Group" and "comment-user-name" in class_name:
                user_controls.append(control)
            elif control_type == "Text":
                visible_texts.append(control)
        except Exception:
            continue
    comments = []
    for content_control in sorted(content_controls, key=lambda item: item.rectangle().top):
        content_rect = content_control.rectangle()
        lineage_classes = []
        container = content_control
        for _ in range(6):
            try:
                candidate = container.parent()
                if candidate is None:
                    break
                candidate_class = (candidate.element_info.class_name or "").strip()
                if candidate_class:
                    lineage_classes.append(candidate_class)
                container = candidate
                if "comment-item" in candidate_class.lower() or "comment_item" in candidate_class.lower():
                    break
            except Exception:
                break
        content_parts = []
        for child in content_control.descendants():
            try:
                if child.element_info.control_type not in {"Text", "Hyperlink"}:
                    continue
                value = re.sub(r"\s+", " ", child.window_text() or "").strip()
                if value:
                    content_parts.append(value)
            except Exception:
                continue
        # Chromium exposes wrapped inline fragments in DOM reading order;
        # sorting by rectangles corrupts links whose bounding box spans two
        # visual lines, so preserve the accessibility enumeration order.
        content = "".join(content_parts).strip()
        if not content:
            continue
        preceding_users = [
            item for item in user_controls
            if 0 <= content_rect.top - item.rectangle().top <= 100
        ]
        user_name = None
        user_top = content_rect.top - 60
        if preceding_users:
            user = min(preceding_users, key=lambda item: content_rect.top - item.rectangle().top)
            user_top = user.rectangle().top
            names = [
                re.sub(r"\s+", " ", child.window_text() or "").strip()
                for child in user.descendants(control_type="Text")
                if re.sub(r"\s+", " ", child.window_text() or "").strip()
            ]
            user_name = names[0] if names else None
        metadata = []
        for text_control in visible_texts:
            try:
                rect = text_control.rectangle()
                value = re.sub(r"\s+", " ", text_control.window_text() or "").strip()
                if value and user_top - 4 <= rect.top < content_rect.top:
                    metadata.append(value)
            except Exception:
                continue
        published_at = next((value for value in metadata if re.fullmatch(
            r"(?:刚刚|昨天|前天|\d+\s*(?:分钟|小时|天|周|个月|年)前|20\d{2}[年/-]\d{1,2}[月/-]\d{1,2}日?)",
            value,
        )), None)
        if published_at:
            parsed_comment_time = parse_search_card_time(published_at)
            if parsed_comment_time is not None:
                published_at = parsed_comment_time.isoformat(sep=" ", timespec="seconds")
        ignored = {user_name, published_at, "作者", "置顶"}
        location = next((value for value in metadata if value not in ignored), None)
        signature = hashlib.sha256(
            f"{user_name or ''}|{content}".encode("utf-8")
        ).hexdigest()
        like_count = reply_count = None
        try:
            container_records = control_element_records(container, include_self=True)
            for value, _control_type, class_name in container_records:
                text = re.sub(r"\s+", " ", value).strip()
                lowered = class_name.lower()
                if like_count is None and any(token in lowered for token in ("like", "praise")):
                    like_count = parse_count(text)
                if reply_count is None and "reply" in lowered:
                    reply_count = parse_count(text)
                if reply_count is None:
                    reply_match = re.search(r"(?:展开|查看|共)?\s*(\d+)\s*条?回复", text)
                    if reply_match:
                        reply_count = int(reply_match.group(1))
        except Exception:
            pass
        comments.append({
            "comment_id": signature,
            "user_url": None,
            "user_id": None,
            "user_name": user_name,
            "content": content,
            "published_at": published_at,
            "location": location,
            "like_count": like_count,
            "reply_count": reply_count,
            "_left": content_rect.left,
            "_top": content_rect.top,
            "_lineage": " ".join(lineage_classes),
        })
    if comments:
        parent_left = min(int(item.get("_left") or 0) for item in comments)
        for item in comments:
            lineage = str(item.get("_lineage") or "").lower()
            indented = int(item.get("_left") or 0) >= parent_left + 12
            semantic_child = any(token in lineage for token in (
                "reply-item", "reply_item", "sub-comment", "sub_comment", "child-comment", "child_comment",
            ))
            item["_depth"] = 1 if indented or semantic_child else 0
    return comments, panel_left


def expand_visible_comment_replies(surface, panel_left: int | None = None) -> int:
    """Expand visible child-reply groups without submitting any new content."""
    pattern = re.compile(r"(?:展开|查看|更多|共)?\s*\d+\s*条?回复|展开回复|查看回复")
    candidates = []
    for control in surface.descendants():
        try:
            value = re.sub(r"\s+", " ", control.window_text() or "").strip()
            rect = control.rectangle()
            if (
                value and pattern.search(value)
                and control.is_visible() and rect.width() > 3 and rect.height() > 3
                and (panel_left is None or rect.left >= panel_left)
            ):
                candidates.append(control)
        except Exception:
            continue
    clicked = 0
    seen_rects = set()
    for control in sorted(candidates, key=lambda item: item.rectangle().top)[:20]:
        try:
            rect = control.rectangle()
            rect_key = (rect.left, rect.top, rect.right, rect.bottom)
            if rect_key in seen_rects:
                continue
            seen_rects.add(rect_key)
            mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
            clicked += 1
            time.sleep(0.2)
        except Exception:
            continue
    return clicked


def merge_parent_child_comments(items: list[dict]) -> list[dict]:
    """Return one ODS row per parent thread, summing parent/child likes."""
    if not items:
        return []
    ordered = sorted(items, key=lambda item: (int(item.get("_sequence") or 0), int(item.get("_top") or 0)))
    roots: dict[str, dict] = {}
    root_order = []
    last_root_id = None
    for item in ordered:
        item_id = str(item.get("comment_id") or hashlib.sha256(
            f"{item.get('user_name') or ''}|{item.get('content') or ''}".encode("utf-8")
        ).hexdigest())
        depth = int(item.get("_depth") or 0)
        root_id = str(item.get("_root_comment_id") or "") or None
        if depth <= 0 or not last_root_id:
            root_id = item_id
            last_root_id = root_id
            if root_id not in roots:
                roots[root_id] = {"parent": dict(item), "children": [], "child_ids": set()}
                root_order.append(root_id)
            continue
        root_id = root_id or last_root_id
        if root_id not in roots:
            # Never discard an orphaned visible reply; degrade it to a parent.
            roots[root_id] = {"parent": dict(item), "children": [], "child_ids": set()}
            root_order.append(root_id)
            last_root_id = root_id
            continue
        if item_id not in roots[root_id]["child_ids"]:
            roots[root_id]["child_ids"].add(item_id)
            roots[root_id]["children"].append(dict(item))

    merged = []
    for root_id in root_order:
        thread = roots[root_id]
        parent = thread["parent"]
        children = thread["children"]
        parent_name = parent.get("user_name") or "未知用户"
        parts = [f"【父评论｜{parent_name}】{parent.get('content') or ''}"]
        for index, child in enumerate(children, start=1):
            child_name = child.get("user_name") or "未知用户"
            parts.append(f"【子评论{index}｜{child_name}】{child.get('content') or ''}")
        known_likes = [
            int(value) for value in [parent.get("like_count"), *(child.get("like_count") for child in children)]
            if value is not None
        ]
        page_reply_count = parent.get("reply_count")
        actual_reply_count = len(children)
        result = {key: value for key, value in parent.items() if not key.startswith("_")}
        result["comment_id"] = parent.get("comment_id") or root_id
        result["content"] = "".join(parts)
        result["like_count"] = sum(known_likes) if known_likes else 0
        result["reply_count"] = max(int(page_reply_count or 0), actual_reply_count)
        merged.append(result)
    return merged


def is_article_caption_panel_entry(content: str | None, article_title: str | None) -> bool:
    """Return True for the author/post summary pinned above real comments."""
    comment_key = normalize_video_identity_text(content)
    title_key = normalize_video_identity_text(article_title)
    if len(title_key) < 6 or not comment_key:
        return False
    if title_key in comment_key:
        return True
    return SequenceMatcher(None, title_key, comment_key).ratio() >= 0.88


def filter_actual_comment_entries(
    items: list[dict],
    article_title: str | None,
    *,
    structured: bool = False,
) -> list[dict]:
    """Remove pinned captions only from ambiguous legacy card fallbacks.

    The 4.1.15 ``comment-content`` group is a real comment.  An author may
    legitimately post a comment identical to the video caption, so content
    similarity must not discard structured entries.
    """
    if structured:
        return list(items)
    return [
        item for item in items
        if not is_article_caption_panel_entry(item.get("content"), article_title)
    ]


def record_comment_capture_result(config, *, status, expected, actual, reason):
    audit = {'event': 'comment_capture_result', 'status': status,
             'expectedCount': expected, 'actualEntries': actual, 'reason': reason,
             'emptyConfirmed': status == 'confirmed_empty'}
    config['_comment_capture_audit'] = audit
    print(json.dumps(audit, ensure_ascii=False), flush=True)


def comment_panel_empty_confirmed(surface):
    # Never infer empty from a timeout, a zero parsed count, or whole-window text.
    for group in surface.descendants(control_type='Group'):
        if not group.is_visible() or 'comment' not in (group.element_info.class_name or '').lower():
            continue
        for text in group.descendants(control_type='Text'):
            if text.is_visible() and (text.window_text() or '').strip() in {
                '暂无评论', '暂时没有评论', '还没有评论', '暂无评论，快来抢沙发吧',
            }:
                return True
    return False


def collect_video_comments(
    web,
    config: dict,
    expected_count: int | None = None,
    article_title: str | None = None,
) -> list[dict]:
    """Open and exhaust the visible comment pane with bounded scrolling."""
    config.pop('_comment_capture_audit', None)
    opened = False
    attempts = min(3, max(1, int(config.get('comment_panel_open_attempts', 2))))
    for attempt in range(1, attempts + 1):
        ensure_capture_not_stopped(config)
        current = get_video_detail_window()
        opened = open_video_comments_panel(current, config)
        if comment_panel_empty_confirmed(get_video_detail_window()):
            # A positive page count conflicts with an empty placeholder.
            status = 'incomplete' if expected_count and expected_count > 0 else 'confirmed_empty'
            record_comment_capture_result(config, status=status, expected=expected_count,
                actual=0, reason='page_count_conflicts_with_empty' if status == 'incomplete' else 'explicit_empty_panel')
            return []
        if opened:
            break
        print(json.dumps({'event': 'comment_panel_retry', 'attempt': attempt,
            'maxAttempts': attempts, 'ready': False}, ensure_ascii=False), flush=True)
    if not opened:
        record_comment_capture_result(config, status='unavailable', expected=expected_count,
            actual=0, reason='panel_not_ready_after_bounded_retry')
        return []
    max_rounds = max(1, int(config.get("comment_scroll_rounds", 20)))
    stable_required = max(1, int(config.get("comment_stable_rounds", 3)))
    wait_seconds = max(0.4, float(config.get("comment_scroll_wait_seconds", 0.8)))
    found: dict[str, dict] = {}
    stable_rounds = 0
    last_root_id = None
    sequence = 0
    for round_index in range(1, max_rounds + 1):
        ensure_capture_not_stopped(config)
        surface = get_video_detail_window()
        new_comments, panel_left = parse_new_comment_panel(surface)
        expanded_replies = expand_visible_comment_replies(surface, panel_left)
        if expanded_replies:
            time.sleep(wait_seconds)
            surface = get_video_detail_window()
            new_comments, panel_left = parse_new_comment_panel(surface)
        before = len(found)
        actual_new_comments = filter_actual_comment_entries(
            new_comments,
            article_title,
            structured=True,
        )
        for item in sorted(actual_new_comments, key=lambda entry: int(entry.get("_top") or 0)):
            item_id = str(item.get("comment_id") or hashlib.sha256(
                f"{item.get('user_name') or ''}|{item.get('content') or ''}".encode("utf-8")
            ).hexdigest())
            if int(item.get("_depth") or 0) <= 0 or not last_root_id:
                last_root_id = item_id
                item["_depth"] = 0
            item["_root_comment_id"] = last_root_id
            item["_sequence"] = sequence
            sequence += 1
            signature = "|".join(str(item.get(key) or "") for key in (
                "comment_id", "user_name", "content", "published_at"
            ))
            found.setdefault(hashlib.sha256(signature.encode("utf-8")).hexdigest(), item)
        root = surface
        candidates = []
        for control_type in ("ListItem", "Group"):
            for control in root.descendants(control_type=control_type):
                try:
                    class_name = (control.element_info.class_name or "").lower()
                    rect = control.rectangle()
                    if "comment" in class_name and control.is_visible() and rect.width() > 20 and rect.height() > 12:
                        candidates.append(control)
                except Exception:
                    continue
        # Prefer leaf-like ListItems. Group fallback covers older clients.
        list_items = [item for item in candidates if item.element_info.control_type == "ListItem"]
        candidates = list_items or candidates
        # Newer WeChat panels expose the pinned video caption through the
        # same comment-content class as real comments.  If that caption is
        # the only structured entry, still parse the surrounding comment
        # cards; otherwise a video showing one real comment exports zero.
        for control in ([] if actual_new_comments else candidates):
            try:
                automation_id = (control.element_info.automation_id or "").strip() or None
                item = parse_comment_card(control_element_records(control, include_self=True), automation_id)
                if not item:
                    continue
                if is_article_caption_panel_entry(item.get("content"), article_title):
                    continue
                item["_depth"] = 0
                item["_sequence"] = sequence
                sequence += 1
                signature = "|".join(str(item.get(key) or "") for key in (
                    "comment_id", "user_name", "content", "published_at"
                ))
                found.setdefault(hashlib.sha256(signature.encode("utf-8")).hexdigest(), item)
            except Exception:
                continue
        if round_index == 1 and expected_count and expected_count > 0 and not found:
            diagnostic_dir = os.environ.get("CHANNELS_COMMENT_DIAGNOSTIC_DIR")
            if diagnostic_dir:
                try:
                    diagnostic_path = Path(diagnostic_dir)
                    diagnostic_path.mkdir(parents=True, exist_ok=True)
                    save_window_screenshot(
                        surface,
                        diagnostic_path / f"comment-panel-{int(time.time())}.png",
                    )
                except Exception:
                    pass
            diagnostics = []
            for control in candidates[:8]:
                try:
                    diagnostics.append({
                        "control_type": control.element_info.control_type or "",
                        "class_name": control.element_info.class_name or "",
                        "automation_id": control.element_info.automation_id or "",
                        "records": [
                            {
                                "text": re.sub(r"\s+", " ", value).strip()[:200],
                                "control_type": control_type,
                                "class_name": class_name,
                            }
                            for value, control_type, class_name
                            in control_element_records(control, include_self=True)[:40]
                        ],
                    })
                except Exception as exc:
                    diagnostics.append({"error": str(exc)})
            print(json.dumps({
                "event": "comment_parse_diagnostic",
                "expected_count": expected_count,
                "structured_entries": len(new_comments),
                "actual_structured_entries": len(actual_new_comments),
                "candidates": diagnostics,
            }, ensure_ascii=False), flush=True)
        stable_rounds = stable_rounds + 1 if len(found) == before else 0
        print(json.dumps({
            "event": "comment_scroll_pass",
            "round": round_index,
            "visible_cards": len(new_comments) + len(candidates),
            "structuredEntries": len(new_comments),
            "fallbackCandidates": len(candidates),
            "candidateCountsMayOverlap": True,
            "unique_comments": len(found),
            "expanded_reply_groups": expanded_replies,
            "stable_rounds": stable_rounds,
        }, ensure_ascii=False), flush=True)
        if expected_count is not None and expected_count > 0 and len(found) >= expected_count:
            break
        if stable_rounds >= stable_required:
            break
        if panel_left is not None:
            rect = surface.rectangle()
            mouse.scroll(
                coords=(int((panel_left + rect.right) / 2), max(rect.top + 80, rect.bottom - 100)),
                wheel_dist=-8,
            )
            time.sleep(wait_seconds)
            continue
        scroll_targets = []
        for control in root.descendants():
            try:
                class_name = (control.element_info.class_name or "").lower()
                rect = control.rectangle()
                if "comment" in class_name and control.is_visible() and rect.width() > 120 and rect.height() > 120:
                    scroll_targets.append(control)
            except Exception:
                continue
        target = max(scroll_targets, key=lambda item: item.rectangle().width() * item.rectangle().height()) if scroll_targets else root
        rect = target.rectangle()
        mouse.scroll(coords=(int((rect.left + rect.right) / 2), max(rect.top + 40, rect.bottom - 60)), wheel_dist=-8)
        time.sleep(wait_seconds)
    incomplete = (not found or (expected_count is not None and len(found) < expected_count))
    record_comment_capture_result(config,
        status='incomplete' if incomplete else 'collected', expected=expected_count,
        actual=len(found), reason='expected_comments_not_fully_observed' if incomplete else 'visible_entries_collected')
    return merge_parent_child_comments(list(found.values()))


def open_video_comments_panel(web, config: dict) -> bool:
    root = active_video_root(web)
    candidates = []
    for control in root.descendants(control_type="Button"):
        try:
            name = (control.window_text() or "").strip()
            class_name = (control.element_info.class_name or "").lower()
            if "评论" in name or "comment" in class_name or "bubble" in class_name:
                rect = control.rectangle()
                if control.is_visible() and rect.width() > 2 and rect.height() > 2:
                    candidates.append(control)
        except Exception:
            continue
    if not candidates:
        return False
    control = candidates[0]
    rect = control.rectangle()
    print(json.dumps({
        "event": "comment_panel_click",
        "button": (control.window_text() or "").strip(),
        "class": control.element_info.class_name or "",
        "rect": [rect.left, rect.top, rect.right, rect.bottom],
        "candidate_count": len(candidates),
    }, ensure_ascii=False), flush=True)
    mouse.click(coords=(int((rect.left + rect.right) / 2), int((rect.top + rect.bottom) / 2)))
    panel_ready = False
    try:
        wait_for(
            lambda: any(
                "comment-content" in (item.element_info.class_name or "").lower()
                and item.is_visible()
                for item in get_video_detail_window().descendants(control_type="Group")
            ),
            timeout=4.0,
        )
        panel_ready = True
    except Exception:
        # The first click can merely reveal the hidden action overlay after a
        # share-menu operation. Re-resolve the active item's comment button
        # and click it once more instead of giving up with a false empty list.
        refreshed_root = active_video_root(get_video_detail_window())
        refreshed_buttons = [
            item for item in refreshed_root.descendants(control_type="Button")
            if "评论" in (item.window_text() or "")
            and item.is_visible()
        ]
        if refreshed_buttons:
            refreshed_rect = refreshed_buttons[0].rectangle()
            mouse.click(coords=(
                int((refreshed_rect.left + refreshed_rect.right) / 2),
                int((refreshed_rect.top + refreshed_rect.bottom) / 2),
            ))
            try:
                wait_for(
                    lambda: any(
                        "comment-content" in (item.element_info.class_name or "").lower()
                        and item.is_visible()
                        for item in get_video_detail_window().descendants(control_type="Group")
                    ),
                    timeout=4.0,
                )
                panel_ready = True
            except Exception:
                pass
        if not panel_ready:
            time.sleep(max(1.0, float(config.get("wait_short", 0.8))))
    print(json.dumps({
        "event": "comment_panel_opened",
        "ready": panel_ready,
        "visible_comment_contents": len([
            item for item in get_video_detail_window().descendants(control_type="Group")
            if "comment-content" in (item.element_info.class_name or "").lower()
            and item.is_visible()
        ]),
    }, ensure_ascii=False), flush=True)
    return panel_ready


def is_search_results_page(web) -> bool:
    return any("搜一搜" in (control.window_text() or "") for control in web.descendants(control_type="Document"))


NO_SEARCH_RESULT_PATTERN = re.compile(
    r"(?:暂无|没有|未找到|无)\s*(?:相关|匹配)?\s*(?:视频|内容|搜索结果|结果)|"
    r"换个关键词试试|未搜索到相关内容"
)


def has_explicit_no_search_results(web) -> bool:
    """Return True only when WeChat explicitly renders an empty-result state.

    A zero-length UIA card list is also produced while Chromium is hydrating or
    when the accessibility tree has detached.  Treating that transient surface
    as EOF silently acknowledged uncollected Redis tasks.
    """
    try:
        for control in web.descendants():
            try:
                text = re.sub(r"\s+", "", (control.window_text() or "").strip())
            except Exception:
                continue
            if text and NO_SEARCH_RESULT_PATTERN.search(text):
                return True
    except Exception:
        return False
    return False


def explicit_empty_status(known_result_count: int) -> str:
    """Checkpoint every first empty result instead of finalizing the keyword.

    WeChat may render the normal empty-result copy while 搜一搜 is temporarily
    restricted.  A new Redis generation also has no local history, so treating
    that first empty page as complete would permanently XACK a keyword that
    actually has videos.  Both variants therefore remain non-terminal and are
    retried from the top on a later worker run.
    """
    return (
        "partial_inconsistent_empty"
        if int(known_result_count or 0) > 0
        else "partial_empty_unverified"
    )


def is_video_detail_page(web) -> bool:
    return any(
        (control.window_text() or "").strip() == "视频号"
        for control in web.descendants(control_type="Document")
    )


def close_stacked_video_tabs_until_search_results(
    surface, *, max_tabs: int = 32, settle_seconds: float = 0.35,
) -> int:
    """Close same-window video tabs until the underlying search tab is active.

    Some WeChat builds open every clicked card as a new tab in the same
    ``Chrome_WidgetWin_0`` window.  Top-level-window cleanup cannot see those
    tabs, so one failed return used to leave each player behind indefinitely.
    """
    closed = 0
    current = surface
    for _attempt in range(max(1, int(max_tabs))):
        if is_search_results_page(current):
            break
        if not is_video_detail_page(current):
            raise SearchWindowUnavailableError(
                "关闭历史视频标签时当前页面既不是视频详情页也不是搜索结果页"
            )
        current.set_focus()
        keyboard.send_keys("^w", pause=0.02)
        closed += 1
        time.sleep(max(0.1, float(settle_seconds)))

        deadline = time.time() + 2.0
        current = None
        while time.time() < deadline:
            windows = web_windows()
            roles = [(detached_web_window_role(win), win) for win in windows]
            searches = [win for role, win in roles if role == "search_results"]
            details = [win for role, win in roles if role == "video_detail"]
            if len(searches) == 1:
                current = searches[0]
                break
            if len(details) == 1:
                current = details[0]
                break
            if not windows:
                break
            time.sleep(0.1)
        if current is not None and is_search_results_page(current):
            break
        if current is None:
            break
    else:
        raise SearchWindowUnavailableError(
            f"连续关闭 {max_tabs} 个视频标签后仍未恢复搜索结果页"
        )

    print(json.dumps({
        "event": "stacked_video_tabs_closed",
        "closed": closed,
        "search_results_ready": bool(current is not None and is_search_results_page(current)),
    }, ensure_ascii=False), flush=True)
    return closed


def return_to_search_results(config: dict | None = None) -> None:
    global SEARCH_LAYOUT
    def settle_after_return() -> None:
        if config is not None:
            jitter_pause(
                config, "jitter_return_list_seconds", (2.0, 5.0),
                "after_return_to_list",
            )
    # Newer WeChat builds open the player as a second top-level Chromium
    # window. Close that window directly while retaining the independent search
    # result window beneath it. This must happen before try_get_search_window(),
    # otherwise the old strict singleton check cannot even perform cleanup.
    closed = cleanup_stale_video_windows(
        reason="return_to_search_results", require_search_peer=True,
    )
    if closed:
        surface = get_search_results_window(wait_seconds=4.0)
        if is_search_results_page(surface) and not has_video_detail_window():
            settle_after_return()
            return
    if has_search_window() and not has_video_detail_window():
        settle_after_return()
        return
    surface = try_get_video_detail_window()
    if surface is None:
        if has_search_window():
            settle_after_return()
            return
        raise SearchWindowUnavailableError("返回搜索结果时既未找到视频详情页，也未找到搜索结果页")
    # When both windows still exist, cleanup should have closed the detail
    # window or raised a bounded error. Reaching here is the legacy same-window
    # navigation path only.
    if has_search_window() and SEARCH_LAYOUT == "detached":
        settle_after_return()
        return
    if SEARCH_LAYOUT == "detached":
        tab_closes = close_stacked_video_tabs_until_search_results(surface)
        if tab_closes == 0:
            backs = [
                control for control in surface.descendants(control_type="Button")
                if (control.window_text() or "").startswith("点按该按钮可后退")
            ]
            if backs:
                backs[0].click_input()
            else:
                surface.set_focus()
                keyboard.send_keys("%{LEFT}", pause=0.02)
    else:
        backs = [
            control for control in surface.descendants(control_type="Button")
            if (control.window_text() or "").startswith("点按该按钮可后退")
            or (control.window_text() or "").strip() in {"返回", "后退"}
        ]
        if backs:
            backs[0].click_input()
        else:
            # The embedded Chromium view keeps normal navigation history even
            # when its back button is omitted from the UIA tree.
            surface.set_focus()
            keyboard.send_keys("%{LEFT}", pause=0.02)
    wait_for(
        has_search_window,
        timeout=8.0,
    )
    cleanup_stale_video_windows(
        reason="return_to_search_results_post_navigation",
        require_search_peer=True,
    )
    surface = get_search_results_window(wait_seconds=4.0)
    if not is_search_results_page(surface) or has_video_detail_window():
        raise SearchWindowUnavailableError(
            "视频采集后未恢复为唯一搜索结果页，已停止继续打开新详情窗口"
        )
    settle_after_return()


def collect_live_keyword(
    keyword: str,
    project: str,
    cutoff: datetime,
    config: dict,
    connection: sqlite3.Connection,
    screenshot_dir: Path,
    max_results: int,
    collect_comments: bool = False,
    stop_file: Path | None = None,
) -> dict:
    config = dict(config)
    config['_capture_stop_file'] = stop_file
    config['force_asr'] = bool(config.get('force_asr')) or os.environ.get('WECHAT_FORCE_ASR') == '1'
    asr_backfill_mode = bool(config.get("asr_backfill_mode", False))
    retry_fingerprint = str(
        config.get("retry_fingerprint")
        or os.environ.get("WECHAT_REDIS_VIDEO_FINGERPRINT")
        or ""
    ).strip()
    remote_store = RedisVideoCheckpointStore()
    # Topic/project metadata can be corrected independently of the crawl
    # checkpoint.  Apply it to already committed rows before any early return
    # so resumed and legacy databases export the same canonical project name.
    connection.execute(
        "UPDATE captured_article SET project=? WHERE keyword=? AND project<>?",
        (project, keyword, project),
    )
    connection.execute(
        "UPDATE captured_comment SET project=? WHERE keyword=? AND project<>?",
        (project, keyword, project),
    )
    connection.commit()
    hydrated_remote = hydrate_remote_video_checkpoints(
        connection, keyword, project, remote_store,
        {retry_fingerprint} if retry_fingerprint else None,
    ) if remote_store.enabled else 0
    if hydrated_remote:
        print(json.dumps({
            "event": "redis_video_checkpoints_hydrated",
            "keyword": keyword,
            "records": hydrated_remote,
        }, ensure_ascii=False), flush=True)
    if STOP_REQUESTED or (stop_file and stop_file.exists()):
        return {"keyword": keyword, "complete": False, "status": "partial_safe_stop"}
    if asr_backfill_mode and retry_fingerprint:
        from precise_video_retry import try_direct_retry
        direct_result = try_direct_retry(sys.modules[__name__], keyword, project,
            retry_fingerprint, config, connection, remote_store, collect_comments)
        if direct_result is not None:
            return direct_result
    # A cross-machine ASR retry often runs days after the metadata scan. New
    # videos have meanwhile shifted the old card far below its original
    # position, and legacy checkpoints did not persist ``result_index``.  Use
    # the already-hydrated exact title as the search query so the one target is
    # normally in the first batch instead of rescanning dozens of unrelated
    # cards from the broad business keyword. The durable keyword/project stay
    # unchanged; this is only a locator query.
    search_query = keyword
    if asr_backfill_mode and retry_fingerprint:
        retry_checkpoint = load_video_checkpoints(connection, keyword).get(
            retry_fingerprint
        ) or {}
        retry_row = retry_checkpoint.get("row") or {}
        retry_title = str(
            retry_row.get("title") or retry_checkpoint.get("title") or ""
        ).strip()
        if not retry_title and retry_checkpoint.get('raw_card_text'):
            recovered = parse_video_card_text(retry_checkpoint['raw_card_text'])
            retry_title = recovered.get('title') or ''
            if retry_title:
                retry_checkpoint.update(title=retry_title,
                    author_name=recovered.get('author_name') or retry_checkpoint.get('author_name'))
                connection.execute('UPDATE video_checkpoint SET title=?, author_name=COALESCE(author_name,?) WHERE keyword=? AND fingerprint=?',
                    (retry_title, recovered.get('author_name'), keyword, retry_fingerprint))
                connection.commit()
                if remote_store.enabled:
                    remote_store.put(retry_fingerprint, retry_checkpoint)
        if not retry_title and not retry_row.get('url'):
            upsert_video_checkpoint(connection, keyword=keyword, project=project,
                fingerprint=retry_fingerprint, status='retry_pending',
                last_error='identity_recovery_required: no saved URL, title or recoverable raw card',
                result_index=retry_checkpoint.get('result_index'),
                retry_count=int(retry_checkpoint.get('retry_count') or 0), remote_store=remote_store)
            return {'keyword':keyword, 'saved':0, 'failed':1, 'examined':0,
                    'status':'partial_identity_recovery_required', 'complete':False}
        normalized_retry_title = normalize_video_identity_text(retry_title)
        if len(normalized_retry_title) >= 8 and normalized_retry_title != "...":
            search_query = retry_title
            print(json.dumps({
                "event": "redis_retry_exact_title_search",
                "keyword": keyword,
                "query": search_query,
                "stored_result_index": retry_checkpoint.get("result_index"),
            }, ensure_ascii=False), flush=True)
    try:
        search_in_web(search_query, config)
    except Exception:
        if search_query == keyword:
            raise
        print(json.dumps({
            "event": "redis_retry_exact_title_search_failed_fallback",
            "keyword": keyword,
            "query": search_query,
        }, ensure_ascii=False), flush=True)
        search_query = keyword
        search_in_web(keyword, config)
    screenshot_dir.mkdir(parents=True, exist_ok=True)
    save_window_screenshot(get_search_window(), screenshot_dir / f"{keyword}_results.png")
    search_origin = search_scan_surface_snapshot(get_search_window(), keyword)
    result_index, old_streak, previous_status = load_task_state(connection, keyword, project)
    known_result_count = int(connection.execute(
        "SELECT COUNT(*) FROM captured_article WHERE keyword=?", (keyword,)
    ).fetchone()[0] or 0) + int(connection.execute(
        "SELECT COUNT(*) FROM video_checkpoint WHERE keyword=?", (keyword,)
    ).fetchone()[0] or 0)
    # Search results are sorted newest-first and new articles can be inserted
    # above a persisted numeric index between runs.  Therefore a fresh process
    # must rebuild the scan from the true top after any partial/interrupted
    # status.  URL-key upserts keep this idempotent while preventing both the
    # visual "start in the middle" problem and missed newly published posts.
    restart_partial_from_top = bool(config.get("restart_partial_from_top", True))
    if restart_partial_from_top and previous_status and not previous_status.startswith("complete_"):
        print(json.dumps({
            "event": "partial_checkpoint_rescan_from_top",
            "keyword": keyword,
            "previous_status": previous_status,
            "previous_next_result_index": result_index,
        }, ensure_ascii=False), flush=True)
        result_index = 0
        old_streak = 0
    # Pagination v2 rechecks legacy complete_exhausted checkpoints because
    # older builds could mistake a non-growing first lazy-load batch for EOF.
    recheck_exhausted = bool(config.get("recheck_legacy_exhausted", True))
    # Legacy builds finalized the first explicit empty page. Always re-open
    # that checkpoint, even when this generation has no local article rows.
    # This changes only the checkpoint of the exact task already being run;
    # it does not reopen Redis tasks or reset per-video attempts/history.
    legacy_explicit_empty_checkpoint = previous_status.startswith(
        "complete_explicit_empty"
    )
    if legacy_explicit_empty_checkpoint:
        print(json.dumps({
            "event": "legacy_explicit_empty_checkpoint_reopened",
            "keyword": keyword,
            "known_result_count": known_result_count,
            "previous_status": previous_status,
        }, ensure_ascii=False), flush=True)
    recheck_annual = legacy_annual_rescan_required(
        previous_status, asr_backfill=asr_backfill_mode,
        recheck_exhausted=recheck_exhausted,
    ) or (bool(os.environ.get('WECHAT_SEARCH_SCAN_RUN_TOKEN')) and not asr_backfill_mode)
    if recheck_annual:
        print(json.dumps({
            'event': 'legacy_annual_checkpoint_rescan_from_top',
            'keyword': keyword, 'previous_status': previous_status,
            'previous_next_result_index': result_index,
        }, ensure_ascii=False), flush=True)
        result_index = 0
        old_streak = 0
    if previous_status.startswith("complete_") and not asr_backfill_mode and not recheck_annual:
        return {
            "keyword": keyword,
            "saved": 0,
            "failed": 0,
            "examined": 0,
            "status": previous_status,
            "complete": True,
        }
    if asr_backfill_mode:
        # Phase two is an intentional second pass. It must scan from the newest
        # card even when phase one already reached complete_cutoff/exhausted.
        result_index = 0
        old_streak = 0
        print(json.dumps({
            "event": "asr_backfill_rescan_from_top",
            "keyword": keyword,
            "existing_transcriptions": len(existing_asr) if "existing_asr" in locals() else 0,
        }, ensure_ascii=False), flush=True)
    asr_queue = None
    if config.get("enable_client_asr", True) and config.get("client_asr_background", False):
        database_row = connection.execute("PRAGMA database_list").fetchone()
        if database_row and database_row[2]:
            retry_checkpoint_for_recovery = (
                load_video_checkpoints(connection, keyword).get(retry_fingerprint) or {}
                if retry_fingerprint else {}
            )
            retry_article_key = str(
                os.environ.get("WECHAT_REDIS_VIDEO_ARTICLE_KEY")
                or retry_checkpoint_for_recovery.get("article_key")
                or ""
            ).strip()
            recovery_article_keys = (
                {retry_article_key} if retry_article_key else set()
            ) if retry_fingerprint else None
            asr_queue = BackgroundAsrQueue(
                Path(database_row[2]), config,
                recovery_keyword=keyword,
                recovery_article_keys=recovery_article_keys,
            )
    existing_article_keys = {
        row[0] for row in connection.execute(
            "SELECT article_key FROM captured_article WHERE keyword=?", (keyword,)
        ).fetchall()
    }
    video_checkpoints = load_video_checkpoints(connection, keyword)
    retry_card_scan_limit = (
        redis_retry_card_scan_limit(config)
        if asr_backfill_mode and retry_fingerprint else None
    )
    if retry_card_scan_limit is not None:
        print(json.dumps({
            "event": "redis_retry_scan_bounded",
            "keyword": keyword,
            "fingerprint": retry_fingerprint,
            "max_cards": retry_card_scan_limit,
            "original_index": os.environ.get("WECHAT_REDIS_VIDEO_RESULT_INDEX"),
        }, ensure_ascii=False), flush=True)
    incomplete_asr_keys = {
        str(item.get("article_key"))
        for item in video_checkpoints.values()
        if item.get("article_key") and str(item.get("asr_status") or "") in {
            "partial_review", "pending", "retry_pending", "running",
        }
    }
    existing_asr = {
        row[0]: (row[1], row[2])
        for row in connection.execute(
            "SELECT article_key, video_content, video_duration FROM captured_article "
            "WHERE keyword=? AND video_content IS NOT NULL AND video_content<>''",
            (keyword,),
        ).fetchall()
        if row[0] not in incomplete_asr_keys
        and not has_untrusted_transcript(row[1])
        and not asr_output_needs_review(row[1], row[2], config=config)
    }
    if retry_fingerprint or config.get('force_asr'):
        existing_asr = {}
    saved = 0
    failed = 0
    processed_this_run = 0
    rest_low, rest_high = configured_range(
        config, "jitter_long_break_every_videos", (8.0, 15.0),
    )
    rest_low_i = max(1, int(round(rest_low)))
    rest_high_i = max(rest_low_i, int(round(rest_high)))
    next_long_break_at = (
        random.randint(rest_low_i, rest_high_i)
        if bool(config.get("jitter_enabled", False)) else None
    )
    first_failed_index = None
    consecutive_position_skips = 0
    seen_fingerprints_this_run: set[str] = set()
    status = "running"
    search_end_proof = None
    scan_started_at_top = result_index == 0 and valid_origin(search_origin, keyword)
    last_published_at: datetime | None = None
    update_task_state(connection, keyword, project, result_index, old_streak, status)

    volume_local_target = None
    if os.environ.get('SOCIAL_VOLUME_ENABLED') == '1':
        baseline = connection.execute(
            'SELECT COUNT(DISTINCT article_key) FROM captured_article WHERE keyword=? AND project=?',
            (keyword, project)).fetchone()[0]
        volume_local_target = volume_export_local_target(baseline,
            int(os.environ['SOCIAL_VOLUME_INITIAL_POSTS']),
            int(os.environ['SOCIAL_VOLUME_TARGET_POSTS']))

    while max_results <= 0 or processed_this_run < max_results:
        if asr_queue is not None:
            asr_queue.raise_if_infrastructure_failed()
        if os.environ.get('SOCIAL_VOLUME_ENABLED') == '1' and not os.environ.get('WECHAT_RECOLLECT_ONCE'):
            target = int(os.environ['SOCIAL_VOLUME_TARGET_POSTS'])
            local_posts = connection.execute(
                "SELECT COUNT(DISTINCT article_key) FROM captured_article WHERE keyword=? AND project=?",
                (keyword,project)).fetchone()[0]
            if local_posts >= volume_local_target:
                # Stop only for export/import. The scheduler MUST independently
                # re-read SQL before quantity acceptance; this is not scan proof.
                status = 'complete_volume_export_pending'
                update_task_state(connection,keyword,project,result_index,old_streak,status)
                break
        if STOP_REQUESTED or (stop_file and stop_file.exists()):
            status = "partial_stopped"
            update_task_state(
                connection, keyword, project, result_index, old_streak, status,
                "graceful stop requested",
            )
            break
        if (
            next_long_break_at is not None
            # Exact ASR recovery is only comparing lightweight search-card
            # identities until it finds one target.  Counting inspected cards
            # as captured videos caused a 30-90 second break every 8-15 cards,
            # making a bounded 36-card lookup appear to hang.  Search submit,
            # page-load, scroll and detail-open jitter still remain active.
            and not retry_fingerprint
            and processed_this_run >= next_long_break_at
        ):
            if not jitter_pause(
                config, "jitter_long_break_seconds", (30.0, 90.0),
                "periodic_video_rest", keyword=keyword,
                processed=processed_this_run,
            ):
                status = "partial_stopped"
                update_task_state(
                    connection, keyword, project, result_index, old_streak,
                    status, "长休息期间收到安全停止请求",
                )
                break
            next_long_break_at = processed_this_run + random.randint(
                rest_low_i, rest_high_i,
            )
        if (
            retry_card_scan_limit is not None
            and processed_this_run >= retry_card_scan_limit
        ):
            if search_query != keyword:
                print(json.dumps({'event':'retry_title_not_found_keyword_fallback', 'keyword':keyword}, ensure_ascii=False), flush=True)
                search_in_web(keyword, config)
                search_query = keyword
                result_index = processed_this_run = 0
                retry_card_scan_limit = max(retry_card_scan_limit, int(config.get('redis_retry_max_cards', 240)))
                continue
            status = "partial_retry_target_not_found"
            detail = (
                f"精准字幕补采已检查 {processed_this_run} 张卡片，"
                "未找到目标视频；保留指纹供后续重试"
            )
            update_task_state(
                connection, keyword, project, result_index, old_streak,
                status, detail,
            )
            save_collection_issue(
                connection, keyword, result_index,
                "redis_retry_target_not_found", "", detail,
                retry_fingerprint,
            )
            print(json.dumps({
                "event": "redis_retry_scan_limit_reached",
                "keyword": keyword,
                "fingerprint": retry_fingerprint,
                "examined": processed_this_run,
                "max_cards": retry_card_scan_limit,
            }, ensure_ascii=False), flush=True)
            break
        if not has_search_window():
            search_in_web(keyword, config)
        web = get_search_window()
        results = search_result_buttons(web)
        # A resumed checkpoint can be several lazy-load batches beyond the
        # initial DOM (for example index 120 while only 36 controls exist).
        # Keep loading until the requested index actually exists. The former
        # single load attempt could grow 36 -> 56 and then still access [120].
        exhausted_before_index = False
        no_growth_cycles = 0
        exhaustion_cycles = max(3, int(config.get("exhaustion_confirmation_cycles", 3)))
        search_end_tracker = SearchEndTracker(keyword, search_origin, scan_started_at_top)
        while result_index >= len(results):
            if (
                result_index == 0
                and not results
                and has_explicit_no_search_results(get_search_window())
            ):
                status = explicit_empty_status(known_result_count)
                error = (
                    f"平台本次显示无结果；本地已有 {known_result_count} 条文章/视频断点；"
                    "已保存为异常断点，后续运行重新采集，不作为完成"
                )
                update_task_state(
                    connection, keyword, project, result_index, old_streak, status, error
                )
                exhausted_before_index = True
                print(json.dumps({
                    "event": (
                        "search_explicit_empty_inconsistent"
                        if known_result_count > 0
                        else "search_explicit_empty_checkpointed"
                    ),
                    "keyword": keyword,
                    "known_result_count": known_result_count,
                }, ensure_ascii=False), flush=True)
                break
            previous_count = len(results)
            try:
                results, grew = load_more_search_results(
                    web, previous_count, precise_retry_load_config(config) if retry_fingerprint else config, keyword=keyword,
                )
            except SearchWindowUnavailableError as exc:
                status = "partial_search_window_unavailable"
                detail = (
                    "滚动加载期间搜一搜窗口消失且有界恢复失败；"
                    "已保留当前索引，后续运行从断点重新扫描：" + str(exc)
                )
                update_task_state(
                    connection, keyword, project, result_index, old_streak,
                    status, detail,
                )
                save_collection_issue(
                    connection, keyword, result_index,
                    "search_window_unavailable", "", detail, None,
                )
                print(json.dumps({
                    "event": "search_load_more_recovery_exhausted",
                    "keyword": keyword,
                    "index": result_index,
                    "status": status,
                    "error": str(exc),
                }, ensure_ascii=False), flush=True)
                exhausted_before_index = True
                break
            if result_index >= len(results) and not grew:
                no_growth_cycles += 1
                if retry_card_scan_limit is not None:
                    status = "partial_retry_target_not_found"
                    detail = (
                        f"精准字幕补采搜索到列表末尾，已检查 {processed_this_run} 张卡片，"
                        "未找到目标视频；不再重建整个关键词页面"
                    )
                    update_task_state(
                        connection, keyword, project, result_index, old_streak,
                        status, detail,
                    )
                    save_collection_issue(
                        connection, keyword, result_index,
                        "redis_retry_target_not_found", "", detail,
                        retry_fingerprint,
                    )
                    print(json.dumps({
                        "event": "redis_retry_search_exhausted_once",
                        "keyword": keyword,
                        "fingerprint": retry_fingerprint,
                        "examined": processed_this_run,
                        "result_count": len(results),
                    }, ensure_ascii=False), flush=True)
                    exhausted_before_index = True
                    break
                explicit_empty = (
                    result_index == 0
                    and not results
                    and has_explicit_no_search_results(get_search_window())
                )
                print(json.dumps({
                    "event": "search_exhaustion_probe",
                    "keyword": keyword,
                    "result_count": len(results),
                    "cycle": no_growth_cycles,
                    "required_cycles": exhaustion_cycles,
                    "last_published_at": (
                        last_published_at.isoformat() if last_published_at else None
                    ),
                    "explicit_empty": explicit_empty,
                }, ensure_ascii=False), flush=True)
                if result_index == 0 and not results and not explicit_empty:
                    # Do not spend 12 scrolls x 3 rebuilds (and then repeat the
                    # whole sequence in ASR phase) on an unhydrated UIA tree.
                    # Fail closed so Redis keeps/retries the keyword while later
                    # unread tasks can continue.
                    status = "partial_result_surface_empty"
                    update_task_state(
                        connection, keyword, project, result_index, old_streak,
                        status,
                        "搜索页未显示卡片，也没有平台明确的无结果提示",
                    )
                    print(json.dumps({
                        "event": "search_result_surface_empty_unconfirmed",
                        "keyword": keyword,
                        "page_titles": search_page_titles(get_search_window()),
                    }, ensure_ascii=False), flush=True)
                    exhausted_before_index = True
                    break
                if explicit_empty:
                    status = explicit_empty_status(known_result_count)
                    error = (
                        f"平台本次显示无结果；本地已有 {known_result_count} 条文章/视频断点；"
                        "已保存为异常断点，后续运行重新采集，不作为完成"
                    )
                    update_task_state(
                        connection, keyword, project, result_index, old_streak, status, error
                    )
                    exhausted_before_index = True
                    break
                if no_growth_cycles >= exhaustion_cycles:
                    endpoint = search_end_tracker.observe(search_scan_surface_snapshot(
                        get_search_window(), keyword, scanned_count=result_index,
                        surface_generation=no_growth_cycles))
                    if endpoint.get('ok'):
                        status = 'complete_exhausted'
                        search_end_proof = search_end_tracker.proof
                    else:
                        status = 'partial_search_end_unverified'
                    update_task_state(
                        connection, keyword, project, result_index, old_streak, status,
                        None if endpoint.get('ok') else endpoint.get('reason')
                    )
                    print(json.dumps({'event': 'search_end_evidence_checked', 'keyword': keyword,
                                      'status': status, 'evidence': endpoint}, ensure_ascii=False), flush=True)
                    exhausted_before_index = True
                    break
                search_end_tracker.observe(search_scan_surface_snapshot(
                    get_search_window(), keyword, scanned_count=result_index,
                    surface_generation=no_growth_cycles))
                # Rebuild the exact same keyword/video/latest page between
                # exhaustion probes. This borrows the公众号 anti-stall pattern:
                # a lazy-load failure in one Chromium surface is not accepted
                # as the end of the yearly list without independent reloads.
                search_in_web(keyword, config)
                web = get_search_window()
                results = search_result_buttons(web)
                print(json.dumps({
                    "event": "search_pagination_surface_rebuilt",
                    "keyword": keyword,
                    "cycle": no_growth_cycles,
                    "result_count_after_rebuild": len(results),
                }, ensure_ascii=False), flush=True)
            # Do not reset the cross-surface exhaustion confirmation merely
            # because a rebuilt page re-hydrated its already-known first
            # batches (12 -> 24 -> ...).  That growth does not reveal a card
            # beyond ``result_index`` and the old reset made every independent
            # probe return to cycle 1 forever.  If genuinely new cards extend
            # past result_index, the while condition exits naturally and the
            # next card is processed.
            web = get_search_window()
        if exhausted_before_index and retry_fingerprint and search_query != keyword and status in {
            'partial_retry_target_not_found', 'partial_explicit_empty',
            'partial_result_surface_empty', 'partial_explicit_empty_inconsistent',
            'partial_empty_search', 'partial_empty_search_inconsistent',
            'partial_inconsistent_empty', 'partial_empty_unverified',
            'complete_exhausted',
        }:
            print(json.dumps({'event':'retry_title_empty_keyword_fallback', 'keyword':keyword}, ensure_ascii=False), flush=True)
            search_in_web(keyword, config)
            search_query = keyword
            result_index = processed_this_run = 0
            status = 'running'
            continue
        if exhausted_before_index:
            break

        result = results[result_index]
        raw_card_text = result.window_text().strip()
        card = parse_video_card_text(raw_card_text)
        title = card["title"] or ""
        card_published_at = parse_search_card_time(card["relative_time"] or "")
        card_fingerprint = normalize_video_card_fingerprint(raw_card_text)
        if card_fingerprint and card_fingerprint not in video_checkpoints:
            canonical_checkpoint = unique_checkpoint_for_stable_card_identity(
                video_checkpoints, card,
            )
            if canonical_checkpoint is not None:
                observed_fingerprint = card_fingerprint
                card_fingerprint, canonical_payload = canonical_checkpoint
                print(json.dumps({
                    "event": "card_countdown_fingerprint_canonicalized",
                    "keyword": keyword,
                    "index": result_index,
                    "observed_fingerprint": observed_fingerprint,
                    "canonical_fingerprint": card_fingerprint,
                    "observed_duration": card.get("duration_seconds"),
                    "stored_duration": (
                        (canonical_payload.get("row") or {}).get("video_duration")
                    ),
                }, ensure_ascii=False), flush=True)
        if card_fingerprint:
            seen_fingerprints_this_run.add(card_fingerprint)
        if asr_backfill_mode and retry_fingerprint and card_fingerprint != retry_fingerprint:
            retry_checkpoint = video_checkpoints.get(retry_fingerprint)
            if retry_checkpoint_matches_card(retry_checkpoint, card):
                print(json.dumps({
                    "event": "redis_retry_card_matched_stable_identity",
                    "keyword": keyword,
                    "index": result_index,
                    "retry_fingerprint": retry_fingerprint,
                    "observed_fingerprint": card_fingerprint,
                    "observed_duration": card.get("duration_seconds"),
                    "stored_duration": ((retry_checkpoint or {}).get("row") or {}).get("video_duration"),
                }, ensure_ascii=False), flush=True)
                # Continue to update the durable checkpoint represented by the
                # Redis message, not a transient hash derived from countdown.
                card_fingerprint = retry_fingerprint
                seen_fingerprints_this_run.add(retry_fingerprint)
            else:
                # A Redis child task represents exactly one failed/incomplete
                # video. Scan card identities only; never open unrelated pages.
                result_index += 1
                processed_this_run += 1
                continue
        checkpoint = video_checkpoints.get(card_fingerprint) if card_fingerprint else None
        checkpoint_status = str((checkpoint or {}).get("status") or "")
        checkpoint_asr = str((checkpoint or {}).get("asr_status") or "")
        checkpoint_row = (checkpoint or {}).get("row") or {}
        checkpoint_asr_is_reusable = (
            os.environ.get('SOCIAL_VOLUME_ENABLED')=='1'
            or checkpoint_asr in {"no_content", "not_applicable"}
            or (
                checkpoint_asr == "completed"
                and bool(str(checkpoint_row.get("video_content") or "").strip())
                and not asr_output_needs_review(
                    checkpoint_row.get("video_content"),
                    checkpoint_row.get("video_duration"),
                    config=config,
                )
            )
        )
        completed_for_phase = checkpoint and (
            checkpoint_status in {"dead_letter","skipped_unresolved"}
            or (
                asr_backfill_mode
                and checkpoint_asr_is_reusable
            )
            or (
                not asr_backfill_mode
                and checkpoint_status in {"completed", "completed_metadata"}
            )
        )
        repair_reason = checkpoint_repair_reason(
            checkpoint, collect_comments=collect_comments)
        if completed_for_phase and not (os.environ.get('WECHAT_RECOLLECT_ONCE') or retry_fingerprint or force_retranscribe(config, checkpoint) or repair_reason):
            failed_index = result_index
            result_index += 1
            processed_this_run += 1
            checkpoint_published = checkpoint.get("published_at")
            try:
                checkpoint_date = datetime.fromisoformat(checkpoint_published) if checkpoint_published else card_published_at
            except (TypeError, ValueError):
                checkpoint_date = card_published_at
            if checkpoint_date is not None:
                last_published_at = checkpoint_date
                old_streak = old_streak + 1 if checkpoint_date < cutoff else 0
            update_task_state(
                connection, keyword, project, result_index, old_streak, "running",
                f"结果索引 {failed_index} 命中视频级完成指纹，仅卡片复核后跳过",
            )
            print(json.dumps({
                "event": "rescan_skip_completed_video_fingerprint",
                "keyword": keyword,
                "index": failed_index,
                "fingerprint": card_fingerprint,
                "checkpoint_status": checkpoint_status,
                "asr_status": checkpoint_asr,
            }, ensure_ascii=False), flush=True)
            if observe_old_date_window(config, old_streak):
                print(json.dumps({'event': 'old_date_window_observed_not_search_end',
                                  'keyword': keyword, 'old_streak': old_streak},
                                 ensure_ascii=False), flush=True)
            continue
        # 搜一搜“视频”流会混排正在进行的直播。直播卡没有普通视频
        # 详情页/分享链接，反复点击只会永久留下 partial_error，阻塞下一
        # 个关键词。它不属于本任务的已发布视频，明确审计后跳过。
        if re.search(r"(?:^|\s)直播中(?:\s|$)", raw_card_text):
            print(json.dumps({
                "event": "live_card_skipped",
                "keyword": keyword,
                "index": result_index,
                "title": title,
            }, ensure_ascii=False), flush=True)
            result_index += 1
            processed_this_run += 1
            update_task_state(
                connection, keyword, project,
                first_failed_index if first_failed_index is not None else result_index,
                old_streak, "running",
            )
            continue
        item_success = False
        last_item_exception = None
        media_cache_before = None
        if (
            config.get("enable_client_asr", True)
            and config.get("enable_wechat_media_cache", True)
            and os.environ.get("WECHAT_CLIENT_ASR", "1") != "0"
        ):
            cache_dirs = discover_cache_dirs(config.get("wechat_media_cache_dirs"))
            media_cache_before = snapshot_cache(cache_dirs)
            print(json.dumps({
                "event": "wechat_media_cache_snapshot",
                "keyword": keyword,
                "index": result_index,
                "cache_dirs": len(cache_dirs),
                "files": len(media_cache_before),
            }, ensure_ascii=False), flush=True)
        item_attempts_this_run = max(1, int(config.get("video_attempts_per_run", 1)))
        for item_attempt in range(1, item_attempts_this_run + 1):
            try:
                # Refresh virtualized UIA elements before every retry. Some
                # WeChat builds expose Invoke but silently ignore it; use a
                # physical click for subsequent attempts.
                refreshed = search_result_buttons(get_search_window())
                if result_index < len(refreshed):
                    result = refreshed[result_index]
                    raw_card_text = result.window_text().strip()
                    card = parse_video_card_text(raw_card_text)
                    title = card["title"] or ""
                    card_published_at = parse_search_card_time(card["relative_time"] or "")
                visible_result = bring_search_result_into_view(result_index, config)
                if visible_result is None:
                    raise ResultPositionError(f"结果索引 {result_index} 无法滚动到可视区域")
                result = visible_result
                raw_card_text = result.window_text().strip()
                card = parse_video_card_text(raw_card_text)
                title = card["title"] or ""
                if (
                    config.get("enable_client_asr", True)
                    and os.environ.get("WECHAT_CLIENT_ASR", "1") != "0"
                ):
                    config["_search_card_audio_enabled"] = enable_search_card_audio(result)
                else:
                    config.pop("_search_card_audio_enabled", None)
                jitter_pause(
                    config, "jitter_detail_open_seconds", (3.0, 7.0),
                    "before_detail_open", keyword=keyword,
                    result_index=result_index,
                )
                click_point = click_visible_search_result(result)
                print(json.dumps({
                    "event": "search_result_clicked",
                    "index": result_index,
                    "point": click_point,
                }, ensure_ascii=False), flush=True)
                try:
                    wait_for(
                        lambda: has_video_detail_window(),
                        timeout=float(config.get("wait_video_detail_open_seconds", 8.0)),
                    )
                except Exception as exc:
                    raise RuntimeError("视频详情页未在时限内打开") from exc
                web = prepare_video_detail(get_video_detail_window(), config, card,
                                           keyword=keyword, result_index=result_index)
                try:
                    wait_for(
                        lambda: pause_active_video_if_running(get_video_detail_window()),
                        timeout=min(
                            6.0,
                            float(config.get("wait_video_detail_ready_seconds", 15.0)),
                        ),
                    )
                except Exception:
                    # Some videos expose action metadata before the native
                    # player control. Continue; the normal identity/link
                    # checks will still reject a mismatched feed item.
                    print(json.dumps({
                        "event": "video_pause_control_unavailable",
                        "keyword": keyword,
                        "index": result_index,
                    }, ensure_ascii=False), flush=True)
                try:
                    wait_for(
                        lambda: len([
                            control for control in active_video_root(get_video_detail_window()).descendants(control_type="Button")
                            if "op-item" in (control.element_info.class_name or "")
                        ]) >= 4,
                        timeout=float(config.get("wait_video_detail_ready_seconds", 15.0)),
                    )
                except Exception as exc:
                    raise RuntimeError("视频详情页已打开，但互动栏未在时限内完成加载") from exc
                web = get_video_detail_window()
                records = active_video_records(web)
                require_video_detail_identity(
                    card, records, keyword=keyword, result_index=result_index,
                )
                detail_duration = resolve_video_detail_duration(
                    web, records, card["duration_seconds"],
                )
                sample_duration_limit = int(
                    os.environ.get("CHANNELS_SAMPLE_MAX_DURATION_SECONDS", "0") or 0
                )
                if sample_duration_limit > 0 and (
                    detail_duration or card["duration_seconds"] or 0
                ) > sample_duration_limit:
                    print(json.dumps({
                        "event": "sample_long_video_skipped",
                        "keyword": keyword,
                        "index": result_index,
                        "duration_seconds": detail_duration or card["duration_seconds"],
                        "sample_limit_seconds": sample_duration_limit,
                    }, ensure_ascii=False), flush=True)
                    return_to_search_results(config)
                    result_index += 1
                    processed_this_run += 1
                    update_task_state(
                        connection, keyword, project, result_index, old_streak,
                        "running", "测试抽样跳过长视频；未写入生产视频检查点",
                    )
                    item_success = True
                    break
                if (
                    detail_duration and card["duration_seconds"]
                    and detail_duration != card["duration_seconds"]
                ):
                    print(json.dumps({
                        "event": "video_duration_corrected_from_detail_clock",
                        "keyword": keyword,
                        "index": result_index,
                        "card_duration_seconds": card["duration_seconds"],
                        "detail_duration_seconds": detail_duration,
                    }, ensure_ascii=False), flush=True)
                published_at = parse_publish_time(records) or card_published_at
                if published_at is not None:
                    last_published_at = published_at
                details = parse_article_details(records)
                if published_at is not None and details["published_at"] is None:
                    details["published_at"] = published_at.isoformat(sep=" ", timespec="seconds")
                like_count, comment_count, favorite_count, share_count = parse_article_metrics(records)
                link = copy_article_link_from_menu(get_video_detail_window(), config)
                if asr_backfill_mode and retry_fingerprint:
                    retry_target_checkpoint = video_checkpoints.get(
                        retry_fingerprint
                    ) or {}
                    retry_target_row = retry_target_checkpoint.get("row") or {}
                    canonical_retry_url = str(
                        retry_target_row.get("url") or ""
                    ).strip()
                    if canonical_retry_url and canonical_retry_url != link:
                        # 视频号 can issue a different /sph/ token for the
                        # same card on a later share. The Redis child task is
                        # an enrichment of the already committed row, not a
                        # new article. Keep its canonical URL/article_key so
                        # the completed transcript is exported and upserted
                        # into MySQL instead of being stranded on a new key.
                        print(json.dumps({
                            "event": "redis_retry_share_url_canonicalized",
                            "keyword": keyword,
                            "observed_url": link,
                            "canonical_url": canonical_retry_url,
                        }, ensure_ascii=False), flush=True)
                        link = canonical_retry_url
                candidate_key = article_key_for_row(ResultRow(
                    project=project, keyword=keyword, title=title, url=link,
                    author_name=card["author_name"] or details["author_name"],
                    published_at=details["published_at"],
                    video_duration=detail_duration or card["duration_seconds"],
                ))
                # Cross-keyword dedupe is project-scoped.  The canonical
                # finder/share identity is checked only after opening the card,
                # avoiding false positives from authors who reuse the same
                # title and duration.  This saves the expensive comments/ASR
                # phase while preserving one row per business topic.
                local_duplicates = connection.execute(
                    "SELECT keyword, url FROM captured_article "
                    "WHERE article_key=? AND project=? AND keyword<>? ORDER BY keyword",
                    (candidate_key, project, keyword),
                ).fetchall()
                remote_duplicate = remote_store.get_global(project, candidate_key) \
                    if remote_store.enabled else None
                reuse_sources = [cp for duplicate in local_duplicates
                    for cp in load_video_checkpoints(connection, duplicate[0]).values()
                    if cp.get('article_key') == candidate_key]
                if remote_duplicate:
                    reuse_sources.append(remote_duplicate)
                reusable = next((cp for cp in reuse_sources if (cp.get('row') or {}).get('keyword') != keyword
                    and reusable_cross_keyword_checkpoint(
                    cp, project=project, article_key=candidate_key, url=link,
                    author=card['author_name'] or details['author_name'],
                    duration=detail_duration or card['duration_seconds'], cutoff=cutoff,
                    collect_comments=collect_comments)), None)
                if (
                    published_at is not None and published_at >= cutoff
                    and reusable is not None
                    and not (asr_backfill_mode and retry_fingerprint)
                ):
                    duplicate_keyword = reusable['row']['keyword']
                    reused = persist_cross_keyword_reuse(connection, checkpoint=reusable,
                        keyword=keyword, project=project, fingerprint=card_fingerprint,
                        result_index=result_index, cutoff=cutoff, collect_comments=collect_comments,
                        remote_store=remote_store)
                    video_checkpoints[card_fingerprint] = reused
                    saved += 1
                    return_to_search_results(config)
                    result_index += 1
                    processed_this_run += 1
                    old_streak = 0
                    update_task_state(
                        connection, keyword, project, result_index, old_streak, "running",
                        f"结果索引 {result_index - 1} 与关键词 {duplicate_keyword} 命中同一视频，已跨关键词去重",
                    )
                    print(json.dumps({
                        "event": "cross_keyword_verified_data_reused",
                        "keyword": keyword,
                        "index": result_index - 1,
                        "duplicate_keyword": duplicate_keyword,
                        "article_key": candidate_key,
                        "url": link,
                    }, ensure_ascii=False), flush=True)
                    item_success = True
                    break
                comments = []
                detail_capture = {"completed": False}

                def collect_detail_fields_while_recording():
                    nonlocal comments, details
                    nonlocal like_count, comment_count, favorite_count, share_count
                    config['_capture_stop_file'] = stop_file
                    config.pop('_comment_capture_audit', None)
                    ensure_capture_not_stopped(config)
                    current_web = get_video_detail_window()
                    current_records = active_video_records(current_web)
                    require_video_detail_identity(
                        card, current_records, keyword=keyword,
                        result_index=result_index,
                    )
                    refreshed_details = parse_article_details(current_records)
                    if published_at is not None and refreshed_details["published_at"] is None:
                        refreshed_details["published_at"] = published_at.isoformat(
                            sep=" ", timespec="seconds",
                        )
                    details = refreshed_details
                    (
                        like_count, comment_count, favorite_count, share_count,
                    ) = parse_article_metrics(current_records)
                    if collect_comments and (comment_count is None or comment_count > 0):
                        comments = collect_video_comments(
                            current_web, config, expected_count=comment_count,
                            article_title=card.get("title"),
                        )
                    detail_capture["completed"] = True
                # A Chromium navigation can expose the menu/link before the
                # article metadata has hydrated. Treating that intermediate
                # state as success silently drops the row because its date is
                # unknown. Re-open it through the normal retry loop instead.
                if published_at is None:
                    raise RuntimeError("文章页面尚未加载发布时间，重新打开重试")

                in_period = published_at is not None and published_at >= cutoff
                if published_at is None:
                    # Keep an unknown date out of the requested date window;
                    # it also must not reset an existing old-article streak.
                    pass
                elif published_at < cutoff:
                    old_streak += 1
                else:
                    old_streak = 0
                    pending_asr = None
                    optional_audio_error = None
                    # Reuse the same canonical share URL/finder-id key used by
                    # phase one and direct-link ASR enrichment. A title-only
                    # fingerprint can miss an existing transcription and cause
                    # an unnecessary full replay in the WeChat client.
                    if candidate_key in existing_asr and not force_retranscribe(config, checkpoint):
                        video_content, video_duration = existing_asr[candidate_key]
                        print(json.dumps({
                            "event": "client_audio_asr_reused",
                            "article_key": candidate_key,
                            "text_length": len(video_content or ""),
                        }, ensure_ascii=False), flush=True)
                    else:
                        video_content, video_duration, optional_audio_error = capture_primary_audio_optional(
                            raw_card_text, web, {**config, '_capture_identity':{
                                'title':title, 'author_name':card.get('author_name') or details.get('author_name')}},
                            duration_seconds=detail_duration or card["duration_seconds"],
                            prompt_context=None,
                            media_cache_before=media_cache_before,
                            during_capture_callback=collect_detail_fields_while_recording,
                            stop_file=stop_file,
                        )
                        if isinstance(video_content, PendingClientAsr):
                            pending_asr = video_content
                            video_content = None
                    # Cache-media reuse and existing transcripts do not start a
                    # loopback recorder, so collect the same fields directly.
                    if not detail_capture["completed"]:
                        collect_detail_fields_while_recording()
                    saved_row = ResultRow(
                        project=project,
                        keyword=keyword,
                        title=title,
                        url=link,
                        # The search-card tail is stable across current WeChat
                        # UI families; detail-page class names are not and may
                        # expose an internal id as the apparent nickname.
                        author_name=card["author_name"] or details["author_name"],
                        tags=card["tags"],
                        published_at=details["published_at"],
                        location=details["location"],
                        content=None,
                        image_count=details["image_count"],
                        video_count=max(1, details["video_count"] or 0),
                        video_content=video_content,
                        video_duration=video_duration,
                        like_count=like_count,
                        comment_count=comment_count,
                        favorite_count=favorite_count,
                        share_count=share_count,
                    )
                    save_row(connection, saved_row)
                    if optional_audio_error:
                        save_collection_issue(connection, keyword, result_index,
                            'optional_audio_failed_volume_policy', title,
                            optional_audio_error, card_fingerprint)
                    if collect_comments:
                        save_comments(connection, saved_row, comments)
                        comment_audit = config.get('_comment_capture_audit') or {}
                        if comment_audit.get('status') in {'unavailable', 'incomplete'}:
                            save_collection_issue(connection, keyword, result_index,
                                'comments_not_fully_observed', title,
                                json.dumps(comment_audit, ensure_ascii=False), card_fingerprint)
                    if asr_backfill_mode:
                        asr_status = (
                            "pending" if pending_asr is not None
                            else "partial_review" if video_content and asr_output_needs_review(
                                video_content, video_duration, config=config,
                            )
                            else "completed" if video_content
                            else "no_content"
                        )
                    else:
                        asr_status = (
                            "pending" if pending_asr is not None
                            else "partial_review" if video_content and asr_output_needs_review(
                                video_content, video_duration, config=config,
                            )
                            else "completed" if video_content
                            else "pending"
                        )
                    if optional_audio_error:
                        asr_status = 'skipped_volume_policy'
                    checkpoint_status = (
                        'completed_metadata' if optional_audio_error else
                        "completed" if asr_status in {"completed", "no_content", "not_applicable"}
                        else "retry_pending" if asr_status == "partial_review"
                        else "completed_metadata"
                    )
                    checkpoint_payload = upsert_video_checkpoint(
                        connection,
                        keyword=keyword,
                        project=project,
                        fingerprint=card_fingerprint,
                        result_index=result_index,
                        status=checkpoint_status,
                        retry_count=0,
                        row=saved_row,
                        comments=comments,
                        asr_status=asr_status,
                        last_error=(
                            optional_audio_error or (
                                "ASR transcript is implausibly short for captured duration"
                                if asr_status == "partial_review" else None)
                        ),
                        remote_store=remote_store,
                    )
                    video_checkpoints[card_fingerprint] = checkpoint_payload
                    if pending_asr is not None and asr_queue is not None:
                        asr_queue.submit(
                            article_key_for_row(saved_row), keyword, pending_asr,
                            fingerprint=card_fingerprint, project=project,
                            remote_store=remote_store,
                        )
                    if video_content:
                        existing_asr[article_key_for_row(saved_row)] = (video_content, video_duration)
                    article_key = article_key_for_row(saved_row)
                    if article_key not in existing_article_keys:
                        existing_article_keys.add(article_key)
                        saved += 1
                return_to_search_results(config)
                result_index += 1
                processed_this_run += 1
                update_task_state(
                    connection, keyword, project,
                    first_failed_index if first_failed_index is not None else result_index,
                    old_streak, "running"
                )
                print(json.dumps({
                    "keyword": keyword,
                    "index": result_index - 1,
                    "title": title,
                    "published_at": published_at.isoformat() if published_at else None,
                    "url": link,
                    "like": like_count,
                    "comment": comment_count,
                    "favorite_or_recommend": favorite_count,
                    "share": share_count,
                    "saved": in_period,
                }, ensure_ascii=False))
                item_success = True
                if asr_backfill_mode and retry_fingerprint:
                    status = "complete_retry_fingerprint"
                    update_task_state(
                        connection, keyword, project, result_index, old_streak, status,
                        f"视频指纹 {retry_fingerprint} 精准补采完成",
                    )
                break
            except (VideoPlaybackUnavailableError, VideoDetailRecoveryBlockedError) as exc:
                last_item_exception = exc
                status = 'partial_stopped' if 'STOP_REQUESTED' in str(exc) else 'partial_error'
                save_collection_issue(connection, keyword, result_index,
                    'video_playback_unavailable', title, str(exc), card_fingerprint)
                update_task_state(connection, keyword, project, result_index,
                                  old_streak, status, str(exc))
                # Infrastructure/player readiness is not a malformed video:
                # do not advance, spend content retries or create a dead letter.
                try:
                    return_to_search_results(config)
                except Exception:
                    pass
                break
            except RedisCheckpointUnavailable:
                # Transport exhaustion is not a failed video. The local
                # checkpoint was already committed: preserve its retry count
                # and cursor and let the outer salvage/exit-75 path handle it.
                print(json.dumps({
                    "event": "article_infrastructure_paused",
                    "keyword": keyword,
                    "index": result_index,
                    "contentAttemptCharged": False,
                    "nextItemAllowed": False,
                    "localCheckpointRetained": True,
                }, ensure_ascii=False), flush=True)
                raise
            except Exception as exc:
                last_item_exception = exc
                failed += 1
                print(json.dumps({
                    "keyword": keyword,
                    "index": result_index,
                    "title": title,
                    "attempt": item_attempt,
                    "error": str(exc),
                }, ensure_ascii=False))
                try:
                    # Do not guard cleanup with a singleton probe.  When two or
                    # more stale players exist that probe is deliberately
                    # ambiguous; skipping cleanup there caused linear window
                    # accumulation after repeated item failures.
                    return_to_search_results(config)
                except Exception:
                    try:
                        cleanup_stale_video_windows(
                            reason="article_exception_fallback",
                            require_search_peer=False,
                        )
                    except Exception:
                        pass
                if isinstance(exc, ResultPositionError) and item_attempt < item_attempts_this_run:
                    try:
                        print(json.dumps({
                            "event": "search_result_position_recovery",
                            "keyword": keyword,
                            "index": result_index,
                            "attempt": item_attempt,
                        }, ensure_ascii=False), flush=True)
                        search_in_web(keyword, config)
                        recovered = search_result_buttons(get_search_window())
                        while result_index >= len(recovered):
                            previous_count = len(recovered)
                            recovered, grew = load_more_search_results(
                                get_search_window(), previous_count, config,
                                keyword=keyword,
                            )
                            if result_index >= len(recovered) and not grew:
                                break
                    except Exception as recovery_exc:
                        print(json.dumps({
                            "event": "search_result_position_recovery_failed",
                            "keyword": keyword,
                            "index": result_index,
                            "error": str(recovery_exc),
                        }, ensure_ascii=False), flush=True)
                if not has_search_window():
                    try:
                        search_in_web(keyword, config)
                    except Exception:
                        pass
                if has_search_window() and is_search_results_page(get_search_window()):
                    refreshed = search_result_buttons(get_search_window())
                    if result_index < len(refreshed):
                        result = refreshed[result_index]
                if isinstance(exc, VideoDetailIdentityMismatchError):
                    # An ultra-short video that has already advanced will fail
                    # identically on every immediate retry.  Cleanup the detail
                    # tab first, then audit/skip it once so it cannot block the
                    # rest of the year or the next keyword.
                    try:
                        cleanup_stale_video_windows(
                            reason="identity_mismatch_fallback",
                            require_search_peer=False,
                        )
                    except Exception:
                        pass
                    if not has_search_window():
                        try:
                            search_in_web(keyword, config)
                        except Exception:
                            pass
                    break
        if isinstance(last_item_exception, (VideoPlaybackUnavailableError, VideoDetailRecoveryBlockedError)):
            break
        if not item_success:
            failed_index = result_index
            safe_keyword = re.sub(r"[^\w-]+", "_", keyword)
            failure_screenshot = screenshot_dir / (
                f"{safe_keyword}_{failed_index}_"
                f"{(card_fingerprint or 'unknown')[:12]}_failed.png"
            )
            try:
                failure_surface = (
                    try_get_video_detail_window(focus=False)
                    or try_get_search_results_window(focus=False)
                )
                if failure_surface is None:
                    raise SearchWindowUnavailableError("失败截图时未找到微信网页窗口")
                save_window_screenshot(failure_surface, failure_screenshot)
                failure_screenshot_value = str(failure_screenshot)
            except Exception:
                failure_screenshot_value = None
            if card_fingerprint:
                failure_status, retry_count = record_video_failure(
                    connection,
                    keyword=keyword,
                    project=project,
                    fingerprint=card_fingerprint,
                    result_index=failed_index,
                    title=title,
                    author_name=card.get("author_name"),
                    raw_card_text=raw_card_text,
                    error=str(last_item_exception or "unknown article error"),
                    screenshot_path=failure_screenshot_value,
                    max_attempts=int(config.get("video_retry_max_attempts", 3)),
                    remote_store=remote_store,
                )
                video_checkpoints = load_video_checkpoints(connection, keyword)
            else:
                # An index plus an empty UIA card is not a stable article
                # identity. Hashing it created retry work that could never be
                # found again and polluted the dead-letter table.
                failure_status, retry_count = "audit_only", 0
                save_collection_issue(
                    connection, keyword, failed_index,
                    "unidentified_result_failure_skipped", title,
                    str(last_item_exception or "unknown article error"), None,
                )
            if isinstance(last_item_exception, VideoDetailIdentityMismatchError):
                save_collection_issue(
                    connection, keyword, failed_index,
                    "detail_identity_mismatch_skipped", title,
                    str(last_item_exception), card_fingerprint,
                )
                result_index += 1
                processed_this_run += 1
                consecutive_position_skips = 0
                update_task_state(
                    connection, keyword, project, result_index, old_streak, "running",
                    f"结果索引 {failed_index} 详情自动跳转，已审计并继续后续视频",
                )
                print(json.dumps({
                    "event": "video_detail_identity_mismatch_skipped",
                    "keyword": keyword,
                    "index": failed_index,
                    "fingerprint": card_fingerprint,
                    "retry_status": failure_status,
                    "retry_count": retry_count,
                }, ensure_ascii=False), flush=True)
                continue
            if isinstance(last_item_exception, ResultPositionError):
                save_collection_issue(
                    connection, keyword, failed_index,
                    "position_unreachable_skipped", title,
                    str(last_item_exception), card_fingerprint,
                )
                consecutive_position_skips += 1
                result_index += 1
                processed_this_run += 1
                update_task_state(
                    connection, keyword, project, result_index, old_streak, "running",
                    f"结果索引 {failed_index} 定位失败，已记录指纹并跳过",
                )
                print(json.dumps({
                    "event": "search_result_position_skipped",
                    "keyword": keyword,
                    "index": failed_index,
                    "fingerprint": card_fingerprint,
                    "consecutive_position_skips": consecutive_position_skips,
                    "retry_status": failure_status,
                    "retry_count": retry_count,
                }, ensure_ascii=False), flush=True)
                if consecutive_position_skips >= max(
                    1, int(config.get("max_consecutive_position_skips", 3))
                ):
                    print(json.dumps({
                        "event": "consecutive_position_failures_rebuild_surface",
                        "keyword": keyword,
                        "count": consecutive_position_skips,
                    }, ensure_ascii=False), flush=True)
                    try:
                        search_in_web(keyword, config)
                        consecutive_position_skips = 0
                    except Exception as recovery_error:
                        status = "partial_position_error"
                        update_task_state(
                            connection, keyword, project, result_index, old_streak, status,
                            f"连续定位失败且搜索页重建失败：{recovery_error}",
                        )
                        break
                continue
            consecutive_position_skips = 0
            if failure_status == "retry_pending" and first_failed_index is None:
                first_failed_index = failed_index
            save_collection_issue(
                connection, keyword, failed_index, "article_retry_exhausted",
                title, str(last_item_exception or "unknown article error"),
                normalize_video_card_fingerprint(raw_card_text),
            )
            # One malformed/unresponsive result must not prevent the keyword
            # from reaching later dates. Keep the earliest failed index as the
            # durable retry checkpoint while advancing locally in this run.
            result_index += 1
            processed_this_run += 1
            status = "running"
            update_task_state(
                connection, keyword, project,
                first_failed_index if first_failed_index is not None else result_index,
                old_streak, status,
                f"结果索引 {failed_index} 本轮采集失败；已进入{failure_status}并继续向后采集",
            )
            print(json.dumps({
                "event": "article_skipped_after_retries",
                "keyword": keyword,
                "index": failed_index,
                "title": title,
                "retry_checkpoint": first_failed_index,
                "retry_status": failure_status,
                "retry_count": retry_count,
            }, ensure_ascii=False), flush=True)
            if asr_backfill_mode and retry_fingerprint:
                # This Redis child owns one exact video only. Once that video
                # has been attempted and durably marked retry/dead-letter,
                # scanning the rest of the keyword cannot make it succeed and
                # only delays the next independent child task.
                status = (
                    "complete_retry_dead_letter"
                    if failure_status == "dead_letter"
                    else "partial_retry_pending"
                )
                update_task_state(
                    connection, keyword, project, result_index, old_streak, status,
                    f"视频指纹 {retry_fingerprint} 本轮补采失败；状态={failure_status}",
                )
                print(json.dumps({
                    "event": "redis_video_retry_attempt_finished",
                    "keyword": keyword,
                    "fingerprint": retry_fingerprint,
                    "status": failure_status,
                    "retry_count": retry_count,
                }, ensure_ascii=False), flush=True)
                break
            continue
        consecutive_position_skips = 0
        if status == "complete_retry_fingerprint":
            break
        if STOP_REQUESTED or (stop_file and stop_file.exists()):
            status = "partial_stopped"
            update_task_state(
                connection, keyword, project, result_index, old_streak, status,
                "graceful stop requested after current article",
            )
            break
        # 最新 does not prove strictly monotonic dates: pinned/interleaved
        # results may put another target-year video after an old one.
        if observe_old_date_window(config, old_streak):
            print(json.dumps({'event': 'old_date_window_observed_not_search_end',
                              'keyword': keyword, 'old_streak': old_streak},
                             ensure_ascii=False), flush=True)

    if status == "running" and max_results > 0 and processed_this_run >= max_results:
        status = "partial_limit"
        update_task_state(
            connection, keyword, project, result_index, old_streak, status,
            f"本次达到安全上限 {max_results} 条，需继续当前关键词",
        )
    if asr_queue is not None:
        # Navigation is non-blocking; the durable queue is drained only at the
        # keyword boundary. Count retry/dead-letter states after the drain.
        asr_queue.close_and_wait()
    # A previously failed video may disappear from the current search result
    # set. Count that as a bounded retry instead of leaving the keyword pending
    # forever with no observable way to reach the card.
    precise_miss=record_precise_search_miss(connection,keyword=keyword,project=project,
        retry_fingerprint=retry_fingerprint,scan_status=status,
        seen_fingerprints=seen_fingerprints_this_run,
        max_attempts=int(config.get('video_retry_max_attempts',3)),remote_store=remote_store)
    if precise_miss:
        print(json.dumps({'event':'redis_retry_verified_search_miss_counted',
            'fingerprint':retry_fingerprint,'status':precise_miss[0],'retry_count':precise_miss[1]},ensure_ascii=False),flush=True)
    current_checkpoints = load_video_checkpoints(connection, keyword)
    for fingerprint, checkpoint in current_checkpoints.items():
        if not should_age_missing_retry(
            fingerprint,
            checkpoint,
            seen_fingerprints_this_run,
            asr_backfill_mode=asr_backfill_mode,
            retry_fingerprint=retry_fingerprint,
        ):
            continue
        row_data = checkpoint.get("row") or {}
        if status not in {'complete_cutoff', 'complete_exhausted'}:
            continue
        missing_index = checkpoint.get('result_index')
        missing_status, missing_attempts = record_video_failure(
            connection,
            keyword=keyword,
            project=project,
            fingerprint=fingerprint,
            result_index=missing_index,
            title=str(row_data.get("title") or checkpoint.get("title") or ""),
            author_name=row_data.get("author_name") or checkpoint.get("author_name"),
            raw_card_text="",
            error="待重试视频在本轮完整搜索结果中未再次出现",
            screenshot_path=None,
            max_attempts=int(config.get("video_retry_max_attempts", 3)),
            remote_store=remote_store,
        )
        save_collection_issue(
            connection, keyword, missing_index, "retry_card_not_found",
            str(row_data.get("title") or ""),
            f"status={missing_status}; attempts={missing_attempts}", fingerprint,
        )
    retry_pending_count = connection.execute(
        "SELECT COUNT(*) FROM video_checkpoint WHERE keyword=? AND status='retry_pending'",
        (keyword,),
    ).fetchone()[0]
    dead_letter_count = connection.execute(
        "SELECT COUNT(*) FROM video_checkpoint WHERE keyword=? AND status='dead_letter'",
        (keyword,),
    ).fetchone()[0]
    if retry_pending_count > 0 and status != "partial_stopped":
        deferred_status = status_with_deferred_video_retries(status, retry_pending_count)
        if deferred_status.startswith("complete_"):
            # The yearly list has reached its real boundary.  Keep the keyword
            # scan complete and let Redis persist the individual failed-video
            # checkpoints in the independent retry stream.  Reclassifying the
            # whole keyword as partial here made XAUTOCLAIM run it from the top
            # indefinitely and starved all later keywords.
            status = deferred_status
            update_task_state(
                connection, keyword, project, result_index, old_streak, status,
                f"年度主扫描已完成；{retry_pending_count} 个视频已转入精准补采队列",
            )
        else:
            status = deferred_status
            update_task_state(
                connection, keyword, project,
                first_failed_index if first_failed_index is not None else result_index,
                old_streak, status,
                f"主扫描尚未到达边界；仍有 {retry_pending_count} 个视频等待后续运行精准重试",
            )
    elif first_failed_index is not None and not status.startswith("complete_") and status != "partial_stopped":
        # The scan was cut short before its time boundary, so the earliest
        # failed index is the durable resume checkpoint. A keyword that already
        # reached its cutoff/exhaustion boundary (or was safely stopped) keeps
        # its own status even when a few articles refused to open; otherwise the
        # worker would retry the same keyword forever and never advance.
        status = "partial_error"
        update_task_state(
            connection, keyword, project, first_failed_index, old_streak, status,
            "已继续采集后续结果，但仍有失败文章待下次从最早失败索引复查",
        )
    elif dead_letter_count > 0 and status.startswith("complete_"):
        status = f"{status}_with_dead_letters"
        update_task_state(
            connection, keyword, project, result_index, old_streak, status,
            f"完成扫描；{dead_letter_count} 个耗尽视频明确跳过，保留失败证据，不计采集成功",
        )
    complete = status.startswith("complete_")
    return {
        "keyword": keyword,
        "saved": saved,
        "failed": failed,
        "examined": processed_this_run,
        "next_result_index": result_index,
        "status": status,
        "complete": complete,
        "searchEndProof": search_end_proof,
        "retry_pending": int(retry_pending_count),
        "dead_letters": int(dead_letter_count),
    }


def export_template(template: Path, output: Path, connection: sqlite3.Connection, keyword: str | None = None) -> None:
    workbook = openpyxl.load_workbook(template)
    sheet = workbook["采集结果"]
    # Keep the supplied layout while adding the requested title field after keyword.
    headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
    if "文章标题" not in headers:
        sheet.insert_cols(3, 1)
        source = sheet.cell(1, 4)
        target = sheet.cell(1, 3)
        if source.has_style:
            from copy import copy
            target._style = copy(source._style)
        target.value = "文章标题"
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    if keyword:
        rows = connection.execute(
            """
            SELECT project, keyword, title, url, like_count, comment_count, favorite_count, share_count
            FROM captured_article WHERE keyword=?
            ORDER BY project, keyword, captured_at, url
            """,
            (keyword,),
        ).fetchall()
    else:
        rows = connection.execute(
            """
            SELECT project, keyword, title, url, like_count, comment_count, favorite_count, share_count
            FROM captured_article
            ORDER BY project, keyword, captured_at, url
            """
        ).fetchall()
    for values in rows:
        sheet.append(values)
    # Make the newly added title and URL columns readable without changing metric semantics.
    sheet.column_dimensions["C"].width = max(sheet.column_dimensions["C"].width or 0, 55)
    sheet.column_dimensions["D"].width = max(sheet.column_dimensions["D"].width or 0, 60)
    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=3, max_col=4):
        for cell in row_cells:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")
        sheet.row_dimensions[row_cells[0].row].height = 32
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = output.with_name(f".{output.name}.{os.getpid()}.tmp{output.suffix}")
    try:
        workbook.save(temporary_output)
        os.replace(temporary_output, output)
    finally:
        temporary_output.unlink(missing_ok=True)


def export_ods_local(
    output_dir: Path,
    connection: sqlite3.Connection,
    keywords: list[str] | None = None,
    article_keys: list[str] | None = None,
) -> dict:
    params: list[str] = []
    conditions = []
    if keywords:
        conditions.append(f"keyword IN ({','.join('?' for _ in keywords)})")
        params.extend(keywords)
    if article_keys:
        conditions.append(f"article_key IN ({','.join('?' for _ in article_keys)})")
        params.extend(article_keys)
    where = "WHERE " + " AND ".join(conditions) if conditions else ""
    rows = connection.execute(
        f"""
        SELECT article_key, project, keyword, title, url, author_name, tags, published_at,
               location, content, image_count, video_count, like_count, comment_count,
               favorite_count, share_count, video_content, video_duration, captured_at
        FROM captured_article {where}
        ORDER BY captured_at, keyword, url
        """,
        params,
    ).fetchall()
    records = []
    seen_urls = set()
    for row in rows:
        (article_key, project, keyword, title, url, author_name, tags, published_at,
         location, content, image_count, video_count, like_count, comment_count,
         favorite_count, share_count, video_content, video_duration, captured_at) = row
        # Keep the historical value in SQLite/backup, but never re-export an
        # identified internal-label contamination as a valid new subtitle.
        if has_untrusted_transcript(video_content):
            video_content = None
        if url in seen_urls:
            continue
        seen_urls.add(url)
        content_type = "视频" if (video_count or 0) else "图文" if (image_count or 0) else "纯文本"
        record = {field: None for field in ODS_FIELDS}
        record.update({
            "作者名字": author_name,
            "文章标题": title,
            "文章内容": None,
            "文章标签": tags,
            "文章时间地点": published_at,
            "文章点赞数": like_count,
            "文章收藏数": favorite_count,
            "文章评论数": comment_count,
            "文章链接": url,
            "文章id": article_key,
            "文章视频内容": video_content,
            "关键词": keyword,
            "作者id": None,
            "作者粉丝数": None,
            "作者点赞数": None,
            "文章分享数": share_count,
            "文章视频时长": video_duration,
            "平台": "微信视频号",
            "文章获取时间": captured_at,
            "项目": project,
            "地点": location,
            "内容类型": content_type,
        })
        records.append(record)
    output_dir.mkdir(parents=True, exist_ok=True)
    jsonl_path = output_dir / "ods_social_media_post_channels.jsonl"
    csv_path = output_dir / "ods_social_media_post_channels.csv"
    atomic_write_text(jsonl_path, "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in records))
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=ODS_FIELDS)
    writer.writeheader()
    writer.writerows(records)
    atomic_write_text(csv_path, "\ufeff" + csv_buffer.getvalue())
    comment_fields = [
        "id", "评论id", "评论用户链接", "评论用户id", "评论用户名字", "评论内容",
        "评论时间", "评论IP", "评论点赞数", "评论回复数", "文章链接", "文章id",
        "关键词", "评论获取时间", "平台", "分类", "项目", "清洗标注",
    ]
    comment_where = where
    comment_rows = connection.execute(
        f"""
        SELECT comment_key, article_key, project, keyword, article_url,
               user_url, user_id, user_name, content, published_at, location,
               like_count, reply_count, captured_at
        FROM captured_comment {comment_where}
        ORDER BY captured_at, keyword, article_url
        """,
        params,
    ).fetchall()
    comments = []
    for item in comment_rows:
        (comment_key, article_key, project, keyword, article_url, user_url, user_id,
         user_name, content, published_at, location, like_count, reply_count,
         captured_at) = item
        comments.append({
            "id": None, "评论id": comment_key, "评论用户链接": user_url, "评论用户id": user_id,
            "评论用户名字": user_name, "评论内容": content, "评论时间": published_at,
            "评论IP": location, "评论点赞数": like_count, "评论回复数": reply_count,
            "文章链接": article_url, "文章id": article_key, "关键词": keyword,
            "评论获取时间": captured_at, "平台": "微信视频号", "分类": None,
            "项目": project, "清洗标注": None,
        })
    comment_jsonl = output_dir / "ods_social_media_comment_channels.jsonl"
    comment_csv = output_dir / "ods_social_media_comment_channels.csv"
    atomic_write_text(comment_jsonl, "".join(json.dumps(record, ensure_ascii=False) + "\n" for record in comments))
    comment_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(comment_buffer, fieldnames=comment_fields)
    writer.writeheader()
    writer.writerows(comments)
    atomic_write_text(comment_csv, "\ufeff" + comment_buffer.getvalue())
    # video_dead_letter has no article_key column.  Its filter therefore only
    # uses the keyword subset; reusing ``params`` here is incorrect whenever
    # export-ods is called with both --keyword and --article-key (the query has
    # one placeholder but receives two values).  Keep an independent parameter
    # list so an ASR worker can export one hydrated article without crashing.
    dead_letter_params: list[str] = list(keywords or [])
    dead_letter_rows = connection.execute(
        """
        SELECT keyword, fingerprint, project, result_index, title, author_name,
               raw_card_text, retry_count, last_error, screenshot_path,
               created_at, updated_at
        FROM video_dead_letter
        """
        + (f" WHERE keyword IN ({','.join('?' for _ in keywords)})" if keywords else "")
        + " ORDER BY updated_at, keyword, result_index",
        dead_letter_params,
    ).fetchall()
    dead_letter_fields = [
        "关键词", "视频指纹", "项目", "结果索引", "标题", "作者名字",
        "原始卡片", "累计失败次数", "最后错误", "截图路径", "首次失败时间", "更新时间",
    ]
    dead_letters = [dict(zip(dead_letter_fields, row)) for row in dead_letter_rows]
    dead_jsonl = output_dir / "channels_video_dead_letter.jsonl"
    dead_csv = output_dir / "channels_video_dead_letter.csv"
    atomic_write_text(dead_jsonl, "".join(json.dumps(item, ensure_ascii=False) + "\n" for item in dead_letters))
    dead_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(dead_buffer, fieldnames=dead_letter_fields)
    writer.writeheader()
    writer.writerows(dead_letters)
    atomic_write_text(dead_csv, "\ufeff" + dead_buffer.getvalue())
    return {
        "records": len(records), "comments": len(comments), "dead_letters": len(dead_letters),
        "jsonl": str(jsonl_path), "csv": str(csv_path),
        "comment_jsonl": str(comment_jsonl), "comment_csv": str(comment_csv),
        "dead_letter_jsonl": str(dead_jsonl), "dead_letter_csv": str(dead_csv),
    }


def backfill_titles_live(keyword: str, config: dict, connection: sqlite3.Connection, screenshot_dir: Path) -> dict:
    """Recover titles for already captured rows by matching each live result URL."""
    targets = {row[0]: row[1] for row in connection.execute(
        "SELECT url, article_key FROM captured_article WHERE keyword=?", (keyword,)
    ).fetchall()}
    search_in_web(keyword, config)
    result_index = 0
    matched = 0
    seen = set()
    while targets:
        web = get_search_window()
        results = search_result_buttons(web)
        if result_index >= len(results):
            if results:
                results[-1].scroll_into_view()
            time.sleep(config.get("wait_page", 3.0))
            results = search_result_buttons(get_search_window())
            if result_index >= len(results):
                break
        result = results[result_index]
        title = result.window_text().strip()
        for attempt in range(3):
            try:
                result.invoke()
                time.sleep(config.get("wait_article", 3.0))
                wait_for(has_video_detail_window, 5.0)
                link = copy_article_link_from_menu(get_video_detail_window(), config)
                if link in targets:
                    connection.execute("UPDATE captured_article SET title=? WHERE keyword=? AND url=?",
                                       (title, keyword, link))
                    connection.commit()
                    matched += 1
                    targets.pop(link, None)
                    print(json.dumps({"keyword": keyword, "index": result_index, "title": title, "url": link, "matched": matched}, ensure_ascii=False))
                return_to_search_results(config)
                break
            except Exception as exc:
                if attempt == 2:
                    print(json.dumps({"keyword": keyword, "index": result_index, "error": str(exc)}, ensure_ascii=False))
                try:
                    if has_video_detail_window():
                        return_to_search_results(config)
                except Exception:
                    pass
        result_index += 1
    return {"keyword": keyword, "matched": matched, "remaining": len(targets), "examined": result_index}


def search_keyword(keyword: str, config: dict, screenshot_dir: Path) -> None:
    """进入 PC 微信搜一搜文章页。结果遍历和链接复制由配置坐标控制。"""
    win = get_wechat_window()
    keyboard.send_keys("^f", pause=0.02)
    time.sleep(config.get("wait_short", 0.8))
    paste_text(keyword)
    time.sleep(config.get("wait_search_suggest", 1.5))
    save_window_screenshot(win, screenshot_dir / f"{keyword}_01_suggest.png")
    if config.get("submit_search_with_enter", False):
        keyboard.send_keys("{ENTER}", pause=0.02)
    else:
        click_relative(win, config["points"]["soso_entry"])
    time.sleep(config.get("wait_page", 3.0))
    if "article_tab" in config["points"]:
        click_relative(win, config["points"]["article_tab"])
        time.sleep(config.get("wait_page", 3.0))
    save_window_screenshot(win, screenshot_dir / f"{keyword}_02_articles.png")


def copy_current_article_link(config: dict) -> str:
    win = get_wechat_window()
    before = get_clipboard()
    click_relative(win, config["points"]["article_more"])
    time.sleep(config.get("wait_short", 0.8))
    click_relative(win, config["points"]["copy_link"])
    time.sleep(config.get("wait_short", 0.8))
    value = get_clipboard().strip()
    if value == before or not value.startswith(("http://", "https://")):
        raise RuntimeError("未从剪贴板取得新的文章链接；请重新校准 article_more/copy_link")
    parsed = urlparse(value)
    is_channels_share = (
        parsed.netloc.lower() == "weixin.qq.com" and parsed.path.startswith("/sph/")
    ) or (
        parsed.netloc.lower() == "channels.weixin.qq.com"
        and parsed.path.startswith("/finder-preview/")
    )
    if not is_channels_share:
        raise RuntimeError(f"复制结果不是视频号分享链接，拒绝保存：{value}")
    return value


def command_ingest(args) -> None:
    connection = init_state(Path(args.state_db))
    row = ResultRow(
        project=args.project,
        keyword=args.keyword,
        title=args.title or "",
        url=args.url,
        like_count=parse_count(args.like),
        comment_count=parse_count(args.comment),
        favorite_count=parse_count(args.favorite),
        share_count=parse_count(args.share),
    )
    save_row(connection, row)
    print(json.dumps(asdict(row), ensure_ascii=False, indent=2))


def ingest_asr_enrichment(
    path: Path, connection: sqlite3.Connection, config: dict | None = None,
) -> dict:
    updated = 0
    examined = 0
    completed = 0
    partial_review = 0
    remote_store = RedisVideoCheckpointStore()
    if not path.exists():
        return {
            "examined": 0, "updated": 0, "completed": 0,
            "partial_review": 0, "path": str(path),
        }
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        examined += 1
        record = json.loads(line)
        text = clean_asr_text(record.get("videoContent"))
        url = str(record.get("url") or "").strip()
        if not text or not url:
            continue
        durations = [
            item.get("duration") for item in (record.get("videos") or [])
            if isinstance(item, dict) and item.get("duration")
        ]
        duration = round(max(durations)) if durations else None
        needs_review = has_untrusted_transcript(text) or asr_output_needs_review(text, duration, config=config)
        asr_status = "partial_review" if needs_review else "completed"
        checkpoint_status = "retry_pending" if needs_review else "completed"
        for key, word in connection.execute('SELECT article_key, keyword FROM captured_article WHERE url=?', (url,)).fetchall():
            stage_candidate(connection, key, word, text, approved=not needs_review,
                            reason='direct enrichment quality review' if needs_review else 'approved direct enrichment')
        if not needs_review:
            cursor = connection.execute(
                "UPDATE captured_article SET video_content=?, "
                "video_duration=COALESCE(video_duration, ?) WHERE url=?",
                (text, duration, url),
            )
            updated += cursor.rowcount
        checkpoint_rows = connection.execute(
            """
            SELECT vc.keyword, vc.fingerprint, vc.project, vc.retry_count,
                   ca.article_key, vc.comments_json
            FROM video_checkpoint vc
            JOIN captured_article ca
              ON ca.article_key=vc.article_key AND ca.keyword=vc.keyword
            WHERE ca.url=?
            """,
            (url,),
        ).fetchall()
        connection.commit()
        for keyword, fingerprint, project, retry_count, article_key, comments_json in checkpoint_rows:
            row = load_captured_result_row(connection, article_key, keyword)
            if row is None:
                continue
            comments = json.loads(comments_json) if comments_json else []
            upsert_video_checkpoint(
                connection,
                keyword=keyword,
                project=project,
                fingerprint=fingerprint,
                status=checkpoint_status,
                retry_count=int(retry_count or 0),
                row=row,
                comments=comments,
                asr_status=asr_status,
                last_error=(
                    "Direct ASR transcript is implausibly short for media duration"
                    if needs_review else None
                ),
                remote_store=remote_store,
            )
            if needs_review:
                partial_review += 1
            else:
                completed += 1
    connection.commit()
    return {
        "examined": examined, "updated": updated, "completed": completed,
        "partial_review": partial_review, "path": str(path),
    }


def collect_keyword_sequence(
    selected: list[tuple[str, str]], collect_one, allow_sample_transition: bool = False,
) -> list[dict]:
    """Run keywords sequentially and gate the transition on complete status.

    Kept separate from UI code so the critical "next keyword" behavior is
    regression-testable without driving a live WeChat account.
    """
    summaries = []
    for keyword, project in selected:
        summary = collect_one(keyword, project)
        summaries.append(summary)
        sample_continuable = {
            "partial_limit", "partial_error", "partial_position_error",
            "partial_retry_pending",
        }
        if not summary.get("complete") and not (
            allow_sample_transition and summary.get("status") in sample_continuable
        ):
            break
    return summaries


def main() -> None:
    parser = argparse.ArgumentParser(description="PC 微信视频号搜一搜采集器")
    sub = parser.add_subparsers(dest="command", required=True)

    probe = sub.add_parser("probe", help="检查微信窗口")

    search = sub.add_parser("search", help="打开一个关键词的视频搜索页")
    search.add_argument("keyword")
    search.add_argument("--config", default="collector_config.json")
    search.add_argument("--screenshots", default="work/screenshots")

    ingest = sub.add_parser("ingest", help="将已核验的链接和计数写入断点库")
    ingest.add_argument("--project", required=True)
    ingest.add_argument("--keyword", required=True)
    ingest.add_argument("--url", required=True)
    ingest.add_argument("--title", default="")
    ingest.add_argument("--like")
    ingest.add_argument("--comment")
    ingest.add_argument("--favorite")
    ingest.add_argument("--share")
    ingest.add_argument("--state-db", default="work/wechat_collect.sqlite")

    ingest_asr = sub.add_parser("ingest-asr-enrichment", help="把链接直解析得到的ASR结果回写断点库")
    ingest_asr.add_argument("--input", required=True)
    ingest_asr.add_argument("--state-db", default="work/wechat_collect.sqlite")
    ingest_asr.add_argument("--config", default="collector_config.json")
    ingest_asr.add_argument("--status-file")

    hydrate_retry = sub.add_parser(
        "hydrate-redis-retry",
        help="将一个Redis视频补采指纹水合到本机SQLite",
    )
    hydrate_retry.add_argument("--state-db", default="work/wechat_collect.sqlite")
    hydrate_retry.add_argument("--keyword", required=True)
    hydrate_retry.add_argument("--fingerprint", required=True)

    export = sub.add_parser("export", help="将断点库导出到用户结果表")
    export.add_argument("--template", required=True)
    export.add_argument("--output", required=True)
    export.add_argument("--state-db", default="work/wechat_collect.sqlite")
    export.add_argument("--keyword")

    export_ods = sub.add_parser("export-ods", help="按32字段ODS结构导出本地JSONL和CSV")
    export_ods.add_argument("--output-dir", required=True)
    export_ods.add_argument("--state-db", default="work/wechat_collect.sqlite")
    export_ods.add_argument("--keyword", action="append")
    export_ods.add_argument("--article-key", action="append")

    backfill = sub.add_parser("backfill-titles", help="从微信搜一搜结果回填已采集文章标题")
    backfill.add_argument("--keyword", required=True)
    backfill.add_argument("--config", default="collector_config.json")
    backfill.add_argument("--state-db", default="work/wechat_collect.sqlite")
    backfill.add_argument("--screenshots", default="work/screenshots")

    keywords = sub.add_parser("keywords", help="预览待采集关键词")
    keywords.add_argument("--input", required=True)
    keywords.add_argument("--limit", type=int)

    live = sub.add_parser("collect-live", help="在已经打开的微信搜一搜网页窗口中批量实采")
    live.add_argument("--input", required=True, help="关键词工作簿")
    live.add_argument("--config", default="collector_config.json")
    live.add_argument("--keyword", action="append", help="只采指定检索词，可重复传入")
    live.add_argument("--keyword-limit", type=int, default=4)
    live.add_argument("--cutoff", default=DEFAULT_START_DATE)
    # Zero means unlimited. Production workers must stop on the requested date
    # boundary or true list exhaustion, not after an arbitrary 100 articles.
    live.add_argument("--max-results", type=int, default=0)
    live.add_argument("--state-db", default="work/wechat_collect.sqlite")
    live.add_argument("--stop-file", help="存在时在当前文章完成后安全停止")
    live.add_argument("--screenshots", default="work/screenshots")
    live.add_argument("--comments", action="store_true", help="采集当前视频号页面公开可见评论（专题二）")
    live.add_argument("--asr-backfill", action="store_true", help="第二阶段：重扫当前关键词并只补充缺失的视频字幕")
    live.add_argument(
        "--retry-fingerprint",
        help="Redis精准补采：只打开与该卡片指纹匹配的视频",
    )
    live.add_argument(
        "--sample-transition", action="store_true",
        help="仅测试：达到--max-results后继续下一个指定关键词；生产任务禁止使用",
    )

    args = parser.parse_args()
    if args.command == "probe":
        win = get_wechat_window()
        rect = win.rectangle()
        print(json.dumps({
            "title": win.window_text(),
            "rect": [rect.left, rect.top, rect.right, rect.bottom],
            "dpi_scale": window_dpi_scale(win),
        }, ensure_ascii=False))
    elif args.command == "search":
        search_keyword(args.keyword, load_config(Path(args.config)), Path(args.screenshots))
    elif args.command == "ingest":
        command_ingest(args)
    elif args.command == "ingest-asr-enrichment":
        connection = init_state(Path(args.state_db))
        result = ingest_asr_enrichment(
            Path(args.input), connection,
            load_config(Path(args.config)),
        )
        if args.status_file:
            atomic_write_text(
                Path(args.status_file),
                json.dumps(result, ensure_ascii=False, indent=2),
            )
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif args.command == "hydrate-redis-retry":
        connection = init_state(Path(args.state_db))
        store = RedisVideoCheckpointStore()
        payload = store.load_all().get(args.fingerprint) if store.enabled else None
        project = str((payload or {}).get("project") or "")
        hydrated = hydrate_remote_video_checkpoints(
            connection, args.keyword, project, store, {args.fingerprint},
        ) if store.enabled else 0
        local_checkpoint = connection.execute(
            "SELECT article_key, title, author_name FROM video_checkpoint "
            "WHERE keyword=? AND fingerprint=?",
            (args.keyword, args.fingerprint),
        ).fetchone()
        print(json.dumps({
            "keyword": args.keyword,
            "fingerprint": args.fingerprint,
            "hydrated": hydrated,
            "article_key": local_checkpoint[0] if local_checkpoint else None,
            "title": local_checkpoint[1] if local_checkpoint else None,
            "author_name": local_checkpoint[2] if local_checkpoint else None,
            "retry_mode": (
                "asr_exact" if local_checkpoint and local_checkpoint[0]
                else "detail_locator"
            ),
        }, ensure_ascii=False, indent=2))
        if hydrated != 1:
            raise SystemExit(4)
    elif args.command == "export":
        connection = init_state(Path(args.state_db))
        export_template(Path(args.template), Path(args.output), connection, args.keyword)
    elif args.command == "export-ods":
        connection = init_state(Path(args.state_db))
        print(json.dumps(export_ods_local(
            Path(args.output_dir), connection, args.keyword, args.article_key,
        ), ensure_ascii=False, indent=2))
    elif args.command == "backfill-titles":
        connection = init_state(Path(args.state_db))
        result = backfill_titles_live(args.keyword, load_config(Path(args.config)), connection, Path(args.screenshots))
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif args.command == "keywords":
        print(json.dumps(load_keywords(Path(args.input), args.limit), ensure_ascii=False, indent=2))
    elif args.command == "collect-live":
        scan_started_at = datetime.now().astimezone().isoformat()
        signal.signal(signal.SIGINT, request_graceful_stop)
        if hasattr(signal, "SIGBREAK"):
            signal.signal(signal.SIGBREAK, request_graceful_stop)
        all_keywords = load_keywords(Path(args.input))
        selected = []
        requested = set(args.keyword or [])
        for keyword, project in all_keywords:
            if requested and keyword not in requested:
                continue
            selected.append((keyword, project))
            if not requested and len(selected) >= args.keyword_limit:
                break
        if requested and {keyword for keyword, _ in selected} != requested:
            missing = requested - {keyword for keyword, _ in selected}
            raise ValueError(f"关键词表中不存在：{sorted(missing)}")
        config = load_config(Path(args.config))
        if args.asr_backfill:
            config["enable_client_asr"] = True
            config["client_asr_background"] = True
            config["asr_backfill_mode"] = True
            if args.retry_fingerprint:
                config["retry_fingerprint"] = args.retry_fingerprint
        connection = init_state(Path(args.state_db))
        cutoff = datetime.fromisoformat(args.cutoff)
        def collect_one(keyword: str, project: str) -> dict:
            return collect_live_keyword(
                keyword=keyword,
                project=project,
                cutoff=cutoff,
                config=config,
                connection=connection,
                screenshot_dir=Path(args.screenshots),
                max_results=args.max_results,
                collect_comments=args.comments,
                stop_file=Path(args.stop_file) if args.stop_file else None,
            )
        summaries = collect_keyword_sequence(
            selected, collect_one, allow_sample_transition=args.sample_transition,
        )
        proof_file = os.environ.get('WECHAT_SEARCH_SCAN_PROOF_FILE')
        if proof_file:
            # Unique parent-provided file/token: never reuse a prior completed
            # keyword report to confirm this process's primary delivery.
            atomic_write_text(Path(proof_file), json.dumps({
                'policy': SEARCH_SCAN_POLICY,
                'runToken': os.environ.get('WECHAT_SEARCH_SCAN_RUN_TOKEN'),
                'taskId': os.environ.get('WECHAT_REDIS_TASK_ID'),
                'topic': os.environ.get('WECHAT_REDIS_TOPIC'),
                'keyword': os.environ.get('WECHAT_REDIS_KEYWORD'),
                'runId': os.environ.get('WECHAT_REDIS_RUN_ID'),
                'generation': os.environ.get('WECHAT_REDIS_GENERATION'),
                'startedAt': scan_started_at, 'finishedAt': datetime.now().astimezone().isoformat(),
                'summaries': summaries,
            }, ensure_ascii=False, indent=2))
        print(json.dumps({"summaries": summaries}, ensure_ascii=False, indent=2))
        # A partial keyword is not considered complete. The worker wrapper
        # will still export/import the rows already committed to SQLite, then
        # retry this same keyword from crawl_task.next_result_index.
        if summaries and any(
            not summary.get("complete")
            and not (
                args.sample_transition
                and summary.get("status") in {
                    "partial_limit", "partial_error", "partial_position_error",
                    "partial_retry_pending",
                }
            )
            for summary in summaries
        ):
            raise SystemExit(4)


if __name__ == "__main__":
    try:
        main()
    except RedisCheckpointUnavailable as error:
        print(json.dumps({"event": "redis_checkpoint_unavailable",
            "errorType": str(error), "localCheckpointRetained": True,
            "remoteSyncConfirmed": False, "exitCode": 75}), flush=True)
        raise SystemExit(75)
